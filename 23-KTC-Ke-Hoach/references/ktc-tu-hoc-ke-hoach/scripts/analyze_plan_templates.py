#!/usr/bin/env python3
"""Extract auditable formatting DNA from KTC XLSX/DOCX planning templates."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX


def color_value(color):
    if color is None:
        return None
    if color.type == "rgb":
        return color.rgb
    if color.type == "indexed":
        idx = color.indexed
        return COLOR_INDEX[idx] if isinstance(idx, int) and idx < len(COLOR_INDEX) else str(idx)
    if color.type == "theme":
        return f"theme:{color.theme}:tint:{color.tint}"
    return color.type


def side_value(side):
    if side is None or side.style is None:
        return None
    return {"style": side.style, "color": color_value(side.color)}


def cell_style(cell):
    return {
        "coordinate": cell.coordinate,
        "value": cell.value,
        "style_id": cell.style_id,
        "font": {
            "name": cell.font.name,
            "size": cell.font.sz,
            "bold": cell.font.b,
            "italic": cell.font.i,
            "color": color_value(cell.font.color),
        },
        "fill": {"type": cell.fill.fill_type, "fg": color_value(cell.fill.fgColor)},
        "alignment": {
            "horizontal": cell.alignment.horizontal,
            "vertical": cell.alignment.vertical,
            "wrap_text": cell.alignment.wrap_text,
            "text_rotation": cell.alignment.text_rotation,
        },
        "border": {
            "left": side_value(cell.border.left),
            "right": side_value(cell.border.right),
            "top": side_value(cell.border.top),
            "bottom": side_value(cell.border.bottom),
        },
        "number_format": cell.number_format,
    }


def analyze_xlsx(path: Path):
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        style_counts = Counter()
        nonempty = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    style_counts[cell.style_id] += 1
                    nonempty.append(cell)
        representative = []
        used_style_ids = set()
        for cell in nonempty:
            if cell.style_id not in used_style_ids:
                representative.append(cell_style(cell))
                used_style_ids.add(cell.style_id)
        page_margins = ws.page_margins
        sheets.append({
            "title": ws.title,
            "dimensions": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
            "print_area": str(ws.print_area),
            "print_title_rows": ws.print_title_rows,
            "sheet_view": {"show_grid_lines": ws.sheet_view.showGridLines, "zoom": ws.sheet_view.zoomScale},
            "page_setup": {
                "orientation": ws.page_setup.orientation,
                "paper_size": ws.page_setup.paperSize,
                "fit_to_width": ws.page_setup.fitToWidth,
                "fit_to_height": ws.page_setup.fitToHeight,
                "horizontal_centered": ws.print_options.horizontalCentered,
                "vertical_centered": ws.print_options.verticalCentered,
                "margins": {k: getattr(page_margins, k) for k in ("left", "right", "top", "bottom", "header", "footer")},
            },
            "column_widths": {k: v.width for k, v in ws.column_dimensions.items() if v.width},
            "row_heights": {str(k): v.height for k, v in ws.row_dimensions.items() if v.height},
            "style_usage": dict(style_counts),
            "representative_styles": representative,
            "header_footer": {
                "odd_header": {"left": ws.oddHeader.left.text, "center": ws.oddHeader.center.text, "right": ws.oddHeader.right.text},
                "odd_footer": {"left": ws.oddFooter.left.text, "center": ws.oddFooter.center.text, "right": ws.oddFooter.right.text},
            },
        })
    return {"file": path.name, "type": "xlsx", "sheets": sheets}


def analyze_docx(path: Path):
    doc = Document(path)
    section_data = []
    for section in doc.sections:
        section_data.append({
            "page_width_cm": round(section.page_width.cm, 2),
            "page_height_cm": round(section.page_height.cm, 2),
            "orientation": str(section.orientation),
            "margins_cm": {
                "top": round(section.top_margin.cm, 2),
                "bottom": round(section.bottom_margin.cm, 2),
                "left": round(section.left_margin.cm, 2),
                "right": round(section.right_margin.cm, 2),
            },
            "header_distance_cm": round(section.header_distance.cm, 2),
            "footer_distance_cm": round(section.footer_distance.cm, 2),
        })
    style_counts = Counter(p.style.name for p in doc.paragraphs if p.text.strip())
    para_samples = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        run = next((r for r in p.runs if r.text.strip()), None)
        para_samples.append({
            "text": p.text[:180],
            "style": p.style.name,
            "alignment": str(p.alignment),
            "left_indent_cm": round(p.paragraph_format.left_indent.cm, 2) if p.paragraph_format.left_indent else None,
            "first_line_indent_cm": round(p.paragraph_format.first_line_indent.cm, 2) if p.paragraph_format.first_line_indent else None,
            "space_before_pt": p.paragraph_format.space_before.pt if p.paragraph_format.space_before else None,
            "space_after_pt": p.paragraph_format.space_after.pt if p.paragraph_format.space_after else None,
            "line_spacing": p.paragraph_format.line_spacing,
            "run": None if run is None else {
                "font": run.font.name,
                "size_pt": run.font.size.pt if run.font.size else None,
                "bold": run.bold,
                "italic": run.italic,
                "underline": bool(run.underline),
            },
        })
    tables = []
    for idx, table in enumerate(doc.tables, 1):
        tables.append({
            "index": idx,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "style": table.style.name if table.style else None,
            "first_rows": [[c.text[:100] for c in row.cells] for row in table.rows[:3]],
        })
    return {
        "file": path.name,
        "type": "docx",
        "sections": section_data,
        "paragraph_style_usage": dict(style_counts),
        "paragraph_samples": para_samples[:40],
        "tables": tables,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--output", required=True)
    args = ap.parse_args()
    reports = []
    for name in args.inputs:
        path = Path(name)
        if path.suffix.lower() == ".xlsx":
            reports.append(analyze_xlsx(path))
        elif path.suffix.lower() == ".docx":
            reports.append(analyze_docx(path))
        else:
            raise SystemExit(f"Unsupported file: {path}")
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(reports, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
