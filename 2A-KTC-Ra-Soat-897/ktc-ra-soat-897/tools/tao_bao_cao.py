#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tao_bao_cao.py – Dựng báo cáo rà soát .docx theo khuôn cố định (Skill ktc-ra-soat-897 v4.5)

Người rà soát (AI) chỉ điền DỮ LIỆU (tệp JSON); script:
  * nhập nguyên kết quả đo thể thức của kiem_tra_the_thuc.py (không cho sửa mức),
  * nhập kết quả cổng hồ sơ của kiem_ke_ho_so.py (từ chối dựng nếu cổng CHẶN),
  * kiểm tính hợp lệ (mã, mức, bằng chứng hai phía, 6 giá trị ma trận…),
  * TỰ TÍNH Bảng 1, Bảng 2, mức độ tuân thủ và NHÃN KẾT LUẬN theo
    references/Quy-Tac-KTC/02-Quy-Tac-Xep-Muc.md,
  * dựng .docx với cấu trúc, định dạng cố định.

Cách dùng:
  python3 tao_bao_cao.py du_lieu.json --cong cong.json --the-thuc tt.json [--the-thuc tt2.json] -o bao_cao.docx
  python3 tao_bao_cao.py du_lieu.json ... --chi-kiem      (chỉ kiểm hợp lệ, không dựng tệp)
Mẫu dữ liệu: assets/Mau-Du-Lieu-Bao-Cao.json
"""
import sys, json, re, argparse, datetime, unicodedata
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PHIEN_BAN = 'ktc-ra-soat-897 v4.5 (27/10/2026)'
TRANG_THAI = 'BẢN NHÁP RÀ SOÁT – CHƯA ĐƯỢC PHÊ DUYỆT CHÍNH THỨC'

# ---- mức cho phép theo mã (02-Quy-Tac-Xep-Muc.md mục 2) ----
MUC_CHO_PHEP = {
    'TT-34': {2, 3}, 'TB-01': {3}, 'TB-02': {4}, 'TB-03': {3},
    'TD-01': {2}, 'TD-02': {2}, 'TD-05': {2}, 'TD-06': {2}, 'TD-03': {3}, 'TD-04': {3}, 'TD-07': {3},
    'CC-01': {1}, 'CC-02': {2}, 'CC-03': {1}, 'CC-04': {1}, 'CC-05': {1, 2}, 'CC-06': {2},
    'TQ-01': {1}, 'TQ-02': {2},
    'DC-01': {1, 2}, 'DC-02': {2}, 'DC-03': {1, 2}, 'DC-04': {2}, 'DC-05': {2}, 'DC-06': {2},
    'ND-01': {1, 2}, 'ND-02': {1}, 'ND-03': {2},
    'VP-01': {2, 3}, 'GY-01': {4}, 'DG-01': {2}, 'DG-02': {2},
}
# mã TT, TB, TD, CC do script đo (bảng 01, 02) – chỉ các mã này được nhập từ script hoặc thêm tay có lý do
MA_SCRIPT = {'TT-%02d' % k for k in list(range(1, 5)) + list(range(10, 25)) + list(range(30, 47)) + [50, 51]} | \
            {'TB-01', 'TB-02', 'TB-03', 'TD-01', 'TD-02', 'TD-03', 'TD-04', 'TD-05', 'CC-01', 'CC-02', 'CC-06'}
NHOM = {'TT': 'Thể thức', 'TB': 'Thể thức', 'TD': 'Thể thức', 'DG': 'Thể thức',
        'CC': 'Nội dung', 'TQ': 'Nội dung', 'DC': 'Nội dung', 'ND': 'Nội dung',
        'VP': 'Chất lượng soạn thảo', 'GY': 'Chất lượng soạn thảo'}
NHOM_LIST = ['Thể thức', 'Nội dung', 'Điều khoản bắt buộc', 'Chất lượng soạn thảo']
MA_TRAN_GT = ['Khớp', 'Mâu thuẫn', 'Hẹp hơn', 'Rộng hơn', 'Văn bản liên quan không quy định', 'Chưa kiểm']
TRANG_THAI_DC = ['Đã sửa đạt', 'Sửa một phần', 'Chưa sửa', 'Không tiếp thu – có giải trình', 'Không còn áp dụng',
                 'Kiến nghị lần trước không còn phù hợp']  # PB-09
NGOAC = re.compile(r'[“"«].+[”"»]', re.S)
NGOAC2 = re.compile(r'[“"«][^”"»]+[”"»]')
TRONG_YEU = ['thẩm quyền', 'thời hạn', 'điều cấm', 'cho phép', 'đối tượng', 'hiệu lực']
TEN_MUC = {1: 'Mức 1 – Bắt buộc sửa', 2: 'Mức 2 – Cần sửa', 3: 'Mức 3 – Nên sửa', 4: 'Mức 4 – Góp ý nâng cao'}


def muc_so(v):
    m = re.search(r'([1-4])', str(v)); return int(m.group(1)) if m else None


def nhom_cua(ma):
    if ma == 'ND-02': return 'Điều khoản bắt buộc'
    return NHOM.get(ma.split('-')[0], 'Nội dung')


TEP_IDS = []  # tên tệp dự thảo/văn bản kèm theo theo cổng hồ sơ – đặt trong kiem() (v4.5, PB5-WORK-14)


def tep_cua(v):
    t = unicodedata.normalize('NFC', v or '')
    hit = [x for x in TEP_IDS if x and x in t]
    return max(hit, key=len) if hit else None


def norm_vt(v):
    """Khóa vị trí chuẩn hóa: tập {tệp, phần văn bản, điều, khoản, điểm, phụ lục} – để “Quy định, Điều 7 khoản 2” và
    “Điều 7, khoản 2 Quy định” được coi là một vị trí (PB2-12); v4.5: vị trí ghi tên tệp thì tệp cũng là một phần của vị trí."""
    tp = tep_cua(v)
    t = unicodedata.normalize('NFC', (v or '').replace(tp, ' ') if tp else (v or '')).lower()
    toks = set()
    if tp: toks.add(('tep', tp))
    for pat, ten in ((r'điều\s*(\d+)', 'dieu'), (r'khoản\s*(\d+)', 'khoan'), (r'điểm\s*([a-zđ])\b', 'diem'), (r'phụ lục\s*([ivxlc\d]+)', 'pl'), (r'đoạn\s*(\d+)', 'doan')):
        for m in re.finditer(pat, t): toks.add((ten, m.group(1)))
    for phan in ('quy định', 'quy chế', 'quyết định', 'kế hoạch', 'điều lệ', 'nội quy'):
        if phan in t: toks.add(('phan', phan))
    return frozenset(toks) if any(k not in ('phan', 'tep') for k, _ in toks) else re.sub(r'\s+', ' ', t.strip())


def trung_vt(a, b):
    """v4.3 (TH-12): hai vị trí trùng hoặc lồng nhau (cha – con): “Điều 7” chứa “Điều 7 khoản 2”."""
    if isinstance(a, frozenset) and isinstance(b, frozenset):
        ta = {x for x in a if x[0] == 'tep'}; tb = {x for x in b if x[0] == 'tep'}
        if ta and tb and ta != tb: return False  # v4.5 (PB5-WORK-14): khác tệp thì không cùng vị trí
        ka = {x for x in a if x[0] not in ('phan', 'tep')}; kb = {x for x in b if x[0] not in ('phan', 'tep')}
        pa = {x for x in a if x[0] == 'phan'}; pb = {x for x in b if x[0] == 'phan'}
        if pa and pb and not (pa & pb): return False
        return bool(ka) and bool(kb) and (ka <= kb or kb <= ka)
    return a == b


def chuan_nv(s_):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', s_ or '').replace('…', ' ').strip(' “”"«».')).strip().lower()


TRAI = [('cấm', 'cho phép'), ('cấm', 'được phép'), ('bắt buộc', 'khuyến khích'), ('tối đa', 'tối thiểu'), ('tăng', 'giảm')]
DV_TG = ('giờ', 'ngày', 'tuần', 'tháng', 'quý', 'năm', 'phút')
SO_CHU = {'một': 1, 'hai': 2, 'ba': 3, 'bốn': 4, 'tư': 4, 'năm': 5, 'sáu': 6, 'bảy': 7, 'tám': 8, 'chín': 9, 'mười': 10, 'mươi': 10}
NEG_RE = r'(?<!\w)(không|chưa|chẳng|cấm|thiếu)(?!\w)'


def dau_hieu(s_):
    """v4.5: chữ ký so sánh của một phía – (i) dãy (số, đơn vị) THEO THỨ TỰ, bỏ số dẫn Điều/khoản/điểm/Chương/Mục/Phụ lục/số hiệu và năm;
    số viết bằng chữ (“năm ngày”) được đổi ra số; (ii) dãy phủ định theo từng vế (tách ở dấu phẩy, chấm phẩy, “và”, “nhưng”, “còn”)."""
    t = chuan_nv(s_)
    so = []
    for m in re.finditer(r'(?<![\w/])(\d+(?:[.,]\d+)?|' + '|'.join(SO_CHU) + r')\s*(' + '|'.join(DV_TG) + r')?(?!\w)', t):
        truoc = t[max(0, m.start() - 12):m.start()]
        if re.search(r'(điều|khoản|điểm|chương|mục|phụ lục|số|năm|lần|ngày)\s*$', truoc) and not m.group(2): continue
        v = m.group(1)
        if v in SO_CHU:
            if not m.group(2): continue  # chữ số chỉ tính khi đi kèm đơn vị (“năm ngày”), tránh nhầm “năm” (năm học)
            v = str(SO_CHU[v])
        so.append((str(int(float(v.replace(',', '.')))) if re.match(r'^\d+([.,]0+)?$', v) else v, m.group(2) or ''))
    ve = [x for x in re.split(r',|;|(?<!\w)(?:và|nhưng|còn|song)(?!\w)', t) if x.strip()]
    pd = [bool(re.search(NEG_RE, x)) for x in ve]
    return so, pd


def trai_nghia(a_, b_):
    """v4.4 (PB4-GPT-07), v4.5 (PB5-GPT-03, PB5-WORK-15): dấu hiệu hai phía không khớp. Trả về mô tả hoặc ''."""
    sa, pa = dau_hieu(a_); sb, pb = dau_hieu(b_)
    if (sa or sb) and sa != sb:
        return 'khác số liệu/đơn vị theo thứ tự (%s / %s)' % (', '.join(' '.join(x).strip() for x in sa) or 'không có', ', '.join(' '.join(x).strip() for x in sb) or 'không có')
    if len(pa) == len(pb) and pa != pb: return 'khác phủ định theo từng vế'
    if len(pa) != len(pb) and sum(pa) != sum(pb): return 'khác số vế phủ định'
    A, B = chuan_nv(a_), chuan_nv(b_)
    def co(s_, w_): return re.search(r'(?<!\w)' + w_ + r'(?!\w)', s_) is not None
    for x, y in TRAI:
        if (co(A, x) and co(B, y) and not co(B, x)) or (co(A, y) and co(B, x) and not co(A, x)):
            return 'trái nghĩa (“%s” / “%s”)' % (x, y)
    return ''


GIAI_TRINH = []


def kiem(d, cong, tts):
    loi = []
    del GIAI_TRINH[:]
    TEP_IDS[:] = sorted(((cong or {}).get('sha256') or {}).keys(), key=len, reverse=True)
    if cong is None:
        loi.append('Thiếu kết quả cổng hồ sơ (--cong). Phải chạy kiem_ke_ho_so.py trước.')
    elif cong.get('ket_luan') != 'DU':
        loi.append('Cổng hồ sơ CHẶN – không được lập báo cáo. Gửi nguyên văn thông báo chặn cho người dùng.')
    if not tts:
        loi.append('Thiếu kết quả đo thể thức (--the-thuc). Phải chạy kiem_tra_the_thuc.py cho dự thảo (và từng tệp phụ lục riêng nếu có).')
    if cong is not None:
        # v4.3 (TH-11): kết quả đo thể thức phải thuộc đúng hồ sơ đã qua cổng (khớp mã băm tệp)
        sha_cong = cong.get('sha256') or {}
        if not sha_cong:
            loi.append('Kết quả cổng hồ sơ không có mã băm tệp (chạy lại kiem_ke_ho_so.py của gói v4.5).')
        for t in tts:
            if not t.get('sha256'):
                loi.append('Kết quả đo thể thức “%s” không có mã băm tệp (chạy lại kiem_tra_the_thuc.py của gói v4.5).' % t.get('_ten_tep', ''))
            elif sha_cong and t['sha256'] not in sha_cong.values():
                loi.append('Kết quả đo thể thức “%s” (tệp %s) không thuộc hồ sơ đã qua cổng – phải đo đúng tệp dự thảo/phụ lục của hồ sơ này.' % (t.get('_ten_tep', ''), t.get('tep', '?')))
        vt_c = cong.get('vai_tro', {})
        da_do = {t.get('sha256') for t in tts}  # v4.4 (PB4-WORK-07): đối chiếu theo mã băm từng tệp, không theo tên tệp
        for ten, v in vt_c.items():
            if v in ('du_thao', 'kem_theo') and ten.lower().endswith('.docx') and sha_cong.get(ten) not in da_do:
                loi.append('Chưa đo thể thức tệp “%s” (%s) – chạy kiem_tra_the_thuc.py và truyền thêm --the-thuc.' % (ten, 'dự thảo' if v == 'du_thao' else 'phụ lục/văn bản kèm theo'))
        # v4.3 (TH-04, quyết định PB2-07 ngày 27/10/2026): mọi tệp “PHẢI XEM ẢNH” phải có kết luận xem ảnh
        xa = {x.get('tep'): x for x in d.get('xem_anh', [])}
        for ten, so_anh, _ in cong.get('anh_can_xem', []):
            x = xa.get(ten)
            if not x or x.get('ket_luan') not in ('so_do', 'van_ban_chup') or int(x.get('so_anh_da_xem', 0)) < so_anh or not (x.get('mo_ta') or '').strip():
                loi.append('Tệp “%s” có %d ảnh “PHẢI XEM ẢNH” nhưng “xem_anh” chưa có kết luận đủ (tep, so_anh_da_xem ≥ %d, ket_luan = so_do | van_ban_chup, mo_ta).' % (ten, so_anh, so_anh))
            elif x.get('ket_luan') == 'van_ban_chup':
                loi.append('Tệp “%s” chứa ảnh văn bản chụp/scan → hồ sơ CHẶN theo G1: không lập báo cáo; gửi thông báo yêu cầu bản chữ.' % ten)
    for i, f in enumerate(d.get('phat_hien', []), 1):
        ma = f.get('ma', '')
        if not re.match(r'^(TT|TB|TD|CC|TQ|DC|ND|VP|GY|DG)-\d{2}$', ma):
            loi.append('Phát hiện %d: mã “%s” không hợp lệ.' % (i, ma)); continue
        if ma.startswith('TT-') and f.get('_nguon') != 'script' and not (f.get('thu_cong_ly_do') or '').strip():
            loi.append('Phát hiện %d (%s): mã TT do script đo; chỉ thêm tay khi script không đo được hạng mục này – phải ghi “thu_cong_ly_do” và số đo cụ thể.' % (i, ma))
        if ma not in MUC_CHO_PHEP and ma not in MA_SCRIPT:
            loi.append('Phát hiện %d: mã “%s” không có trong bảng mã 02-Quy-Tac-Xep-Muc.md – không tự đặt mã mới.' % (i, ma)); continue
        m = muc_so(f.get('muc'))
        cho = MUC_CHO_PHEP.get(ma, {2} if ma.startswith('TT-') else None)
        if m is None: loi.append('Phát hiện %d (%s): thiếu mức.' % (i, ma))
        elif cho and m not in cho:
            loi.append('Phát hiện %d (%s): mức %d không đúng bảng mã (cho phép: %s).' % (i, ma, m, ', '.join('Mức %d' % x for x in sorted(cho))))
        if m in (1, 2):
            if not (f.get('nguyen_van_du_thao') or '').strip():
                loi.append('Phát hiện %d (%s, Mức %d): thiếu nguyên văn/số đo của dự thảo.' % (i, ma, m))
            if not (f.get('nguyen_van_can_cu') or '').strip() or not (f.get('can_cu') or '').strip():
                loi.append('Phát hiện %d (%s, Mức %d): thiếu căn cứ hoặc nguyên văn căn cứ.' % (i, ma, m))
    for i, r in enumerate(d.get('ma_tran', []), 1):
        kq = r.get('ket_qua', '')
        base = kq if kq in MA_TRAN_GT else ('Chưa kiểm' if kq.startswith('Chưa kiểm') else ('Khớp' if kq == 'Khớp – có điều khoản thay thế' else None))
        if base is None:
            loi.append('Ma trận dòng %d: kết quả “%s” ngoài 6 giá trị cho phép.' % (i, kq))
        if base in ('Khớp', 'Mâu thuẫn', 'Hẹp hơn', 'Rộng hơn'):
            if not (r.get('du_thao_nguyen_van') or '').strip() or not (r.get('vblq_nguyen_van') or '').strip():
                loi.append('Ma trận dòng %d (%s): thiếu nguyên văn hai phía.' % (i, kq))
        if base in ('Mâu thuẫn', 'Hẹp hơn', 'Rộng hơn'):
            mp = (r.get('ma_phat_hien') or '').split('#')[0]
            can = {'Mâu thuẫn': 'DC-01', 'Hẹp hơn': 'DC-02', 'Rộng hơn': 'DC-03'}[base]
            if mp != can:  # v4.3 (TH-06): đúng mã theo loại kết quả
                loi.append('Ma trận dòng %d (%s): phải chuyển thành phát hiện %s ở Phần II (trường “ma_phat_hien” = “%s”, hiện “%s”).' % (i, kq, can, can, mp or 'trống'))
            elif not any(f.get('ma') == can and muc_so(f.get('muc')) in (1, 2) for f in d.get('phat_hien', [])):
                loi.append('Ma trận dòng %d (%s): không có phát hiện %s Mức 1–2 tương ứng ở Phần II.' % (i, kq, can))
            elif len((cong or {}).get('cac_ban_du_thao') or []) > 1 and not tep_cua(r.get('du_thao_vi_tri')):  # v4.5 (PB5-WORK-14)
                loi.append('Ma trận dòng %d (%s): hồ sơ có %d bản – “du_thao_vi_tri” phải ghi tên tệp như cổng hồ sơ ghi (ví dụ “%s – Điều 7 khoản 2”).' % (
                    i, kq, len(cong['cac_ban_du_thao']), cong['cac_ban_du_thao'][0]))
            elif not any(f.get('ma') == can and muc_so(f.get('muc')) in (1, 2) and trung_vt(norm_vt(r.get('du_thao_vi_tri')), norm_vt(f.get('vi_tri')))
                         for f in d.get('phat_hien', [])):  # v4.4 (PB4-GPT-06, PB4-WORK-10)
                loi.append('Ma trận dòng %d (%s, “%s”): phát hiện %s ở Phần II phải đặt CÙNG vị trí (hoặc vị trí cha – con) với dòng ma trận; hiện các phát hiện %s ở: %s.' % (
                    i, kq, r.get('du_thao_vi_tri'), can, can, '; '.join(f.get('vi_tri', '') for f in d.get('phat_hien', []) if f.get('ma') == can)[:200]))
    # PB-11 / B2 đề xuất 6, 7: “Khớp” phải trích nguyên văn trong ngoặc kép và nêu yếu tố khớp
    vt_loi = [norm_vt(f.get('vi_tri')) for f in d.get('phat_hien', []) if f.get('vi_tri')]
    for i, r in enumerate(d.get('ma_tran', []), 1):
        kq = r.get('ket_qua', '')
        if kq.startswith('Khớp') or kq in ('Mâu thuẫn', 'Hẹp hơn', 'Rộng hơn'):
            if not NGOAC.search(r.get('du_thao_nguyen_van') or '') or not NGOAC.search(r.get('vblq_nguyen_van') or ''):
                loi.append('Ma trận dòng %d (%s): nguyên văn hai phía phải đặt trong ngoặc kép “…”.' % (i, kq))
        if kq == 'Khớp – có điều khoản thay thế':  # v4.5: ngoại lệ bảng 02 mục 4 – nguyên văn điều khoản thay thế thay cho yeu_to_khop
            if not NGOAC.search(r.get('nguyen_van_thay_the') or ''):
                loi.append('Ma trận dòng %d (Khớp – có điều khoản thay thế): phải ghi “nguyen_van_thay_the” – nguyên văn điều khoản thay thế/sửa đổi trong ngoặc kép.' % i)
        elif kq.startswith('Khớp'):
            ytk = r.get('yeu_to_khop')
            thieu = [k for k in ('chu_the', 'hanh_vi', 'muc_thoi_han') if not isinstance(ytk, dict) or not (ytk.get(k) or '').strip()]
            if thieu:
                loi.append('Ma trận dòng %d (Khớp): “yeu_to_khop” phải là bảng {chu_the, hanh_vi, muc_thoi_han} có đủ 3 yếu tố của hai bên (thiếu: %s).' % (i, ', '.join(thieu)))
            elif any(re.search(r'trùng từ khóa|tương tự|gần giống|na ná', str(v), re.I) for v in ytk.values()):
                loi.append('Ma trận dòng %d (Khớp): yếu tố khớp chỉ là “trùng từ khóa/tương tự” – không đủ để ghi Khớp.' % i)
            elif any(len(NGOAC2.findall(str(v))) < 2 for v in ytk.values()):  # v4.3 (TH-12)
                loi.append('Ma trận dòng %d (Khớp): mỗi yếu tố khớp (chủ thể, hành vi, mức/thời hạn) phải trích HAI PHÍA trong ngoặc kép (dự thảo “…”; văn bản liên quan “…”).' % i)
            else:
                gt = r.get('giai_trinh_khop') or {}
                for k_, v_ in ytk.items():  # v4.4–v4.5: hai phía khác số/đơn vị theo thứ tự, khác phủ định, trái nghĩa thì không phải Khớp
                    ly = trai_nghia(*[x.strip('“”"«»') for x in NGOAC2.findall(str(v_))[:2]])
                    if not ly: continue
                    if isinstance(gt, dict) and len((gt.get(k_) or '').strip()) >= 20:  # quyết định F2: giải trình, in ở Bảng 3
                        GIAI_TRINH.append(dict(dong=i, yeu_to=k_, dau_hieu=ly, giai_trinh=gt[k_].strip()))
                        continue
                    loi.append('Ma trận dòng %d (Khớp), yếu tố “%s”: hai phía %s – không được ghi Khớp; xem lại thành Mâu thuẫn/Hẹp hơn/Rộng hơn (DC-01…03), hoặc nếu đọc nghĩa vẫn khớp thì ghi “giai_trinh_khop”: {“%s”: “lý do ≥ 20 ký tự”} (quyết định F2 – in ở Bảng 3).' % (i, k_, ly, k_))
            if any(trung_vt(norm_vt(r.get('du_thao_vi_tri')), v) for v in vt_loi):
                loi.append('Ma trận dòng %d: vị trí “%s” ghi Khớp nhưng lại có phát hiện ở Phần II.' % (i, r.get('du_thao_vi_tri')))
    for f in d.get('dat', []):
        if f.get('nhom') in ('Nội dung', 'Điều khoản bắt buộc') and not (f.get('vi_tri') or '').strip():
            loi.append('Mục Đạt “%s” (nhóm %s): thiếu vị trí trong dự thảo (“vi_tri”).' % (f.get('noi_dung', '')[:40], f.get('nhom')))
        if f.get('vi_tri') and any(trung_vt(norm_vt(f.get('vi_tri')), v) for v in vt_loi):
            loi.append('Mục Đạt “%s” (vị trí %s) trùng vị trí có phát hiện ở Phần II – một vị trí không vừa lỗi vừa Đạt.' % (f.get('noi_dung', '')[:40], f.get('vi_tri')))
    # PB-12: cấu trúc
    vong = (cong or {}).get('vong', 1)
    if d.get('che_do') == 'ĐC':
        if vong < 2:
            loi.append('Chế độ ĐC nhưng cổng hồ sơ ghi lần rà soát %s – chế độ ĐC chỉ dùng từ lần 2.' % vong)
        hq = [r for r in d.get('hoi_quy', []) if (r.get('ma_cu') or '').strip() and (r.get('ket_qua') or '').strip()]
        if not hq or len(hq) < len(d.get('hoi_quy', [])):
            loi.append('Chế độ ĐC: II.C Kiểm hồi quy (“hoi_quy”) phải có ít nhất một dòng và mọi dòng có đủ “ma_cu”, “ket_qua”.')
    n_uu = len(d.get('uu_tien', []))
    can = min(3, len(d.get('phat_hien', [])))
    if n_uu > 5 or n_uu < can:
        loi.append('Phần VI: nêu %d–5 vấn đề ưu tiên (hiện %d).' % (can, n_uu))
    for i, r in enumerate(d.get('doi_chieu', []), 1):
        if r.get('trang_thai') in ('Sửa một phần', 'Chưa sửa'):
            mp = (r.get('ma_phat_hien') or '').split('#')[0]
            if not mp or not any(f.get('ma') == mp for f in d.get('phat_hien', [])):
                loi.append('Đối chiếu dòng %d (%s): phần còn lại phải ghi thành phát hiện ở Phần II (trường “ma_phat_hien”).' % (i, r.get('trang_thai')))
        if r.get('trang_thai') == 'Kiến nghị lần trước không còn phù hợp' and not (r.get('can_cu_dung') or '').strip():
            loi.append('Đối chiếu dòng %d: “Kiến nghị lần trước không còn phù hợp” phải nêu căn cứ đúng (“can_cu_dung”).' % i)
    for i, r in enumerate(d.get('doi_chieu', []), 1):
        if r.get('trang_thai') not in TRANG_THAI_DC:
            loi.append('Đối chiếu dòng %d: trạng thái “%s” không hợp lệ.' % (i, r.get('trang_thai')))
        if r.get('trang_thai') in ('Đã sửa đạt', 'Sửa một phần', 'Chưa sửa') and not (r.get('nguyen_van_hien_hanh') or '').strip():
            loi.append('Đối chiếu dòng %d: thiếu nguyên văn/số đo bản hiện hành.' % i)
    if d.get('che_do') == 'ĐC' and not d.get('doi_chieu'):
        loi.append('Chế độ ĐC nhưng chưa có bảng đối chiếu việc sửa (“doi_chieu”).')
    for f in d.get('dat', []):
        if not (f.get('nguyen_van') or f.get('so_do') or '').strip():
            loi.append('Mục Đạt “%s”: thiếu nguyên văn/số đo – không được ghi Đạt.' % f.get('noi_dung', '')[:50])
    for k in ('ten_du_thao', 'don_vi_soan_thao', 'loai_van_ban', 'nguoi_kiem_tra'):
        if not d.get(k): loi.append('Thiếu trường “%s”.' % k)
    kieu = d.get('kieu_bao_cao', 'chinh_thuc')
    if kieu not in KIEU:
        loi.append('“kieu_bao_cao” chỉ nhận: %s.' % ', '.join(KIEU))
    tv = (cong or {}).get('tac_vu', 'ra-soat')
    if (tv == 'so-sanh') != (kieu == 'so_sanh'):
        loi.append('Kiểu báo cáo “%s” không khớp tác vụ của cổng hồ sơ (“%s”).' % (kieu, tv))
    if kieu == 'so_sanh':
        if not ((d.get('so_sanh') or {}).get('ket_luan') or '').strip():
            loi.append('So sánh hai bản: thiếu “so_sanh.ket_luan” (bản nào hợp lý hơn, vì sao).')
        if not d.get('_ss'):
            loi.append('So sánh hai bản: thiếu bảng thay đổi (--so-sanh ss.json do so_sanh_ban.py tạo).')
        else:  # v4.5 (PB5-WORK-16): bảng thay đổi phải của đúng hai bản đã qua cổng
            hai = {(cong or {}).get('sha256', {}).get(x) for x in ((cong or {}).get('cac_ban_du_thao') or [])}
            ss_ = d['_ss']
            if not ss_.get('sha256_cu') or not ss_.get('sha256_moi'):
                loi.append('So sánh hai bản: bảng thay đổi không có mã băm hai tệp (chạy lại so_sanh_ban.py của gói v4.5).')
            elif {ss_['sha256_cu'], ss_['sha256_moi']} != hai:
                loi.append('So sánh hai bản: bảng thay đổi (--so-sanh) lập cho cặp tệp khác với hai bản đã qua cổng hồ sơ (%s) – chạy lại so_sanh_ban.py cho đúng hai bản.' % '; '.join((cong or {}).get('cac_ban_du_thao') or []))
    if d.get('nhan_ket_luan'):
        loi.append('Không điền “nhan_ket_luan” – nhãn do script tự tính.')
    # v4.3 (TH-05), v4.4 (PB4-GPT-02, PB4-WORK-06; quyết định E2 của Hiệu trưởng ngày 27/10/2026): bỏ qua dương tính giả phải chỉ đúng
    # từng phát hiện (tệp + mã + vị trí) script báo, lý do cụ thể; Mức 1 phải có nguyên văn ĐÚNG LÀ CHỮ của đoạn script báo (≥ 15 ký tự)
    for b in d.get('bo_qua_the_thuc', []):
        k = (b.get('ma'), b.get('vi_tri', ''))
        if not (b.get('vi_tri') or '').strip():
            loi.append('Bỏ qua %s: phải ghi “vi_tri” đúng từng phát hiện của script (không bỏ cả mã ở mọi vị trí).' % b.get('ma')); continue
        if len(tts) > 1 and not (b.get('tep') or '').strip():
            loi.append('Bỏ qua %s tại “%s”: hồ sơ có %d tệp đã đo – phải ghi “tep” (tên tệp như cổng hồ sơ ghi) để không bỏ nhầm phát hiện của tệp khác.' % (k + (len(tts),))); continue
        tt_ = tep_khop(tts, b.get('tep'))
        if (b.get('tep') or '').strip() and tt_ is None:
            loi.append('Bỏ qua %s tại “%s”: “tep” = “%s” không chỉ đúng MỘT tệp đã đo (ghi tên tệp như cổng hồ sơ ghi, ví dụ “A/Ho_so.docx”).' % (k + (b.get('tep'),))); continue
        hit = [x for t in ([tt_] if tt_ is not None else tts) for x in t.get('loi', []) if (x['ma'], x['vi_tri']) == k]
        if not hit:
            loi.append('Bỏ qua %s tại “%s”%s: không khớp phát hiện nào script đã báo.' % (k + ((' (tệp %s)' % b['tep']) if b.get('tep') else '',)))
        elif len((b.get('ly_do') or '').strip()) < 20:
            loi.append('Bỏ qua %s tại “%s”: lý do phải cụ thể (nêu vì sao là dương tính giả).' % k)
        elif muc_so(hit[0].get('muc')) == 1:
            nv = chuan_nv(b.get('nguyen_van') or '')
            if not NGOAC.search(b.get('nguyen_van') or '') or len(nv) < 15:
                loi.append('Bỏ qua %s Mức 1 tại “%s”: phải có “nguyen_van” trích trong ngoặc kép (≥ 15 ký tự) từ chính đoạn script báo.' % k)
            elif not hit[0].get('doan_goc'):
                loi.append('Bỏ qua %s Mức 1 tại “%s”: kết quả đo không có nguyên văn đoạn (chạy lại kiem_tra_the_thuc.py của gói v4.5).' % k)
            elif nv not in chuan_nv(hit[0]['doan_goc']):
                loi.append('Bỏ qua %s Mức 1 tại “%s”: “nguyen_van” không phải chữ của đoạn script báo – không được dùng làm bằng chứng dương tính giả.' % k)
    return loi


def tep_khop(tts, ten):
    """v4.5 (PB5-WORK-06): tệp đo thể thức mà “tep” chỉ tới – tên như cổng hồ sơ ghi (_id); tên ngắn chỉ nhận khi duy nhất."""
    ten = (ten or '').strip()
    if not ten: return None
    a_ = [t for t in tts if t.get('_id') == ten]
    if len(a_) == 1: return a_[0]
    b_ = [t for t in tts if t.get('tep') == ten]
    return b_[0] if len(b_) == 1 else None


def nhap_the_thuc(tts, bo_qua):
    out, dat, chua, loai = [], [], [], []
    def bq_cua(t, x):
        for b in bo_qua:
            if not (b.get('vi_tri') or '').strip() or (b.get('ma'), b.get('vi_tri')) != (x['ma'], x['vi_tri']): continue
            if (b.get('tep') or '').strip() and tep_khop(tts, b.get('tep')) is not t: continue
            if not (b.get('tep') or '').strip() and len(tts) > 1: continue  # v4.4 (PB4-WORK-06): nhiều tệp thì phải chỉ đúng tệp
            return b
        return None
    for t in tts:
        nm = t.get('_id') or t.get('tep') or t.get('_ten_tep', '')
        for x in t.get('loi', []):
            b = bq_cua(t, x)
            if b:
                loai.append(dict(tep=nm, ma=x['ma'], vi_tri=x['vi_tri'], muc=x['muc'], nguyen_van=b.get('nguyen_van', ''), ly_do=b.get('ly_do', '')))
                continue
            out.append(dict(ma=x['ma'], _nguon='script', vi_tri=(nm + ' – ' if len(tts) > 1 else '') + x['vi_tri'],
                            mo_ta=x['hang_muc'], nguyen_van_du_thao='đo được ' + x['do'],
                            can_cu=x['can_cu'], nguyen_van_can_cu='Yêu cầu: ' + x['yeu_cau'], kien_nghi='Sửa thành: ' + x['yeu_cau'],
                            muc=x['muc'], chac_chan='chắc chắn'))
        for x in t.get('dat', []):
            dat.append(dict(nhom='Thể thức', noi_dung=x['hang_muc'], so_do=x['do'], nguon=x['can_cu']))
        for x in t.get('chua_kiem', []):
            chua.append(dict(noi_dung=x['hang_muc'], ly_do=x['ly_do'], nguoi_xac_minh='Người soạn thảo kiểm tra bằng Word'))
    return out, dat, chua, loai


KIEU = {'chinh_thuc': 'rà soát chính thức', 'nhanh': 'rà soát nhanh', 'so_sanh': 'so sánh hai bản'}


def tinh(ph, ma_tran, da_ban_hanh=False, kieu='chinh_thuc'):
    dem = {m: 0 for m in (1, 2, 3, 4)}
    for f in ph: dem[muc_so(f['muc'])] += 1
    ck_trong_yeu = [r for r in ma_tran if r.get('ket_qua', '').startswith('Chưa kiểm') and any(k in (r.get('chu_de', '') or '').lower() for k in TRONG_YEU)]
    if dem[1]: nhan, qt = 'Chưa đủ điều kiện ban hành', 'Quy tắc 1: còn %d lỗi Mức 1' % dem[1]
    elif dem[2] or ck_trong_yeu: nhan, qt = 'Cần chỉnh sửa trước khi trình ký', 'Quy tắc 2: 0 Mức 1; %d Mức 2%s' % (dem[2], ('; %d chủ đề trọng yếu “Chưa kiểm”' % len(ck_trong_yeu)) if ck_trong_yeu else '')
    elif dem[3]: nhan, qt = 'Có thể ban hành sau khi chỉnh sửa nhỏ', 'Quy tắc 3: 0 Mức 1, 0 Mức 2; %d Mức 3' % dem[3]
    else: nhan, qt = 'Có thể ban hành ngay', 'Quy tắc 4: 0 Mức 1, 0 Mức 2, 0 Mức 3'
    if da_ban_hanh:  # văn bản đã ban hành: không gán nhãn trình ký (02 mục 5)
        nhan, qt = 'Kiến nghị sửa đổi, bổ sung', 'văn bản đã ban hành – %d Mức 1, %d Mức 2, %d Mức 3, %d Mức 4' % (dem[1], dem[2], dem[3], dem[4])
    if kieu in ('nhanh', 'so_sanh'):  # PB2-21: (b) và (c) không gán nhãn điều kiện trình ký
        nhan, qt = ('Rà soát nhanh – không gán nhãn trình ký' if kieu == 'nhanh' else 'So sánh hai bản – không gán nhãn trình ký'), \
                   'không áp dụng Quy tắc 1–4 của 02-Quy-Tac-Xep-Muc mục 5'
    tuan_thu = 'Chưa đạt' if dem[1] else ('Đạt' if dem[2] else ('Tốt' if dem[3] else 'Rất tốt'))
    nhom = {g: 0 for g in NHOM_LIST}
    for f in ph: nhom[f['nhom']] += 1
    return dem, nhan, qt, tuan_thu, nhom


# ---------------------------------------------------------------- docx
def shade(cell, hexcolor='D9EAF7'):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexcolor); tcPr.append(s)


def para(doc, text='', bold=False, size=14, align=None, italic=False, color=None, before=6, after=6):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align in ('j', None):  # v4.2: mọi đoạn lời văn của báo cáo dàn đều, thụt 1 cm (quyết định HT 26/10 số 5; TT-23, TT-24)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    return p


def table(doc, header, rows, widths=None, size=11):
    t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; shade(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ''; r = cs[i].paragraphs[0].add_run(str(v if v is not None else '')); r.font.size = Pt(size)
            cs[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        t.autofit = False
        tot = sum(widths); widths = [w * 16.0 / tot for w in widths]
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i] * 567)))
        for i, col in enumerate(t.columns): col.width = Cm(widths[i])
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
        tblPr = t._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def dung(d, cong, ph, dat, chua, out):
    kieu = d.get('kieu_bao_cao', 'chinh_thuc')
    dem, nhan, qt, tuan_thu, nhom = tinh(ph, d.get('ma_tran', []), d.get('da_ban_hanh'), kieu)
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7); s.orientation = WD_ORIENT.PORTRAIT
        # chuẩn Trường (quyết định của Hiệu trưởng 26/10/2026, điểm 5): lề 2/2/3/2 cm; header/footer 1,27 cm; số trang giữa lề trên, cỡ 14, ẩn trang 1
        s.top_margin = s.bottom_margin = Cm(2); s.left_margin = Cm(3); s.right_margin = Cm(2)
        s.header_distance = s.footer_distance = Cm(1.27)
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f1 = OxmlElement('w:fldSimple'); f1.set(qn('w:instr'), 'PAGE')
        rr = OxmlElement('w:r'); rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '28'); rpr.append(sz); rr.append(rpr)
        tt_ = OxmlElement('w:t'); tt_.text = '2'; rr.append(tt_); f1.append(rr); hp._p.append(f1)
    st.paragraph_format.space_before = Pt(6); st.paragraph_format.space_after = Pt(6)
    para(doc, 'TRƯỜNG CAO ĐẲNG KON TUM', True, 13, 'c', before=0)
    tieu = {'chinh_thuc': 'BÁO CÁO RÀ SOÁT DỰ THẢO VĂN BẢN TRƯỚC KHI TRÌNH KÝ', 'nhanh': 'BÁO CÁO RÀ SOÁT NHANH DỰ THẢO VĂN BẢN',
            'so_sanh': 'BÁO CÁO SO SÁNH HAI BẢN DỰ THẢO VĂN BẢN'}[d.get('kieu_bao_cao', 'chinh_thuc')]
    para(doc, tieu, True, 15, 'c')
    para(doc, d['ten_du_thao'], True, 14, 'c')
    para(doc, TRANG_THAI, True, 13, 'c', color=(192, 0, 0))
    # Kết luận điều hành
    chan = [f for f in ph if muc_so(f['muc']) == 1]
    kl = 'KẾT LUẬN: %s. Số phát hiện: %d Mức 1, %d Mức 2, %d Mức 3, %d Mức 4 (%s).' % (nhan.upper(), dem[1], dem[2], dem[3], dem[4], qt)
    if chan: kl += ' Mục chặn trình ký: ' + '; '.join('%s (%s)' % (f['ma'], f['vi_tri'][:60]) for f in chan[:8]) + '.'
    if dem[1]: kl += ' Văn bản chưa nên phát hành cho đến khi các lỗi bắt buộc được khắc phục.'
    para(doc, kl, True, 14, 'j', color=(192, 0, 0))

    para(doc, 'PHẦN I. THÔNG TIN CHUNG', True)
    he = {'chinh_quyen': 'Văn bản hành chính (chính quyền) – TB 597/TB-CĐKT, QĐ 1355/QĐ-CĐKT, NĐ 30/2020/NĐ-CP', 'dang': 'Văn bản của Đảng – QĐ 208-QĐ/TW, HD 05-HD/VPTW, QĐ 399-QĐ/TW'}.get(d.get('he'), d.get('he', ''))
    table(doc, ['Mục', 'Nội dung'], [
        ['1. Dự thảo', d['ten_du_thao']], ['2. Loại văn bản', d['loai_van_ban']], ['3. Đơn vị soạn thảo', d['don_vi_soan_thao']],
        ['4. Hệ văn bản', he], ['5. Tác vụ, chế độ', '%s – chế độ %s' % (KIEU.get(d.get('kieu_bao_cao', 'chinh_thuc')), d.get('che_do', 'ĐL'))],
        ['6. Lần rà soát', str(d.get('lan_ra_soat', cong.get('vong', 1)))], ['7. Ngày rà soát', d.get('ngay_ra_soat', datetime.date.today().strftime('%d/%m/%Y'))],
        ['8. Phiên bản Skill', PHIEN_BAN], ['9. Phạm vi', d.get('pham_vi', 'Toàn bộ dự thảo, văn bản kèm theo và phụ lục')],
        ['10. Ghi chú', d.get('ghi_chu', '')]], [4.5, 12.5])
    para(doc, 'I.7. Kiểm kê hồ sơ (kết quả kiem_ke_ho_so.py: ĐỦ ĐIỀU KIỆN RÀ SOÁT)', True)
    vt = cong.get('vai_tro', {})
    table(doc, ['#', 'Tệp', 'Vai trò', 'Tình trạng'], [[i, t[0], {'du_thao': 'Dự thảo hiện hành', 'du_thao_cu': 'Dự thảo lần trước', 'bao_cao_ra_soat': 'Báo cáo rà soát lần trước', 'bc_tiep_thu': 'Báo cáo tiếp thu, giải trình', 'kem_theo': 'Phụ lục/văn bản kèm theo', 'lien_quan': 'Văn bản liên quan'}.get(vt.get(t[0], ''), '—'), t[2]] for i, t in enumerate(cong.get('tep', []), 1)], [1, 9, 4, 3])
    table(doc, ['#', 'Văn bản được dự thảo đề cập', 'Đoạn', 'Nguồn toàn văn'], [[i, x.get('ten') or ('%s số %s' % (x['loai'], x['so'])), x['doan'], '%s – %s' % (x['trang_thai'], x['nguon'])] for i, x in enumerate(cong.get('van_ban_de_cap', []), 1)], [1, 8, 1.5, 6.5])
    if cong.get('xac_nhan_kem_thu_cong'):  # v4.5 – quyết định F1 của Hiệu trưởng ngày 27/10/2026
        para(doc, 'Văn bản kèm theo do người rà soát xác nhận thủ công (trích yếu không khớp hoặc không có – quyết định F1):', True, 13, color=(192, 0, 0))
        table(doc, ['Bản', 'Văn bản kèm theo dự thảo ban hành', 'Tệp được xác nhận', 'Tiêu đề tệp', 'Lý do xác nhận'],
              [[x.get('ban'), x.get('van_ban_kem_theo'), x.get('tep'), x.get('tieu_de_tep'), x.get('ly_do')] for x in cong['xac_nhan_kem_thu_cong']], [2.5, 4, 2.5, 4, 4], 10)
    if d.get('xem_anh'):
        table(doc, ['Tệp có ảnh ≥ 1/4 trang', 'Số ảnh đã xem', 'Kết luận', 'Mô tả'], [[x.get('tep'), x.get('so_anh_da_xem'), {'so_do': 'Sơ đồ/hình minh họa – nhận'}.get(x.get('ket_luan'), x.get('ket_luan')), x.get('mo_ta')] for x in d['xem_anh']], [5, 2.5, 4, 5.5])
    para(doc, 'I.8. Nội dung chưa kiểm được', True)
    ck = chua + d.get('chua_kiem', [])
    if ck: table(doc, ['#', 'Nội dung', 'Lý do', 'Ai cần xác minh'], [[i, x['noi_dung'], x['ly_do'], x.get('nguoi_xac_minh', '')] for i, x in enumerate(ck, 1)], [1, 6, 6, 4])
    else: para(doc, 'Không có.')

    para(doc, 'PHẦN II. KẾT QUẢ RÀ SOÁT CHI TIẾT', True)
    if d.get('che_do') == 'ĐC':
        para(doc, 'II.A. Theo dõi các mục của báo cáo rà soát lần trước', True)
        table(doc, ['Mã cũ', 'Nội dung lần trước', 'Vị trí bản hiện hành', 'Nguyên văn/số đo bản hiện hành', 'Trạng thái', 'Mức còn lại'],
              [[r.get('ma_cu'), r.get('noi_dung_cu'), r.get('vi_tri_hien_hanh'), (r.get('nguyen_van_hien_hanh') or '') + (('\nCăn cứ đúng: ' + r['can_cu_dung']) if r.get('can_cu_dung') else ''), r.get('trang_thai'), (r.get('muc_con_lai') or '') + ((' → ' + r['ma_phat_hien']) if r.get('ma_phat_hien') else '')] for r in d['doi_chieu']], [1.6, 3.6, 2.6, 4.6, 2.4, 2.2], 10)
        para(doc, 'II.B. Danh sách phát hiện hiện hành (tồn tại và mới)', True)
    rows = []
    for i, f in enumerate(ph, 1):
        rows.append([i, f['ma'], f['nhom'], f['vi_tri'], f['mo_ta'] + ('\nDự thảo: ' + f['nguyen_van_du_thao'] if f.get('nguyen_van_du_thao') else ''),
                     f.get('can_cu', '') + ('\n' + f['nguyen_van_can_cu'] if f.get('nguyen_van_can_cu') else ''), f.get('kien_nghi', ''),
                     'Mức %d\n%s' % (muc_so(f['muc']), f.get('chac_chan', 'chắc chắn'))])
    if rows: table(doc, ['STT', 'Mã', 'Nhóm', 'Vị trí', 'Mô tả vấn đề', 'Căn cứ nhận xét', 'Kiến nghị sửa', 'Mức độ'], rows, [0.9, 1.3, 1.7, 2.3, 4.0, 3.2, 2.6, 1.4], 10)
    else: para(doc, 'Không phát hiện vấn đề.')
    if d.get('che_do') == 'ĐC' and d.get('hoi_quy'):
        para(doc, 'II.C. Kiểm hồi quy', True)
        table(doc, ['Mã cũ', 'Kết quả kiểm lại'], [[r.get('ma_cu'), r.get('ket_qua')] for r in d['hoi_quy']], [3, 14])
    if kieu == 'so_sanh':
        para(doc, 'II.D. Bảng thay đổi giữa hai bản (so_sanh_ban.py)', True)
        ssd = d.get('_ss') or {}
        tk = ssd.get('thong_ke', {})
        para(doc, 'Bản cũ: %s; bản mới: %s; %s sửa, %s thêm, %s xóa; vị trí bị ảnh hưởng: %s.' % (
            str(ssd.get('cu', '')).rsplit('/', 1)[-1], str(ssd.get('moi', '')).rsplit('/', 1)[-1], tk.get('Sửa', 0), tk.get('Thêm', 0), tk.get('Xóa', 0), ', '.join(ssd.get('vi_tri', [])) or 'không có'), align='j')
        rows_ss = [[k, r['loai'], r.get('doan_cu') or '', r.get('doan_moi') or '', r.get('vi_tri', ''), r.get('noi_dung', '')[:700]] for k, r in enumerate(ssd.get('thay_doi', []), 1)]
        if rows_ss: table(doc, ['#', 'Loại', 'Đoạn cũ', 'Đoạn mới', 'Vị trí', 'Nội dung thay đổi'], rows_ss, [0.8, 1.3, 1.3, 1.3, 2.3, 10], 10)
        para(doc, 'Kết luận so sánh: ' + d['so_sanh']['ket_luan'], True, align='j')

    para(doc, 'PHẦN III. CÁC NỘI DUNG KHÔNG PHÁT HIỆN VẤN ĐỀ', True)
    allok = dat + d.get('dat', [])
    for g in NHOM_LIST:
        items = [x for x in allok if x.get('nhom') == g]
        para(doc, '%s (%d mục)' % (g, len(items)), True, 14)
        if items: table(doc, ['#', 'Nội dung', 'Số đo / nguyên văn', 'Nguồn đối chiếu'], [[i, x['noi_dung'], x.get('so_do') or x.get('nguyen_van'), x.get('nguon', '')] for i, x in enumerate(items, 1)], [1, 5, 7, 4], 10)
        else: para(doc, 'Không có mục Đạt kèm bằng chứng.', italic=True)

    para(doc, 'PHẦN IV. BẢNG TỔNG HỢP', True)
    para(doc, 'Bảng 1. Số phát hiện theo nhóm', True, 14)
    table(doc, ['Nhóm', 'Số phát hiện'], [[g, nhom[g]] for g in NHOM_LIST] + [['Tổng cộng', sum(nhom.values())]], [10, 4])
    para(doc, 'Bảng 2. Số phát hiện theo mức', True, 14)
    table(doc, ['Mức', 'Số phát hiện'], [[TEN_MUC[m], dem[m]] for m in (1, 2, 3, 4)] + [['Tổng cộng', sum(dem.values())]], [10, 4])
    l1 = [x for x in d.get('_da_loai', []) if muc_so(x['muc']) == 1]
    if l1:  # v4.4 – quyết định E2 của Hiệu trưởng ngày 27/10/2026: Mức 1 của script đã loại phải in rõ để người ký xem lại
        para(doc, 'Bảng 2a. Phát hiện Mức 1 của script đã loại khỏi kết luận (người rà soát xác định là báo oan) – người ký xem lại', True, 14, color=(192, 0, 0))
        table(doc, ['#', 'Tệp', 'Mã', 'Vị trí', 'Nguyên văn chứng minh', 'Lý do loại'], [[i, x['tep'], x['ma'], x['vi_tri'][:120], x['nguyen_van'], x['ly_do']] for i, x in enumerate(l1, 1)], [0.8, 2.6, 1.3, 4, 4.3, 4], 10)
    para(doc, 'Bảng 3. Ma trận đối chiếu chéo', True, 14)
    mt = d.get('ma_tran', [])
    if mt:
        table(doc, ['STT', 'Chủ đề', 'Dự thảo: vị trí + nguyên văn', 'Văn bản liên quan + nguyên văn', 'Kết quả', 'Phát hiện'],
              [[i, r.get('chu_de'), '%s: %s' % (r.get('du_thao_vi_tri', ''), r.get('du_thao_nguyen_van', '')), '%s: %s' % (r.get('vblq', ''), r.get('vblq_nguyen_van', '')), r.get('ket_qua') + (('\nĐiều khoản thay thế: ' + r['nguyen_van_thay_the']) if r.get('nguyen_van_thay_the') else ''), r.get('ma_phat_hien', '')] for i, r in enumerate(mt, 1)], [0.9, 2.4, 5, 5, 2, 1.7], 10)
        if GIAI_TRINH:  # v4.5 – quyết định F2: “Khớp” mà script thấy dấu hiệu khác nhau, AI giải trình vẫn khớp
            para(doc, 'Giải trình “Khớp” khi script phát hiện dấu hiệu khác nhau giữa hai phía (quyết định F2) – người ký xem lại:', True, 13, color=(192, 0, 0))
            table(doc, ['Dòng', 'Yếu tố', 'Dấu hiệu script phát hiện', 'Giải trình'], [[x['dong'], x['yeu_to'], x['dau_hieu'], x['giai_trinh']] for x in GIAI_TRINH], [1.2, 2.3, 5.5, 8], 10)
        tk = {g: sum(1 for r in mt if r.get('ket_qua', '').startswith(g)) for g in MA_TRAN_GT}
        para(doc, 'Tổng hợp ma trận: ' + '; '.join('%s: %d' % (g, n) for g, n in tk.items()) + '.', size=14, align='j')
    else: para(doc, 'Không có văn bản liên quan cùng chủ đề để đối chiếu (đã xác nhận tại I.7).', italic=True)

    para(doc, 'PHẦN V. ĐÁNH GIÁ TỔNG THỂ', True)
    para(doc, 'Mức độ tuân thủ pháp luật: %s (gán theo số lượng Mức 1–3, 02-Quy-Tac-Xep-Muc.md mục 5).' % tuan_thu, align='j')
    para(doc, 'Mức độ sẵn sàng phát hành: %s – %s.' % (nhan, qt), True, align='j')

    para(doc, 'PHẦN VI. KẾT LUẬN', True)
    for i, x in enumerate(d.get('uu_tien', []), 1): para(doc, '%d. %s' % (i, x), align='j')
    if d.get('rui_ro'): para(doc, 'Rủi ro nếu không sửa: ' + d['rui_ro'], align='j')
    if d.get('kien_nghi_chung'): para(doc, 'Kiến nghị chung: ' + d['kien_nghi_chung'], align='j')

    para(doc, 'KIỂM TRA CHẤT LƯỢNG CUỐI CÙNG', True)
    kc = [['Nhãn kết luận do script tính từ Bảng 2', 'Đạt (tự động)'], ['Mọi Mức 1–2 có bằng chứng hai phía', 'Đạt (script kiểm)'],
          ['Ma trận chỉ dùng 6 giá trị; “Khớp” có nguyên văn hai phía', 'Đạt (script kiểm)']]
    kc += [[x['hang_muc'], x['ket_qua']] for x in d.get('kiem_tra_chat_luong', [])]
    kc += [['Bỏ qua phát hiện của script: %s – %s' % (b['ma'], b.get('vi_tri', '')), 'Lý do: ' + b.get('ly_do', '')] for b in d.get('bo_qua_the_thuc', [])]
    table(doc, ['Hạng mục kiểm', 'Kết quả'], kc, [10, 7])
    if d.get('kha_thi'):
        para(doc, 'PHẦN VII. PHÂN TÍCH TÍNH KHẢ THI VÀ RỦI RO (theo yêu cầu)', True)
        para(doc, d['kha_thi'], align='j')

    para(doc, 'THÔNG TIN TRÁCH NHIỆM', True)
    table(doc, ['Trường', 'Nội dung'], [
        ['Nguồn dữ liệu đã dùng', '; '.join(d.get('tai_lieu_da_dung', [])) or '; '.join(t[0] for t in cong.get('tep', []))],
        ['Người kiểm tra', d['nguoi_kiem_tra']],
        ['Trạng thái phê duyệt', TRANG_THAI + '. Kết quả AI chỉ có giá trị tham khảo; Trưởng/Phó đơn vị phê duyệt trước khi dùng chính thức.']], [5, 12])
    doc.core_properties.title = 'Báo cáo rà soát – ' + d['ten_du_thao'][:120]
    doc.core_properties.author = 'Skill ' + PHIEN_BAN
    doc.save(out)
    return dem, nhan, qt, tuan_thu


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('du_lieu'); ap.add_argument('--cong'); ap.add_argument('--the-thuc', action='append', default=[])
    ap.add_argument('-o', '--out', default='bao_cao_ra_soat.docx'); ap.add_argument('--chi-kiem', action='store_true')
    ap.add_argument('--so-sanh', default=None, help='JSON do so_sanh_ban.py tạo (bắt buộc với kieu_bao_cao = so_sanh)')
    a = ap.parse_args()
    d = json.load(open(a.du_lieu, encoding='utf-8'))
    if a.so_sanh: d['_ss'] = json.load(open(a.so_sanh, encoding='utf-8'))
    cong = json.load(open(a.cong, encoding='utf-8')) if a.cong else None
    tts = []
    for p in a.the_thuc:
        t = json.load(open(p, encoding='utf-8')); t['_ten_tep'] = p.rsplit('/', 1)[-1].replace('.json', ''); tts.append(t)
    sha_id = {}
    for k_, v_ in ((cong or {}).get('sha256') or {}).items(): sha_id.setdefault(v_, []).append(k_)
    for t in tts:  # v4.4 (PB4-WORK-07): tên tệp theo cổng hồ sơ (đường dẫn tương đối khi trùng tên), xác định bằng mã băm
        ids = sha_id.get(t.get('sha256'), [])
        t['_id'] = (t.get('tep') if t.get('tep') in ids else ids[0]) if ids else t.get('tep', t['_ten_tep'])
    sp, sdat, schua, sloai = nhap_the_thuc(tts, d.get('bo_qua_the_thuc', []))
    d['_da_loai'] = sloai
    ph = sp + d.get('phat_hien', [])
    for f in ph: f['nhom'] = nhom_cua(f['ma'])
    ph.sort(key=lambda f: (muc_so(f.get('muc')) or 9, NHOM_LIST.index(f['nhom']), f['ma']))  # thứ tự cố định
    loi = kiem(dict(d, phat_hien=ph), cong, tts)
    if loi:
        print('DỮ LIỆU BÁO CÁO CHƯA HỢP LỆ – không dựng tệp:')
        for x in loi: print(' -', x)
        sys.exit(1)
    dem, nhan, qt, tt = tinh(ph, d.get('ma_tran', []), d.get('da_ban_hanh'), d.get('kieu_bao_cao', 'chinh_thuc'))[:4] if a.chi_kiem else dung(d, cong, ph, sdat, schua, a.out)
    n_l1 = sum(1 for x in sloai if muc_so(x['muc']) == 1)
    print('HỢP LỆ. Nhãn: %s (%s). Mức 1/2/3/4 = %d/%d/%d/%d. Tuân thủ: %s.%s%s' % (nhan, qt, dem[1], dem[2], dem[3], dem[4], tt,
          (' Đã loại %d phát hiện Mức 1 của script (in ở Bảng 2a).' % n_l1) if n_l1 else '', '' if a.chi_kiem else ' Đã ghi: ' + a.out))


if __name__ == '__main__':
    main()
