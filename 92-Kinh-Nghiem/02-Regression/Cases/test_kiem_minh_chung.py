# -*- coding: utf-8 -*-
"""Hoi quy kiem_minh_chung.py — DL-20260919-006. Tu tao tep theo dung tieu de cua
24-KTC-Theo-doi-CV/01. Bo du lieu van hanh (sheet Nhiem vu + Minh chung) va Master Task Register."""
import datetime as dt
import os
import sys
import tempfile

import openpyxl

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, os.path.join(GOC, "29-Cong-Cu"))
import kiem_minh_chung as k  # noqa: E402

sai = []


def kiem(dk, ten):
    print(("  OK " if dk else "  X  ") + ten)
    if not dk:
        sai.append(ten)


T = tempfile.mkdtemp(prefix="ktc_mc_")
that = os.path.join(T, "bien-ban-that.pdf")
open(that, "wb").write(b"%PDF")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Nhiệm vụ"
ws.append(["Mã nhiệm vụ", "Mã KH nguồn", "Kỳ kế hoạch", "Trục TB 817", "Nhóm nhiệm vụ", "Tên nhiệm vụ",
           "Đơn vị chủ trì", "Đơn vị phối hợp", "Sản phẩm/kết quả", "Hạn baseline", "Hạn hiện hành", "Trạng thái"])
for ma, tt in (("NV1", "Hoàn thành"), ("NV2", "Hoàn thành"), ("NV3", "Đang thực hiện"), ("NV4", "Hoàn thành")):
    ws.append([ma, "", "Quý III/2026", "Trục 1", "", f"Nhiệm vụ {ma}", "K-KTCN", "", "", dt.date(2026, 9, 30),
               dt.date(2026, 9, 30), tt])
m = wb.create_sheet("Minh chứng")
m.append(["Mã minh chứng", "Mã nhiệm vụ", "Loại minh chứng", "Mô tả", "Ngày phát sinh", "Liên kết Drive",
          "Tình trạng xác minh", "Người xác minh", "Ngày xác minh", "Ghi chú"])
m.append(["MC1", "NV1", "Văn bản", "", dt.date(2026, 9, 10), that, "Chưa xác minh", "", "", ""])
m.append(["MC2", "NV2", "Văn bản", "", dt.date(2026, 10, 5), os.path.join(T, "khong-co.pdf"), "Đã xác minh", "", "", ""])
m.append(["MC3", "NV3", "Ảnh", "", None, "https://drive.google.com/file/d/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/view",
          "Chưa xác minh", "", "", ""])
m.append(["MC4", "NV9", "Ảnh", "", dt.date(2026, 9, 1), that, "Chưa xác minh", "", "", ""])
m.append(["MC5", "NV3", "Ảnh", "", dt.date(2026, 9, 1), "", "Chưa xác minh", "", "", ""])
cn = wb.create_sheet("Cập nhật tiến độ")
cn.append(["Mã cập nhật", "Mã nhiệm vụ", "Ngày cập nhật", "Trạng thái", "% hoàn thành", "Kết quả đã thực hiện",
           "Khó khăn/vướng mắc", "Hành động tiếp theo", "Đề xuất hạn mới", "Liên kết minh chứng", "Người cập nhật"])
cn.append(["CN1", "NV3", dt.date(2026, 9, 15), "Hoàn thành", 100, "", "", "", "", "", ""])
p = os.path.join(T, "theo-doi.xlsx")
wb.save(p)

nv, mc = k.doc_tep(p)
kq = k.kiem(nv, mc)
ma = lambda code, vt: any(x[1] == code and vt in x[2] for x in kq)  # noqa: E731
kiem(len(nv) == 4 and len(mc) == 6, f"đọc cột theo tên tiêu đề: 4 nhiệm vụ, 5 minh chứng + 1 dòng cập nhật (được {len(nv)}, {len(mc)})")
kiem(nv["NV3"]["trang_thai"] == "Đang thực hiện",
     "sheet “Cập nhật tiến độ” KHÔNG bị đọc nhầm thành sheet nhiệm vụ (không ghi đè trạng thái)")
kiem(ma("MC01", "NV4"), "MC01: NV4 hoàn thành không có minh chứng → Mức 1")
kiem(not ma("MC01", "NV1"), "MC01: NV1 có minh chứng thật → không báo (thử ngược)")
kiem(not ma("MC01", "NV3"), "MC01: NV3 đang thực hiện → không đòi minh chứng")
kiem(any(x[1] == "MC03" and x[0] == 2 and "NV2" in x[2] for x in kq), "MC03: đường dẫn không tồn tại → Mức 2")
kiem(any(x[1] == "MC03" and x[0] is None and "CẦN MỞ QUA GOOGLE DRIVE" in x[3] for x in kq),
     "MC03: URL Drive → chuyển mở qua Google Drive, không tự kết luận")
kiem(ma("MC04", "NV2") and any("sau hạn" in x[3] for x in kq), "MC04: minh chứng phát sinh sau hạn")
kiem(any(x[1] == "MC04" and "thiếu ngày" in x[3] and "MC3" in x[2] for x in kq), "MC04: thiếu ngày phát sinh")
kiem(ma("MC05", "NV2"), "MC05: “Đã xác minh” thiếu người/ngày → xác minh hình thức")
kiem(ma("MC06", "NV9"), "MC06: minh chứng trỏ tới nhiệm vụ không tồn tại")
kiem(any(x[1] == "MC07" and "NV1" in x[2] and "NV9" in x[2] for x in kq), "MC07: một liên kết dùng cho nhiều nhiệm vụ")
kiem(any(x[1] == "MC02" and "MC5" in x[2] for x in kq), "MC02: minh chứng không có liên kết")

# Master Task Register: cot Minh_Chung tren cung dong
wb2 = openpyxl.Workbook()
w2 = wb2.active
w2.append(["Task_ID", "Ten_Nhiem_Vu", "Don_Vi_Chu_Tri", "Han_Hoan_Thanh", "Trang_Thai", "Minh_Chung"])
w2.append(["(bắt buộc)", "", "", "", "", ""])
w2.append(["KTC-2026-Q3-00001", "A", "P-THHC", dt.date(2026, 9, 30), "Đã hoàn thành", that])
w2.append(["KTC-2026-Q3-00002", "B", "P-THHC", dt.date(2026, 9, 30), "Đã hoàn thành", None])
p2 = os.path.join(T, "mtr.xlsx")
wb2.save(p2)
nv2, mc2 = k.doc_tep(p2)
kq2 = k.kiem(nv2, mc2)
kiem(len(nv2) == 2, "Master Task Register: bỏ dòng mô tả “(bắt buộc)”")
kiem([x[2] for x in kq2 if x[1] == "MC01"] == ["KTC-2026-Q3-00002"], "Master Task Register: MC01 đúng nhiệm vụ thiếu Minh_Chung")

print(f"\n{'ĐẠT' if not sai else 'KHÔNG ĐẠT'}: {len(sai)} ca sai")
sys.exit(1 if sai else 0)
