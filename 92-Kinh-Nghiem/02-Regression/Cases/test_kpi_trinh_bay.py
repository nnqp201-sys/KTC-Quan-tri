# -*- coding: utf-8 -*-
"""Hoi quy trinh bay sheet KPI (lenh sua 25/9/2026: che chu L1, thieu cot C–F L2) — kpi_mau.py, kpi_danh_gia.py.

Du lieu gia dinh ("Nguyen Van A"), thu muc tam. Viet TRUOC khi sua, chay thay that bai, roi moi sua (quy tac them phep
kiem phai co ca thu nguoc). T9 ghim tong diem truoc khi sua: sua trinh bay khong duoc doi so nao.
"""
import ctypes
import math
import os
import subprocess
import sys
import tempfile

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
CC = os.path.join(GOC, "29-Cong-Cu")
sys.path.insert(0, CC)
import kpi_calc as kc      # noqa: E402
import kpi_danh_gia as kd  # noqa: E402
import kpi_mau as km       # noqa: E402
import validate_plan as vp  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")
sai = []
T = tempfile.mkdtemp(prefix="kpi_tb_")


def kiem(dk, ten):
    print(("  OK " if dk else "  X  ") + ten)
    if not dk:
        sai.append(ten)


def dv(truc, nd, sl=1, muc="Trung bình", **k):
    d = {"truc": truc, "noi_dung": nd, "cap_trinh": "Trưởng đơn vị", "muc_do": muc, "san_pham": "Báo cáo",
         "so_luong": sl, "thoi_han": "30/11/2026", "minh_chung": "Văn bản trên hệ thống Office"}
    d.update(k)
    return d


def lap(nhom, ds, ten):
    kh = {"ca_nhan": {"ho_ten": "Nguyễn Văn A", "don_vi": "Đơn vị X"}, "dau_viec": ds}
    kh = dict(kh, dau_viec=kc.quy_doi_ke_hoach(kh, "muc-do")["dau_viec"], phuong_an="muc-do")
    p = os.path.join(T, f"{nhom}_{ten}.xlsx")
    tb = km.ghi_ke_hoach(nhom, kh, p, "III", "2026")
    return p, tb


def ma(p, nhom):
    return [x["ma"] for x in vp.kiem(p, nhom, "muc-do")["loi"]]


def goi(fn):
    try:
        return fn()
    except Exception as e:     # ham chua co (truoc khi sua) -> ca that bai, khong dung bo thu
        return e


DAI = ("Theo dõi, tổng hợp, báo cáo, đánh giá tất cả các mặt hoạt động của Trường để làm báo cáo hằng tháng, quý, "
       "sáu tháng, chín tháng và công tác năm học; phối hợp các đơn vị thuộc Trường thu thập số liệu, đối chiếu kế hoạch, "
       "tổng hợp kết quả thực hiện nhiệm vụ trọng tâm theo sáu Trục và đề xuất giải pháp khắc phục hạn chế tồn tại.")[:300]

print("== T1. Nội dung 300 ký tự (hanh-chinh): dòng KPI đủ cao cho cột B (công thức trỏ Ke Hoach) ==")
p, _ = lap("hanh-chinh", [dv(1, DAI)], "t1")
wb = km.mo(p)
ct = km.cau_truc(wb)
kp = wb["KPI"]
rk = ct["kpi_dong"][ct["truc"][1][0]]
rong = goi(lambda: km.do_rong_cot(kp))
rong_b = rong.get(2, 32.3) if isinstance(rong, dict) else 32.3
co = float(kp[f"B{rk}"].font.sz or 12)
can = math.ceil(len(DAI) * 1.2 / rong_b) * co * 1.3
cao = kp.row_dimensions[rk].height or 15
kiem(cao >= can - 0.01, f"KPI dòng {rk}: cao {cao:g} pt ≥ ước {can:g} pt ({len(DAI)} ký tự, cột B rộng {rong_b:g})")

print("== T2. Độ rộng cột nằm giữa nhóm column_dimensions gộp (min–max) ==")
kiem(isinstance(rong, dict) and abs(rong.get(4, 0) - 8.7109375) < 1e-6 and abs(rong.get(21, 0) - 30.0) < 1e-6,
     f"KPI cột D (giữa nhóm C–F) = {rong.get(4) if isinstance(rong, dict) else rong}; cột U (giữa nhóm S–Y) = "
     f"{rong.get(21) if isinstance(rong, dict) else '-'}")

print("== T3. Truyền đủ người chỉ đạo, phối hợp, đơn vị tham mưu -> KPI!C–E; KPI!F là công thức ==")
p, _ = lap("hanh-chinh", [dv(1, "Việc có đủ trường", nguoi_chi_dao="Phó Hiệu trưởng X", nguoi_phoi_hop="Các phòng",
                            don_vi_tham_muu="Phòng Y")], "t3")
wb = km.mo(p)
ct = km.cau_truc(wb)
r = ct["truc"][1][0]
rk = ct["kpi_dong"][r]
kp = wb["KPI"]
kiem((kp[f"C{rk}"].value, kp[f"D{rk}"].value, kp[f"E{rk}"].value) == ("Phó Hiệu trưởng X", "Các phòng", "Phòng Y"),
     f"KPI!C/D/E = {kp[f'C{rk}'].value!r}, {kp[f'D{rk}'].value!r}, {kp[f'E{rk}'].value!r}")
kiem(kp[f"F{rk}"].value == f"='Ke Hoach'!E{r}", f"KPI!F{rk} = {kp[f'F{rk}'].value!r}")
kiem(kp[f"C{rk}"].alignment.wrap_text and kp[f"C{rk}"].alignment.vertical == "top", "C–F xuống dòng, căn trên")
kiem("KH17" not in ma(p, "hanh-chinh") and "KH18" not in ma(p, "hanh-chinh"), "ca ngược: đủ trường -> không KH17/KH18")

print("== T4. Không truyền người phối hợp -> để trống + KH17; mặc định C = cấp trình, E = đơn vị ==")
p, _ = lap("hanh-chinh", [dv(1, "Việc thiếu người phối hợp")], "t4")
wb = km.mo(p)
ct = km.cau_truc(wb)
rk = ct["kpi_dong"][ct["truc"][1][0]]
kp = wb["KPI"]
kiem(kp[f"D{rk}"].value in (None, "") and "KH17" in ma(p, "hanh-chinh"), "KPI!D trống, có KH17 (không tự bịa)")
kiem(kp[f"C{rk}"].value == "Trưởng đơn vị" and kp[f"E{rk}"].value == "Đơn vị X", "mặc định C = cấp trình, E = đơn vị")
wb2 = km.mo(p)
wb2["KPI"]["D4"].value = None                       # ca nguoc: mau KHONG co cot "Nguoi phoi hop" -> khong KH17
p2 = os.path.join(T, "t4_khong_cot.xlsx")
wb2.save(p2)
kiem("KH17" not in ma(p2, "hanh-chinh"), "ca ngược: mẫu không có cột Người phối hợp -> không KH17")
wb3 = km.mo(p)
wb3["KPI"][f"F{rk}"].value = "='Ke Hoach'!E99"      # tro dong trong -> KH18
p3 = os.path.join(T, "t4_f_sai.xlsx")
wb3.save(p3)
kiem("KH18" in ma(p3, "hanh-chinh"), "KPI!F trỏ dòng trống của Ke Hoach -> KH18 (LỖI)")

print("== T5. Nội dung quá dài (> 409 pt) -> KH19, không cắt im lặng ==")
p, tb = lap("hanh-chinh", [dv(1, DAI * 12)], "t5")
kiem(any("KH19" in str(x) for x in (tb or [])) and "KH19" in ma(p, "hanh-chinh"), "cảnh báo KH19 ở cả kpi_mau và validate_plan")
wb = km.mo(p)
ct = km.cau_truc(wb)
kiem((wb["KPI"].row_dimensions[ct["kpi_dong"][ct["truc"][1][0]]].height or 0) <= 409, "chiều cao không vượt 409 pt")

print("== T6. Sau danh-gia: dòng KPI đủ cao cho sản phẩm thực tế (K) và minh chứng (S) ==")
p, _ = lap("hanh-chinh", [dv(1, "Việc ngắn")], "t6")
khd = kd.doc_ke_hoach(p)
bh = os.path.join(T, "t6_bh.xlsx")
kd.sinh_bang_hoi(khd, bh)
from openpyxl import load_workbook  # noqa: E402
w = load_workbook(bh)
for row in w["A-Tieu-chi-chung"].iter_rows(min_row=5):
    if len(str(row[0].value or "")) > 2:
        row[3].value = row[2].value
K = "Đã ban hành: " + ", ".join(f"{i}/BC-CĐKT" for i in range(300, 340))
for row in w["B-KPI"].iter_rows(min_row=5):
    if str(row[0].value or "").startswith("B"):
        row[7].value, row[8].value, row[9].value, row[10].value, row[11].value = K, 1, 100, 100, "Hoàn thành đúng hạn"
        row[13].value = "Văn bản trên hệ thống Office, sổ theo dõi văn bản đi, biên bản họp giao ban tháng 7, 8, 9"
for row in w["C-Dieu-kien"].iter_rows(min_row=5):
    m_ = str(row[0].value or "")
    if m_[:1] in "CD" and m_[1:].isdigit():
        row[2].value = {"C07": "Có", "C15": kd.MUC[1]}.get(m_, "Đạt" if m_[0] == "D" else "Không")
w.save(bh)
tl = kd.doc_bang_hoi(bh)
ra = os.path.join(T, "t6_TDG.xlsx")
kd.ghi_ket_qua(khd, tl, kd.tinh(khd, tl), ra)
wb = km.mo(ra)
ct = km.cau_truc(wb)
rk = ct["kpi_dong"][ct["truc"][1][0]]
rong = goi(lambda: km.do_rong_cot(wb["KPI"]))
rk_ = rong.get(11, 8.7) if isinstance(rong, dict) else 8.7
can = min(409, math.ceil(len(K) * 1.2 / rk_) * float(wb["KPI"][f"K{rk}"].font.sz or 12) * 1.3)
cao = wb["KPI"].row_dimensions[rk].height or 15
kiem(cao >= can - 0.01, f"KPI dòng {rk}: cao {cao:g} pt ≥ ước cho cột K {can:g} pt")

print("== T7. Tệp ra đang bị khóa (như Excel đang mở) -> thông báo tiếng Việt, mã ≠ 0, không traceback ==")
khoa = os.path.join(T, "dang_mo.xlsx")
open(khoa, "wb").write(b"x")
k32 = ctypes.windll.kernel32 if os.name == "nt" else None
h = k32.CreateFileW(khoa, 0x80000000, 0, None, 3, 0x80, None) if k32 else None   # GENERIC_READ, share=0 (doc quyen)
import json  # noqa: E402
js = os.path.join(T, "t7.json")
json.dump({"ca_nhan": {"ho_ten": "Nguyễn Văn A", "don_vi": "Đơn vị X"}, "dau_viec": [dv(1, "Việc")]},
          open(js, "w", encoding="utf-8"), ensure_ascii=False)
r1 = subprocess.run([sys.executable, os.path.join(CC, "kpi_mau.py"), "--nhom", "hanh-chinh", "--json", js,
                     "--phuong-an", "muc-do", "--ra", khoa], capture_output=True, text=True, encoding="utf-8")
r2 = subprocess.run([sys.executable, os.path.join(CC, "kpi_danh_gia.py"), "danh-gia", "--ke-hoach", p, "--bang-hoi", bh,
                     "--ra", khoa], capture_output=True, text=True, encoding="utf-8")
if k32:
    k32.CloseHandle(h)
for ten, r in (("kpi_mau", r1), ("kpi_danh_gia", r2)):
    kiem(r.returncode != 0 and "Traceback" not in r.stderr and "đang mở" in r.stdout,
         f"{ten}: mã {r.returncode}, thông báo: {r.stdout.strip()[-90:]!r}")

print("== T8. Cả 6 nhóm: C–F đúng theo tiêu đề cột của mẫu ==")
for nhom in km.NHOM:
    p, _ = lap(nhom, [dv(1, "Việc A"), dv(4, "Việc Trục 4", nguoi_phoi_hop="Chi bộ")], "t8")
    wb = km.mo(p)
    ct = km.cau_truc(wb)
    kp = wb["KPI"]
    ok = True
    for n in (1, 4):
        r = ct["truc"][n][0]
        rk = ct["kpi_dong"][r]
        ok &= kp[f"C{rk}"].value == "Trưởng đơn vị" and kp[f"E{rk}"].value == "Đơn vị X" \
            and kp[f"F{rk}"].value == f"='Ke Hoach'!E{r}"
    kiem(ok and "KH18" not in ma(p, nhom), f"{nhom}: C, E, F đúng ở Trục 1 và 4; không KH18")

print("== T9. Không số nào đổi: tổng điểm, A, B, từng Trục, mức theo điểm = trước khi sửa ==")
CHUAN = {"truong-pho-don-vi": (74.282353, 24.0, 50.282353), "bo-mon": (72.54902, 24.0, 48.54902),
         "nha-giao": (72.54902, 24.0, 48.54902), "giao-vu": (78.184314, 24.0, 54.184314),
         "hanh-chinh": (78.184314, 24.0, 54.184314), "ho-tro": (83.521569, 24.0, 59.521569)}
B = [(4, 100, 90, "Vượt mức"), (1, 80, 100, "Hoàn thành đúng hạn"), (2, 100, 60, "Hoàn thành chậm tiến độ"),
     (1, 100, 100, "Hoàn thành đúng hạn")]
for nhom, (tong, a, b) in CHUAN.items():
    p, _ = lap(nhom, [dv(1, "Tổng hợp báo cáo công tác tháng", 3), dv(1, "Xây dựng kế hoạch quý", 1, "Cao"),
                      dv(2, "Góp ý quy chế", 2, "Thấp"), dv(4, "Sinh hoạt chi bộ")], "t9")
    khd = kd.doc_ke_hoach(p)
    bh = os.path.join(T, f"{nhom}_t9_bh.xlsx")
    kd.sinh_bang_hoi(khd, bh)
    w = load_workbook(bh)
    for row in w["A-Tieu-chi-chung"].iter_rows(min_row=5):
        if len(str(row[0].value or "")) > 2:
            row[3].value = (row[2].value or 0) * 0.8
    i = 0
    for row in w["B-KPI"].iter_rows(min_row=5):
        if str(row[0].value or "").startswith("B"):
            row[8].value, row[9].value, row[10].value, row[11].value = B[i]
            i += 1
    for row in w["C-Dieu-kien"].iter_rows(min_row=5):
        m_ = str(row[0].value or "")
        if m_[:1] in "CD" and m_[1:].isdigit():
            nd = km._kd(row[1].value)
            row[2].value = ({"C07": "Có", "C08": "Có", "C13": "Chưa có", "C15": kd.MUC[1]}.get(m_, "Không")
                            if m_[0] == "C" else ("Đạt" if "bang kiem" in nd else (100 if "gio giang" in nd else "Có")))
    w.save(bh)
    r = kd.tinh(khd, kd.doc_bang_hoi(bh))
    kiem(abs(r["tong"] - tong) < 1e-5 and abs(r["diem_a"] - a) < 1e-5 and abs(r["diem_b"] - b) < 1e-5,
         f"{nhom}: tổng {r['tong']:.6f} = {tong}")

print()
print("ĐẠT: 0 ca sai" if not sai else f"CÓ {len(sai)} CA SAI: {sai}")
sys.exit(1 if sai else 0)
