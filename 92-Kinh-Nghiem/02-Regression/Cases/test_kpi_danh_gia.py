# -*- coding: utf-8 -*-
"""Hoi quy cho skill ktc-kpi-tu-danh-gia (kpi_danh_gia.py + cau_truc_danh_gia trong kpi_mau.py) — lenh 25/9/2026.

Thu tren CA 6 mau vi tri (khong chi truong-pho-don-vi — mau VC-Hanh-Chinh lech dong). Moi phep kiem co ca nguoc.
Du lieu thu la gia dinh ("Nguyen Van A"), chay trong thu muc tam — khong dung du lieu ca nhan that.
Neu may co Excel: mo tep ket qua bang Excel (COM qua PowerShell), tinh lai cong thuc, so tong diem voi Python.
"""
import copy
import os
import shutil
import subprocess
import sys
import tempfile

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, os.path.join(GOC, "29-Cong-Cu"))
import kpi_calc as kc      # noqa: E402
import kpi_danh_gia as kd  # noqa: E402
import kpi_mau as km       # noqa: E402
import validate_plan as vp  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")
sai = []
T = tempfile.mkdtemp(prefix="kpi_dg_")


def kiem(dk, ten):
    print(("  OK " if dk else "  X  ") + ten)
    if not dk:
        sai.append(ten)


def dv(truc, nd, sl=2, muc="Trung bình"):
    return {"truc": truc, "noi_dung": nd, "cap_trinh": "Trưởng đơn vị", "muc_do": muc, "san_pham": "Báo cáo",
            "so_luong": sl, "thoi_han": "30/11/2026", "minh_chung": "Văn bản trên hệ thống Office"}


def lap(nhom, ten="kh.xlsx", ds=None):
    ds = ds or [dv(1, "Tổng hợp báo cáo công tác tháng"), dv(1, "Xây dựng kế hoạch quý", 1, "Cao"),
                dv(2, "Góp ý dự thảo quy chế", 1), dv(3, "Cập nhật dữ liệu lên hệ thống số", 1),
                dv(4, "Sinh hoạt chi bộ, học tập nghị quyết", 1), dv(5, "Tham gia hoạt động văn hóa", 1),
                dv(6, "Tham gia diễn tập phòng cháy", 1)]
    kh = {"ca_nhan": {"ho_ten": "Nguyễn Văn A", "don_vi": "Đơn vị X", "chuc_vu_chinh_quyen": "Viên chức"},
          "dau_viec": ds}
    kh = dict(kh, dau_viec=kc.quy_doi_ke_hoach(kh, "muc-do")["dau_viec"], phuong_an="muc-do")
    p = os.path.join(T, nhom + "_" + ten)
    km.ghi_ke_hoach(nhom, kh, p, "IV", "2026")
    return p


def dien(khd, bh, A=None, B=None, C=None):
    """Dien bang hoi nhu nguoi dung: A = diem moi tieu chi (mac dinh = toi da), B = (L, N, P, ket qua) theo SL,
    C = tra loi muc C (mac dinh: khong dac thu, du dieu kien)."""
    from openpyxl import load_workbook
    wb = load_workbook(bh)
    ws = wb["A-Tieu-chi-chung"]
    for row in ws.iter_rows(min_row=5):
        ma = str(row[0].value or "")
        if len(ma) > 2 and ma.startswith("A"):
            row[3].value = (A or {}).get(ma, row[2].value)
    ws = wb["B-KPI"]
    for row in ws.iter_rows(min_row=5):
        ma = str(row[0].value or "")
        if ma.startswith("B"):
            L, N, P, kq = (B or {}).get(ma, (row[4].value, 100, 100, "Hoàn thành đúng hạn"))
            row[8].value, row[9].value, row[10].value, row[11].value = L, N, P, kq
    ws = wb["C-Dieu-kien"]
    mac_dinh = {"C07": "Có", "C08": "Có", "C09": "Không", "C10": "Không", "C11": "Không", "C12": "Không",
                "C13": "Chưa có", "C14": "Không", "C15": kd.MUC[1]}
    for row in ws.iter_rows(min_row=5):
        ma = str(row[0].value or "")
        if ma[:1] in ("C", "D") and ma[1:].isdigit():
            if ma.startswith("D"):
                nd = km._kd(row[1].value)
                v = "Đạt" if "bang kiem" in nd else (100 if "gio giang" in nd else "Có")
            else:
                v = mac_dinh.get(ma, "Không")
            row[2].value = (C or {}).get(ma, v)
    wb.save(bh)
    return kd.doc_bang_hoi(bh)


def excel_tong(p, sheet, o):
    """Tinh lai bang Excel that (COM) va doc gia tri o; None neu khong co Excel."""
    ps = (f"$x=New-Object -ComObject Excel.Application;$x.DisplayAlerts=$false;"
          f"$w=$x.Workbooks.Open('{os.path.abspath(p)}');$x.CalculateFull();"
          f"$v=$w.Worksheets.Item('{sheet}').Range('{o}').Value2;$w.Close($false);$x.Quit();"
          f"[Console]::Out.Write([string]$v)")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=120)
        return float(r.stdout.strip().replace(",", ".")) if r.stdout.strip() else None
    except Exception:
        return None


print("== 1. cau_truc_danh_gia — ca 6 mau ==")
for nhom in km.NHOM:
    d = km.cau_truc_danh_gia(km.mo(km.tep_mau(nhom)))
    ok = (len(d["a"]) == 3 and abs(sum(g["diem_max"] for g in d["a"]) - 30) < 1e-9
          and all(g["diem_max"] >= 5 for g in d["a"]) and sorted(d["b"]) == [1, 2, 3, 4, 5, 6]
          and abs(sum(b["diem_max"] for b in d["b"].values()) - 70) < 1e-9 and d["dieu_kien"]["muc"]
          and d["de_xuat"] and len(d["ca_nhan"]) == 5)
    kiem(ok, f"{nhom}: 3 nhóm A = 30 (≥5/nhóm), 6 Trục = 70, khối II ({len(d['dieu_kien']['muc'])} điều kiện), III")
    kiem(km.nhan_nhom(km.mo(km.tep_mau(nhom))) == nhom, f"{nhom}: nhận đúng nhóm từ tiêu đề mẫu (không cần --nhom)")
d1 = km.cau_truc_danh_gia(km.mo(km.tep_mau("truong-pho-don-vi")))
d2 = km.cau_truc_danh_gia(km.mo(km.tep_mau("hanh-chinh")))
kiem(d1["b"][1]["dong"] != d2["b"][1]["dong"] and d1["dieu_kien"]["cot_kq"] != d2["dieu_kien"]["cot_kq"],
     "hai mẫu lệch dòng và lệch cột kết quả điều kiện — dò động, không ghi cứng")

print("== 2. Ca ngược cấu trúc: thiếu khối -> dừng, không đoán ==")
for moi, ten in ((lambda dg: [c for c in dg["B"] if str(c.value or "").startswith("Trục (3)")][0], "Trục (3)"),
                 (lambda dg: [c for c in dg["A"] if str(c.value or "").startswith("III.")][0], "III")):
    wb = km.mo(km.tep_mau("hanh-chinh"))
    moi(km.sheet_danh_gia(wb)).value = "đã xóa"
    try:
        km.cau_truc_danh_gia(wb)
        kiem(False, f"xóa {ten} -> phải báo LoiCauTruc")
    except km.LoiCauTruc as e:
        kiem(ten.split()[0] in str(e) or "III" in str(e), f"xóa {ten} -> LoiCauTruc: {str(e)[:80]}")

print("== 3. Giai đoạn 1: xóa số thực tế ví dụ của mẫu (Known-Issues #12), ẩn dòng trống, KH16 ==")
p = lap("hanh-chinh")
wb = km.mo(p)
kiem(wb["KPI"]["L9"].value is None and wb["KPI"]["N9"].value is None, "kế hoạch xuất ra không còn L9=4/N9=100 của mẫu")
kiem(km.mo(km.tep_mau("hanh-chinh"))["KPI"]["L9"].value == 4, "ca ngược: mẫu gốc vẫn có L9=4 (không sửa assets)")
kiem(wb["Ke Hoach"].row_dimensions[16].hidden and not wb["Ke Hoach"].row_dimensions[14].hidden,
     "dòng trống trong khối Trục bị ẩn, dòng có việc không ẩn")
kiem(wb["Ke Hoach"]["B14"].alignment.wrap_text, "ô nội dung đầu việc bật xuống dòng")
kiem(not [x for x in vp.kiem(p, "hanh-chinh", "muc-do")["loi"] if x["ma"] == "KH16"], "KH16 không báo kế hoạch sạch")
kiem([x for x in vp.kiem(km.tep_mau("hanh-chinh"), "hanh-chinh", "muc-do")["loi"] if x["ma"] == "KH16"],
     "ca ngược: KH16 bắt số ví dụ trong mẫu gốc")

print("== 4. Trọn quy trình trên cả 6 nhóm (sinh bảng hỏi -> dừng khi trống -> điền -> đánh giá -> xuất) ==")
for nhom in km.NHOM:
    p = lap(nhom)
    khd = kd.doc_ke_hoach(p)
    bh = os.path.join(T, nhom + "_bh.xlsx")
    kd.sinh_bang_hoi(khd, bh, "IV", "2026")
    try:
        kd.tinh(khd, kd.doc_bang_hoi(bh))
        kiem(False, f"{nhom}: bảng hỏi trống -> phải dừng")
    except kd.Dung:
        pass
    tl = dien(khd, bh)
    kq = kd.tinh(khd, tl)
    ra = os.path.join(T, nhom + "_TDG.xlsx")
    kd.ghi_ket_qua(khd, tl, kq, ra, "IV", "2026")
    dg = km.cau_truc_danh_gia(km.mo(ra))
    w = km.mo(ra)
    s = w[dg["sheet"]]
    g1 = dg["a"][0]
    ok = (abs(kq["tong"] - 100) < 1e-6 and kq["muc_theo_diem"] == kd.MUC[0] and khd["nhom"] == nhom
          and s[f"{g1['cot_diem']}{g1['tieu_chi'][0][0]}"].value == g1["tieu_chi"][0][3]
          and kd.MUC[1] in str(s.cell(dg["de_xuat"], 1).value)
          and "SUMPRODUCT" in str(w["KPI"]["R8"].value) and w["KPI"]["L9"].value == 2)
    kiem(ok, f"{nhom}: đủ điểm = {kd.cat2(kq['tong'])}, mức theo điểm HTXS, ghi mục A/III, KPI R chặn trần")
    if nhom in ("hanh-chinh", "truong-pho-don-vi"):
        v = excel_tong(ra, dg["sheet"], f"{dg['b'][1]['cot_dat']}{dg['tong']}")
        if v is None:
            print("  -- bỏ qua so với Excel: máy không có Excel")
        else:
            kiem(abs(v - kq["tong"]) < 1e-6, f"{nhom}: Excel tính lại tổng = {v:g}, khớp Python {kq['tong']:g}")
    try:
        kd.ghi_ket_qua(khd, tl, kq, p)
        kiem(False, "ca ngược: ghi đè kế hoạch đã duyệt phải bị chặn")
    except kd.Dung:
        pass

print("== 5. Ngưỡng điểm (không làm tròn có lợi) ==")
for t, m in ((89.99, 1), (90, 0), (69.99, 2), (70, 1), (49.99, 3), (50, 2)):
    kiem(kc.xep_loai_theo_diem(t)["muc_theo_diem"] == kd.MUC[m], f"{t} -> {kd.MUC[m]}")
kiem(kd.cat2(89.996) == "89,99" and kd.cat2(90) == "90,00", "hiển thị cắt 89,996 -> 89,99 (không làm tròn lên 90)")

print("== 6. Vượt 100% bị chặn trần theo chỉ tiêu; công thức gốc của mẫu thì không ==")
p = lap("hanh-chinh", "vuot.xlsx")
khd = kd.doc_ke_hoach(p)
bh = os.path.join(T, "vuot_bh.xlsx")
kd.sinh_bang_hoi(khd, bh)
v1 = [v for v in khd["viec"] if v["truc"] == 1]
tl = dien(khd, bh, B={v1[0]["ma"]: (6, 100, 100, "Vượt mức"), v1[1]["ma"]: (1, 50, 50, "Hoàn thành chậm tiến độ")})
kq = kd.tinh(khd, tl)
t1 = kq["truc"][1]
kiem(t1["pt_mau"] > 100 and t1["pt"] <= 100 + 1e-9, f"Trục 1: công thức mẫu {kd.cat2(t1['pt_mau'])}% -> chặn "
     f"{kd.cat2(t1['pt'])}% (việc vượt không bù cho việc thiếu)")
kiem(any("Đ11.6" in x for x in kq["canh_bao"]), "cảnh báo nêu Đ11.6 và chênh lệch với mẫu")
ra = os.path.join(T, "vuot_TDG.xlsx")
kd.ghi_ket_qua(khd, tl, kq, ra)
dg = km.cau_truc_danh_gia(km.mo(ra))
v = excel_tong(ra, dg["sheet"], f"{dg['b'][1]['cot_pt']}{dg['b'][1]['dong']}")
if v is not None:
    kiem(abs(v - t1["pt"]) < 1e-6, f"Excel tính lại % Trục 1 = {v:.4f} = Python {t1['pt']:.4f} (công thức chặn trần)")

print("== 7. Điều kiện thiếu nguồn -> 'Thiếu dữ liệu', không tự cho Đạt ==")
p = lap("nha-giao", "tn.xlsx")
khd = kd.doc_ke_hoach(p)
bh = os.path.join(T, "tn_bh.xlsx")
kd.sinh_bang_hoi(khd, bh)
d_bk = [f"D{tt}" for _, tt, nd, _, _ in khd["dg"]["dieu_kien"]["muc"] if "bang kiem" in km._kd(nd)][0]
kq = kd.tinh(khd, dien(khd, bh, C={d_bk: "Chưa có kết quả"}))
kiem(kq["dieu_kien"]["tong_hop"][kd.MUC[1]].startswith("Thiếu"), "bảng kiểm 'Chưa có kết quả' -> điều kiện HTT: Thiếu dữ liệu")
kq = kd.tinh(khd, dien(khd, bh, C={d_bk: "Đạt"}))
kiem(kq["dieu_kien"]["tong_hop"][kd.MUC[1]] == "Đạt", "ca ngược: đủ dữ liệu, đều đạt -> điều kiện HTT: Đạt")
d_gg = [f"D{tt}" for _, tt, nd, _, _ in khd["dg"]["dieu_kien"]["muc"] if "gio giang" in km._kd(nd)][0]
kq = kd.tinh(khd, dien(khd, bh, C={d_gg: 60}))
kiem(kq["dieu_kien"]["tong_hop"][kd.MUC[0]] == "Không đạt" and kq["dieu_kien"]["tong_hop"][kd.MUC[2]] == "Đạt",
     "giờ giảng 60%: không đạt điều kiện HTXS/HTT, đạt điều kiện HT (≥50%)")
kq = kd.tinh(khd, dien(khd, bh, C={d_bk: "Không đạt"}))
kiem(any("Bảng kiểm" in x for x in kq["dieu_kien"]["khong_hoan_thanh"]), "bảng kiểm Không đạt -> trường hợp Đ19.1d")

kq = kd.tinh(khd, dien(khd, bh))
kiem(kq["dieu_kien"]["tong_hop"][kd.MUC[0]] == "Không đạt", "không có việc vượt mức -> điều kiện HTXS (≥30% vượt mức) Không đạt")
kq = kd.tinh(khd, dien(khd, bh, B={v["ma"]: (v["so_luong"] + 1, 100, 100, "Vượt mức") for v in khd["viec"]}))
kiem(kq["dieu_kien"]["tong_hop"][kd.MUC[0]] == "Đạt", "ca ngược: mọi việc vượt mức, đủ dữ liệu -> điều kiện HTXS Đạt")

print("== 8. Dừng, không tự chấm ==")
for C, ten in (({"C01": "Có"}, "đào tạo tập trung ≥ 2 tháng (Đ21.6a)"), ({"C05": "Có"}, "≥1/2 quý (Đ21.4)")):
    try:
        kd.tinh(khd, dien(khd, bh, C=C))
        kiem(False, f"{ten} -> phải dừng")
    except kd.Dung as e:
        kiem("Đ21" in str(e), f"{ten} -> dừng, không chấm")
try:
    kd.tinh(khd, dien(khd, bh, A={"A1a": 2.5}))
    kiem(False, "điểm tiêu chí vượt tối đa -> phải dừng")
except kd.Dung as e:
    kiem("A1a" in str(e), "điểm tiêu chí vượt tối đa -> dừng")
tl = dien(khd, bh)
tl["A_muc"]["A1"] = "Mức 3"
try:
    kd.tinh(khd, tl)
    kiem(False, "mức nhóm chọn không khớp tổng điểm -> phải dừng")
except kd.Dung as e:
    kiem("Đ10.5" in str(e), "mức nhóm chọn không khớp tổng điểm nhóm -> dừng [Đ10.5]")
tl = dien(khd, bh)
tl["A_muc"]["A1"] = "Mức 1"
kiem(abs(kd.tinh(khd, tl)["tong"] - 100) < 1e-6, "ca ngược: mức nhóm chọn khớp -> chấm bình thường")
tl = dien(khd, bh)
tl["B"][khd["viec"][0]["ma"]]["trong_tam"] = "Mức 3"
try:
    kd.tinh(khd, tl)
    kiem(False, "trọng tâm Mức 3 mà hoàn thành 100% -> phải dừng")
except kd.Dung as e:
    kiem("Đ18" in str(e), "nhiệm vụ trọng tâm: % không khớp mức Đ18 -> dừng")

print("== 9. Cảnh báo cấu trúc mẫu (nhóm A < 5 điểm) là lỗi mẫu, không phải lỗi người dùng ==")
khd2 = copy.deepcopy(khd)
g3 = khd2["dg"]["a"][2]
g3["tieu_chi"] = g3["tieu_chi"][:2]
g3["diem_max"] = sum(x[3] for x in g3["tieu_chi"])
tl = dien(khd, bh)
kq = kd.tinh(khd2, tl)
kiem(any("Cấu trúc mẫu sai" in x and "Đ10.4" in x for x in kq["canh_bao"]), f"nhóm 3 = {g3['diem_max']:g} điểm < 5 "
     "-> cảnh báo cấu trúc mẫu")
kiem(not any("Cấu trúc mẫu sai" in x for x in kd.tinh(khd, tl)["canh_bao"]), "ca ngược: mẫu đúng -> không cảnh báo")

print("== 10. Người đứng đầu không cao hơn tập thể; tự đề xuất cao hơn mức theo điểm ==")
p = lap("truong-pho-don-vi", "ql.xlsx")
khd = kd.doc_ke_hoach(p)
bh = os.path.join(T, "ql_bh.xlsx")
kd.sinh_bang_hoi(khd, bh)
kq = kd.tinh(khd, dien(khd, bh, C={"C12": "Có", "C13": kd.MUC[1], "C15": kd.MUC[0]}))
kiem(any("Đ14.4" in x for x in kq["canh_bao"]), "đứng đầu: HTXS theo điểm > tập thể HTT -> cảnh báo Đ14.4")
kq = kd.tinh(khd, dien(khd, bh, C={"C12": "Không", "C13": kd.MUC[1]}))
kiem(not any("Đ14.4, Đ19.4]" in x and "cao hơn" in x for x in kq["canh_bao"]), "ca ngược: cấp phó -> không áp")
kq = kd.tinh(khd, dien(khd, bh, A={f"A1{k}": 0 for k in "abcdđeghik"}, C={"C15": kd.MUC[0]}))
kiem(kq["muc_theo_diem"] == kd.MUC[1] and any("CAO HƠN" in x for x in kq["canh_bao"]),
     f"tự đề xuất HTXS khi điểm {kd.cat2(kq['tong'])} -> cảnh báo")

print("== 11. Bảo mật: tệp tự đánh giá không vào git (QĐ 1923 Đ23.2) ==")
thu = os.path.join(GOC, "30-Ket-Qua", "2099-01-01", "KPI-ca-nhan", "TDG-KPI-Q4-2026_Nguyen-Van-A.xlsx")
r = subprocess.run(["git", "-C", GOC, "check-ignore", "-q", thu])
kiem(r.returncode == 0, "30-Ket-Qua/<ngày>/KPI-ca-nhan/TDG-*.xlsx nằm trong .gitignore")
os.makedirs(os.path.dirname(thu), exist_ok=True)
open(thu, "w").write("x")
out = subprocess.run(["git", "-C", GOC, "add", "-A", "--dry-run", "30-Ket-Qua/2099-01-01"], capture_output=True,
                     text=True).stdout
shutil.rmtree(os.path.join(GOC, "30-Ket-Qua", "2099-01-01"))
kiem("KPI-ca-nhan" not in out, "git add -A --dry-run không đưa tệp tự đánh giá vào git")
r = subprocess.run(["git", "-C", GOC, "check-ignore", "-q", os.path.join(GOC, "30-Ket-Qua", "2099-01-01", "Bao-cao",
                                                                          "x.xlsx")])
kiem(r.returncode != 0, "ca ngược: thư mục kết quả khác vẫn được git theo dõi")

shutil.rmtree(T, ignore_errors=True)
print()
print("ĐẠT: 0 ca sai" if not sai else f"CÓ {len(sai)} CA SAI: {sai}")
sys.exit(1 if sai else 0)
