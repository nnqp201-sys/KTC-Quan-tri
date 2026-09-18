# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "29-Cong-Cu"))
from docx import Document
from ktc_trackchanges import (TrackChanges, kiem_tra, nhat_ky_sua_doi,
                              doi_chieu_goc, BO, BO_SUNG, DIEU_CHINH)

TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Fixtures")
os.makedirs(TMP, exist_ok=True)
GOC = os.path.join(TMP, "goc.docx")
SUA = os.path.join(TMP, "sua.docx")

# --- dung file goc ---
d = Document()
d.add_paragraph("Muc 3. Ket qua tuyen sinh")
d.add_paragraph("Nam 2026 Nha truong tuyen duoc 2.000 hoc sinh, dat 80% chi tieu.")
d.add_paragraph("Phong da tham muu cho Lanh dao Truong ban hanh ke hoach tuyen sinh.")
t = d.add_table(rows=3, cols=3)
for i, r in enumerate(t.rows):
    for j, c in enumerate(r.cells):
        c.text = f"O{i}{j}"
d.save(GOC)
print("Goc:", len(Document(GOC).paragraphs), "doan,", len(Document(GOC).tables[0].rows), "hang bang")

# --- sua co dau vet ---
tc = TrackChanges(GOC)
tc.thay("2.000 hoc sinh", "2.150 hoc sinh")                 # dieu chinh
tc.thay("80%", "86%")                                        # dieu chinh
tc.xoa_cum("da tham muu cho Lanh dao Truong ")               # bo
tc.them_sau("Muc 3", "3.1. Cong tac tu van huong nghiep duoc trien khai som.")
tc.xoa_hang(0, 1)                                            # bo hang bang
out = tc.luu(SUA)
print("\n--- Nhat ky phien lam viec ---")
print(tc.bao_cao())

# --- validate ---
kq = kiem_tra(SUA)
print("--- Validator ---")
print("DAT:", kq["dat"])
for e in kq["loi"]:      print("  LOI:", e)
for w in kq["canh_bao"]: print("  CB :", w)
print("  thong ke:", kq["thong_ke"])

# --- doc lai nhat ky tu file (nhu the do nguoi khac sua) ---
print("\n--- nhat_ky_sua_doi() doc lai tu file ---")
print(nhat_ky_sua_doi(SUA))

# --- doi chieu goc ---
print("--- doi_chieu_goc ---")
print(doi_chieu_goc(GOC, SUA))

# --- kiem tra Word mo duoc: reload ---
try:
    Document(SUA); print("\nReload OK — file hop le voi python-docx")
except Exception as ex:
    print("\nReload LOI:", ex)

# --- chan ghi de file goc ---
try:
    TrackChanges(GOC).luu(GOC); print("LOI: khong chan ghi de!")
except ValueError as ex:
    print("Chan ghi de goc: OK -", ex)

# --- ca thu: soan lai tu dau roi trinh bay nhu ban sua (NT-2 cam) ---
gia = os.path.join(TMP, "gia_mao.docx")
g = Document()
g.add_paragraph("Muc 3. Ket qua tuyen sinh")
g.add_paragraph("Nam 2026 tuyen 2.150 em, vuot chi tieu de ra.")
g.add_paragraph("Ke hoach tuyen sinh da duoc ban hanh.")
g.save(gia)
print("\n--- Ca thu: soan lai tu dau ---")
print(doi_chieu_goc(GOC, gia))

# ===================================================================
# Ca thu nguoc cho 3 phuong thuc them 14/9/2026:
#   xoa_doan() · xoa_tu_den() · xoa_bang()
# Ly do ton tai: xoa_cum() chi gach TEXT, dau doan van con — chap nhan
# thay doi xong con lai hang loat doan RONG. Ba ham moi phai danh dau ca
# dau doan (<w:pPr><w:rPr><w:del/></w:rPr>). Neu dat <w:del> sai vi tri
# trong pPr thi Word bao "unreadable content" — phai kiem bang validator.
# ===================================================================
print("\n" + "=" * 60)
print("CA THU NGUOC — xoa nguyen doan / khoang doan / bang")
print("=" * 60)

G2 = os.path.join(TMP, "goc2.docx")
S2 = os.path.join(TMP, "sua2.docx")
d2 = Document()
for t_ in ["Phan I", "I. THUC TRANG", "Noi dung 1", "Noi dung 2",
           "II. KET QUA", "Noi dung 3", "Phan II", "III. DE XUAT", "Giu lai dong nay"]:
    d2.add_paragraph(t_)
tb = d2.add_table(rows=3, cols=2)
for i, r in enumerate(tb.rows):
    for j, c in enumerate(r.cells):
        c.text = f"B{i}{j}"
d2.save(G2)

tc2 = TrackChanges(G2)
n_khoang = tc2.xoa_tu_den("Phan I", "Phan II")     # bo 6 doan dau
n_doan = tc2.xoa_doan("III. DE XUAT")              # bo dung 1 doan
n_bang = tc2.xoa_bang(0)
tc2.luu(S2)

loi = []
def kiem2(ten, dk, mo_ta=""):
    print(f"  {'✓' if dk else '✗'} {ten:52s} {mo_ta}")
    if not dk:
        loi.append(ten)

kiem2("xoa_tu_den bo dung 6 doan (Phan I -> truoc Phan II)", n_khoang == 6, f"-> {n_khoang}")
kiem2("xoa_doan bo dung 1 doan", n_doan == 1, f"-> {n_doan}")
kiem2("xoa_bang bo dung 3 hang", n_bang == 3, f"-> {n_bang}")

kq2 = kiem_tra(S2)
kiem2("validator OOXML dat (dat <w:del> dung cho trong pPr)", kq2["dat"], str(kq2["loi"])[:60])

# Noi dung sau khi CHAP NHAN: chi con 'Phan II' va 'Giu lai dong nay'
from ktc_trackchanges import _text_day_du
con = [_text_day_du(p._element, chap_nhan=True).strip()
       for p in Document(S2).paragraphs]
con = [x for x in con if x]
kiem2("sau khi chap nhan chi con 2 doan", len(con) == 2, f"-> {con}")
kiem2("khong con doan RONG sot lai", "" not in con)

# Sau khi TU CHOI: phai tro ve nguyen ban 9 doan
tuchoi = [_text_day_du(p._element, chap_nhan=False).strip()
          for p in Document(S2).paragraphs]
tuchoi = [x for x in tuchoi if x]
kiem2("sau khi tu choi tro ve du 9 doan", len(tuchoi) == 9, f"-> {len(tuchoi)}")

# Moc khong ton tai -> phai bao loi, khong im lang bo qua
for ham, arg in (("xoa_doan", ("KHONG CO CUM NAY",)),
                 ("xoa_tu_den", ("KHONG CO", "Phan II"))):
    try:
        getattr(TrackChanges(G2), ham)(*arg)
        kiem2(f"{ham}: moc sai phai bao loi", False, "khong bao loi")
    except ValueError:
        kiem2(f"{ham}: moc sai bao ValueError", True)

if loi:
    print(f"\nTHAT BAI: {len(loi)} phep kiem")
    for x in loi:
        print("   ✗", x)
    sys.exit(1)
print("\nDAT — 3 phuong thuc xoa hoat dong dung va bao loi dung cho")
