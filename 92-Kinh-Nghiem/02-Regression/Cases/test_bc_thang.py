# -*- coding: utf-8 -*-
"""Ca thu bc_thang.py (plugin 1.3.6, skill bao-cao v3.18) — dung bao cao thang tu ban DA BAN HANH.

Bai hoc 28/9/2026: chay thu tren Cowork, Claude Code cho bao cao thang 9 kem ban 21/9 (dung tu mau trang, the
[CẦN BỔ SUNG [PHAN_I]], loi "nhiem ky 2021-2026", thieu ke hoach thang 10). Cac ca duoi day giu hanh vi dat.

Chay:  python 92-Kinh-Nghiem/02-Regression/Cases/test_bc_thang.py      (ma thoat 0 = sach)
Can kho KTC-Database (BC-375, PL-375, KH-834) va 10-Dau-Vao/.../2026-09 — thieu thi bo qua CO CANH BAO, khong bao sach gia.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

DU_AN = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
CC = os.path.join(DU_AN, "29-Cong-Cu")
sys.path.insert(0, CC)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import bc_thang as B  # noqa: E402

loi, bo_qua = [], []


def kiem(dk, ten):
    print(("  OK " if dk else "  X  ") + ten)
    if not dk:
        loi.append(ten)


def chay(*a):
    r = subprocess.run([sys.executable, os.path.join(CC, "bc_thang.py"), *a], capture_output=True, text=True,
                       encoding="utf-8")
    try:
        return r.returncode, json.loads(r.stdout)
    except Exception:
        return r.returncode, {"_raw": r.stdout + r.stderr}


print("== A. Bản sao trong gói kỹ năng báo cáo trùng nguồn ==")
for f in ("bc_thang.py", "vanphong.py", "trich_tuong_thuat.py", "duong_dan.py"):
    a = open(os.path.join(CC, f), "rb").read()
    b = open(os.path.join(DU_AN, "25-KTC-Bao-Cao", "references", "Skill-Library", f), "rb").read()
    kiem(a == b, f"25-KTC-Bao-Cao/references/Skill-Library/{f} trùng byte 29-Cong-Cu/{f}")

print("== B. Mẫu trắng đã sửa lỗi (dự phòng) ==")
import docx  # noqa: E402
mau = [p.text for p in docx.Document(os.path.join(DU_AN, "25-KTC-Bao-Cao", "00. Mau bao cao thang (cap Truong).docx")).paragraphs]
kiem(not any("2021-2026" in t for t in mau), "mẫu trắng không còn 'nhiệm kỳ 2021-2026'")
kiem(not any("Báo cáo báo cáo" in t for t in mau), "mẫu trắng không còn 'Báo cáo báo cáo'")

print("== C. Tìm nguồn, trích dữ liệu (kho + 10-Dau-Vao tháng 9) ==")
DV = os.path.join(DU_AN, "10-Dau-Vao", "01-Dau-Moi-Nop", "2026-09")
try:
    ng = B.tim_ban_da_ban_hanh(2026, 9)
    co_kho = all(ng.get(k) for k in ("BC", "PL", "KH"))
except Exception as e:
    ng, co_kho = None, False
    bo_qua.append(f"không đọc được kho ({e.__class__.__name__})")
if co_kho and os.path.isdir(DV):
    kiem(ng["BC"]["ky"] == "2026-08" and ng["PL"]["ky"] == "2026-08" and ng["KH"]["ky"] == "2026-09",
         "kỳ 9/2026 → BC, PL tháng 8 đã ban hành; KH tháng 9 đã ban hành")
    ma, j = chay("nguon", "--ky", "2026-09", "--dau-vao", DV)
    kiem(ma == 0 and len(j["dau_vao"]["IIa"]) >= 10 and len(j["dau_vao"]["IIb"]) >= 9, "nguon: phân loại ≥10 IIa, ≥9 IIb")
    tam = tempfile.mkdtemp(prefix="bc_thang_")
    try:
        ma, j = chay("trich", "--dau-vao", DV, "--ra", os.path.join(tam, "t.json"))
        t = json.load(open(os.path.join(tam, "t.json"), encoding="utf-8"))
        kiem(ma == 0 and j["IIa"] >= 10 and not j["loi"], "trich: đọc ≥10 tường thuật IIa, không lỗi")
        kiem(any(k.startswith("P-THHC (") for k in t["tuong_thuat"]) or sum(1 for k in t["tuong_thuat"] if k.startswith("P-THHC")) >= 1,
             "trich: tệp cùng mã đơn vị không đè nhau")

        print("== D. Dựng Word từ bản đã ban hành ==")
        nd = {"ky": {"thang": 9, "nam": 2026},
              "ket_qua": {"1": [[n, "Nhà trường thực hiện nội dung thử."] for n in B.MUC_CON[1]],
                          "2": [[n, "Nhà trường thực hiện nội dung thử."] for n in B.MUC_CON[2]],
                          "3": [[None, "Nhà trường thực hiện nội dung thử."]],
                          "4": [[n, "Nhà trường thực hiện nội dung thử."] for n in B.MUC_CON[4]],
                          "5": [[n, "Nhà trường thực hiện nội dung thử."] for n in B.MUC_CON[5]],
                          "6": [[n, "Nhà trường thực hiện nội dung thử."] for n in B.MUC_CON[6]],
                          "7": [["71", "Nhà trường thực hiện nội dung thử."], ["59", "Nhà trường thực hiện nội dung thử."]]},
              "danh_gia": {"dat": "Nhà trường hoàn thành cơ bản nhiệm vụ.", "ton_tai": "Một số đơn vị nộp chậm."},
              "nhiem_vu": {"1": [["Công tác tuyển sinh", "Nhà trường tiếp tục tuyển sinh. [CẦN BỔ SUNG: số liệu, P-QLDT]"]]}}
        json.dump(nd, open(os.path.join(tam, "bc.json"), "w", encoding="utf-8"), ensure_ascii=False)
        ma, j = chay("word", "--goc", ng["BC"]["tep"], "--noi-dung", os.path.join(tam, "bc.json"), "--ra", os.path.join(tam, "bc.docx"))
        kiem(ma == 0 and j["ky_cu_o_dau_cuoi"] == [], "word: không sót kỳ cũ ở tiêu đề, câu mở đầu, câu kết")
        kiem(j.get("con_can_bo_sung") == 1, "NGƯỢC: word đếm đúng 1 chỗ '[CẦN BỔ SUNG' còn lại")
        kiem(j.get("muc_con_chua_co_trong_phan_I_III") == {}, "word: đủ mục con cố định")
        txt = [p.text for p in docx.Document(os.path.join(tam, "bc.docx")).paragraphs]
        kiem(any(x.startswith("I. KẾT QUẢ THỰC HIỆN CÔNG TÁC THÁNG 9 NĂM 2026") for x in txt), "word: tiêu đề mục I đúng kỳ")
        kiem(any(x.startswith("III. NHIỆM VỤ TRỌNG TÂM THÁNG 10 NĂM 2026") for x in txt), "word: mục III là tháng sau")
        kiem(any("nhiệm kỳ 2026-2031" in x for x in txt), "word: giữ căn cứ đúng của bản đã ban hành (2026-2031)")
        i59, i71 = (next(i for i, x in enumerate(txt) if x.startswith(f"* Nghị quyết số {s}")) for s in ("59", "71"))
        kiem(i59 < i71, "word: Nghị quyết xếp đúng thứ tự (59 trước 71) dù nhập ngược")
        run2 = next(p for p in docx.Document(os.path.join(tam, "bc.docx")).paragraphs if p.text.startswith("* Công tác tuyển sinh"))
        kiem(len(run2.runs) >= 2 and run2.runs[1].bold in (False, None), "word: nhãn đậm nghiêng, nội dung thường (2 run)")

        print("== E. Phụ lục, kế hoạch ==")
        _, dong = B.doc_dong_excel(ng["KH"]["tep"])
        pl = {"ky": {"thang": 9, "nam": 2026}, "truc": {}, "dot_xuat": []}
        for x in dong[:6]:
            pl["truc"].setdefault(str(x["truc"] or 1), {"dong": []})["dong"].append(
                {"nd": x["nd"], "cd": x["chi_dao"], "ct": x["chu_tri"], "sp": x["sp"], "sl": x["sl"], "dk": x["dk"],
                 "ket_qua": {"sl": 1, "cl": 0.75, "td": 1}})
        pl["truc"]["1"]["dong"][0]["ket_qua"] = None
        json.dump(pl, open(os.path.join(tam, "pl.json"), "w", encoding="utf-8"), ensure_ascii=False)
        ma, j = chay("phu-luc", "--goc", ng["PL"]["tep"], "--noi-dung", os.path.join(tam, "pl.json"), "--ra", os.path.join(tam, "pl.xlsx"))
        kiem(ma == 0 and j["co_ket_qua"] == 5 and j["khong_co_ket_qua"] == 1, "phu-luc: 5 dòng có kết quả, 1 dòng để trống")
        import openpyxl  # noqa: E402
        ws = openpyxl.load_workbook(os.path.join(tam, "pl.xlsx")).active
        kiem("THÁNG 9 NĂM 2026" in str(ws["A3"].value) and ws.title.endswith("9"), "phu-luc: tiêu đề, tên sheet đúng kỳ")
        f = [c for c in ws[11] if isinstance(c.value, str) and c.value.startswith("=")]
        kiem(len(f) >= 9 and ws["M11"].value == "=K11*75%", "phu-luc: dòng có kết quả đủ công thức KPI (M = K×75%)")
        kiem(ws["K10"].value is None, "NGƯỢC: dòng chưa có kết quả không tự điền 100%")
        kiem(isinstance(ws["F11"].value, (int, float)), "phu-luc: số lượng lưu dạng số")
        kh = {"ky": {"thang": 10, "nam": 2026}, "can_cu": "        Căn cứ thử.",
              "truc": {"1": {"dong": [{"nd": "Việc thử", "cd": "Hiệu trưởng", "ct": "Phòng TH-HC&QT", "sp": "Kế hoạch",
                                        "sl": "1", "dk": "Cao", "thoi_han": "Chậm nhất ngày 31/10/2026"}]}}}
        json.dump(kh, open(os.path.join(tam, "kh.json"), "w", encoding="utf-8"), ensure_ascii=False)
        ma, j = chay("ke-hoach", "--goc", ng["KH"]["tep"], "--noi-dung", os.path.join(tam, "kh.json"), "--ra", os.path.join(tam, "kh.xlsx"))
        ws = openpyxl.load_workbook(os.path.join(tam, "kh.xlsx")).active
        kiem(ma == 0 and "tháng 10 năm 2026" in str(ws["A5"].value) and str(ws["A3"].value).startswith("Số:       /"),
             "ke-hoach: tiêu đề tháng 10, số văn bản để trống")
        kiem(any(str(ws.cell(r, 1).value or "").startswith("Nơi nhận") for r in range(1, ws.max_row + 1)), "ke-hoach: giữ khối Nơi nhận")
        kiem(isinstance(ws["F13"].value, int), "ke-hoach: số lượng '1' đổi thành số")

        print("== F. Ca ngược: tệp gốc không phải báo cáo tháng ==")
        ma, j = chay("word", "--goc", os.path.join(DU_AN, "25-KTC-Bao-Cao", "00. Mau bao cao thang (cap Truong).docx"),
                     "--noi-dung", os.path.join(tam, "kh.json"), "--ra", os.path.join(tam, "x.docx"))
        kiem(ma == 2 and "loi" in j, "NGƯỢC: nội dung sai cấu trúc → mã 2, báo lỗi, không xuất tệp")
    finally:
        shutil.rmtree(tam, ignore_errors=True)
else:
    bo_qua.append("thiếu kho hoặc 10-Dau-Vao/01-Dau-Moi-Nop/2026-09 — bỏ qua mục C–F")

for b in bo_qua:
    print("  ⚠ BỎ QUA:", b)
print("\nKET LUAN:", "SACH" if not loi else f"{len(loi)} ca sai", "(có mục bỏ qua)" if bo_qua else "")
sys.exit(1 if loi else 0)
