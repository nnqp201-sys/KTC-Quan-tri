# -*- coding: utf-8 -*-
"""Hoi quy kiem_the_thuc.py + hook ktc_the_thuc_hook.py — DL-20260919-003.

Ca THU NGUOC (biet chac sai) va ca DUNG — LL-20260914-001. Tu tao tep, khong phu thuoc Drive;
neu Drive co san thi do them 03-Templates(1) that (mau dung phai dat).
"""
import json
import os
import subprocess
import sys
import tempfile

GOC = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
sys.path.insert(0, os.path.join(GOC, "29-Cong-Cu"))
import docx  # noqa: E402
import openpyxl  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
import kiem_the_thuc as k  # noqa: E402

sai = []


def kiem(dk, ten):
    print(("  OK " if dk else "  X  ") + ten)
    if not dk:
        sai.append(ten)


def ma(p):
    return {x[1] for x in k.kiem_tep(p)}


TMP = tempfile.mkdtemp(prefix="ktc_tt_")

# 1. Document() rong — loi da mac 18/9/2026 (kho Letter)
p_rong = os.path.join(TMP, "rong.docx")
d = docx.Document()
d.add_paragraph("Nội dung báo cáo " * 10)
d.save(p_rong)
m = ma(p_rong)
kiem("TT01" in m, "TT01 Document() rỗng bị bắt khổ Letter")
kiem("TT02" in m, "TT02 lề mặc định (trái 3,18 / 2,54 cm) bị bắt")
kiem("TT03" in m, "TT03 phông Calibri mặc định bị bắt")

# 2. Van ban dung chuan
p_dung = os.path.join(TMP, "dung.docx")
d = docx.Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
st = d.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(14)
t = d.add_table(rows=2, cols=2)
for o, (txt, co) in zip([t.cell(0, 0), t.cell(0, 1), t.cell(1, 0), t.cell(1, 1)],
                        [("UBND TỈNH QUẢNG NGÃI", 13), ("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", 13),
                         ("TRƯỜNG CAO ĐẲNG KON TUM", 13), ("Độc lập - Tự do - Hạnh phúc", 14)]):
    r = o.paragraphs[0].add_run(txt)
    r.font.size, r.bold = Pt(co), True
for _ in range(3):
    d.add_paragraph("Nhà trường triển khai nhiệm vụ trọng tâm theo kế hoạch đã ban hành, bảo đảm tiến độ. " * 2)
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
for p in d.paragraphs:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
d.save(p_dung)
kq = k.kiem_tep(p_dung)
kiem(not [x for x in kq if x[0] <= 2], f"văn bản đúng chuẩn không có Mức 1–2 (được: {kq})")

# 3. Ca thu nguoc tung loi
d = docx.Document(p_dung)
d.tables[0].cell(0, 0).paragraphs[0].runs[0].text = "UBND TỈNH KON TUM"
d.tables[0].cell(0, 1).paragraphs[0].runs[0].font.size = Pt(12)
d.paragraphs[0].runs[0].font.name = ".VnTime"
for p in d.paragraphs:
    for r in p.runs:
        r.font.size = Pt(13)
p_sai = os.path.join(TMP, "sai.docx")
d.save(p_sai)
m = ma(p_sai)
kiem("TT08" in m, "TT08 cơ quan chủ quản “UBND TỈNH KON TUM” bị bắt")
kiem("TT09" in m, "TT09 Quốc hiệu cỡ 12 bị bắt")
kiem("TT03" in m, "TT03 .VnTime bị bắt")
kiem("TT04" in m, "TT04 nội dung cỡ 13 bị bắt")

# 4. Bien the ten TNR -> Muc 3, khong phai Muc 2
d = docx.Document(p_dung)
d.paragraphs[0].runs[0].font.name = "TimesNewRomanPSMT"
p_bt = os.path.join(TMP, "bienthe.docx")
d.save(p_bt)
kq = k.kiem_tep(p_bt)
kiem(any(x[1] == "TT03b" and x[0] == 3 for x in kq) and not any(x[1] == "TT03" for x in kq),
     "TT03b biến thể TimesNewRomanPSMT chỉ Mức 3")

# 5. xlsx
p_x = os.path.join(TMP, "kh.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
from openpyxl.styles import Font  # noqa: E402
for i in range(1, 6):
    for j in range(1, 9):
        c = ws.cell(i, j, f"ô {i}-{j}")
        c.font = Font(name="Arial", sz=11)
wb.save(p_x)
m = ma(p_x)
kiem("TX02" in m, "TX02 xlsx phông Arial bị bắt (skill xlsx cho phép, KTC thì không)")
for row in ws.iter_rows():
    for c in row:
        c.font = Font(name="Times New Roman", sz=14)
ws.page_setup.paperSize = 9
ws.page_setup.orientation = "landscape"
wb.save(p_x)
kiem(not [x for x in k.kiem_tep(p_x) if x[0] <= 2], "xlsx TNR, A4 ngang đạt")

# 6. tao_tu_mau: .dotx -> .docx mo duoc, giu le
p_mau = os.path.join(TMP, "mau.dotx")
raw = open(p_dung, "rb").read()
open(p_mau, "wb").write(k._doi_kieu(raw, b"wordprocessingml.document.main+xml",
                                     b"wordprocessingml.template.main+xml"))
p_tu_mau = k.tao_tu_mau(p_mau, os.path.join(TMP, "tu_mau.docx"))
d2 = docx.Document(p_tu_mau)
kiem(abs(d2.sections[0].left_margin.cm - 3) < 0.01, "tao_tu_mau giữ lề trái 3 cm của mẫu")

# 7. Hook: trong du an KTC, tep sai -> ma 2; tep dung -> 0; ngoai du an -> 0
HOOK = os.path.join(GOC, "31-Plugin", "scripts", "ktc_the_thuc_hook.py")
du_an = os.path.join(TMP, "du_an")
os.makedirs(os.path.join(du_an, "90-Nhat-Ky-Van-Hanh"))
os.makedirs(os.path.join(du_an, "30-Ket-Qua"))
env = dict(os.environ, USERPROFILE=TMP, HOME=TMP)
os.makedirs(os.path.join(TMP, ".claude"), exist_ok=True)


def chay_hook(tep, cwd):
    vao = json.dumps({"tool_name": "Write", "tool_input": {"file_path": tep}, "cwd": cwd})
    return subprocess.run([sys.executable, HOOK], input=vao.encode(), capture_output=True, env=env)


import shutil  # noqa: E402
t_sai = shutil.copy(p_sai, os.path.join(du_an, "30-Ket-Qua", "sai.docx"))
t_dung = shutil.copy(p_dung, os.path.join(du_an, "30-Ket-Qua", "dung.docx"))
r = chay_hook(t_sai, du_an)
kiem(r.returncode == 2 and "TT08" in r.stderr.decode("utf-8", "ignore"), "hook: tệp sai trong dự án KTC → mã 2, báo TT08")
kiem(chay_hook(t_dung, du_an).returncode == 0, "hook: tệp đúng → mã 0")
kiem(chay_hook(p_sai, TMP).returncode == 0, "hook: ngoài dự án KTC → không can thiệp")

# 8. Mau that tren Drive (neu co)
try:
    import duong_dan
    mau_dir = os.path.join(duong_dan.ktc_database(canh_bao_ban_cu=False), "03-Templates(1)")
except Exception:
    mau_dir = None
if mau_dir and os.path.isdir(mau_dir):
    loi_mau = {}
    for f in sorted(os.listdir(mau_dir)):
        if f.endswith((".dotx", ".xltx")) and not f.startswith("10-"):   # 10-Giay-moi: loi mau da biet
            nang = [x for x in k.kiem_tep(os.path.join(mau_dir, f)) if x[0] <= 2]
            if nang:
                loi_mau[f] = nang
    kiem(not loi_mau, f"15 mẫu 03-Templates(1) (trừ 10-Giay-moi) không bị báo Mức 1–2: {loi_mau}")
    kiem("TT08" in ma(os.path.join(mau_dir, "10-Giay-moi.dotx")), "lỗi đã biết của 10-Giay-moi vẫn bị bắt")
else:
    print("  -- bỏ qua ca mẫu thật: không đọc được Drive")

print(f"\n{'ĐẠT' if not sai else 'KHÔNG ĐẠT'}: {len(sai)} ca sai")
sys.exit(1 if sai else 0)
