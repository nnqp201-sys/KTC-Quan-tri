# -*- coding: utf-8 -*-
"""Ca thử công cụ office_ca_nhan.py (Skill 38, KTC-Bao-Cao v3.21) — dữ liệu tổng hợp, không dùng dữ liệu Office thật."""
import datetime as dt
import json
import os
import sys
import tempfile

import openpyxl

GOC = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.path.insert(0, os.path.join(GOC, "25-KTC-Bao-Cao", "references", "Skill-Library"))
import office_ca_nhan as oc  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
loi = []


def kt(ten, dk):
    print(("  ✓ " if dk else "  ✗ ") + ten)
    if not dk:
        loi.append(ten)


def tep_vb(p):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["DANH MỤC VĂN BẢN ĐÃ XỬ LÝ"]); ws.append(["Nguồn: thử"])
    ws.append(["STT", "Số ký hiệu", "Loại văn bản", "Trích yếu", "Ngày soạn thảo", "Ngày ban hành", "Người soạn thảo", "Độ mật"])
    ws.append([1, "10/KH-CĐKT", "", "Kế hoạch A", dt.datetime(2026, 9, 1), dt.datetime(2026, 9, 2), "Trần Văn A", ""])
    ws.append([2, "10/KH-CĐKT", "", "Kế hoạch A (trùng)", "01/09/2026", "02/09/2026", "Trần  Văn A", ""])
    ws.append([3, "11/BC-CĐKT", "", "Báo cáo B", "30/06/2026", "30/06/2026", "Trần Văn A", ""])          # ngoài kỳ
    ws.append([4, "12/QĐ-CĐKT", "", "Quyết định C", "05/09/2026", "06/09/2026", "Lê Thị B", ""])           # người khác
    ws.append([5, "13/BC-CĐKT", "", "Nội dung mật", "05/09/2026", "06/09/2026", "Trần Văn A", "Mật"])     # có độ mật
    ws.append([6, "", "", "Dự thảo D", "07/09/2026", None, "Trần Văn A", ""])                              # chưa ban hành
    wb.save(p)


def tep_cv(p):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(4):
        ws.append(["tiêu đề"])
    ws.append(["STT", "Mức độ công việc", "Ngày giao việc", "Ngày hết hạn", "Chủ trì", "Tên công việc", "Tình trạng", "Kết quả thực hiện"])
    ws.append([1, "Thường", "08/09/2026", "11/09/2026", "Trần Văn A", "Tham mưu kế hoạch X", "Chưa hoàn thành - Quá hạn", ""])
    ws.append([2, "Thường", "08/09/2026", "20/09/2026", "Lê Thị B", "Chủ trì, phối hợp với phòng Thử tham mưu Y", "Hoàn thành", "Đã xong"])
    ws.append([3, "Thường", "08/09/2026", "20/09/2026", "Lê Thị B", "Việc không liên quan", "Hoàn thành", "Đã xong"])
    wb.save(p)


with tempfile.TemporaryDirectory() as d:
    v, c, r = os.path.join(d, "vb.xlsx"), os.path.join(d, "cv.xlsx"), os.path.join(d, "ra.xlsx")
    tep_vb(v); tep_cv(c)
    ma = oc.main(["--nguoi", "Trần Văn A", "--don-vi", "phòng Thử", "--tu", "01/07/2026", "--den", "30/09/2026", "--ra", r, v, c])
    kt("mã thoát 0", ma == 0)
    wb = openpyxl.load_workbook(r)
    kt("đủ 4 sheet", wb.sheetnames == ["Van-ban-da-xu-ly", "Cong-viec", "Tong-hop", "Nguon"])
    ws = wb["Van-ban-da-xu-ly"]
    skh = [ws.cell(i, 2).value for i in range(4, ws.max_row + 1)]
    kt("loại trùng, ngoài kỳ, người khác, có độ mật: còn 2 văn bản", len(skh) == 2)
    kt("không có trích yếu của văn bản có độ mật", all("mật" not in str(ws.cell(i, 4).value) for i in range(4, ws.max_row + 1)))
    kt("suy loại từ ký hiệu KH", any(ws.cell(i, 3).value == "Kế hoạch" for i in range(4, ws.max_row + 1)))
    kt("cờ văn bản chưa ban hành", any("Chưa có ngày ban hành" in str(ws.cell(i, 14).value) for i in range(4, ws.max_row + 1)))
    kt("cột đề xuất để trống (không tự đặt Trục, hệ số)", all(ws.cell(i, j).value in (None, "") for i in range(4, ws.max_row + 1) for j in (9, 10, 11, 12)))
    ws = wb["Cong-viec"]
    vt = [ws.cell(i, 7).value for i in range(4, ws.max_row + 1)]
    kt("công việc: 1 chủ trì + 1 phối hợp suy từ nội dung", len(vt) == 2 and "Chủ trì" in vt and any("cần xác nhận" in x for x in vt))
    kt("cờ quá hạn", any("Quá hạn" in str(ws.cell(i, 12).value) for i in range(4, ws.max_row + 1)))
    # vai trò chỉ phối hợp
    oc.main(["--nguoi", "Trần Văn A", "--don-vi", "phòng Thử", "--tu", "01/07/2026", "--den", "30/09/2026", "--vai-tro", "phoi-hop", "--ra", r, v, c])
    wb = openpyxl.load_workbook(r)
    kt("lọc vai trò phối hợp: 0 văn bản, 1 công việc", wb["Van-ban-da-xu-ly"].max_row == 3 and wb["Cong-viec"].max_row == 4)
    # bản chụp json: đọc thiếu phải cảnh báo
    j = os.path.join(d, "chup.json")
    json.dump({"nguon": "thử", "thoi_diem_doc": "x", "so_dong_trang_bao": {"van_ban": 5},
               "van_ban": [{"so_ky_hieu": "20/TB-CĐKT", "trich_yeu": "T", "ngay_ban_hanh": "03/09/2026", "nguoi_soan": "Trần Văn A"}]},
              open(j, "w", encoding="utf-8"), ensure_ascii=False)
    _, _, gc = oc.doc_json(j)
    kt("bản chụp đọc thiếu → cảnh báo ĐỌC THIẾU", any("ĐỌC THIẾU" in x for x in gc))
    # tệp sai cấu trúc: không báo sạch
    s = os.path.join(d, "sai.xlsx")
    wb = openpyxl.Workbook(); wb.active.append(["a", "b"]); wb.save(s)
    kt("tệp không nhận ra cấu trúc → mã thoát 1 (THIEU_DU_LIEU)", oc.main(["--nguoi", "A", "--tu", "01/07/2026", "--den", "30/09/2026", "--ra", r, s]) == 1)
    kt("khoảng thời gian ngược → mã thoát 2", oc.main(["--nguoi", "A", "--tu", "30/09/2026", "--den", "01/07/2026", "--ra", r, v]) == 2)

print("\nKẾT QUẢ:", "ĐẠT" if not loi else f"HỎNG {len(loi)}")
sys.exit(1 if loi else 0)
