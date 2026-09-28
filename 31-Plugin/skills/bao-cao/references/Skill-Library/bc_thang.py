# -*- coding: utf-8 -*-
"""Dung BAO CAO THANG cap Truong tu BAN DA BAN HANH (Nguyen tac 7) — bo cong cu chung, khong ghi cung ky, ten tep.

Tong quat hoa quy trinh da cho ket qua dat ngay 21/9/2026 (build_bc2/build_xl + kich ban thang 9). Claude SOAN NOI DUNG
(tuong thuat cap Truong, chu the "Nha truong", nhan muc con co dinh); cong cu lo phan co hoc: tim ban da ban hanh, trich
du lieu don vi, mo ban da ban hanh roi thay noi dung giu nguyen the thuc, dung lai cong thuc KPI, tu kiem.

    python bc_thang.py nguon    --ky 2026-09 [--dau-vao <thu muc nop>]      # ban da ban hanh + tep don vi (JSON)
    python bc_thang.py trich    --dau-vao <thu muc nop> --ra trich.json      # tuong thuat IIa theo Truc + dong IIb, Ib
    python bc_thang.py word     --goc <BC da ban hanh .docx> --noi-dung bc.json --ra <.docx>
    python bc_thang.py phu-luc  --goc <PL da ban hanh .xlsx> --noi-dung pl.json --ra <.xlsx>
    python bc_thang.py ke-hoach --goc <KH thang da ban hanh .xlsx> --noi-dung kh.json --ra <.xlsx>

Cau truc JSON noi dung: xem `references/Skill-Library/37-Quy-Trinh-Bao-Cao-Thang-Tu-Ban-Da-Ban-Hanh.md` (skill bao-cao).
Moi lenh dung san pham tu kiem va in JSON: so cho con `[CẦN BỔ SUNG`, ky cu con sot o phan dau/cuoi, vi pham van phong
cap Truong, so cong thuc — ma thoat 0 (xuat duoc) / 2 (thieu dau vao, loi cau truc tep goc).
"""
import argparse
import copy
import datetime as dt
import glob
import json
import os
import re
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
# Ban trong goi ky nang (skills/bao-cao/references/Skill-Library/) dung them cong cu o scripts/ goc plugin
_GOC_PLUGIN = os.path.abspath(os.path.join(HERE, "..", "..", "..", "..", "scripts"))
if os.path.isdir(_GOC_PLUGIN):
    sys.path.append(_GOC_PLUGIN)

TEN_TRUC = {
    1: "Thực hiện mục tiêu phát triển kinh tế - xã hội và nhiệm vụ chính trị",
    2: "Hoàn thiện thể chế, đẩy mạnh phân cấp, phân quyền gắn với kiểm tra, giám sát",
    3: "Thúc đẩy phát triển KH-CN, đổi mới sáng tạo và chuyển đổi số",
    4: "Xây dựng Đảng và hệ thống chính trị; phòng, chống tham nhũng, tiêu cực",
    5: "Phát triển văn hóa, con người, bảo đảm an sinh xã hội, nâng cao đời sống Nhân dân",
    6: "Củng cố quốc phòng, an ninh, giữ vững ổn định chính trị - xã hội, "
       "nâng cao hiệu quả đối ngoại và hội nhập quốc tế",
}
# Danh muc muc con CO DINH theo mau bao cao thang cap Truong (Skill 33 BUOC 0B) — Truc 3 khong co muc con
MUC_CON = {
    1: ["Công tác tuyển sinh", "Công tác đào tạo", "Công tác khảo thí", "Công tác bảo đảm chất lượng",
        "Công tác kế hoạch, tổng hợp", "Công tác tổ chức, cán bộ"],
    2: ["Về thể chế", "Công tác Kiểm tra, giám sát"],
    3: [],
    4: ["Công tác xây dựng Đảng", "Chấp hành kỷ cương hành chính", "Công tác Đảng, Công đoàn, Đoàn Thanh niên"],
    5: ["Công tác quản lý cơ sở vật chất", "Công tác Tài chính", "Công tác an sinh giáo dục", "Công tác truyền thông"],
    6: ["Về Quốc phòng - An ninh", "Về hoạt động Đối ngoại và Hợp tác", "Về hoạt động hợp tác phát triển"],
}
TEN_NQ = {
    "59": "Nghị quyết số 59-NQ/TW ngày 24/01/2025 hội nhập quốc tế trong tình hình mới",
    "66": "Nghị quyết số 66-NQ/TW ngày 30/4/2025 đổi mới công tác xây dựng và thi hành pháp luật",
    "68": "Nghị quyết số 68-NQ/TW ngày 04/5/2025 phát triển kinh tế tư nhân",
    "79": "Nghị quyết số 79-NQ/TW ngày 06/01/2026 phát triển kinh tế nhà nước",
    "70": "Nghị quyết số 70-NQ/TW ngày 20/8/2025 bảo đảm an ninh năng lượng quốc gia",
    "71": "Nghị quyết số 71-NQ/TW ngày 22/8/2025 về đột phá phát triển giáo dục và đào tạo",
    "72": "Nghị quyết số 72-NQ/TW ngày 09/9/2025 bảo vệ, chăm sóc sức khỏe nhân dân",
    "80": "Nghị quyết số 80-NQ/TW ngày 07/01/2026 phát triển văn hóa Việt Nam",
}
THU_TU_NQ = ["59", "66", "68", "79", "70", "71", "72", "80"]
DIEM = {"Khó và phức tạp": 200, "Cao": 150, "Trung bình": 120, "Thấp": 100}
CAN_BO_SUNG = "[CẦN BỔ SUNG"


def _kd(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d").replace("Đ", "D").lower()


def _ky(s):
    m = re.fullmatch(r"(20\d\d)-(\d{1,2})", s.strip())
    if not m:
        raise SystemExit("--ky phải dạng YYYY-MM, ví dụ 2026-09")
    return int(m.group(1)), int(m.group(2))


def _truoc(y, m):
    return (y, m - 1) if m > 1 else (y - 1, 12)


def _sau(y, m):
    return (y, m + 1) if m < 12 else (y + 1, 1)


# ================================================================ nguon
def _thang_tep(p):
    t = _kd(os.path.basename(p))
    m = re.search(r"thang-(\d{1,2})", t)
    y = re.search(r"(20\d\d)", t[m.end():] if m else t) if m else None
    return (int(y.group(1)), int(m.group(1))) if (m and y) else None


def tim_ban_da_ban_hanh(y, m):
    """BC, PL thang truoc (ky N-1) va KH thang N (ke hoach cung ky) trong kho — uu tien 04-Good-Documents."""
    import duong_dan
    db = duong_dan.ktc_database()
    tat = glob.glob(os.path.join(db, "0[24]-*", "**", "*.*"), recursive=True)

    def loai(p):
        b = _kd(os.path.basename(p))
        if b.startswith("~$"):
            return None
        if b.endswith(".docx") and b.startswith("bc-") and "ket-qua" in b and "thang" in b:
            return "BC"
        if b.endswith(".xlsx") and b.startswith("pl-") and "thang" in b:
            return "PL"
        if b.endswith(".xlsx") and b.startswith("kh-") and "cong-tac-thang" in b:
            return "KH"
        return None
    ds = {"BC": [], "PL": [], "KH": []}
    for p in tat:
        k = loai(p)
        tm = _thang_tep(p) if k else None
        if tm:
            ds[k].append((tm, 0 if "04-Good-Documents" in p else 1, p))
    ra = {}
    for k, dich in (("BC", _truoc(y, m)), ("PL", _truoc(y, m)), ("KH", (y, m))):
        hop = sorted([x for x in ds[k] if x[0] <= dich], key=lambda x: (x[0], -x[1]))
        ra[k] = {"tep": hop[-1][2], "ky": f"{hop[-1][0][0]}-{hop[-1][0][1]:02d}",
                 "dung_ky": hop[-1][0] == dich} if hop else None
    # nguon bo tro cho ke hoach thang sau va dau moi chua nop (Skill 37 muc 2.3): CTCT nam, ke hoach quy, ket luan giao ban
    def ngay_ten(p):
        m = re.search(r"_(20\d{6})_", os.path.basename(p))
        return m.group(1) if m else "0"
    ctct = [p for p in tat if re.search(r"ctct|chuong-trinh-cong-tac|chuong trinh cong tac", _kd(os.path.basename(p)))
            and str(y) in os.path.basename(p) and p.lower().endswith((".xlsx", ".docx"))]
    quy = [p for p in tat if re.search(r"ke-hoach-cong-tac-quy|ke hoach cong tac quy", _kd(os.path.basename(p)))
           and p.lower().endswith((".xlsx", ".docx"))]
    gb = [p for p in tat if re.search(r"giao-ban|giao ban|tbkl", _kd(os.path.basename(p))) and p.lower().endswith(".docx")]
    ra["bo_tro"] = {"ctct_nam": sorted(ctct), "ke_hoach_quy": sorted(quy),
                    "ket_luan_giao_ban_gan_nhat": sorted(gb, key=ngay_ten)[-6:]}
    ra["kho"] = db
    return ra


def _tieu_de_xlsx(p):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return ""
    t = []
    for ws in wb.worksheets[:2]:
        for r in ws.iter_rows(min_row=1, max_row=8, values_only=True):
            t += [str(c) for c in r if c]
    return _kd(" ".join(t))


def phan_loai_dau_vao(thu_muc):
    """Tep don vi nop: IIa (tuong thuat .docx), IIb (ket qua .xlsx), Ib (ke hoach .xlsx) — theo ten thu muc, roi noi dung."""
    try:
        from doi_soat_so_lieu import ma_don_vi
    except Exception:
        ma_don_vi = None
    ra = {"IIa": [], "IIb": [], "Ib": [], "khac": []}
    for p in sorted(glob.glob(os.path.join(thu_muc, "**", "*.*"), recursive=True)):
        b = os.path.basename(p)
        if b.startswith(("~$", ".")) or not b.lower().endswith((".docx", ".xlsx")):
            continue
        duong = _kd(os.path.relpath(p, thu_muc))
        if b.lower().endswith(".docx"):
            k = "IIa"
        elif re.search(r"\biib\b|pl ?iib|ket qua|\biic\b", duong):
            k = "IIb"
        elif re.search(r"\bib\b|phu luc ib|ke hoach", duong):
            k = "Ib"
        else:
            td = _tieu_de_xlsx(p)
            k = "IIb" if ("iib" in td or "ket qua" in td) else "Ib" if ("ib" in td or "ke hoach" in td) else "khac"
        ma = None
        if ma_don_vi:
            try:
                ma = ma_don_vi(p)
            except Exception:
                ma = None
        ra[k].append({"tep": p, "ma": ma or "?" + os.path.splitext(b)[0]})
    return ra


# ================================================================ trich
def _g(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.day:02d}/{v.month:02d}/{v.year}"
    return " ".join(str(v).split())


def doc_dong_excel(p):
    """Dong nhiem vu cua Phu luc IIb/Ib (sheet dau khong phai 'quy'): muc I/II, Truc, cac cot A..P."""
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = next((w for w in wb.worksheets if "quy" not in _kd(w.title).split()), wb.worksheets[0])
    h = next((i for i, r in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1)
              if any(c and "NỘI DUNG CÔNG VIỆC" in str(c).upper() for c in r)), None)
    if h is None:
        return ws.title, []
    out, muc, truc = [], "I", 0
    for ri, r in enumerate(ws.iter_rows(min_row=h + 1, values_only=True), h + 1):
        r = list(r) + [None] * 20
        stt, nd = _g(r[0]), _g(r[1])
        if not nd or _kd(nd).startswith(("ghi chu", "noi nhan", "muc do")) or re.match(r"^\(\d+\)", stt):
            continue
        if stt in ("I", "II") and "nhiệm vụ" in nd.lower():
            muc, truc = stt, 0
            continue
        mt = re.search(r"Trục\s*\(?(\d)\)?", nd[:14])
        if mt and re.fullmatch(r"\d", stt):
            truc = int(mt.group(1))
            continue
        if not any(r[2:5]) and len(nd) < 15:
            continue
        out.append({"hang": ri, "muc": muc, "truc": truc, "stt": stt, "nd": nd, "chi_dao": _g(r[2]),
                    "chu_tri": _g(r[3]), "sp": _g(r[4]), "sl": _g(r[5]), "dk": _g(r[6]),
                    "cot_8_16": [_g(x) for x in r[7:16]]})
    return ws.title, out


def trich(thu_muc):
    import trich_tuong_thuat as T
    pl = phan_loai_dau_vao(thu_muc)
    ra = {"thu_muc": thu_muc, "tuong_thuat": {}, "iib": {}, "ib": {}, "loi": []}
    for x in pl["IIa"]:
        try:
            d = T.doc_mot(x["tep"])
            d["tep"] = os.path.basename(x["tep"])
            khoa = x["ma"] if x["ma"] not in ra["tuong_thuat"] else f"{x['ma']} ({os.path.splitext(d['tep'])[0]})"
            ra["tuong_thuat"][khoa] = d          # 2 tep cung ma (vd Ban Truyen thong va Phong TH-HC&QT) — giu ca hai
        except Exception as e:
            ra["loi"].append(f"{x['tep']}: {e.__class__.__name__} {e}")
    for k, dich in (("IIb", "iib"), ("Ib", "ib")):
        for x in pl[k]:
            try:
                sheet, dong = doc_dong_excel(x["tep"])
                ra[dich].setdefault(x["ma"], []).append({"tep": os.path.basename(x["tep"]), "sheet": sheet, "dong": dong})
            except Exception as e:
                ra["loi"].append(f"{x['tep']}: {e.__class__.__name__} {e}")
    ra["thong_ke"] = {"IIa": len(ra["tuong_thuat"]), "IIb": len(ra["iib"]), "Ib": len(ra["ib"]),
                      "khong_phan_loai": [os.path.basename(x["tep"]) for x in pl["khac"]]}
    return ra


# ================================================================ word
def _qn(t):
    from docx.oxml.ns import qn
    return qn(t)


def _dat_chu_doan(p, text):
    """Thay chu ca doan, giu dinh dang run dau."""
    rs = p.runs
    if not rs:
        p.add_run(text)
        return
    rs[0].text = text
    for r in rs[1:]:
        r.text = ""


class DungWord:
    """Mo BC da ban hanh: giu doan dau (quoc hieu, can cu) va doan cuoi (Tren day la, noi nhan, chu ky), thay phan than."""

    def __init__(self, goc):
        from docx import Document
        self.doc = Document(goc)
        ps = self.doc.paragraphs
        self.i_dau = next(i for i, p in enumerate(ps) if p.text.strip().startswith("I. KẾT QUẢ"))
        self.i_cuoi = next(i for i, p in enumerate(ps) if p.text.strip().startswith("Trên đây là"))
        self.m_h1 = ps[self.i_dau]._element
        self.m_h2 = ps[self.i_dau + 1]._element
        mau = next(p for p in ps[self.i_dau + 2:self.i_cuoi] if p.text.strip().startswith("*") and len(p.runs) >= 2)
        self.m_doan = mau._element
        self.neo = ps[self.i_cuoi]._element
        for p in ps[self.i_dau:self.i_cuoi]:
            p._element.getparent().remove(p._element)

    def _dat_run(self, r, text, thuong):
        for t in r.findall(_qn("w:t")):
            r.remove(t)
        if thuong:
            rpr = r.find(_qn("w:rPr"))
            if rpr is None:
                rpr = r.makeelement(_qn("w:rPr"), {})
                r.insert(0, rpr)
            for tag in ("w:b", "w:bCs", "w:i", "w:iCs"):
                for e in rpr.findall(_qn(tag)):
                    rpr.remove(e)
                e = rpr.makeelement(_qn(tag), {})
                e.set(_qn("w:val"), "0")
                rpr.append(e)
        t = r.makeelement(_qn("w:t"), {})
        t.text = text
        t.set(_qn("xml:space"), "preserve")
        r.append(t)

    def doan(self, nhan, noi_dung, sao=True):
        el = copy.deepcopy(self.m_doan)
        rs = el.findall(_qn("w:r"))
        for r in rs[2:]:
            el.remove(r)
        self._dat_run(rs[0], f"* {nhan}: " if sao else f"{nhan}: ", thuong=False)
        self._dat_run(rs[1], noi_dung, thuong=True)
        self.neo.addprevious(el)

    def _mot_run(self, mau, text, thuong=False):
        el = copy.deepcopy(mau)
        rs = el.findall(_qn("w:r"))
        for r in rs[1:]:
            el.remove(r)
        self._dat_run(rs[0], text, thuong)
        self.neo.addprevious(el)

    def h1(self, s):
        self._mot_run(self.m_h1, s)

    def h2(self, s):
        self._mot_run(self.m_h2, s)

    def than(self, s):
        self._mot_run(self.m_doan, s, thuong=True)


def _phan(b, d, loai):
    """d: {"1": [[nhan, noi dung], ...], ..., "7": [[so NQ hoac ten, noi dung], ...]}"""
    for t in range(1, 7):
        b.h2(f"{t}. {TEN_TRUC[t]}")
        for muc in d.get(str(t), []):
            nhan, nd = (muc[0], muc[1]) if isinstance(muc, (list, tuple)) else (None, muc)
            if not nd:
                continue
            (b.than(nd) if not nhan else b.doan(nhan, nd))
    nq = d.get("7", [])
    if nq:
        b.h2("7. Kết quả thực hiện các Nghị quyết của Bộ Chính trị" if loai == "kq"
             else "7. Kế hoạch triển khai các Nghị quyết của Bộ Chính trị")
        thu_tu = {s: i for i, s in enumerate(THU_TU_NQ)}
        for so, nd in sorted(nq, key=lambda x: thu_tu.get(str(x[0]), 99)):
            if nd:
                b.doan(TEN_NQ.get(str(so), str(so)), nd)


def dung_word(goc, nd, ra):
    from docx import Document
    y, m = nd["ky"]["nam"], nd["ky"]["thang"]
    y2, m2 = _sau(y, m) if "ky_sau" not in nd else (nd["ky_sau"]["nam"], nd["ky_sau"]["thang"])
    b = DungWord(goc)
    cum = f"tháng {m} và kế hoạch công tác tháng {m2} năm {y2}"
    for i, p in enumerate(b.doc.paragraphs):
        t = p.text
        if i > 12 and not t.startswith("Trên đây là"):
            continue
        t2 = re.sub(r"tháng \d{1,2} và kế hoạch công tác tháng \d{1,2} năm 20\d\d", cum, t)
        if t2 != t:
            _dat_chu_doan(p, t2)
    for tb in b.doc.tables[:1]:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    t = p.text
                    # 1.3.6: ban trong kho co so hieu di dang "Số375BC-CĐKT" (thieu ":" va "/") — chay that 28/9 giu nham so 375
                    t2 = re.sub(r"^(\s*)Số\s*:?\s*\d+\s*/?\s*(?=[A-ZĐ]{2,}-)", r"\1Số:      /", t)
                    t2 = re.sub(r"ngày\s*\d{1,2}\s*tháng\s*\d{1,2}\s*năm", f"ngày     tháng {nd.get('thang_lap', m)} năm", t2)
                    if t2 != t:
                        _dat_chu_doan(p, t2)
    b.h1(f"I. KẾT QUẢ THỰC HIỆN CÔNG TÁC THÁNG {m} NĂM {y}")
    _phan(b, nd["ket_qua"], "kq")
    b.h1("II. ĐÁNH GIÁ CHUNG")
    b.doan("Kết quả đạt được", nd["danh_gia"]["dat"], sao=False)
    b.doan("Tồn tại, hạn chế", nd["danh_gia"]["ton_tai"], sao=False)
    b.h1(f"III. NHIỆM VỤ TRỌNG TÂM THÁNG {m2} NĂM {y2}")
    _phan(b, nd["nhiem_vu"], "kh")
    os.makedirs(os.path.dirname(os.path.abspath(ra)), exist_ok=True)
    b.doc.save(ra)
    # ---- tu kiem
    d = Document(ra)
    txt = [p.text for p in d.paragraphs]
    y0, m0 = _truoc(y, m)
    dau_cuoi = [t for t in txt[:12]] + [t for t in txt if t.startswith("Trên đây là")]
    try:
        from vanphong import kiem_tra
        vp = [(t[:60], l) for t in txt for l in kiem_tra(t)]
    except Exception:
        vp = None
    thieu_muc = {t: [n for n in MUC_CON[t] if not any(x.startswith(f"* {n}:") for x in txt)] for t in (1, 2, 4, 5, 6)}
    return {"tep": ra, "so_doan": sum(1 for t in txt if t.strip()),
            "con_can_bo_sung": sum(t.count(CAN_BO_SUNG) for t in txt),
            "ky_cu_o_dau_cuoi": [t[:80] for t in dau_cuoi if re.search(rf"tháng {m0}\b.*năm", t)],
            "muc_con_chua_co_trong_phan_I_III": {k: v for k, v in thieu_muc.items() if v},
            "vi_pham_van_phong": (len(vp) if vp is not None else "chưa kiểm (thiếu vanphong.py)"),
            "vi_du_vi_pham": (vp[:5] if vp else [])}


# ================================================================ xlsx chung
def _so(v):
    """So luong: '1' -> 1, '2.0' -> 2 (Excel SUM bo qua o chu); de nguyen neu khong phai so."""
    if isinstance(v, str):
        t = v.strip().replace(",", ".")
        try:
            f = float(t)
            return int(f) if f == int(f) else f
        except ValueError:
            return v or None
    return v


def _hang_tieu_de(ws, max_r=20):
    for r in range(1, max_r + 1):
        if any(c.value and "NỘI DUNG CÔNG VIỆC" in str(c.value).upper() for c in ws[r]):
            return r
    raise ValueError("Không thấy hàng tiêu đề 'NỘI DUNG CÔNG VIỆC' trong tệp gốc")


def _dong_bat_dau(ws, r_h):
    for r in range(r_h + 1, r_h + 8):
        if str(ws.cell(r, 1).value or "").strip() == "I":
            return r
    raise ValueError("Không thấy hàng nhóm 'I' sau hàng tiêu đề")


def _mau_hang(ws, r, ncol):
    return [copy.copy(ws.cell(r, c)._style) for c in range(1, ncol + 1)], ws.row_dimensions[r].height


def _gop_hang(ws, r):
    """Cac vung gop nam tren mot hang (cot dau, cot cuoi)."""
    return [(m.min_col, m.max_col) for m in ws.merged_cells.ranges if m.min_row == r == m.max_row]


def _cong_thuc_mau(ws, r, cot):
    from openpyxl.formula.translate import Translator
    ra = {}
    for c in cot:
        v = ws.cell(r, c).value
        if isinstance(v, str) and v.startswith("="):
            ra[c] = (v, ws.cell(r, c).coordinate, Translator)
    return ra


def _dich(ct, dich_toi):
    v, goc, T = ct
    return T(v, origin=goc).translate_formula(dich_toi)


def _tnr(ws, den_hang, ncol):
    from openpyxl.styles import Font
    for row in ws.iter_rows(min_row=1, max_row=den_hang, max_col=ncol):
        for c in row:
            if c.font and c.font.name != "Times New Roman":
                fo = copy.copy(c.font)
                fo.name = "Times New Roman"
                c.font = fo


def _cot_doi_soat(ws, r_h, col, rows, rong=46):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    ws.cell(r_h, col).value = "Ghi chú đối soát (xóa trước khi ban hành)"
    ws.cell(r_h, col)._style = copy.copy(ws.cell(r_h, col - 1)._style)
    ws.column_dimensions[get_column_letter(col)].width = rong
    for r in rows:
        c = ws.cell(r, col)
        if c.value:
            c.font = Font(name="Times New Roman", size=11)
            c.alignment = Alignment(wrap_text=True, vertical="center")


# ================================================================ phu luc
def dung_phu_luc(goc, nd, ra):
    """nd: {"ky":{"thang","nam"}, "truc":{"1":{"ten"?, "dong":[dong]}}, "dot_xuat":[dong]}
    dong: {"nd","cd","ct","sp","sl","dk","ket_qua":{"sl":0..1,"cl":0..1,"td":0..1}|null,"doi_soat":"..."}"""
    import openpyxl
    y, m = nd["ky"]["nam"], nd["ky"]["thang"]
    wb = openpyxl.load_workbook(goc)
    ws = wb.active
    ws.title = re.sub(r"\d{1,2}$", str(m), ws.title) if re.search(r"\d{1,2}$", ws.title) else f"BC Kết quả tháng {m}"
    for r in range(1, 5):
        for c in ws[r]:
            if isinstance(c.value, str):
                c.value = re.sub(r"THÁNG \d{1,2} NĂM 20\d\d", f"THÁNG {m} NĂM {y}", c.value)
    r_h = _hang_tieu_de(ws)
    r_dau = _dong_bat_dau(ws, r_h)
    NC = 16
    s_nhom, s_truc, s_data = (_mau_hang(ws, r_dau + k, NC) for k in (0, 1, 2))
    gop_nhom, gop_truc = _gop_hang(ws, r_dau), _gop_hang(ws, r_dau + 1)
    ct_data = _cong_thuc_mau(ws, r_dau + 2, [8, 9, 10, 12, 14, 16])      # H I J L N P
    cot_tong_truc = [c for c in range(6, NC + 1) if str(ws.cell(r_dau + 1, c).value or "").startswith("=SUM")]
    r_ii = next((r for r in range(r_dau, ws.max_row + 1) if str(ws.cell(r, 1).value or "").strip() == "II"), None)
    cot_tong_ii = [c for c in range(6, NC + 1) if r_ii and str(ws.cell(r_ii, c).value or "").startswith("=SUM")]
    tieu_I = ws.cell(r_dau, 2).value
    tieu_II = ws.cell(r_ii, 2).value if r_ii else "Các nhiệm vụ đột xuất, phát sinh khác ngoài kế hoạch"
    ten_truc_goc = {}
    for r in range(r_dau, ws.max_row + 1):
        a, bv = str(ws.cell(r, 1).value or "").strip(), ws.cell(r, 2).value
        if re.fullmatch(r"[1-6]", a) and bv and "Trục" in str(bv)[:10]:
            ten_truc_goc[int(a)] = bv
    # 1.3.6: tieu de nhom/Truc cua ban goc co the mang ky cu ("... tháng 8") — chay that 28/9 sot o muc II
    doi_ky = lambda v: re.sub(r"(tháng|THÁNG)\s+\d{1,2}\b", lambda k: f"{k.group(1)} {m}", v) if isinstance(v, str) else v
    tieu_I, tieu_II = doi_ky(tieu_I), doi_ky(tieu_II)
    ten_truc_goc = {k: doi_ky(v) for k, v in ten_truc_goc.items()}
    for mg in [x for x in list(ws.merged_cells.ranges) if x.min_row >= r_dau]:
        ws.unmerge_cells(str(mg))
    ws.delete_rows(r_dau, ws.max_row - r_dau + 1)
    from openpyxl.utils import get_column_letter as L
    cur = [r_dau - 1]
    tk = {"dong": 0, "co_ket_qua": 0, "cong_thuc": 0}

    def dat(mau, vals, gop=()):
        cur[0] += 1
        r = cur[0]
        for c in range(1, NC + 1):
            ws.cell(r, c)._style = copy.copy(mau[0][c - 1])
        if mau[1]:
            ws.row_dimensions[r].height = mau[1]
        for c, v in vals.items():
            ws.cell(r, c).value = v
        for c1, c2 in gop:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        return r

    def dong(stt, x):
        r = cur[0] + 1
        v = {1: stt, 2: x["nd"], 3: x.get("cd"), 4: x.get("ct"), 5: x.get("sp"), 6: _so(x.get("sl")), 7: x.get("dk")}
        for c, ctm in ct_data.items():
            if c in (8, 9, 10):
                v[c] = _dich(ctm, f"{L(c)}{r}")
        kq = x.get("ket_qua")
        if kq:
            v[11] = f"=F{r}*{round(float(kq.get('sl', 1)) * 100, 2):g}%"
            v[13] = f"=K{r}*{round(float(kq.get('cl', 1)) * 100, 2):g}%"
            v[15] = f"=K{r}*{round(float(kq.get('td', 1)) * 100, 2):g}%"
            for c in (12, 14, 16):
                if c in ct_data:
                    v[c] = _dich(ct_data[c], f"{L(c)}{r}")
            tk["co_ket_qua"] += 1
        if x.get("doi_soat"):
            v[17] = x["doi_soat"]
        tk["dong"] += 1
        tk["cong_thuc"] += sum(1 for k, s in v.items() if isinstance(s, str) and s.startswith("="))
        return dat(s_data, v)

    dat(s_nhom, {1: "I", 2: tieu_I}, gop_nhom)
    for t in range(1, 7):
        tr = nd["truc"].get(str(t), {})
        rt = dat(s_truc, {1: str(t), 2: tr.get("ten") or ten_truc_goc.get(t) or f"Trục ({t}) - {TEN_TRUC[t]}"}, gop_truc)
        dau = cur[0] + 1
        for k, x in enumerate(tr.get("dong", []), 1):
            dong(f"{t}.{k}", x)
        for c in cot_tong_truc:
            ws.cell(rt, c).value = f"=SUM({L(c)}{dau}:{L(c)}{max(cur[0], dau)})"
    ds2 = nd.get("dot_xuat", [])
    r2 = dat(s_nhom, {1: "II", 2: tieu_II}, gop_nhom)
    dau = cur[0] + 1
    for k, x in enumerate(ds2, 1):
        dong(str(k), x)
    for c in cot_tong_ii:
        ws.cell(r2, c).value = f"=SUM({L(c)}{dau}:{L(c)}{max(cur[0], dau)})"
    _tnr(ws, cur[0], NC + 1)
    _cot_doi_soat(ws, r_h, NC + 1, range(r_dau, cur[0] + 1))
    ws.print_area = f"A1:{L(NC)}{cur[0]}"
    os.makedirs(os.path.dirname(os.path.abspath(ra)), exist_ok=True)
    wb.save(ra)
    return {"tep": ra, "hang_cuoi": cur[0], **tk, "khong_co_ket_qua": tk["dong"] - tk["co_ket_qua"]}


# ================================================================ ke hoach
def dung_ke_hoach(goc, nd, ra):
    """nd: {"ky":{"thang","nam"}, "so"?, "ngay"?, "can_cu": "...", "truc":{"1":{"ten"?,"dong":[dong]}}, "chuyen_tiep":[dong]}
    dong: {"nd","cd","ct","sp","sl","dk","thoi_han","ghi_chu","doi_soat"}"""
    import openpyxl
    y, m = nd["ky"]["nam"], nd["ky"]["thang"]
    wb = openpyxl.load_workbook(goc)
    ws = wb.active
    ws.title = f"KH tháng {m}"
    r_h = _hang_tieu_de(ws)
    r_dau = _dong_bat_dau(ws, r_h)
    NC = 11
    r_nn = next((r for r in range(r_dau, ws.max_row + 1)
                 if str(ws.cell(r, 1).value or "").strip().startswith("Nơi nhận")), None)
    if r_nn is None:
        raise ValueError("Không thấy khối 'Nơi nhận' ở cuối kế hoạch gốc")
    s_nhom, s_truc, s_data = (_mau_hang(ws, r_dau + k, NC) for k in (0, 1, 2))
    gop_nhom, gop_truc = _gop_hang(ws, r_dau), _gop_hang(ws, r_dau + 1)
    ct_data = _cong_thuc_mau(ws, r_dau + 2, [9, 10])
    duoi = {c: (ws.cell(r_nn, c).value, copy.copy(ws.cell(r_nn, c)._style)) for c in range(1, NC + 1)
            if ws.cell(r_nn, c).value is not None}
    gop_duoi = _gop_hang(ws, r_nn)
    cao_duoi = ws.row_dimensions[r_nn].height
    tieu_I = ws.cell(r_dau, 2).value
    r_ii = next((r for r in range(r_dau, r_nn) if str(ws.cell(r, 1).value or "").strip() == "II"), None)
    tieu_II = ws.cell(r_ii, 2).value if r_ii else "Các nhiệm vụ từ tháng trước chuyển sang"
    ten_truc_goc = {}
    for r in range(r_dau, r_nn):
        a, bv = str(ws.cell(r, 1).value or "").strip(), ws.cell(r, 2).value
        if re.fullmatch(r"[1-6]", a) and bv and "Trục" in str(bv)[:10]:
            ten_truc_goc[int(a)] = bv
    # phan dau
    for r in range(1, r_h):
        for c in ws[r]:
            v = c.value
            if not isinstance(v, str):
                continue
            if v.strip().startswith("Số:"):
                c.value = nd.get("so", "Số:       /KH-CĐKT")
            elif v.strip().startswith("Quảng Ngãi, ngày"):
                c.value = nd.get("ngay", f"Quảng Ngãi, ngày     tháng {_truoc(y, m)[1]} năm {_truoc(y, m)[0]}")
            elif v.startswith("KẾ HOẠCH"):
                c.value = f"KẾ HOẠCH\ncông tác tháng {m} năm {y}"
            elif re.fullmatch(r"\s*công tác tháng \d{1,2} năm 20\d\d\s*", v):
                c.value = None
            elif "Căn cứ" in v[:30] and nd.get("can_cu"):
                c.value = nd["can_cu"]
    for mg in [x for x in list(ws.merged_cells.ranges) if x.min_row >= r_dau]:
        ws.unmerge_cells(str(mg))
    ws.delete_rows(r_dau, ws.max_row - r_dau + 1)
    from openpyxl.utils import get_column_letter as L
    cur = [r_dau - 1]
    tk = {"dong": 0, "chuyen_tiep": 0}

    def dat(mau, vals, gop=()):
        cur[0] += 1
        r = cur[0]
        for c in range(1, NC + 1):
            ws.cell(r, c)._style = copy.copy(mau[0][c - 1])
        if mau[1]:
            ws.row_dimensions[r].height = mau[1]
        for c, v in vals.items():
            ws.cell(r, c).value = v
        for c1, c2 in gop:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        return r

    def dong(stt, x):
        r = cur[0] + 1
        v = {1: stt, 2: x["nd"], 3: x.get("cd"), 4: x.get("ct"), 5: x.get("sp"), 6: _so(x.get("sl")), 7: x.get("dk"),
             8: x.get("thoi_han"), 11: x.get("ghi_chu")}
        for c, ctm in ct_data.items():
            v[c] = _dich(ctm, f"{L(c)}{r}")
        if 9 not in ct_data:
            v[9] = DIEM.get(x.get("dk"))
            v[10] = (v[9] / 100) if v[9] else None
        if x.get("doi_soat"):
            v[12] = x["doi_soat"]
        tk["dong"] += 1
        return dat(s_data, v)

    dat(s_nhom, {1: "I", 2: tieu_I}, gop_nhom)
    for t in range(1, 7):
        tr = nd["truc"].get(str(t), {})
        dat(s_truc, {1: str(t), 2: tr.get("ten") or ten_truc_goc.get(t) or f"Trục ({t}) - {TEN_TRUC[t]}"}, gop_truc)
        for k, x in enumerate(tr.get("dong", []), 1):
            dong(f"{t}.{k}", x)
    ct2 = nd.get("chuyen_tiep", [])
    if ct2:
        dat(s_nhom, {1: "II", 2: tieu_II}, gop_nhom)
        for k, x in enumerate(ct2, 1):
            dong(str(k), x)
            tk["chuyen_tiep"] += 1
    cur[0] += 2
    r = cur[0]
    for c, (v, st) in duoi.items():
        ws.cell(r, c).value = v
        ws.cell(r, c)._style = st
    for c1, c2 in gop_duoi:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    if cao_duoi:
        ws.row_dimensions[r].height = cao_duoi
    _tnr(ws, cur[0], NC + 1)
    _cot_doi_soat(ws, r_h, NC + 1, range(r_dau, cur[0]), rong=48)
    ws.print_area = f"A1:{L(NC)}{cur[0]}"
    os.makedirs(os.path.dirname(os.path.abspath(ra)), exist_ok=True)
    wb.save(ra)
    return {"tep": ra, "hang_cuoi": cur[0], **tk}


# ================================================================ main
def main(argv):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="lenh", required=True)
    a = sub.add_parser("nguon")
    a.add_argument("--ky", required=True)
    a.add_argument("--dau-vao")
    a = sub.add_parser("trich")
    a.add_argument("--dau-vao", required=True)
    a.add_argument("--ra", required=True)
    for ten in ("word", "phu-luc", "ke-hoach"):
        a = sub.add_parser(ten)
        a.add_argument("--goc", required=True)
        a.add_argument("--noi-dung", required=True)
        a.add_argument("--ra", required=True)
    a = ap.parse_args(argv)
    try:
        if a.lenh == "nguon":
            y, m = _ky(a.ky)
            ra = {"ky": a.ky, "ban_da_ban_hanh": tim_ban_da_ban_hanh(y, m)}
            if a.dau_vao:
                pl = phan_loai_dau_vao(a.dau_vao)
                ra["dau_vao"] = {k: [{"ma": x["ma"], "tep": os.path.relpath(x["tep"], a.dau_vao)} for x in v]
                                 for k, v in pl.items()}
        elif a.lenh == "trich":
            d = trich(a.dau_vao)
            json.dump(d, open(a.ra, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            ra = {"ra": a.ra, **d["thong_ke"], "loi": d["loi"]}
        else:
            nd = json.load(open(a.noi_dung, encoding="utf-8"))
            f = {"word": dung_word, "phu-luc": dung_phu_luc, "ke-hoach": dung_ke_hoach}[a.lenh]
            ra = f(a.goc, nd, a.ra)
    except (FileNotFoundError, ValueError, StopIteration, KeyError) as e:
        print(json.dumps({"loi": f"{e.__class__.__name__}: {e}"}, ensure_ascii=False))
        return 2
    print(json.dumps(ra, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
