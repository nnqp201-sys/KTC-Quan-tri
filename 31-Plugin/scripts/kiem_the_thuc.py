# -*- coding: utf-8 -*-
"""Kiem the thuc, ky thuat trinh bay san pham .docx/.xlsx — DL-20260919-003.

Chuan: 20-Chuan-Chung/18-Chuan-The-Thuc-San-Pham.md — so do that cua 16 mau
KTC-Database/03-Templates(1) (.dotx/.xltx), NĐ 30/2020 Phu luc I, TB 597/TB-CĐKT (qua 897,
Checklist 05-Hinh-Thuc va 08-Quy-Uoc-Rieng-CDKT muc 4).

Do tren BYTE THAT cua tep (python-docx/openpyxl): phong va co chu duoc giai theo chuoi
run -> style ky tu -> style doan (ke ca base_style) -> docDefaults -> theme.
Chi PHAT HIEN va GOI Y muc — khong sua tep. He A (hanh chinh). Khong ap cho van ban Dang (he B).

Chay:  python 29-Cong-Cu/kiem_the_thuc.py <tep .docx|.xlsx> [...]
       python 29-Cong-Cu/kiem_the_thuc.py --tao <mau .dotx> <dich .docx>   (tao ban lam viec tu mau)
Ma thoat: 1 neu co goi y Muc 1 hoac Muc 2, nguoc lai 0.
"""
import collections
import io
import os
import re
import sys
import zipfile

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PHONG = "Times New Roman"
# (min, max) cm — NĐ 30 Phu luc I; mau 03-Templates(1) dung 2/2/3/2
LE = {"trên": (2.0, 2.5), "dưới": (2.0, 2.5), "trái": (3.0, 3.5), "phải": (1.5, 2.0)}
EPS = 0.05


# ------------------------------------------------------------------ mo tep mau
def _doi_kieu(du_lieu: bytes, tu: bytes, sang: bytes) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(du_lieu)) as zi, zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(tu, sang)
            zo.writestr(it, data)
    return b.getvalue()


def mo_mau(p):
    """Mo .dotx (hoac .docx) thanh docx.Document — python-docx khong tu mo .dotx."""
    import docx
    raw = open(p, "rb").read()
    if p.lower().endswith(".dotx"):
        raw = _doi_kieu(raw, b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
    return docx.Document(io.BytesIO(raw))


def tao_tu_mau(mau, dich):
    """Tao ban lam viec .docx/.xlsx tu mau .dotx/.xltx — giu nguyen style, le, bang the thuc."""
    raw = open(mau, "rb").read()
    if mau.lower().endswith(".dotx"):
        raw = _doi_kieu(raw, b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
    elif mau.lower().endswith(".xltx"):
        raw = _doi_kieu(raw, b"spreadsheetml.template.main+xml", b"spreadsheetml.sheet.main+xml")
    os.makedirs(os.path.dirname(os.path.abspath(dich)), exist_ok=True)
    open(dich, "wb").write(raw)
    return dich


# ------------------------------------------------------------------ giai phong/co chu
class GiaiDocx:
    def __init__(self, d):
        self.d = d
        root = d.styles.element
        rd = root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
        self.mac_dinh_phong = self._phong_rpr(rd)
        sz = rd.find(f"{W}sz") if rd is not None else None
        self.mac_dinh_co = int(sz.get(f"{W}val")) / 2 if sz is not None else 10.0
        self.theme = self._doc_theme()

    def _doc_theme(self):
        """Phong theme ma Word THAT SU hien thi. settings themeFontLang = vi-VN -> Word dung phong
        script="Viet" cua theme (vd. major Viet = Times New Roman) thay vi <a:latin> (Calibri Light).
        Bo qua buoc nay tung bao nham hang tram doan “Calibri Light” tren van ban da ban hanh (19/9/2026)."""
        viet = bool(re.search(r'<w:themeFontLang [^>]*w:val="vi', self.d.settings.element.xml))
        kq = {"major": None, "minor": None}
        for rel in self.d.part.rels.values():
            if rel.reltype.endswith("/theme"):
                x = rel.target_part.blob.decode("utf-8", "ignore")
                for loai, the in (("major", "majorFont"), ("minor", "minorFont")):
                    m = re.search(f"<a:{the}>(.*?)</a:{the}>", x, re.S)
                    if not m:
                        continue
                    khoi = m.group(1)
                    v = re.search(r'script="Viet" typeface="([^"]*)"', khoi) if viet else None
                    latin = re.search(r'<a:latin typeface="([^"]*)"', khoi)
                    kq[loai] = v.group(1) if v else (latin.group(1) if latin else None)
        return kq

    def _phong_rpr(self, rpr):
        if rpr is None:
            return None
        f = rpr.find(f"{W}rFonts")
        if f is None:
            return None
        if f.get(f"{W}ascii"):
            return f.get(f"{W}ascii")
        t = f.get(f"{W}asciiTheme")
        if t:
            return "theme:" + ("major" if "major" in t.lower() else "minor")
        return None

    def _xu_theme(self, v):
        if v and v.startswith("theme:"):
            return self.theme.get(v[6:]) or v
        return v

    def _style_chain(self, st):
        while st is not None:
            yield st
            st = st.base_style

    def phong(self, run, para):
        v = run.font.name or self._phong_rpr(run._element.rPr)
        if not v and run.style is not None:
            for s in self._style_chain(run.style):
                v = s.font.name or self._phong_rpr(s.element.rPr)
                if v:
                    break
        if not v:
            for s in self._style_chain(para.style):
                v = s.font.name or self._phong_rpr(s.element.rPr)
                if v:
                    break
        return self._xu_theme(v or self.mac_dinh_phong)

    def co(self, run, para):
        if run.font.size:
            return run.font.size.pt
        if run.style is not None:
            for s in self._style_chain(run.style):
                if s.font.size:
                    return s.font.size.pt
        for s in self._style_chain(para.style):
            if s.font.size:
                return s.font.size.pt
        return self.mac_dinh_co

    def can_le(self, para):
        if para.paragraph_format.alignment is not None:
            return para.paragraph_format.alignment
        for s in self._style_chain(para.style):
            if s.paragraph_format.alignment is not None:
                return s.paragraph_format.alignment
        return None


def _doan_trong_bang(d):
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p


# ------------------------------------------------------------------ kiem docx
def kiem_docx(d):
    """Tra ve list[(muc, ma, mo_ta)]."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    kq = []
    g = GiaiDocx(d)

    # TT01 kho giay, TT02 le — moi section
    for i, s in enumerate(d.sections, 1):
        w, h = s.page_width.cm, s.page_height.cm
        ngang = w > h
        a4 = abs(min(w, h) - 21.0) <= 0.1 and abs(max(w, h) - 29.7) <= 0.1
        if not a4:
            kq.append((2, "TT01", f"Section {i}: khổ {w:.1f}×{h:.1f} cm — phải A4 (21×29,7)"))
        le = {"trên": s.top_margin.cm, "dưới": s.bottom_margin.cm,
              "trái": s.left_margin.cm, "phải": s.right_margin.cm}
        if not ngang:  # trang ngang (phu luc bang) khong ap le doc
            sai = [f"{k} {round(v, 2)}" for k, v in le.items()
                   if not (LE[k][0] - EPS <= round(v, 2) <= LE[k][1] + EPS)]
            if sai:
                kq.append((2, "TT02", f"Section {i}: lề ngoài khoảng NĐ 30 ({', '.join(sai)} cm). "
                           "Mẫu 03-Templates(1): trên 2 · dưới 2 · trái 3 · phải 2 cm"))

    doan = list(d.paragraphs)
    tat_ca = doan + list(_doan_trong_bang(d))

    # TT03 phong chu, TT05 mau chu
    sai_phong = collections.Counter()
    mau_chu = 0
    for p in tat_ca:
        for r in p.runs:
            if not r.text.strip():
                continue
            f = g.phong(r, p)
            if f != PHONG:
                sai_phong[f or "(không xác định)"] += 1
            c = r.font.color
            if c is not None and c.type is not None and c.rgb is not None and str(c.rgb) not in ("000000",):
                mau_chu += 1
    # Bien the ten cua chinh TNR ("Times New Roman Bold", "TimesNewRomanPSMT" — thuong do chep tu PDF):
    # hien thi van la TNR tren may co phong chuan, nhung may khac co the thay phong -> Muc 3.
    bien_the = {k: v for k, v in sai_phong.items() if re.sub(r"[\s\-]", "", k or "").lower().startswith("timesnewroman")}
    khac = {k: v for k, v in sai_phong.items() if k not in bien_the}
    if khac:
        kq.append((2, "TT03", f"{sum(khac.values())} đoạn chữ không phải {PHONG}: "
                   + ", ".join(f"{k}×{v}" for k, v in collections.Counter(khac).most_common(4))
                   + (" — .VnTime là phông TCVN3 cũ, phải chuyển Unicode" if any(".Vn" in (k or "") for k in khac) else "")))
    if bien_the:
        kq.append((3, "TT03b", f"{sum(bien_the.values())} đoạn chữ dùng biến thể tên phông "
                   + ", ".join(bien_the) + f" — đặt lại đúng tên “{PHONG}”"))
    if mau_chu:
        kq.append((3, "TT05", f"{mau_chu} đoạn chữ có màu khác đen (NĐ 30: màu đen)"))

    # Doan noi dung: ngoai bang, dai >= 60 ky tu
    noi_dung = [p for p in doan if len(p.text.strip()) >= 60]
    co = collections.Counter()
    for p in noi_dung:
        for r in p.runs:
            if r.text.strip():
                co[g.co(r, p)] += len(r.text)
    if co:
        chinh = co.most_common(1)[0][0]
        if chinh != 14:
            kq.append((2, "TT04", f"Cỡ chữ phần nội dung chủ yếu là {chinh:g} — TB 597: cỡ 14"))
        lech = {k: v for k, v in co.items() if k != chinh and v > 0.05 * sum(co.values())}
        if lech:
            kq.append((3, "TT04b", "Cỡ chữ lời văn không thống nhất: "
                       + ", ".join(f"cỡ {k:g} ({v} ký tự)" for k, v in lech.items())))
        khong_deu = sum(1 for p in noi_dung if g.can_le(p) not in (WD_ALIGN_PARAGRAPH.JUSTIFY, None)
                        and g.can_le(p) != WD_ALIGN_PARAGRAPH.CENTER)
        if khong_deu > 0.2 * len(noi_dung):
            kq.append((3, "TT06", f"{khong_deu}/{len(noi_dung)} đoạn nội dung không dàn đều hai lề"))
        gian = [p.paragraph_format.line_spacing for p in noi_dung]
        qua = sum(1 for x in gian if isinstance(x, float) and x > 1.5 + 1e-6)
        hep = sum(1 for x in gian if x is not None and not isinstance(x, float) and x.pt < 15 - 0.1)
        if qua or hep:
            kq.append((3, "TT07", f"Cách dòng ngoài khoảng single/15pt → 1,5 lines: {qua} đoạn >1,5; "
                       f"{hep} đoạn exactly <15pt"))

    # TT08 phan dau the thuc (bang dau tien)
    dau = "\n".join(p.text for p in (list(_doan_trong_bang(d))[:12] if d.tables else doan[:8]))
    dau_up = dau.upper()
    la_vb_hanh_chinh = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in dau_up or "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM" in dau_up
    if la_vb_hanh_chinh:
        if "TRƯỜNG CAO ĐẲNG KON TUM" not in dau_up:
            kq.append((2, "TT08", "Phần đầu thiếu tên đơn vị ban hành “TRƯỜNG CAO ĐẲNG KON TUM”"))
        if "UBND TỈNH KON TUM" in dau_up or "ỦY BAN NHÂN DÂN TỈNH KON TUM" in dau_up:
            kq.append((2, "TT08", "Cơ quan chủ quản ghi “UBND TỈNH KON TUM” — văn bản mới dùng “UBND TỈNH QUẢNG NGÃI”"))
        elif "QUẢNG NGÃI" not in dau_up:
            kq.append((2, "TT08", "Phần đầu thiếu cơ quan chủ quản “UBND TỈNH QUẢNG NGÃI”"))
        for p in _doan_trong_bang(d):
            t = p.text.strip().upper()
            rs = [r for r in p.runs if r.text.strip()]
            if not rs:
                continue
            if t.startswith("CỘNG H") and "VIỆT NAM" in t:
                cs = {g.co(r, p) for r in rs}
                if cs != {13.0}:
                    kq.append((2, "TT09", f"Quốc hiệu cỡ {sorted(cs)} — TB 597: cỡ 13, đậm"))
            if t.startswith("ĐỘC LẬP"):
                cs = {g.co(r, p) for r in rs}
                if cs != {14.0}:
                    kq.append((2, "TT09", f"Tiêu ngữ cỡ {sorted(cs)} — TB 597: cỡ 14, đậm"))
                if any(r.font.underline for r in rs):
                    kq.append((3, "TT09", "Tiêu ngữ dùng Underline — TB 597: kẻ đường bằng Draw"))

    # TT10 "Noi nhan"
    for p in tat_ca:
        if p.text.strip().startswith("Nơi nhận"):
            rs = [r for r in p.runs if r.text.strip()]
            cs = {g.co(r, p) for r in rs}
            if cs and cs != {12.0}:
                kq.append((3, "TT10", f"“Nơi nhận” cỡ {sorted(cs)} — TB 597: cỡ 12, nghiêng, đậm"))
            break

    # TT11 so trang: co truong PAGE o header
    if len(noi_dung) > 25:
        co_page = any("PAGE" in (s.header._element.xml if s.header is not None else "") for s in d.sections)
        if not co_page:
            kq.append((3, "TT11", "Văn bản nhiều trang nhưng không có số trang (canh giữa lề trên, ẩn trang 1)"))
    kq.sort(key=lambda x: x[0])
    return kq


# ------------------------------------------------------------------ kiem xlsx
def kiem_xlsx(wb):
    kq = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or ws.max_row < 2:
            continue
        ps = ws.page_setup
        if ps.paperSize not in (None, 9, "9"):
            kq.append((2, "TX01", f"[{ws.title}] khổ giấy in mã {ps.paperSize} — phải A4 (mã 9)"))
        f = collections.Counter()
        co = collections.Counter()
        for row in ws.iter_rows(max_row=min(ws.max_row, 400)):
            for c in row:
                if c.value is not None and str(c.value).strip():
                    f[c.font.name] += 1
                    co[c.font.sz] += 1
        sai = {k: v for k, v in f.items() if k != PHONG}
        if sai:
            kq.append((2, "TX02", f"[{ws.title}] {sum(sai.values())} ô không phải {PHONG}: "
                       + ", ".join(f"{k}×{v}" for k, v in list(sai.items())[:4])))
        if co:
            chinh = co.most_common(1)[0][0]
            if chinh is not None and not (12 <= chinh <= 14):
                kq.append((3, "TX03", f"[{ws.title}] cỡ chữ chủ yếu {chinh:g} — mẫu 05B dùng 12–14"))
        if ws.max_column > 6 and ps.orientation != "landscape" and not ps.fitToWidth:
            kq.append((4, "TX04", f"[{ws.title}] bảng {ws.max_column} cột in dọc, không đặt vừa trang — "
                       "mẫu 05B: A4 ngang"))
    kq.sort(key=lambda x: x[0])
    return kq


def kiem_tep(p):
    if p.lower().endswith((".docx", ".dotx")):
        return kiem_docx(mo_mau(p))
    if p.lower().endswith((".xlsx", ".xltx")):
        import warnings
        import openpyxl
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return kiem_xlsx(openpyxl.load_workbook(p))
    raise ValueError(f"Không kiểm được loại tệp: {p}")


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    if argv[0] == "--tao":
        print("Đã tạo", tao_tu_mau(argv[1], argv[2]))
        return 0
    xau = False
    for p in argv:
        kq = kiem_tep(p)
        print(f"\n=== {p} — {len(kq)} gợi ý ===")
        if not kq:
            print("  ✓ Đạt chuẩn thể thức theo các phép kiểm TT/TX.")
        for muc, ma, mo_ta in kq:
            print(f"  [Mức {muc}] {ma}: {mo_ta}")
        xau |= any(x[0] <= 2 for x in kq)
    return 1 if xau else 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
