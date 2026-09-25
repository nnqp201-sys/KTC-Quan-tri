# -*- coding: utf-8 -*-
"""Cau truc 6 mau Ke hoach + KPI Quy (03-Templates/03-12) va dien ke hoach vao BAN SAO cua mau.

Mau (6 nhom vi tri, CV 694/CĐKT-TCCB muc I.2; QD 1923 D15.1b):
  truong-pho-don-vi   Truong/Pho phong, khoa                     (vien chuc quan ly)
  bo-mon              Truong/Pho bo mon, Phong Kham da khoa      (vien chuc quan ly)
  nha-giao            Nhom 1 — Nha giao truc tiep giang day cac Khoa
  giao-vu             Nhom 2 — Giao vu khoa
  hanh-chinh          Nhom 3 — Vien chuc, NLD lam viec theo che do hanh chinh
  ho-tro              Nhom 4 — Nhan vien ho tro, phuc vu

Moi mau: sheet "Ke Hoach" (6 Truc x 20 dong, noi cong thuc sang sheet "KPI"), "KPI" (so luong quy doi, KPI 3 chieu),
"Danh gia"/"Danh Gia" (diem toi da tung Truc, nhom tieu chi chung). Cau truc DOC TU MAU, khong ghi cung dong.

KHONG ghi de mau: luon chep sang tep ra roi moi dien. Tep ra chuan hoa phong chu ve Times New Roman (mau goc co o
Calibri — the thuc Muc 2, skill the-thuc); co, dam, can le giu nguyen.
"""
import glob
import os
import re
import shutil
import unicodedata
import warnings

HERE = os.path.dirname(os.path.abspath(__file__))
DU_AN = os.path.dirname(HERE)

NHOM = {
    "truong-pho-don-vi": ("Trưởng/Phó phòng, khoa", "Truong-Pho-Truong-Cac-Don-Vi", True),
    "bo-mon": ("Trưởng/Phó bộ môn, Phòng Khám đa khoa", "VCQL-Bo-Mon-Va-Tuong-Duong", True),
    "nha-giao": ("Nhóm 1 — Nhà giáo trực tiếp giảng dạy", "Nha-Giao-Giang-Day-Cac-Khoa", False),
    "giao-vu": ("Nhóm 2 — Giáo vụ khoa", "Giao-Vu-Khoa", False),
    "hanh-chinh": ("Nhóm 3 — Viên chức hành chính", "VC-Hanh-Chinh", False),
    "ho-tro": ("Nhóm 4 — Nhân viên hỗ trợ, phục vụ", "NV-Ho-Tro-Phuc-Vu", False),
}
LA_MA = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6}
PHONG = "Times New Roman"


def _kd(s):
    t = unicodedata.normalize("NFD", str(s or "").lower())
    return " ".join("".join(c for c in t if unicodedata.category(c) != "Mn").replace("đ", "d").split())


def thu_muc_mau():
    """assets/ cua skill (khi chay trong goi) hoac 28-KTC-KPI/assets (du an)."""
    for d in (os.path.join(HERE, "..", "assets"), os.path.join(DU_AN, "28-KTC-KPI", "assets"),
              os.path.join(HERE, "..", "skills", "kpi-lap-ke-hoach", "assets")):   # ban chep o scripts/ goc plugin
        if glob.glob(os.path.join(d, "Mau-KeHoach-DanhGia_*.xlsx")):
            return os.path.abspath(d)
    raise FileNotFoundError("Không thấy 6 mẫu Kế hoạch trong assets/ — dừng, hỏi người dùng.")


def tep_mau(nhom):
    if nhom not in NHOM:
        raise ValueError(f"Nhóm vị trí '{nhom}' không có; chọn một trong: {', '.join(NHOM)}")
    ds = glob.glob(os.path.join(thu_muc_mau(), f"Mau-KeHoach-DanhGia_{NHOM[nhom][1]}_*.xlsx"))
    if len(ds) != 1:
        raise FileNotFoundError(f"Không xác định được đúng 1 mẫu cho nhóm '{nhom}' ({len(ds)} tệp).")
    return ds[0]


def nhan_nhom(wb):
    """Doan nhom tu tieu de sheet Danh gia (vd 'DOI VOI VIEN CHUC HANH CHINH'). Khong chac -> None."""
    ws = sheet_danh_gia(wb)
    t = _kd(" ".join(str(ws.cell(r, 1).value or "") for r in range(1, 6)))
    for k, dau in (("giao-vu", "giao vu"), ("ho-tro", "ho tro"), ("nha-giao", "nha giao"),
                   ("hanh-chinh", "hanh chinh"), ("bo-mon", "bo mon"),
                   ("truong-pho-don-vi", "truong/pho cac don vi")):   # tieu de mau: "TRUONG/PHO CAC DON VI"
        if dau in t:
            return k
    return None


def sheet_danh_gia(wb):
    return next(w for w in wb.worksheets if _kd(w.title).startswith("danh gia"))


def cau_truc(wb):
    """{'truc': {1: (dong_dau, dong_cuoi)} tren 'Ke Hoach', 'diem_truc': {1: 45,...}, 'nhom_a': [13,12,5],
        'kpi_dong': {dong Ke Hoach: dong KPI}, 'cot_minh_chung': 'S' | None}"""
    kh, kp, dg = wb["Ke Hoach"], wb["KPI"], sheet_danh_gia(wb)
    kpi_dong, dau_truc = {}, []
    for r in kp.iter_rows(min_row=7):
        a = str(r[0].value or "").strip()
        if a in LA_MA:
            dau_truc.append((r[0].row, LA_MA[a]))
        m = re.match(r"='Ke Hoach'!B(\d+)$", str(r[1].value or ""))
        if m:
            kpi_dong[int(m.group(1))] = r[0].row
    truc = {}
    for i, (d, n) in enumerate(dau_truc):
        het = dau_truc[i + 1][0] if i + 1 < len(dau_truc) else 10 ** 6
        ks = [k for k, v in kpi_dong.items() if d < v < het]
        if ks:
            truc[n] = (min(ks), max(ks))
    diem_truc, nhom_a = {}, []
    for r in dg.iter_rows():
        b = str(r[1].value or "")
        m = re.match(r"Trục \((\d)\)", b)
        if m and len(r) > 5 and isinstance(r[5].value, (int, float)):
            diem_truc[int(m.group(1))] = float(r[5].value)
        if str(r[0].value or "").strip() in ("1", "2", "3") and str(r[3].value or "").startswith("=SUM(D"):
            a_, b_ = map(int, re.findall(r"D(\d+):D(\d+)", r[3].value)[0])
            nhom_a.append(sum(float(dg.cell(i, 4).value or 0) for i in range(a_, b_ + 1)))
    cot_mc = None
    for c in kp[3]:
        if "minh chứng" in str(c.value or "").lower():
            cot_mc = c.column_letter
    return {"truc": truc, "diem_truc": diem_truc, "nhom_a": nhom_a, "kpi_dong": kpi_dong, "cot_minh_chung": cot_mc,
            "kpi_truc": {n: d for d, n in dau_truc}}


class LoiCauTruc(ValueError):
    """Sheet Danh gia khong do duoc cau truc — DUNG, bao nguoi dung; khong doan (lenh 25/9/2026 muc 12)."""


def cau_truc_danh_gia(wb):
    """Do DONG sheet Danh gia/Danh Gia (6 mau lech so dong va cot — khong ghi cung theo mau nao).
    Tra ve {'sheet', 'a': [{'so','dong','cot_max','cot_diem','tieu_chi':[(dong, ky_hieu, noi_dung, diem_max)]}],
    'tong_a': dong, 'b': {truc: {'dong','cot_pt','cot_max','cot_dat','diem_max','cong_thuc'}}, 'tong_b', 'tong',
    'dieu_kien': {'dong', 'cot_kq', 'cot_ghi_chu', 'muc': [(dong, tt, noi_dung, goi_y, ghi_chu)]},
    'de_xuat': dong muc III, 'ca_nhan': {nhan: dong}}. Thieu khoi nao -> LoiCauTruc."""
    from openpyxl.utils import get_column_letter
    dg = sheet_danh_gia(wb)
    o = lambda r, c: dg.cell(r, c).value  # noqa: E731
    kq = {"sheet": dg.title, "a": [], "b": {}, "ca_nhan": {}}
    # --- Muc A: hang tieu de "TT | TIEU CHI DANH GIA | ... Diem toi da | Diem cham"; nhom = dong cot A la 1/2/3
    #     co o "Diem toi da" = SUM(<cot>a:<cot>b) (cac tieu chi con a, b, c...)
    hang_a = next((r for r in range(1, dg.max_row + 1) if str(o(r, 1) or "").strip() == "TT"
                   and _kd(o(r, 2)).startswith("tieu chi danh gia")), None)
    if hang_a:
        tde = {_kd(o(hang_a, c)): get_column_letter(c) for c in range(1, dg.max_column + 1) if o(hang_a, c)}
        cmax = next((v for k, v in tde.items() if k.startswith("diem toi da")), None)
        cdiem = next((v for k, v in tde.items() if k.startswith("diem cham")), None)
        for r in range(hang_a + 1, dg.max_row + 1):
            a = str(o(r, 1) or "").strip()
            m = re.match(r"=SUM\(([A-Z]+)(\d+):\1(\d+)\)$", str(dg[f"{cmax}{r}"].value or "")) if cmax else None
            if a == str(len(kq["a"]) + 1) and m and m.group(1) == cmax and len(kq["a"]) < 3:
                tu, den = int(m.group(2)), int(m.group(3))
                tc = [(i, str(o(i, 1) or "").strip(), str(o(i, 2) or "").strip(),
                       float(dg[f"{cmax}{i}"].value or 0)) for i in range(tu, den + 1)]
                kq["a"].append({"so": int(a), "dong": r, "ten": str(o(r, 2) or "").strip(), "cot_max": cmax,
                                "cot_diem": cdiem, "tieu_chi": tc,
                                "diem_max": sum(x[3] for x in tc)})
    # --- Muc B: dong cot B bat dau "Truc (n)" co diem toi da so (cot co tieu de "Điểm tối đa")
    hang_b = None
    for r in range(1, dg.max_row + 1):
        if str(o(r, 1) or "").strip() == "TT" and "Tiêu chí/Nội dung" in str(o(r, 2) or ""):
            hang_b = r
    if hang_b:
        tieu_de = {_kd(o(hang_b, c)): get_column_letter(c) for c in range(1, dg.max_column + 1) if o(hang_b, c)}
        cot_pt = next((v for k, v in tieu_de.items() if k.startswith("diem kpi")), None)
        cot_max = next((v for k, v in tieu_de.items() if k.startswith("diem toi da")), None)
        cot_dat = next((v for k, v in tieu_de.items() if k.startswith("diem dat")), None)
        for r in range(hang_b + 1, dg.max_row + 1):
            m = re.match(r"Trục \((\d)\)", str(o(r, 2) or ""))
            if m and cot_max and isinstance(dg[f"{cot_max}{r}"].value, (int, float, str)):
                try:
                    dmax = float(dg[f"{cot_max}{r}"].value)
                except (TypeError, ValueError):
                    continue
                kq["b"][int(m.group(1))] = {"dong": r, "cot_pt": cot_pt, "cot_max": cot_max, "cot_dat": cot_dat,
                                            "diem_max": dmax, "cong_thuc": dg[f"{cot_pt}{r}"].value}
    # --- Dong tong, dieu kien (II), de xuat (III), thong tin ca nhan
    for r in range(1, dg.max_row + 1):
        a, b = str(o(r, 1) or "").strip(), _kd(o(r, 2))
        if b == "tong diem a + b":
            kq["tong"] = r
        elif b == "tong diem nhom b":
            kq["tong_b"] = r
        elif b in ("tong diem", "tong diem nhom a") and "tong_a" not in kq:
            kq["tong_a"] = r
        elif a == "II." and "dieu kien bat buoc" in b:
            kq["dieu_kien"] = {"dong": r, "muc": []}
        elif a.startswith("III.") and "tu de xuat" in _kd(a):
            kq["de_xuat"] = r
        for nhan in ("Họ và tên", "Chức vụ Đảng", "Chức vụ chính quyền", "Chức vụ đoàn thể", "Đơn vị công tác"):
            if a.startswith(nhan) and nhan not in kq["ca_nhan"]:
                kq["ca_nhan"][nhan] = r
    dk = kq.get("dieu_kien")
    if dk:
        tde = dk["dong"] + 1
        hdr = {_kd(o(tde, c)): get_column_letter(c) for c in range(1, dg.max_column + 1) if o(tde, c)}
        dk["cot_kq"] = next((v for k, v in hdr.items() if k.startswith("ket qua")), None)
        dk["cot_ghi_chu"] = next((v for k, v in hdr.items() if k.startswith("ghi chu")), None)
        het = kq.get("de_xuat", dg.max_row + 1)
        for r in range(tde + 1, het):
            tt = str(o(r, 1) or "").strip()
            if re.fullmatch(r"\d+", tt) and o(r, 2):
                goi_y = dg[f"{dk['cot_kq']}{r}"].value if dk["cot_kq"] else None
                gc = dg[f"{dk['cot_ghi_chu']}{r}"].value if dk["cot_ghi_chu"] else None
                dk["muc"].append((r, tt, str(o(r, 2)).strip(), goi_y, gc))
    thieu = [t for t, ok in (("mục A (3 nhóm tiêu chí chung)", len(kq["a"]) == 3),
                             ("mục B (6 Trục có điểm tối đa)", sorted(kq["b"]) == [1, 2, 3, 4, 5, 6]),
                             ("cột Điểm KPI (%) / Điểm tối đa / Điểm đạt", hang_b and all(
                                 kq["b"].get(1, {}).get(k) for k in ("cot_pt", "cot_max", "cot_dat"))),
                             ("dòng Tổng điểm A + B", "tong" in kq),
                             ("khối II. Điều kiện bắt buộc", dk and dk.get("cot_kq") and dk["muc"]),
                             ("dòng III. Tự đề xuất mức xếp loại", "de_xuat" in kq)) if not ok]
    if thieu:
        raise LoiCauTruc(f"Sheet '{dg.title}': không dò được " + "; ".join(thieu) +
                         " — dừng, báo người dùng; không đoán cấu trúc.")
    return kq


def mo(p):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(p)


def dong_vi_du(nhom):
    """{(sheet, o): gia tri} cua cac o du lieu CO SAN trong mau o vung dau viec — dong vi du chua xoa (known issue 4)."""
    wb = mo(tep_mau(nhom))
    ct = cau_truc(wb)
    kh = wb["Ke Hoach"]
    vd = {}
    for d, c in ct["truc"].values():
        for r in range(d, c + 1):
            for col in "BCDEFGHIJ":
                v = kh[f"{col}{r}"].value
                if v not in (None, ""):
                    vd[f"{col}{r}"] = v
    return vd


O_THUC_TE = ("K", "L", "N", "P")   # sheet KPI: san pham thuc te, SL thuc te, % chat luong, % tien do (o nhap)


def xoa_thuc_te(kp, r):
    """Xoa o NHAP so thuc te cua mot dong viec tren sheet KPI; giu nguyen o cong thuc."""
    for col in O_THUC_TE:
        v = kp[f"{col}{r}"].value
        if v is not None and not str(v).startswith("="):
            kp[f"{col}{r}"].value = None


def xuong_dong(sh, r, cot):
    """Bat wrap_text cho cac o co noi dung va nang chieu cao dong theo so dong chu uoc tinh (co chu x 1,3 pt)."""
    import math
    from copy import copy
    dong = 1
    co = 12
    for col in cot:
        c = sh[f"{col}{r}"]
        if c.value in (None, "") or str(c.value).startswith("="):
            continue
        al = copy(c.alignment)
        al.wrap_text = True
        c.alignment = al
        co = max(co, float(c.font.sz or 12)) if c.font else co
        rong = sh.column_dimensions[col].width or 10
        ky_tu_dong = max(1.0, rong * 1.1 * 12 / co)      # ~ so ky tu vua mot dong cua cot
        dong = max(dong, sum(math.ceil(max(1, len(p)) / ky_tu_dong) for p in str(c.value).split("\n")))
    h = round(dong * co * 1.3 + 3, 1)
    if (sh.row_dimensions[r].height or 0) < h:
        sh.row_dimensions[r].height = h


def thuc_te_vi_du(nhom):
    """{'dong N': {o: gia tri}} — so thuc te VI DU co san trong mau o sheet KPI (Known-Issues-Bieu-Mau #12)."""
    wb = mo(tep_mau(nhom))
    ct = cau_truc(wb)
    kp = wb["KPI"]
    kq = {}
    for r in ct["kpi_dong"].values():
        o = {f"{col}{r}": kp[f"{col}{r}"].value for col in O_THUC_TE
             if kp[f"{col}{r}"].value not in (None, "") and not str(kp[f"{col}{r}"].value).startswith("=")}
        if o:
            kq[f"dòng {r}"] = o
    return kq


def ghi_ke_hoach(nhom, kh, ra, quy=None, nam=None):
    """Dien ke hoach vao BAN SAO mau. kh: {"ca_nhan":{ho_ten,ngay_sinh,chuc_vu_dang,chuc_vu_chinh_quyen,
    chuc_vu_doan_the,don_vi}, "dau_viec":[{truc,noi_dung,cap_trinh,muc_do,san_pham,so_luong,thoi_han,he_so,
    minh_chung,ghi_chu}], "phuong_an":..., "trang_thai_he_so":...}  (he_so da tinh bang kpi_calc).
    Tra ve danh sach thong bao."""
    goc = tep_mau(nhom)
    if os.path.abspath(ra) == os.path.abspath(goc) or os.path.abspath(os.path.dirname(ra)) == thu_muc_mau():
        raise ValueError("Không ghi vào thư mục mẫu assets/ — chọn nơi lưu khác.")
    os.makedirs(os.path.dirname(os.path.abspath(ra)), exist_ok=True)
    shutil.copyfile(goc, ra)
    wb = mo(ra)
    ct = cau_truc(wb)
    ws, kp = wb["Ke Hoach"], wb["KPI"]
    tb = []
    # 1. Xoa dong vi du trong vung dau viec (giu cot A = STT). Sheet KPI cung co SO THUC TE vi du o dong viec dau
    #    (L=4, N=100, P=100 — Known-Issues-Bieu-Mau #12): xoa o nhap lieu, giu cong thuc.
    for d, c in ct["truc"].values():
        for r in range(d, c + 1):
            for col in "BCDEFGHIJ":
                ws[f"{col}{r}"].value = None
            xoa_thuc_te(kp, ct["kpi_dong"][r])
            if ct["cot_minh_chung"]:
                kp[f"{ct['cot_minh_chung']}{ct['kpi_dong'][r]}"].value = None
    # 2. Thong tin ca nhan
    cn = kh.get("ca_nhan", {})
    nhan = {4: ("Họ và tên", "ho_ten", "Ngày sinh", "ngay_sinh"), 5: ("Chức vụ Đảng", "chuc_vu_dang"),
            6: ("Chức vụ chính quyền", "chuc_vu_chinh_quyen"), 7: ("Chức vụ đoàn thể", "chuc_vu_doan_the"),
            8: ("Đơn vị công tác", "don_vi")}
    for r, t in nhan.items():
        if len(t) == 4:
            ws[f"B{r}"].value = f"{t[0]}: {cn.get(t[1]) or '…'}    {t[2]}: {cn.get(t[3]) or '…'}"
        else:
            ws[f"B{r}"].value = f"{t[0]}: {cn.get(t[1]) or '…'}"
    # 3. Ky (quy, nam) trong tieu de
    if quy and nam:
        for sh in (ws, kp):
            for r in range(1, 4):
                for c in sh[r]:
                    if isinstance(c.value, str):
                        c.value = re.sub(r"QUÝ\s+[IVX]+\s+NĂM\s+\d{4}", f"QUÝ {quy} NĂM {nam}", c.value)
    # 4. Dau viec
    dem = {n: 0 for n in ct["truc"]}
    pa = kh.get("phuong_an")
    for dv in kh.get("dau_viec", []):
        n = int(dv["truc"])
        if n not in ct["truc"]:
            raise ValueError(f"Trục {n} không có trong mẫu.")
        d, c = ct["truc"][n]
        r = d + dem[n]
        if r > c:
            raise ValueError(f"Trục {n} vượt {c - d + 1} dòng của mẫu — gộp bớt hoặc hỏi Phòng TCCB&CTHSSV "
                             "(không tự chèn dòng làm lệch công thức sheet KPI).")
        dem[n] += 1
        ghi_chu = [dv.get("ghi_chu") or ""]
        hs = dv.get("he_so")
        ws[f"B{r}"].value = dv.get("noi_dung")
        ws[f"C{r}"].value = dv.get("cap_trinh")
        ws[f"D{r}"].value = dv.get("muc_do")
        ws[f"E{r}"].value = dv.get("san_pham")
        ws[f"F{r}"].value = dv.get("so_luong")
        ws[f"G{r}"].value = dv.get("thoi_han")
        ws[f"I{r}"].value = hs
        if pa == "muc-do" and hs is not None:
            ws[f"H{r}"].value = round(float(hs) * 100)     # Diem cham = He so x 100 [QD 1923 PL I cot (9)(10)]
        elif hs is not None:
            ghi_chu.append(f"Hệ số theo phương án {pa}")
        mc = dv.get("minh_chung")
        if mc:
            if ct["cot_minh_chung"]:
                kp[f"{ct['cot_minh_chung']}{ct['kpi_dong'][r]}"].value = mc
            else:
                ghi_chu.append(f"Minh chứng: {mc}")
        ws[f"J{r}"].value = "; ".join(x for x in ghi_chu if x) or None
    if pa and pa != "muc-do":
        tb.append(f"Hệ số theo phương án '{pa}': {kh.get('trang_thai_he_so', '')}")
    # 4b. Trinh bay (Known-Issues-Bieu-Mau #10, #11 — phien 24–25/9 phai va tay): xuong dong + chieu cao dong
    #     cho dong da ghi; AN (khong xoa) dong trong trong khoi 20 dong/Truc, ca "Ke Hoach" lan "KPI".
    an = 0
    for d, c in ct["truc"].values():
        for r in range(d, c + 1):
            if ws[f"B{r}"].value in (None, ""):
                ws.row_dimensions[r].hidden = True
                kp.row_dimensions[ct["kpi_dong"][r]].hidden = True
                an += 1
            else:
                xuong_dong(ws, r, "BCDEGJ")
                xuong_dong(kp, ct["kpi_dong"][r], "B")
    if an:
        tb.append(f"Đã ẩn {an} dòng trống trong khối đầu việc (không xóa — bỏ ẩn được khi cần thêm việc)")
    # 5. The thuc: phong chu Times New Roman cho moi o co noi dung (mau goc con o Calibri)
    from copy import copy
    doi = 0
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font is not None and cell.font.name != PHONG:
                    f = copy(cell.font)
                    f.name = PHONG
                    cell.font = f
                    doi += 1
    if doi:
        tb.append(f"Đã chuẩn hóa {doi} ô sang {PHONG} (mẫu gốc còn phông khác — Known-Issues-Bieu-Mau.md)")
    wb.save(ra)
    return tb


def main(argv):
    """Mot lenh tron quy trinh: tinh he so (kpi_calc) -> dien mau -> kiem (validate_plan).
    python scripts/kpi_mau.py --nhom hanh-chinh --json ke_hoach.json --phuong-an muc-do --quy IV --nam 2026 --ra <tep.xlsx>
    Ma thoat: 2 = loi dau vao (dung, hoi nguoi dung) · 1 = da xuat nhung con LOI · 0 = sach."""
    import argparse
    import json
    import sys
    sys.path.insert(0, HERE)
    import kpi_calc as kc
    import validate_plan as vp
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--nhom", required=True, choices=list(NHOM))
    ap.add_argument("--json", required=True)
    ap.add_argument("--phuong-an")
    ap.add_argument("--quy")
    ap.add_argument("--nam")
    ap.add_argument("--ra", required=True)
    a = ap.parse_args(argv)
    if a.phuong_an not in kc.PHUONG_AN:
        print(f"✗ Chưa chọn phương án hệ số ({', '.join(kc.PHUONG_AN)}) — Câu hỏi mở số 1, không có mặc định. "
              "Hỏi người dùng.")
        return 2
    kh = json.load(open(a.json, encoding="utf-8"))
    qd = kc.quy_doi_ke_hoach(kh, a.phuong_an)
    if qd["loi"]:
        print("✗ Không tính được hệ số — dừng, hỏi người dùng:")
        for x in qd["loi"]:
            print(f"  dòng {x['dong']} '{x['noi_dung']}': {x['loi']}")
        return 2
    kh = dict(kh, dau_viec=qd["dau_viec"], phuong_an=a.phuong_an, trang_thai_he_so=qd["trang_thai"])
    try:
        tb = ghi_ke_hoach(a.nhom, kh, a.ra, a.quy, a.nam)
    except (ValueError, FileNotFoundError) as e:
        print(f"✗ {e}")
        return 2
    print(f"✓ Đã xuất {a.ra}")
    print(f"  Phương án hệ số: {a.phuong_an} — {qd['trang_thai']}")
    for x in tb:
        print(f"  · {x}")
    for d in qd["dau_viec"]:
        for c in d.get("canh_bao") or []:
            print(f"  ⚠ {d.get('noi_dung', '')[:50]}: {c}")
    kq = vp.kiem(a.ra, a.nhom, a.phuong_an)
    for x in sorted(kq["loi"], key=lambda x: (x["muc"] != "LOI", x["ma"])):
        print(f"  [{x['muc']:8s}] {x['ma']} {x['vi_tri']}: {x['noi_dung']}  [{x['can_cu']}]")
    return 1 if any(x["muc"] == "LOI" for x in kq["loi"]) else 0


if __name__ == "__main__":
    import sys
    sys.exit(main(sys.argv[1:]))
