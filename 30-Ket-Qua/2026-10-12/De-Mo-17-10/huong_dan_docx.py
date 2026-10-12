# -*- coding: utf-8 -*-
"""Dựng tệp Word: Hướng dẫn chạy thử và quay 04 thao tác (để người trình bày mở bên cạnh khi quay video)."""
import os, sys
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tao_du_lieu_gia_dinh as G
sys.stdout.reconfigure(encoding='utf-8')
D = os.path.dirname(os.path.abspath(__file__))


def bang(d, td, hang, rong):
    t = d.add_table(rows=1, cols=len(td)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    G.rong(t, rong)
    for j, x in enumerate(td):
        G.o_bang(t.cell(0, j), [(x, True, False, 13)])
    for h in hang:
        r = t.add_row()
        for j, x in enumerate(h):
            G.o_bang(r.cells[j], [(x, False, False, 13)], can=WD_ALIGN_PARAGRAPH.CENTER if j == 0 and len(x) < 6 else WD_ALIGN_PARAGRAPH.LEFT)
    G.rong(t, rong)
    G.doan(d, '', sau=0)


d = docx.Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
st.element.rPr.rFonts.set(docx.oxml.ns.qn('w:eastAsia'), 'Times New Roman')
G.doan(d, 'HƯỚNG DẪN', dam=True, can='giua', sau=0)
G.doan(d, 'Chạy thử và quay 04 thao tác trình diễn ngày 17/10/2026', dam=True, can='giua', sau=0); G.ke(d)
G.doan(d, 'Mọi thứ còn lại đã chuẩn bị sẵn trong thư mục De-Mo-17-10. Chỉ còn 02 việc cần người làm trên Claude Cowork: Chạy thử và quay màn hình.')
G.doan(d, 'I. ĐÃ CÓ SẴN', dam=True)
bang(d, ['Nội dung', 'Vị trí'], [
 ['Bộ dữ liệu giả định D1 - D4', 'Thư mục Du-lieu-gia-dinh (KHÔNG kéo thư mục _Dap-an-KHONG-dua-vao-cong-cu vào Cowork)'],
 ['Kết quả chạy trước của cả 03 phần (mở khi trình diễn, kể cả khi mất mạng)', 'Thư mục Ket-qua-chay-truoc'],
 ['Giọng thuyết minh 04 đoạn (giọng nữ tiếng Việt do AI đọc)', 'Thư mục Video-thao-tac, các tệp T1 đến T4 (.mp3)'],
 ['Công cụ ghép video', 'Video-thao-tac/ghep_video.py'],
 ['Trang trình chiếu 11 trang', 'Thư mục Trinh-chieu'],
], (7.0, 9.0))
G.doan(d, 'II. BƯỚC 1 - CHẠY THỬ TRÊN CLAUDE COWORK (khoảng 30 phút)', dam=True)
G.doan(d, 'Mở Claude Cowork, kết nối thư mục Du-lieu-gia-dinh. Mỗi thao tác mở một cuộc hội thoại mới (riêng T2 làm tiếp trong cuộc hội thoại T1). Kéo tệp vào, rồi sao chép nguyên văn câu lệnh ở cột bên phải:')
bang(d, ['Thao tác', 'Kéo vào', 'Câu lệnh (sao chép nguyên văn)'], [
 ['T1', 'Thư mục D1-Ke-hoach-quy-IV-2026', 'Kiểm tra hồ sơ kế hoạch quý của các đơn vị trong thư mục này, rồi tổng hợp thành dự thảo Kế hoạch công tác quý IV năm 2026 của cơ quan'],
 ['T2', 'Tệp trong thư mục D2-He-thong-QLVB', 'Từ danh mục văn bản này, liệt kê văn bản chỉ đạo, kế hoạch năm, quý, thông báo kết luận, kế hoạch chuyên đề trong kỳ; nhiệm vụ nào cấp trên đã giao mà chưa đơn vị nào đưa vào kế hoạch?'],
 ['T3', 'Thư mục D3-Bao-cao-thang-9-2026', 'Lập báo cáo kết quả công tác tháng 9 năm 2026 của cơ quan từ báo cáo các đơn vị, đối chiếu với kế hoạch tháng, phát triển từ báo cáo tháng 8 đã ban hành'],
 ['T4', 'Tệp Du-thao_Chuong-trinh-cong-tac-thang-10-2026_van-ban-Dang_an-danh.docx (trong thư mục D4)', 'Rà soát chính thức dự thảo này theo KTC-Ra-Soat-897. Đây là văn bản của cơ quan khác, không áp quy ước riêng của Trường'],
], (1.8, 5.0, 9.2))
G.doan(d, 'So kết quả với tệp DAP-AN.md trong thư mục _Dap-an-KHONG-dua-vao-cong-cu. Công cụ bỏ sót hoặc nêu thêm là bình thường: Ghi lại để nói đúng thực tế khi trình diễn, để bộ công cụ tự chỉ ra điểm chưa tốt.')
G.doan(d, 'III. BƯỚC 2 - QUAY PHẦN ĐẦU MỖI THAO TÁC (mỗi đoạn 30 - 45 giây)', dam=True)
for x in ['1. Đặt độ phân giải màn hình 1920 x 1080; phóng chữ Claude Desktop lên 125%; tắt thông báo của Zalo, thư điện tử.',
          '2. Mở Snipping Tool, chọn Ghi (Record), quét vùng cửa sổ Claude Desktop, bấm Bắt đầu. Cách khác: Bấm tổ hợp phím Windows + Alt + R.',
          '3. Kéo tệp vào, dán câu lệnh, bấm gửi; để công cụ làm việc khoảng 20 - 30 giây (thấy công cụ đọc tệp, gọi kỹ năng) rồi dừng quay. Không cần nói khi quay.',
          '4. Lưu vào thư mục Video-thao-tac/quay với đúng tên: T1.mp4, T2.mp4, T3.mp4, T4.mp4.']:
    G.doan(d, x)
G.doan(d, 'IV. BƯỚC 3 - GHÉP VIDEO (01 lệnh, khoảng 01 phút)', dam=True)
G.doan(d, 'Nhắn Claude Code: “ghép 4 video thao tác”; hoặc mở cửa sổ lệnh tại thư mục De-Mo-17-10 và chạy: python Video-thao-tac/ghep_video.py')
G.doan(d, 'Kết quả nằm ở thư mục Video-thao-tac/xuat: 04 video có giọng thuyết minh và dòng chữ “Dữ liệu giả định phục vụ trình diễn”. Đoạn quay ngắn hơn lời đọc thì hình cuối được giữ đến hết lời; dài hơn thì giữ tối đa 45 giây.')
G.doan(d, 'V. KHI TRÌNH DIỄN', dam=True)
G.doan(d, 'Với mỗi nội dung trọng tâm: Chiếu video thao tác (30 - 45 giây), sau đó mở tệp tương ứng trong thư mục Ket-qua-chay-truoc và nêu các điểm công cụ phát hiện được, kể cả điểm công cụ nêu chưa đúng.')
bang(d, ['Phần', 'Video', 'Tệp kết quả mở sau video'], [
 ['1', 'T1, T2', 'Phan-1_Tong-hop-ke-hoach-quy-IV-2026_ket-qua-chay-truoc.xlsx (05 trang tính)'],
 ['2', 'T3', 'Phan-2_Doi-chieu-ke-hoach-bao-cao-thang-9-2026_ket-qua-chay-truoc.xlsx; Phan-2_Du-thao-Bao-cao-thang-9-2026_Co-quan-A_ket-qua-chay-truoc.docx'],
 ['3', 'T4', 'BaoCao_RaSoat-897-v2.25_Chuong-trinh-cong-tac-thang-10-2026_van-ban-Dang_an-danh.docx (25 phát hiện)'],
], (1.5, 2.5, 12.0))
ra = os.path.join(D, '00-HUONG-DAN-CHAY-THU-VA-QUAY.docx'); d.save(ra); print(ra)
