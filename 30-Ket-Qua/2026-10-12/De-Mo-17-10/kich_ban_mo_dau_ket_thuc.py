# -*- coding: utf-8 -*-
"""Dựng tệp Word: Kịch bản lời trình bày trực tiếp phần MỞ ĐẦU (trước khi chiếu video) và KẾT THÚC (sau video, chuyển sang thảo luận)
tại buổi làm việc ngày 17/10/2026. Căn cứ đối chiếu trong kho KTC-Database: Luật Trí tuệ nhân tạo (kho 01-01), Quyết định số 23-QĐ/BCĐ (kho 01-03)."""
import os, sys
import docx
from docx.shared import Cm, Pt
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import tao_du_lieu_gia_dinh as G
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
sys.stdout.reconfigure(encoding='utf-8')
D = os.path.dirname(os.path.abspath(__file__))


def bang(d, td, hang, rong):
    t = d.add_table(rows=1, cols=len(td)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, x in enumerate(td):
        G.o_bang(t.cell(0, j), [(x, True, False, 13)])
    for h in hang:
        r = t.add_row()
        for j, x in enumerate(h):
            G.o_bang(r.cells[j], [(x, False, False, 13)], can=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    G.rong(t, rong)
    G.doan(d, '', sau=0)


def loi(d, x):
    """Lời nói: In nghiêng, thụt đầu dòng, để người trình bày đọc liền mạch."""
    G.doan(d, x, nghieng=True)


d = docx.Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
st.element.rPr.rFonts.set(docx.oxml.ns.qn('w:eastAsia'), 'Times New Roman')
G.doan(d, 'KỊCH BẢN LỜI TRÌNH BÀY TRỰC TIẾP', dam=True, can='giua', sau=0)
G.doan(d, 'Phần mở đầu và phần kết thúc, buổi làm việc ngày 17/10/2026', dam=True, can='giua', sau=0); G.ke(d)
G.doan(d, 'Người trình bày: Đồng chí Nguyễn Ngọc Quang Phục. Thời lượng: Mở đầu khoảng 02 phút 30 giây; video 24 phút 32 giây; kết thúc '
          'khoảng 03 phút. Chữ in nghiêng là lời nói gợi ý, người trình bày nói theo cách của mình; chữ đứng trong ngoặc vuông là ghi chú, '
          'không đọc. Chỗ ghi [...] điền theo thành phần Đoàn công tác thực tế.')

G.doan(d, 'I. PHẦN MỞ ĐẦU (khoảng 02 phút 30 giây, trước khi chiếu video)', dam=True)
G.doan(d, '[Trang trình chiếu 1 đang hiện trên màn hình. Đứng, nhìn về phía Đoàn công tác, không đọc từ màn hình.]')
loi(d, 'Kính thưa đồng chí [...], Trưởng đoàn công tác của Đảng ủy Ủy ban nhân dân tỉnh; kính thưa các đồng chí trong Đoàn; kính thưa các '
       'đồng chí Lãnh đạo Trường. Thay mặt phòng Tổng hợp - Hành chính và Quản trị, tôi xin báo cáo phần ứng dụng trí tuệ nhân tạo trong '
       'công tác tham mưu, tổng hợp của Trường.')
loi(d, 'Trước khi trình chiếu, tôi xin phép nói ngắn ba điều.')
loi(d, 'Thứ nhất, về định hướng. Nghị quyết số 57-NQ/TW của Bộ Chính trị đặt chuyển đổi số là đột phá. Luật Trí tuệ nhân tạo có hiệu lực '
       'từ ngày 01/3/2026 khuyến khích ưu tiên ứng dụng trí tuệ nhân tạo trong quản lý, điều hành của cơ quan nhà nước, đồng thời nêu rõ '
       'nguyên tắc: Trí tuệ nhân tạo phục vụ con người, không thay thế thẩm quyền và trách nhiệm của con người. Trong các cơ quan đảng, '
       'tháng 8 vừa qua đã có Bộ chỉ số đánh giá năng lực số và mức độ sẵn sàng chuyển đổi số. Toàn bộ cách làm của Trường đi theo đúng '
       'nguyên tắc đó: Công cụ đọc, đối chiếu, dự thảo; người có thẩm quyền quyết định.')
loi(d, 'Thứ hai, về cách làm. Trường không đầu tư phần mềm riêng. Chúng tôi dùng một công cụ trí tuệ nhân tạo phổ thông, kết hợp với ba '
       'thứ do chính Trường xây dựng: Kho dữ liệu nền gồm văn bản của cấp trên và văn bản đã ban hành của Trường; bộ quy tắc rà soát văn '
       'bản; và bộ công cụ quản trị nhiệm vụ. Người sử dụng hằng ngày là chuyên viên của phòng, không phải chuyên gia công nghệ.')
loi(d, 'Thứ ba, về cách theo dõi phần trình chiếu. Video dài khoảng 24 phút, gồm 04 phần. Ba phần đầu dùng dữ liệu giả định để trình diễn '
       'đúng 03 nội dung Đoàn đề nghị; văn bản của Đảng đưa vào rà soát đã được ẩn danh. Phần thứ tư là những việc thật Phòng đã làm. Lời '
       'đọc trong video do trí tuệ nhân tạo tạo ra, chúng tôi xin nói rõ để bảo đảm minh bạch. Video cũng nêu công khai những chỗ công cụ '
       'làm chưa đúng, vì đó chính là kinh nghiệm cần chia sẻ.')
loi(d, 'Sau phần trình chiếu, tôi xin dành thời gian để các đồng chí trao đổi; nếu Đoàn đề nghị, chúng tôi chạy trực tiếp một thao tác. '
       'Xin mời các đồng chí theo dõi.')
G.doan(d, '[Ra hiệu cho đồng chí Đỗ Thành Công bấm chạy tệp A_VIDEO-TOAN-BO. Trong khi chiếu: Ngồi cùng phía Đoàn, ghi lại câu hỏi phát '
          'sinh; nếu cần dừng, dừng ở các mốc chuyển phần (04:57, 09:29, 12:23, 15:44, 23:19).]')

G.doan(d, 'II. PHẦN KẾT THÚC (khoảng 03 phút, sau video)', dam=True)
G.doan(d, '[Video dừng ở trang cuối “Trân trọng cảm ơn”. Chuyển về trang trình chiếu 10 hoặc giữ trang cuối.]')
loi(d, 'Kính thưa các đồng chí, qua phần trình chiếu, tôi xin rút ra 05 kinh nghiệm từ thực tế sử dụng tại Phòng.')
loi(d, 'Một, trí tuệ nhân tạo có thể viết sai mà rất trôi chảy. Một quyết định của Trường đi từ tra cứu đến khi ký khoảng 10 ngày; chính '
       'người có nghiệp vụ đã phát hiện 02 lỗi căn cứ do công cụ đưa vào. Vì vậy, mọi căn cứ phải đối chiếu với văn bản gốc, và ô người '
       'kiểm tra luôn là người thật.')
loi(d, 'Hai, đầu vào quyết định đầu ra. Các đồng chí vừa thấy: Thiếu dự thảo kế hoạch, công cụ nêu tới 09 việc chưa ai nhận, thực tế chỉ có '
       '02. Công cụ chỉ tốt khi được đưa đủ hồ sơ.')
loi(d, 'Ba, giao việc cho công cụ như giao cho một chuyên viên mới: Nói rõ mục đích, mẫu biểu, căn cứ và sản phẩm cần có. Trường tóm lại '
       'thành ba thứ phải có: Dữ liệu đầu vào, yêu cầu sản phẩm đầu ra, và cách xử lý.')
loi(d, 'Bốn, nền móng trước, công cụ sau. Kho văn bản đã ban hành tốt của chính cơ quan giúp công cụ viết đúng văn phong, đúng thể thức; '
       'không có kho thì công cụ chỉ viết chung chung.')
loi(d, 'Năm, phải có quy định rõ: Loại tài liệu nào không được đưa vào công cụ, tài khoản nào được dùng, ai là người kiểm tra. Nếu không, '
       'mỗi người sẽ tự dùng một cách và không kiểm soát được.')
loi(d, 'Từ những kinh nghiệm đó, chúng tôi xin đề xuất để bắt đầu: Lập kho dữ liệu nền; chọn 01 đến 02 việc lặp lại hằng tháng như tổng '
       'hợp kế hoạch, báo cáo để làm trước; và ban hành quy định sử dụng ngay từ đầu. Bắt đầu nhỏ, làm thật, có người kiểm tra.')
loi(d, 'Trường Cao đẳng Kon Tum sẵn sàng chia sẻ, hỗ trợ khi nhận được thông tin đầu vào và yêu cầu sản phẩm đầu ra của cơ quan. Video, '
       'kịch bản và bộ dữ liệu mẫu hôm nay, Trường xin gửi Đoàn làm tư liệu để các đồng chí xem lại sau buổi làm việc.')
loi(d, 'Phần báo cáo của Phòng xin dừng tại đây. Kính mời các đồng chí trao đổi; nếu các đồng chí muốn xem trực tiếp một thao tác, chúng '
       'tôi xin thực hiện ngay. Trân trọng cảm ơn các đồng chí.')

G.doan(d, 'III. NGÂN HÀNG KINH NGHIỆM THỰC TẾ (dùng khi trả lời câu hỏi)', dam=True)
bang(d, ['TT', 'Kinh nghiệm', 'Dẫn chứng có thể nói'], [
 ['1', 'Công cụ có thể bịa căn cứ, số liệu với văn phong rất thuyết phục', 'Người có nghiệp vụ bắt được 02 lỗi căn cứ trước khi ký; rà soát phát hiện kế hoạch ghi sai số hiệu quyết định (1299 thay vì 1229)'],
 ['2', 'Đầu vào thiếu thì kết luận lệch', 'Đối chiếu văn bản chỉ đạo khi thiếu dự thảo kế hoạch: 09 việc so với 02 việc thực tế'],
 ['3', 'Bảng tính do công cụ dựng phải kiểm tra lại', 'Lượt chạy thật: 06 dòng bảng tính mất dữ liệu'],
 ['4', 'Cùng một văn bản, mỗi lượt chạy có thể khác nhau', 'Rà soát nhanh: 06 phát hiện, có 01 sai; rà soát đầy đủ: 25 phát hiện. Rà soát chính thức phải chạy đủ quy trình'],
 ['5', 'Công cụ tốt nhất ở việc lặp lại, có mẫu', 'Báo cáo tháng, kế hoạch tháng, rà soát thể thức; thông báo kết luận dựng ngay trong cuộc họp'],
 ['6', 'Đo được độ tin cậy thì mới dám dùng', 'Kiểm thử định kỳ trên tệp cài sẵn 30 lỗi: Đúng 24, sai mức 02, bỏ sót 04'],
 ['7', 'Công cụ biết dừng lại hỏi khi thiếu thông tin là dấu hiệu dùng đúng cách', 'Tự đánh giá KPI: Thiếu sản phẩm đầu ra thì công cụ hỏi lại, không tự điền'],
 ['8', 'Minh bạch nội dung do trí tuệ nhân tạo tạo ra', 'Luật Trí tuệ nhân tạo yêu cầu bảo đảm người dùng nhận biết nội dung do hệ thống tạo ra (Điều 11); Trường nói rõ lời đọc trong video do trí tuệ nhân tạo tạo'],
], (1.2, 6.0, 8.8))

G.doan(d, 'IV. LƯU Ý KHI TRÌNH BÀY', dam=True)
for x in ['- Nói, không đọc màn hình; mỗi kinh nghiệm một câu chính, một dẫn chứng.',
          '- Dùng đúng tên: Luật Trí tuệ nhân tạo (không nêu số hiệu Luật khi nói); Nghị quyết số 57-NQ/TW; Bộ chỉ số đánh giá năng lực số và mức '
          'độ sẵn sàng chuyển đổi số trong các cơ quan đảng (Quyết định số 23-QĐ/BCĐ ngày 17/8/2026).',
          '- Không nêu chi phí, hạ tầng, tính năng chung của nền tảng: Phần này do phòng QLKHCN&HTPT trình bày.',
          '- Câu hỏi chưa có thông tin: Ghi nhận, hẹn trả lời bằng văn bản; không trả lời ước đoán.',
          '- Tập nói trọn hai phần (khoảng 05 phút) ít nhất 02 lần trước buổi tổng duyệt ngày 16/10/2026.']:
    G.doan(d, x, thut=False)

G.doan(d, 'Nguồn đã đối chiếu: Luật Trí tuệ nhân tạo, Điều 4 khoản 2, Điều 5 khoản 5, Điều 11 khoản 4, Điều 35 (kho KTC-Database 01-01); '
          'Quyết định số 23-QĐ/BCĐ ngày 17/8/2026 của Ban Chỉ đạo chuyển đổi số trong các cơ quan đảng, phần căn cứ dẫn Nghị quyết số 57-NQ/TW '
          'ngày 22/12/2024 (kho 01-03); số liệu trình diễn: Bộ kết quả chạy thật ngày 12/10/2026 và Phương án bản 7.', nghieng=True)
# số trang: Giữa lề trên, ẩn ở trang 1
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
s.different_first_page_header_footer = True
hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for loai, chu in (('begin', None), (None, 'PAGE'), ('end', None)):
    r = hp.add_run(); r.font.size = Pt(13); r.font.name = 'Times New Roman'
    if loai:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), loai)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = chu
    r._r.append(e)
ra = os.path.join(D, 'KICH-BAN-MO-DAU-KET-THUC_17.10.2026.docx'); d.save(ra); print(ra)
