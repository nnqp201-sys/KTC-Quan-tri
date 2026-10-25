# -*- coding: utf-8 -*-
"""Doc TOAN VAN van ban (.docx, .pdf co lop chu, .xlsx phu luc) thanh danh sach dong danh so.

Thu vien dung chung cho quy trinh trich nhiem vu (skill ktc-trich-nhiem-vu):
- Giu dung thu tu doan van va bang trong .docx; dung lai so tu dong cua Word ("1.", "1.1.", "- ")
  vi so do khong nam trong text.
- Dong bang ghi dang "[Bảng b.r] o1 | o2 | ...", dong phu luc Excel "[<sheet> r] o1 | o2 | ...".
- PDF: noi cac dong bi ngat giua cau; PDF anh quet (khong lop chu) -> tra ve [].

Khong suy dien noi dung — chi doc. Thay cho phan doc cua trich_nhiem_vu.py (da bo, 25/10/2026).
"""
import os
import re
import unicodedata

NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def khong_dau(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('đ', 'd').replace('Đ', 'D').lower()


# Ma don vi chuan (20-Chuan-Chung/13-Bang-Ma-Don-Vi.md) — 11 ma cap mot
MA_DON_VI = {
    'P-TCCB': 'Phòng Tổ chức cán bộ và Công tác học sinh, sinh viên',
    'P-QLDT': 'Phòng Quản lý Đào tạo và Bảo đảm chất lượng',
    'P-THHC': 'Phòng Tổng hợp - Hành chính và Quản trị',
    'P-QLKH': 'Phòng Quản lý khoa học công nghệ và Hợp tác phát triển',
    'P-TCKT': 'Phòng Tài chính - Kế toán',
    'K-KHCB': 'Khoa các Khoa học cơ bản',
    'K-SUPH': 'Khoa Sư phạm',
    'K-KTNL': 'Khoa Kinh tế và Nông Lâm',
    'K-KTCN': 'Khoa Kỹ thuật và Công nghệ',
    'K-YDUOC': 'Khoa Y - Dược',
    'K-DTSHLX': 'Khoa Đào tạo và Sát hạch lái xe',
}


def _dinh_dang_so(doc):
    """numId -> {ilvl: numFmt} tu numbering.xml."""
    kq = {}
    try:
        goc = doc.part.numbering_part.element
    except (KeyError, AttributeError, NotImplementedError):
        return kq
    tr = {}
    for a in goc.findall(NS + 'abstractNum'):
        lv = {}
        for l in a.findall(NS + 'lvl'):
            f = l.find(NS + 'numFmt')
            lv[int(l.get(NS + 'ilvl'))] = f.get(NS + 'val') if f is not None else ''
        tr[a.get(NS + 'abstractNumId')] = lv
    for n in goc.findall(NS + 'num'):
        a = n.find(NS + 'abstractNumId')
        if a is not None:
            kq[n.get(NS + 'numId')] = tr.get(a.get(NS + 'val'), {})
    return kq


def _tien_to_so(p, fmt, dem):
    """Them so tu dong cua Word vao dau doan (neu co va chua go tay)."""
    t = p.text.strip()
    ppr = p._p.pPr
    num = ppr.find(NS + 'numPr') if ppr is not None else None
    if num is None and p.style is not None:  # danh so khai bao trong kieu (List Number...)
        sp = p.style.element.find(NS + 'pPr')
        num = sp.find(NS + 'numPr') if sp is not None else None
    if num is None or re.match(r'^([ivx]+|\d+(\.\d+)*|[a-zđ])[.)]\s', t, re.I) or re.match(r'^[-–+•*]', t):
        return t
    nid = num.find(NS + 'numId').get(NS + 'val') if num.find(NS + 'numId') is not None else ''
    lv = int(num.find(NS + 'ilvl').get(NS + 'val')) if num.find(NS + 'ilvl') is not None else 0
    kieu = fmt.get(nid, {}).get(lv, 'decimal' if lv == 0 else 'bullet')
    if kieu == 'bullet':
        return '- ' + t
    dem[(nid, lv)] = dem.get((nid, lv), 0) + 1
    for k in [k for k in dem if k[0] == nid and k[1] > lv]:
        del dem[k]
    if kieu in ('upperRoman', 'lowerRoman'):
        return ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X'][min(dem[(nid, lv)], 10) - 1] + '. ' + t
    if kieu in ('lowerLetter', 'upperLetter'):
        return 'abcdefghij'[min(dem[(nid, lv)], 10) - 1] + ') ' + t
    return '.'.join(str(dem.get((nid, i), 1)) for i in range(lv + 1)) + '. ' + t


def doc_docx(duong):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(duong)
    fmt, dem, kq, so_bang = _dinh_dang_so(doc), {}, [], 0
    for el in doc.element.body.iterchildren():
        if el.tag == NS + 'p':
            t = _tien_to_so(Paragraph(el, doc), fmt, dem)
            if t:
                kq.append(t)
        elif el.tag == NS + 'tbl':
            so_bang += 1
            for r, hang in enumerate(Table(el, doc).rows, 1):
                o, truoc = [], None
                for c in hang.cells:  # o gop ngang lap lai -> bo trung
                    if c._tc is not truoc:
                        o.append(re.sub(r'\s+', ' ', c.text).strip())
                    truoc = c._tc
                if any(o):
                    kq.append('[Bảng %d.%d] %s' % (so_bang, r, ' | '.join(o)))
    return kq


def doc_pdf(duong):
    from pypdf import PdfReader
    dong = []
    for trang in PdfReader(duong).pages:
        dong += [x.strip() for x in (trang.extract_text() or '').splitlines() if x.strip()]
    kq = []
    for x in dong:
        moi = re.match(r'^([IVX]+\.|\d+(\.\d+)*[.)]|[-–+•*]|[a-zđ]\)|[A-ZĐ ]{6,}$)', x)
        if kq and not moi and not re.search(r'[.:;]$', kq[-1]):
            kq[-1] += ' ' + x
        else:
            kq.append(x)
    return kq if sum(len(x) for x in kq) >= 200 else []


def doc_xlsx(duong):
    import openpyxl
    wb = openpyxl.load_workbook(duong, data_only=True, read_only=True)
    kq = []
    for ws in wb.worksheets:
        for r, hang in enumerate(ws.iter_rows(values_only=True), 1):
            o = [re.sub(r'\s+', ' ', str(v)).strip() for v in hang if v is not None and str(v).strip()]
            if o:
                kq.append('[%s %d] %s' % (ws.title, r, ' | '.join(o)))
    return kq


def doc_toan_van(duong):
    """Tra ve danh sach dong. Loi dinh dang -> ValueError co ly do ro."""
    duoi = os.path.splitext(duong)[1].lower()
    if duoi == '.docx':
        return doc_docx(duong)
    if duoi == '.pdf':
        try:
            return doc_pdf(duong)
        except ImportError:
            raise ValueError('cần thư viện pypdf (pip install pypdf) để đọc PDF')
    if duoi in ('.xlsx', '.xlsm'):
        return doc_xlsx(duong)
    raise ValueError('không đọc được định dạng %s' % duoi)
