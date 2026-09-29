# -*- coding: utf-8 -*-
"""Kiem the thuc, ky thuat trinh bay san pham .docx/.xlsx — DL-20260919-003.

Chuan: 20-Chuan-Chung/18-Chuan-The-Thuc-San-Pham.md — so do that cua 16 mau
KTC-Database/03-Templates(1) (.dotx/.xltx), NĐ 30/2020 Phu luc I, TB 597/TB-CĐKT (qua 897,
Checklist 05-Hinh-Thuc va 08-Quy-Uoc-Rieng-CDKT muc 4).

Do tren BYTE THAT cua tep (python-docx/openpyxl): phong va co chu duoc giai theo chuoi
run -> style ky tu -> style doan (ke ca base_style) -> docDefaults -> theme.
Chi PHAT HIEN va GOI Y muc — khong sua tep. He A (hanh chinh). Khong ap cho van ban Dang (he B).

Chay:  python 29-Cong-Cu/kiem_the_thuc.py <tep .docx|.xlsx> [...]
       python 29-Cong-Cu/kiem_the_thuc.py --tao <mau .dotx> <dich .docx>   (tao ban lam viec tu mau)
       python 29-Cong-Cu/kiem_the_thuc.py --khung <dich .docx> [TB|KH|BC|TTr|QĐ|GM|HD|CTr|BB]
           (khong doc duoc KTC-Database: tao tu khung dung tu mau 03A-Thong-bao — KHONG dung bang tieu de tay)
Ma thoat: 1 neu co goi y Muc 1 hoac Muc 2, nguoc lai 0.
"""
import collections
import io
import os
import re
import sys
import zipfile

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PHONG = "Times New Roman"
# (min, max) cm — NĐ 30 Phu luc I; mau 03-Templates(1) dung 2/2/3/2
LE = {"trên": (2.0, 2.5), "dưới": (2.0, 2.5), "trái": (3.0, 3.5), "phải": (1.5, 2.0)}
EPS = 0.05


# ------------------------------------------------------------------ mo tep mau
def _doi_kieu(du_lieu: bytes, tu: bytes, sang: bytes) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(du_lieu)) as zi, zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(tu, sang)
            zo.writestr(it, data)
    return b.getvalue()


def mo_mau(p):
    """Mo .dotx (hoac .docx) thanh docx.Document — python-docx khong tu mo .dotx."""
    import docx
    raw = open(p, "rb").read()
    if p.lower().endswith(".dotx"):
        raw = _doi_kieu(raw, b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
    return docx.Document(io.BytesIO(raw))


def tao_tu_mau(mau, dich):
    """Tao ban lam viec .docx/.xlsx tu mau .dotx/.xltx — giu nguyen style, le, bang the thuc."""
    raw = open(mau, "rb").read()
    if mau.lower().endswith(".dotx"):
        raw = _doi_kieu(raw, b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
    elif mau.lower().endswith(".xltx"):
        raw = _doi_kieu(raw, b"spreadsheetml.template.main+xml", b"spreadsheetml.sheet.main+xml")
    os.makedirs(os.path.dirname(os.path.abspath(dich)), exist_ok=True)
    open(dich, "wb").write(raw)
    return dich


# Khung the thuc VBHC dung tu mau 03A-Thong-bao (28/9/2026; thay khung ghep tu TB 1060 — hien "1" o trang 1,
# duong ke trich yeu lech) — dung khi KHONG doc duoc KTC-Database (Cowork, Chat,
# tai khoan thanh vien) thay cho dung bang tieu de bang tay (28/9/2026).
TEN_KHUNG = "Khung-the-thuc-VBHC.docx"
LOAI_VB = {"TB": "THÔNG BÁO", "KH": "KẾ HOẠCH", "BC": "BÁO CÁO", "TTr": "TỜ TRÌNH", "QĐ": "QUYẾT ĐỊNH",
           "GM": "GIẤY MỜI", "HD": "HƯỚNG DẪN", "CTr": "CHƯƠNG TRÌNH", "BB": "BIÊN BẢN"}


def tim_khung():
    g = os.path.dirname(os.path.abspath(__file__))
    for p in (os.path.join(g, "..", "assets"), os.path.join(g, "..", "skills", "the-thuc", "assets"),
              os.path.join(g, "..", "27-KTC-The-Thuc", "assets")):
        f = os.path.join(p, TEN_KHUNG)
        if os.path.isfile(f):
            return os.path.normpath(f)
    raise FileNotFoundError(f"Không thấy {TEN_KHUNG} (skill the-thuc/assets)")


def tao_khung(dich, loai="TB", nam=None):
    """Tao tep lam viec tu khung: dat loai van ban (ten loai, ky hieu Số: /<loai>-CĐKT) va nam."""
    import datetime
    import docx
    if loai not in LOAI_VB:
        raise ValueError(f"Loại văn bản {loai!r} — chọn một trong {', '.join(LOAI_VB)}")
    d = docx.Document(tim_khung())
    nam = nam or datetime.date.today().year
    for p in list(d.paragraphs) + list(_doan_trong_bang(d)):
        for r in p.runs:
            if r.text.strip() == "THÔNG BÁO":
                r.text = r.text.replace("THÔNG BÁO", LOAI_VB[loai])
            elif "/TB-CĐKT" in r.text:
                r.text = r.text.replace("/TB-CĐKT", f"/{loai}-CĐKT")
            elif "năm 2026" in r.text and "ngày" in p.text:
                r.text = r.text.replace("năm 2026", f"năm {nam}")
    os.makedirs(os.path.dirname(os.path.abspath(dich)), exist_ok=True)
    d.save(dich)
    return dich


# ------------------------------------------------------------------ giai phong/co chu
class GiaiDocx:
    def __init__(self, d):
        self.d = d
        root = d.styles.element
        rd = root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
        self.mac_dinh_phong = self._phong_rpr(rd)
        sz = rd.find(f"{W}sz") if rd is not None else None
        self.mac_dinh_co = int(sz.get(f"{W}val")) / 2 if sz is not None else 10.0
        self.theme = self._doc_theme()

    def _doc_theme(self):
        """Phong theme ma Word THAT SU hien thi. settings themeFontLang = vi-VN -> Word dung phong
        script="Viet" cua theme (vd. major Viet = Times New Roman) thay vi <a:latin> (Calibri Light).
        Bo qua buoc nay tung bao nham hang tram doan “Calibri Light” tren van ban da ban hanh (19/9/2026)."""
        viet = bool(re.search(r'<w:themeFontLang [^>]*w:val="vi', self.d.settings.element.xml))
        kq = {"major": None, "minor": None}
        for rel in self.d.part.rels.values():
            if rel.reltype.endswith("/theme"):
                x = rel.target_part.blob.decode("utf-8", "ignore")
                for loai, the in (("major", "majorFont"), ("minor", "minorFont")):
                    m = re.search(f"<a:{the}>(.*?)</a:{the}>", x, re.S)
                    if not m:
                        continue
                    khoi = m.group(1)
                    v = re.search(r'script="Viet" typeface="([^"]*)"', khoi) if viet else None
                    latin = re.search(r'<a:latin typeface="([^"]*)"', khoi)
                    kq[loai] = v.group(1) if v else (latin.group(1) if latin else None)
        return kq

    def _phong_rpr(self, rpr):
        if rpr is None:
            return None
        f = rpr.find(f"{W}rFonts")
        if f is None:
            return None
        if f.get(f"{W}ascii"):
            return f.get(f"{W}ascii")
        t = f.get(f"{W}asciiTheme")
        if t:
            return "theme:" + ("major" if "major" in t.lower() else "minor")
        return None

    def _xu_theme(self, v):
        if v and v.startswith("theme:"):
            return self.theme.get(v[6:]) or v
        return v

    def _style_chain(self, st):
        while st is not None:
            yield st
            st = st.base_style

    def phong(self, run, para):
        v = run.font.name or self._phong_rpr(run._element.rPr)
        if not v and run.style is not None:
            for s in self._style_chain(run.style):
                v = s.font.name or self._phong_rpr(s.element.rPr)
                if v:
                    break
        if not v:
            for s in self._style_chain(para.style):
                v = s.font.name or self._phong_rpr(s.element.rPr)
                if v:
                    break
        return self._xu_theme(v or self.mac_dinh_phong)

    def co(self, run, para):
        if run.font.size:
            return run.font.size.pt
        if run.style is not None:
            for s in self._style_chain(run.style):
                if s.font.size:
                    return s.font.size.pt
        for s in self._style_chain(para.style):
            if s.font.size:
                return s.font.size.pt
        return self.mac_dinh_co

    def can_le(self, para):
        if para.paragraph_format.alignment is not None:
            return para.paragraph_format.alignment
        for s in self._style_chain(para.style):
            if s.paragraph_format.alignment is not None:
                return s.paragraph_format.alignment
        return None


def _doan_trong_bang(d):
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p


def _co_duong_ke(o):
    """O bang co duong ke (hinh Draw, VML line hoac vien duoi doan)."""
    x = o._tc.xml
    return "<w:drawing" in x or "<v:line" in x or "<v:shape" in x or re.search(r"<w:pBdr>.*?<w:bottom ", x, re.S)


def _kiem_bang_tieu_de(d, g):
    """TT12-TT15: hinh hoc bang tieu de (28/9/2026 — TB do Cowork dung tay: bang 16 cm chia doi, quoc hieu va dia danh
    xuong dong, thieu duong ke duoi ten Truong; cac phep TT08-TT09 cu van bao 'dat').
    So do van ban da ban hanh (TB 1060, 1092, TB phuong an A1): bang 17,25-17,5 cm; cot trai 7,0-7,5; cot phai 9,75-10,3."""
    kq = []
    if not d.tables:
        return kq
    t = d.tables[0]
    o = [c for r in t.rows for c in r.cells]
    chu = lambda c: " ".join(p.text for p in c.paragraphs).upper()  # noqa: E731
    phai = [c for c in o if "CỘNG H" in chu(c) and "VIỆT NAM" in chu(c)]
    trai = [c for c in o if "TRƯỜNG CAO ĐẲNG KON TUM" in chu(c)]
    if not phai or not trai:
        return kq
    grid = [c.w.cm for c in t._tbl.tblGrid.gridCol_lst if c.w is not None]
    rong = lambda c: c.width.cm if c.width is not None else None  # noqa: E731
    w_trai = rong(trai[0]) or (grid[0] if grid else None)
    w_phai = rong(phai[0]) or (grid[-1] if grid else None)
    if w_trai and w_phai:
        # Quet 434 van ban kho 02, 04 (28/9/2026): van ban that co cot phai 9,13-9,28 cm van hien thi dung -> nguong 9,0
        if w_phai < 9.0 or w_trai > w_phai:
            kq.append((2, "TT12", f"Bảng tiêu đề chia cột {w_trai:.2f} + {w_phai:.2f} cm — quốc hiệu, dòng địa danh sẽ xuống "
                       "dòng. Văn bản đã ban hành: cột trái 7–7,5 cm, cột phải 9,75–10,3 cm (dựng bằng --khung)"))
        elif w_trai + w_phai < 16.0:
            kq.append((3, "TT12", f"Bảng tiêu đề rộng {w_trai + w_phai:.2f} cm — văn bản đã ban hành: 17,25–17,5 cm"))
    da_xet, doan_bang = set(), []  # ca bang: dong "Số", "ngày" thuong o hang 2 (khung, TB 1060)
    for c in o:
        if c._tc not in da_xet:
            da_xet.add(c._tc)
            doan_bang += c.paragraphs
    for p in doan_bang:
        t_up = p.text.strip().upper()
        rs = [r for r in p.runs if r.text.strip()]
        if not rs:
            continue
        if t_up.startswith(("UBND", "ỦY BAN NHÂN DÂN", "UỶ BAN NHÂN DÂN")) and any(r.bold for r in rs):
            kq.append((2, "TT13", "Tên cơ quan chủ quản “UBND TỈNH QUẢNG NGÃI” in đậm — NĐ 30: in hoa, đứng, không đậm"))
        if ", NGÀY" in t_up:
            if any(r.bold for r in rs):
                kq.append((2, "TT15", "Dòng địa danh, ngày tháng in đậm — NĐ 30: nghiêng, không đậm, cỡ 13–14"))
            if not all(r.italic or (p.style is not None and p.style.font.italic) for r in rs):
                kq.append((3, "TT15", "Dòng địa danh, ngày tháng không nghiêng — NĐ 30: chữ nghiêng"))
    # Duong ke thuong la hinh noi neo o o "Số"/"ngày" cung cot (TB 1092, 1056) — xet ca cot
    def cot(c):
        for r in t.rows:
            for j, x in enumerate(r.cells):
                if x._tc is c._tc:
                    return [rr.cells[j] for rr in t.rows if j < len(rr.cells)]
        return [c]
    if not any(_co_duong_ke(x) for x in cot(trai[0])):
        kq.append((2, "TT14", "Thiếu đường kẻ dưới tên cơ quan ban hành “TRƯỜNG CAO ĐẲNG KON TUM” "
                   "(NĐ 30: đường kẻ ngang, nét liền, bằng 1/3–1/2 dòng chữ)"))
    if not any(_co_duong_ke(x) for x in cot(phai[0])):
        kq.append((2, "TT14", "Thiếu đường kẻ dưới tiêu ngữ “Độc lập - Tự do - Hạnh phúc”"))
    return kq


# ------------------------------------------------------------------ anh xa bo quy tac 897
# Nguon (KHONG chep, chi anh xa ma kiem): KTC-Ra-Soat-897-v2-Cai-tien/references/Checklist/
#   01-The-Thuc.md muc 1 (9 thanh phan, quy dinh chung NĐ 30) · 08-Quy-Uoc-Rieng-CDKT.md muc 4 (TB 597 so cung),
#   muc 5.1 (KT./TL./TUQ.), muc 5.4 (khong hoc ham, hoc vi). 01 muc 6: sai co chu, kieu chu, vi tri -> Muc 2.
HOC_HAM = re.compile(r"^(GS|PGS|TS|ThS|BS|BSCKI+|CKI+|DS|KS|CN|NCS|TTƯT|NGƯT|NGND)\.?\s", re.I)


def _thuoc_tinh(r, p, ten):
    """bold/italic that: run -> style ky tu -> style doan (ca base_style)."""
    v = getattr(r.font, ten)
    if v is not None:
        return v
    for st in ([r.style] if r.style is not None else []) + [p.style]:
        while st is not None:
            x = getattr(st.font, ten)
            if x is not None:
                return x
            st = st.base_style
    return False


def _kieu(g, p):
    rs = [r for r in p.runs if r.text.strip()]
    if not rs:
        return None
    return ({g.co(r, p) for r in rs}, all(_thuoc_tinh(r, p, "bold") for r in rs),
            any(_thuoc_tinh(r, p, "bold") for r in rs), all(_thuoc_tinh(r, p, "italic") for r in rs))


def _kiem_thanh_phan_897(d, g):
    """TT17 (co, kieu chu tung thanh phan), TT18 (duong ke duoi trich yeu), TT19 (ky thay, hoc ham)."""
    kq = []

    def sai(ten, p, co=None, dam=None, nghieng=None, nguon="01 mục 1 · 08 mục 4"):
        k = _kieu(g, p)
        if k is None:
            return
        cs, dam_het, dam_co, ngh = k
        loi = []
        if co is not None and cs != {float(co)}:
            loi.append(f"cỡ {'/'.join(f'{c:g}' for c in sorted(cs))} (chuẩn {co})")
        if dam is True and not dam_het:
            loi.append("chưa đậm")
        if dam is False and dam_co:
            loi.append("không được đậm")
        if nghieng is True and not ngh:
            loi.append("chưa nghiêng")
        if loi:
            kq.append((2, "TT17", f"{ten}: {', '.join(loi)} — 897 Checklist {nguon}"))

    # Phan dau (bang dau tien)
    if d.tables:
        for p in _doan_trong_bang_dau(d):
            t = p.text.strip()
            tu = t.upper()
            if tu.startswith(("UBND", "ỦY BAN NHÂN DÂN", "UỶ BAN NHÂN DÂN")):
                sai("Tên cơ quan chủ quản", p, co=13)
            elif tu == "TRƯỜNG CAO ĐẲNG KON TUM":
                sai("Tên đơn vị ban hành", p, co=13, dam=True)
            elif re.match(r"Số\s*:", t):
                sai("Số, ký hiệu", p, co=13, dam=False)
                m = re.match(r"Số\s*:(\s*)/", t)
                if m and len(m.group(1)) < 6:
                    kq.append((3, "TT17", f"Sau “Số:” chỉ để trống {len(m.group(1))} ký tự — tối thiểu 6 (897 Checklist 01 mục 1.4)"))
                if re.match(r"Số\s*:\s*[1-9]/", t):
                    kq.append((3, "TT17", "Số nhỏ hơn 10 phải thêm số 0 phía trước (897 Checklist 01 mục 1.4)"))
            elif ", ngày" in t:
                sai("Địa danh, ngày tháng", p, co=14)

    # Ten loai, trich yeu, can cu (ngoai bang)
    ps = [p for p in d.paragraphs]
    i_loai = next((i for i, p in enumerate(ps[:12]) if p.text.strip() in LOAI_VB.values()), None)
    if i_loai is not None:
        sai("Tên loại văn bản", ps[i_loai], co=14, dam=True)
        ty, co_ke = [], False
        for p in ps[i_loai + 1:i_loai + 6]:
            t = p.text.strip()
            x = p._p.xml
            co_ke |= "<w:drawing" in x or "<w:pict" in x or bool(re.search(r"<w:pBdr>.*?<w:bottom ", x, re.S))
            if not t:
                if ty:
                    break
                continue
            if t.startswith("Căn cứ") or (p.paragraph_format.first_line_indent or 0) > 0:
                break
            ty.append(p)
        for p in ty:
            sai("Trích yếu", p, co=14, dam=True)
        if ty and not co_ke:
            kq.append((2, "TT18", "Thiếu đường kẻ ngang dưới trích yếu (dài 1/3–1/2 dòng chữ) — 897 Checklist 01 mục 1.6"))
    for p in ps:
        if p.text.strip().startswith("Căn cứ"):
            sai("Căn cứ", p, co=14, nghieng=True)

    # Nguoi ky, noi nhan (bang co "Nơi nhận")
    for t in d.tables[1:]:
        o = [c for r in t.rows for c in r.cells]
        if not any(p.text.strip().startswith("Nơi nhận") for c in o for p in c.paragraphs):
            continue
        # khoi chu ky co the chia 2 hang (mau 03A: chuc vu hang 1, ho ten hang 2) -> gom moi o khong phai "Noi nhan"
        ky, da = [], set()
        for c in o:
            if c._tc in da:
                continue
            da.add(c._tc)
            dps = [p for p in c.paragraphs if p.text.strip()]
            if dps and not dps[0].text.strip().startswith("Nơi nhận"):
                ky += dps
        for c in [c for c in o if c.paragraphs and c.paragraphs[0].text.strip().startswith("Nơi nhận")][:1] + [None]:
            if c is None:
                dps = ky
            else:
                dps = [p for p in c.paragraphs if p.text.strip()]
            if not dps:
                continue
            if dps[0].text.strip().startswith("Nơi nhận"):
                sai("Từ “Nơi nhận”", dps[0], co=12, dam=True, nghieng=True)
                for p in dps[1:]:
                    sai("Danh sách nơi nhận", p, co=11)  # khong xet dam: mau 09 dam rieng dau "-" dong Luu
                if not re.match(r"-\s*Lưu\s*:\s*VT", dps[-1].text.strip()):
                    kq.append((3, "TT17", "Dòng cuối nơi nhận phải là “- Lưu: VT, <đơn vị soạn thảo>.” (897 Checklist 01 mục 1.9)"))
                continue
            chuc = [p for p in dps if p.text.strip().isupper()]
            for p in chuc:
                sai("Quyền hạn, chức vụ người ký", p, co=14, dam=True)
                if re.match(r"(K/T|T/L|TU/Q)", p.text.strip()):
                    kq.append((2, "TT19", f"“{p.text.strip()[:12]}” — hệ hành chính dùng KT./TL./TUQ. có dấu chấm (897 Checklist 08 mục 5.1)"))
            if chuc and dps[-1] not in chuc:
                ten = dps[-1]
                sai("Họ tên người ký", ten, co=14, dam=True)
                if HOC_HAM.match(ten.text.strip()):
                    kq.append((2, "TT19", f"Ghi học hàm, học vị trước họ tên người ký “{ten.text.strip()}” (897 Checklist 01 mục 1.8)"))
    return kq


def _doan_trong_bang_dau(d):
    da, kq = set(), []
    for r in d.tables[0].rows:
        for c in r.cells:
            if c._tc not in da:
                da.add(c._tc)
                kq += c.paragraphs
    return kq


# ------------------------------------------------------------------ kiem docx
def kiem_docx(d):
    """Tra ve list[(muc, ma, mo_ta)]."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import unicodedata
    # Mau 03-Templates(1)/03A luu chu dang NFD ("â" + dau nang to hop) -> moi phep so chu ("Nơi nhận", "Căn cứ"...)
    # truot va phep kiem im lang bo qua (28/9/2026). Chuan hoa NFC trong bo nho truoc khi do (khong ghi tep).
    for t in d.element.body.iter(f"{W}t"):
        if t.text:
            t.text = unicodedata.normalize("NFC", t.text)
    kq = []
    g = GiaiDocx(d)

    # TT01 kho giay, TT02 le — moi section
    for i, s in enumerate(d.sections, 1):
        w, h = s.page_width.cm, s.page_height.cm
        ngang = w > h
        a4 = abs(min(w, h) - 21.0) <= 0.1 and abs(max(w, h) - 29.7) <= 0.1
        if not a4:
            kq.append((2, "TT01", f"Section {i}: khổ {w:.1f}×{h:.1f} cm — phải A4 (21×29,7)"))
        le = {"trên": s.top_margin.cm, "dưới": s.bottom_margin.cm,
              "trái": s.left_margin.cm, "phải": s.right_margin.cm}
        if not ngang:  # trang ngang (phu luc bang) khong ap le doc
            sai = [f"{k} {round(v, 2)}" for k, v in le.items()
                   if not (LE[k][0] - EPS <= round(v, 2) <= LE[k][1] + EPS)]
            if sai:
                kq.append((2, "TT02", f"Section {i}: lề ngoài khoảng NĐ 30 ({', '.join(sai)} cm). "
                           "Mẫu 03-Templates(1): trên 2 · dưới 2 · trái 3 · phải 2 cm"))

    doan = list(d.paragraphs)
    tat_ca = doan + list(_doan_trong_bang(d))

    # TT03 phong chu, TT05 mau chu
    sai_phong = collections.Counter()
    mau_chu = 0
    for p in tat_ca:
        for r in p.runs:
            if not r.text.strip():
                continue
            f = g.phong(r, p)
            if f != PHONG:
                sai_phong[f or "(không xác định)"] += 1
            c = r.font.color
            if c is not None and c.type is not None and c.rgb is not None and str(c.rgb) not in ("000000",):
                mau_chu += 1
    # Bien the ten cua chinh TNR ("Times New Roman Bold", "TimesNewRomanPSMT" — thuong do chep tu PDF):
    # hien thi van la TNR tren may co phong chuan, nhung may khac co the thay phong -> Muc 3.
    bien_the = {k: v for k, v in sai_phong.items() if re.sub(r"[\s\-]", "", k or "").lower().startswith("timesnewroman")}
    khac = {k: v for k, v in sai_phong.items() if k not in bien_the}
    if khac:
        kq.append((2, "TT03", f"{sum(khac.values())} đoạn chữ không phải {PHONG}: "
                   + ", ".join(f"{k}×{v}" for k, v in collections.Counter(khac).most_common(4))
                   + (" — .VnTime là phông TCVN3 cũ, phải chuyển Unicode" if any(".Vn" in (k or "") for k in khac) else "")))
    if bien_the:
        kq.append((3, "TT03b", f"{sum(bien_the.values())} đoạn chữ dùng biến thể tên phông "
                   + ", ".join(bien_the) + f" — đặt lại đúng tên “{PHONG}”"))
    if mau_chu:
        kq.append((3, "TT05", f"{mau_chu} đoạn chữ có màu khác đen (NĐ 30: màu đen)"))

    # Doan noi dung: ngoai bang, dai >= 60 ky tu
    noi_dung = [p for p in doan if len(p.text.strip()) >= 60]
    co = collections.Counter()
    for p in noi_dung:
        for r in p.runs:
            if r.text.strip():
                co[g.co(r, p)] += len(r.text)
    if co:
        chinh = co.most_common(1)[0][0]
        if chinh != 14:
            kq.append((2, "TT04", f"Cỡ chữ phần nội dung chủ yếu là {chinh:g} — TB 597: cỡ 14"))
        lech = {k: v for k, v in co.items() if k != chinh and v > 0.05 * sum(co.values())}
        if lech:
            kq.append((3, "TT04b", "Cỡ chữ lời văn không thống nhất: "
                       + ", ".join(f"cỡ {k:g} ({v} ký tự)" for k, v in lech.items())))
        khong_deu = sum(1 for p in noi_dung if g.can_le(p) not in (WD_ALIGN_PARAGRAPH.JUSTIFY, None)
                        and g.can_le(p) != WD_ALIGN_PARAGRAPH.CENTER)
        if khong_deu > 0.2 * len(noi_dung):
            kq.append((3, "TT06", f"{khong_deu}/{len(noi_dung)} đoạn nội dung không dàn đều hai lề"))
        gian = [p.paragraph_format.line_spacing for p in noi_dung]
        qua = sum(1 for x in gian if isinstance(x, float) and x > 1.5 + 1e-6)
        hep = sum(1 for x in gian if x is not None and not isinstance(x, float) and x.pt < 15 - 0.1)
        if qua or hep:
            kq.append((3, "TT07", f"Cách dòng ngoài khoảng single/15pt → 1,5 lines: {qua} đoạn >1,5; "
                       f"{hep} đoạn exactly <15pt"))

    # TT08 phan dau the thuc (bang dau tien)
    dau = "\n".join(p.text for p in (list(_doan_trong_bang(d))[:12] if d.tables else doan[:8]))
    dau_up = dau.upper()
    la_vb_hanh_chinh = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in dau_up or "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM" in dau_up
    if la_vb_hanh_chinh:
        if "TRƯỜNG CAO ĐẲNG KON TUM" not in dau_up:
            kq.append((2, "TT08", "Phần đầu thiếu tên đơn vị ban hành “TRƯỜNG CAO ĐẲNG KON TUM”"))
        if "UBND TỈNH KON TUM" in dau_up or "ỦY BAN NHÂN DÂN TỈNH KON TUM" in dau_up:
            kq.append((2, "TT08", "Cơ quan chủ quản ghi “UBND TỈNH KON TUM” — văn bản mới dùng “UBND TỈNH QUẢNG NGÃI”"))
        elif "QUẢNG NGÃI" not in dau_up:
            kq.append((2, "TT08", "Phần đầu thiếu cơ quan chủ quản “UBND TỈNH QUẢNG NGÃI”"))
        for p in _doan_trong_bang(d):
            t = p.text.strip().upper()
            rs = [r for r in p.runs if r.text.strip()]
            if not rs:
                continue
            if t.startswith("CỘNG H") and "VIỆT NAM" in t:
                cs = {g.co(r, p) for r in rs}
                if cs != {13.0}:
                    kq.append((2, "TT09", f"Quốc hiệu cỡ {sorted(cs)} — TB 597: cỡ 13, đậm"))
            if t.startswith("ĐỘC LẬP"):
                cs = {g.co(r, p) for r in rs}
                if cs != {14.0}:
                    kq.append((2, "TT09", f"Tiêu ngữ cỡ {sorted(cs)} — TB 597: cỡ 14, đậm"))
                if any(r.font.underline for r in rs):
                    kq.append((3, "TT09", "Tiêu ngữ dùng Underline — TB 597: kẻ đường bằng Draw"))
        kq += _kiem_bang_tieu_de(d, g)
        kq += _kiem_thanh_phan_897(d, g)

    # TT16 doan "Can cu" ngoai bang phai thut dau dong nhu doan noi dung (TB 1060, 1092 da ban hanh: 1,25 cm)
    can_cu = [p for p in doan if p.text.strip().startswith("Căn cứ")]
    thut = [p for p in noi_dung if not p.text.strip().startswith("Căn cứ")
            and (p.paragraph_format.first_line_indent or 0) > 0]
    khong_thut = [p for p in can_cu if not (p.paragraph_format.first_line_indent or 0) > 0]
    if khong_thut and thut:
        kq.append((3, "TT16", f"{len(khong_thut)} đoạn “Căn cứ…” không thụt đầu dòng trong khi đoạn nội dung có thụt — "
                   "văn bản đã ban hành: căn cứ thụt 1,25 cm, nghiêng"))

    # TT10 "Noi nhan"
    for p in tat_ca:
        if p.text.strip().startswith("Nơi nhận"):
            rs = [r for r in p.runs if r.text.strip()]
            cs = {g.co(r, p) for r in rs}
            if cs and cs != {12.0}:
                kq.append((3, "TT10", f"“Nơi nhận” cỡ {sorted(cs)} — TB 597: cỡ 12, nghiêng, đậm"))
            break

    # TT11b so trang hien o trang 1 (897 Checklist 05 muc 4.3) — ban TB 28/9/2026 dung tu ghep TB 1060 hien "1"
    #   (ban do la CHU "1" co dinh trong header trang dau, khong phai truong PAGE) -> xet ca truong PAGE lan chu
    s0 = d.sections[0]
    tp = s0._sectPr.find(f"{W}titlePg")
    co_titlepg = tp is not None and tp.get(f"{W}val") not in ("0", "false")
    dau_trang1 = s0.first_page_header if co_titlepg else s0.header
    if not dau_trang1.is_linked_to_previous or co_titlepg:
        x1 = dau_trang1._element.xml
        chu1 = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", x1)).strip()
        if "PAGE" in x1 or re.search(r"\d", chu1):
            kq.append((2, "TT11b", f"Trang thứ nhất hiện số trang/chữ ở đầu trang (“{chu1 or 'PAGE'}”) — 897 Checklist "
                       "05 mục 4.3: không hiển thị số trang thứ nhất"))

    # TT11 so trang: co truong PAGE o header
    if len(noi_dung) > 25:
        co_page = any("PAGE" in (s.header._element.xml if s.header is not None else "") for s in d.sections)
        if not co_page:
            kq.append((3, "TT11", "Văn bản nhiều trang nhưng không có số trang (canh giữa lề trên, ẩn trang 1)"))
    kq.sort(key=lambda x: x[0])
    return kq


# ------------------------------------------------------------------ kiem xlsx
def kiem_xlsx(wb):
    kq = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or ws.max_row < 2:
            continue
        ps = ws.page_setup
        if ps.paperSize not in (None, 9, "9"):
            kq.append((2, "TX01", f"[{ws.title}] khổ giấy in mã {ps.paperSize} — phải A4 (mã 9)"))
        f = collections.Counter()
        co = collections.Counter()
        for row in ws.iter_rows(max_row=min(ws.max_row, 400)):
            for c in row:
                if c.value is not None and str(c.value).strip():
                    f[c.font.name] += 1
                    co[c.font.sz] += 1
        sai = {k: v for k, v in f.items() if k != PHONG}
        if sai:
            kq.append((2, "TX02", f"[{ws.title}] {sum(sai.values())} ô không phải {PHONG}: "
                       + ", ".join(f"{k}×{v}" for k, v in list(sai.items())[:4])))
        if co:
            chinh = co.most_common(1)[0][0]
            if chinh is not None and not (12 <= chinh <= 14):
                kq.append((3, "TX03", f"[{ws.title}] cỡ chữ chủ yếu {chinh:g} — mẫu 05B dùng 12–14"))
        if ws.max_column > 6 and ps.orientation != "landscape" and not ps.fitToWidth:
            kq.append((4, "TX04", f"[{ws.title}] bảng {ws.max_column} cột in dọc, không đặt vừa trang — "
                       "mẫu 05B: A4 ngang"))
    kq.sort(key=lambda x: x[0])
    return kq


def kiem_tep(p):
    if p.lower().endswith((".docx", ".dotx")):
        return kiem_docx(mo_mau(p))
    if p.lower().endswith((".xlsx", ".xltx")):
        import warnings
        import openpyxl
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return kiem_xlsx(openpyxl.load_workbook(p))
    raise ValueError(f"Không kiểm được loại tệp: {p}")


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    if argv[0] == "--tao":
        print("Đã tạo", tao_tu_mau(argv[1], argv[2]))
        return 0
    if argv[0] == "--khung":
        print("Đã tạo", tao_khung(argv[1], argv[2] if len(argv) > 2 else "TB"))
        return 0
    xau = False
    for p in argv:
        kq = kiem_tep(p)
        print(f"\n=== {p} — {len(kq)} gợi ý ===")
        if not kq:
            print("  ✓ Đạt chuẩn thể thức theo các phép kiểm TT/TX.")
        for muc, ma, mo_ta in kq:
            print(f"  [Mức {muc}] {ma}: {mo_ta}")
        xau |= any(x[0] <= 2 for x in kq)
    return 1 if xau else 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
