# -*- coding: utf-8 -*-
"""Cap nhat CSDL theo doi nhiem vu tu ket qua Tra cuu van ban cong iOffice Quang Ngai.

Thay phan XU LY cua quy trinh "Theo doi Ke hoach & Ket luan giao ban — iOffice" (buoc 3, 5, 6, 7
va phan doi ten/ghi nhan file cua buoc 4). Script KHONG truy cap cong, KHONG dang nhap: nguoi dung
tra cuu iOffice bang tay roi luu ket qua (Excel xuat ra, CSV, hoac trang HTML luu bang Ctrl+S).

Luat phan loai: import DUNG 10-Dau-Vao/01-Dau-Moi-Nop/phanloai.py — khong chep, khong sua luat.

Chay:
  python 29-Cong-Cu/cap_nhat_csdl.py --thu <ket-qua-KH> <ket-qua-TB>   # xem truoc, khong ghi gi
  python 29-Cong-Cu/cap_nhat_csdl.py <ket-qua-KH> <ket-qua-TB>         # ghi that (co sao luu)
  python 29-Cong-Cu/cap_nhat_csdl.py --nhan-file "C:/Users/nnqp2/Downloads"   # nhan file da tai
  python 29-Cong-Cu/cap_nhat_csdl.py                                    # chi doi chieu, khop bieu mau

Ma thoat: 0 binh thuong; 1 neu loi dau vao (thieu cot, sai cau truc CSDL).
"""
import argparse
import csv
import difflib
import html
import os
import re
import shutil
import sys
import unicodedata
from datetime import datetime

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
THU_MUC = os.path.join(GOC, '10-Dau-Vao', '01-Dau-Moi-Nop')
TEN_CSDL = 'CSDL-Theo-doi-nhiem-vu_20260929_v1.xlsx'
BIEU_MAU = os.path.join('2026-09', 'QLĐT.xlsx')
THU_MUC_FILE = '2026-09'
NGAY_TU = datetime(2026, 9, 28)

COT_CSDL = ['STT', 'Số ký hiệu', 'Loại', 'Ngày ban hành', 'Trích yếu', 'Đơn vị/người xử lý chính',
            'Độ khẩn', 'Nhóm', 'Lý do phân loại', 'Trục TB817', 'Nhiệm vụ tương ứng trong biểu mẫu',
            'Sản phẩm', 'Tình trạng', 'File đã tải', 'Doc ID', 'Ngày ghi nhận']
COT_NHAT_KY = ['Thời điểm chạy', 'Số VB quét', 'VB mới', 'Nhiệm vụ', 'Vận hành', 'Cần xem lại',
               'File đã tải', 'Ghi chú']
MAU = {'Nhiệm vụ': 'E2EFDA', 'Vận hành đào tạo': 'F2F2F2', 'Cần xem lại': 'FFF2CC',
       'Chờ phân loại': 'DDEBF7'}
VANG = 'FFF2CC'
# Nguong khop bieu mau — co y dat cao: thieu con hon khop sai
DIEM_KHOP = 0.55
CACH_BIET = 0.15

sys.path.insert(0, THU_MUC)
from phanloai import phan_loai  # noqa: E402


def khong_dau(s):
    s = unicodedata.normalize('NFD', str(s or ''))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return s.replace('đ', 'd').replace('Đ', 'D').lower()


def khoa_so(so):
    """Khoa chong trung so ky hieu: bo khoang trang, khong dau, chu hoa."""
    return re.sub(r'\s+', '', khong_dau(so)).upper()


def doc_ngay(v):
    if isinstance(v, datetime):
        return v
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(v or ''))
    return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def loai_vb(so):
    s = khoa_so(so)
    return 'TB' if '/TB-' in s else 'KH' if '/KH-' in s else ''


# ---------- Doc dau vao ----------

TRUONG = {  # ten truong -> cac dau hieu ten cot (khong dau)
    'trich_yeu': ['trich yeu'],
    'so_di': ['so di'],
    'so_ky_hieu': ['so ky hieu', 'so, ky hieu', 'so/ky hieu'],
    'ngay': ['ngay van ban', 'ngay ban hanh', 'ngay vb'],
    'do_khan': ['do khan'],
    'dv_ban_hanh': ['don vi ban hanh', 'co quan ban hanh'],
    'xu_ly': ['xu ly chinh', 'nguoi xu ly'],
    'doc_id': ['doc id', 'docid', 'ma van ban'],
    'ten_file': ['ten file', 'file dinh kem'],  # ten tep goc, nhieu tep cach nhau bang ' | '
}
BAT_BUOC = ['trich_yeu', 'so_ky_hieu', 'ngay']


def doc_html(duong_dan):
    """Trang ket qua luu tu iOffice: moi tr co onclick=showDocDetail(<id>,...)."""
    with open(duong_dan, encoding='utf-8', errors='replace') as f:
        noi_dung = f.read()
    ds = []
    for m in re.finditer(r'<tr\b([^>]*)>(.*?)</tr>', noi_dung, re.S | re.I):
        id_m = re.search(r'showDocDetail\(\s*[\'"]?(\d+)', m.group(1))
        if not id_m:
            continue
        td = [html.unescape(re.sub(r'<[^>]+>', ' ', t)) for t in
              re.findall(r'<td\b[^>]*>(.*?)</td>', m.group(2), re.S | re.I)]
        td = [re.sub(r'\s+', ' ', t).strip() for t in td]
        if len(td) < 13:
            continue
        ds.append({'doc_id': id_m.group(1), 'trich_yeu': td[2], 'so_di': td[3], 'so_ky_hieu': td[4],
                   'ngay': td[7], 'do_khan': td[10], 'dv_ban_hanh': td[11], 'xu_ly': td[12]})
    return ds


def doc_bang(hang):
    """Bang Excel/CSV: tim dong tieu de theo ten cot, khong theo vi tri."""
    hang = [list(h) for h in hang]
    for i, h in enumerate(hang[:15]):
        ten = [khong_dau(c) for c in h]
        cot = {}
        for truong, dau_hieu in TRUONG.items():
            for j, t in enumerate(ten):
                if any(d in t for d in dau_hieu):
                    cot.setdefault(truong, j)
        if all(t in cot for t in BAT_BUOC):
            ds = []
            for h2 in hang[i + 1:]:
                ban_ghi = {t: (h2[j] if j < len(h2) else None) for t, j in cot.items()}
                if ban_ghi.get('so_ky_hieu'):
                    ds.append({k: str(v).strip() if v is not None and not isinstance(v, datetime)
                               else v for k, v in ban_ghi.items()})
            return ds
    raise ValueError('không tìm thấy dòng tiêu đề có đủ cột Trích yếu, Số ký hiệu, Ngày văn bản')


def doc_dau_vao(duong_dan):
    duoi = os.path.splitext(duong_dan)[1].lower()
    if duoi in ('.html', '.htm'):
        return doc_html(duong_dan)
    if duoi == '.csv':
        with open(duong_dan, encoding='utf-8-sig', newline='') as f:
            return doc_bang(csv.reader(f))
    if duoi in ('.xlsx', '.xlsm'):
        wb = openpyxl.load_workbook(duong_dan, data_only=True, read_only=True)
        return doc_bang(wb.worksheets[0].iter_rows(values_only=True))
    raise ValueError('định dạng không hỗ trợ: %s (dùng .xlsx, .csv, .html)' % duoi)


# ---------- CSDL ----------

def dong_tieu_de(ws):
    for r in range(1, 12):
        if str(ws.cell(r, 1).value or '').strip() == 'STT':
            return r
    raise ValueError('sheet "%s" không có dòng tiêu đề bắt đầu bằng "STT"' % ws.title)


def chuan_hoa_cau_truc(wb, ghi_chu):
    """Bo sung cot 'Loai' va doi ten cot nhat ky (da duoc nguoi dung dong y 25/10/2026)."""
    ws = wb['CSDL']
    r0 = dong_tieu_de(ws)
    ten = [ws.cell(r0, c).value for c in range(1, ws.max_column + 1)]
    if 'Loại' not in ten:
        vi_tri = ten.index('Số ký hiệu') + 2
        gop = [g for g in ws.merged_cells.ranges if g.min_row < r0]
        for g in gop:
            ws.unmerge_cells(str(g))
        ws.insert_cols(vi_tri)
        for g in gop:
            ws.merge_cells(start_row=g.min_row, start_column=g.min_col,
                           end_row=g.max_row, end_column=g.max_col + 1)
        mau_td = ws.cell(r0, vi_tri - 1)
        o = ws.cell(r0, vi_tri, 'Loại')
        o._style = mau_td._style
        # Dich do rong cot sang phai mot cot, cot moi rong 7
        rong = {k: v.width for k, v in ws.column_dimensions.items()}
        for c in range(ws.max_column, vi_tri, -1):
            l = openpyxl.utils.get_column_letter(c)
            truoc = openpyxl.utils.get_column_letter(c - 1)
            if truoc in rong:
                ws.column_dimensions[l].width = rong[truoc]
        ws.column_dimensions[openpyxl.utils.get_column_letter(vi_tri)].width = 7
        for r in range(r0 + 1, ws.max_row + 1):
            so = ws.cell(r, vi_tri - 1).value
            if so:
                o = ws.cell(r, vi_tri, loai_vb(so))
                o._style = ws.cell(r, vi_tri - 1)._style
        ws.auto_filter.ref = 'A%d:%s%d' % (r0, openpyxl.utils.get_column_letter(ws.max_column),
                                           ws.max_row)
        ghi_chu.append('Chuẩn hóa: thêm cột "Loại"')
    for r in range(1, r0):
        v = ws.cell(r, 1).value
        if v and 'Chỉ Kế hoạch (KH-CĐKT)' in str(v):
            ws.cell(r, 1).value = str(v).replace('Chỉ Kế hoạch (KH-CĐKT)',
                                                 'Kế hoạch (KH-CĐKT) và Thông báo kết luận (TB-CĐKT)')
    nk = wb['Nhật ký chạy']
    if [c.value for c in nk[1]][:8] != COT_NHAT_KY:
        for j, t in enumerate(COT_NHAT_KY, 1):
            nk.cell(1, j).value = t
        ghi_chu.append('Chuẩn hóa: đổi tên cột Nhật ký')
    ten = [ws.cell(r0, c).value for c in range(1, ws.max_column + 1)]
    thieu = [t for t in COT_CSDL if t not in ten]
    if thieu:
        raise ValueError('sheet CSDL thiếu cột: ' + ', '.join(thieu))
    return r0, {t: ten.index(t) + 1 for t in COT_CSDL}


def cac_dong(ws, r0, cot):
    for r in range(r0 + 1, ws.max_row + 1):
        if ws.cell(r, cot['Số ký hiệu']).value:
            yield r


def mo_ta_ngan(trich_yeu, so_tu=8):
    s = khong_dau(trich_yeu)
    s = re.sub(r'^(ke hoach|thong bao)\s+', '', s)
    tu = re.findall(r'[a-z0-9]+', s)[:so_tu]
    return '-'.join(t.capitalize() if i == 0 else t for i, t in enumerate(tu))


def ten_file_chuan(so, trich_yeu, duoi):
    m = re.match(r'\s*(\d+)', str(so))
    return '%s-%s-CĐKT_%s%s' % (m.group(1) if m else 'x', loai_vb(so), mo_ta_ngan(trich_yeu), duoi)


def tim_file(thu_muc, so):
    """File trong thu muc co ten bat dau bang '<so>-<KH|TB>' hoac chua '<so>/<KH|TB>'-dang tuong tu."""
    m = re.match(r'\s*(\d+)', str(so))
    if not m or not os.path.isdir(thu_muc):
        return []
    mau = re.compile(r'(^|\D)%s\D{0,3}%s' % (m.group(1), loai_vb(so).lower()))
    return sorted(f for f in os.listdir(thu_muc)
                  if os.path.splitext(f)[1].lower() in ('.docx', '.doc', '.pdf')
                  and not f.startswith('~$') and mau.search(khong_dau(f)))


UU_TIEN = {'.docx': 0, '.doc': 1, '.pdf': 2}


def chon_chinh(ds):
    """Tep chinh de trich noi dung: .docx truoc, roi .doc, .pdf; tep phu (bao cao ra soat 897, phieu,
    to trinh, phu luc) xep sau; cung loai thi ten ngan nhat."""
    phu = re.compile(r'ra.?soat|897|phieu|to.?trinh|phu.?luc|^bc[-_ ]')
    return min(ds, key=lambda f: (UU_TIEN.get(os.path.splitext(f)[1].lower(), 9),
                                  bool(phu.search(khong_dau(f.split('__')[-1]))), len(f), f))


# ---------- Khop bieu mau QLDT ----------

def tu_khoa(s):
    bo = {'va', 'cua', 'cho', 'cac', 'nam', 'tai', 'theo', 've', 'trong', 'ke', 'hoach', 'thuc',
          'hien', 'trien', 'khai', 'to', 'chuc', 'nhung', 'voi', 'thong', 'bao', 'the', 'nhu'}
    return {t for t in re.findall(r'[a-z0-9]+', khong_dau(s)) if t not in bo and len(t) > 1}


def diem(a, b):
    ta, tb = tu_khoa(a), tu_khoa(b)
    if not ta or not tb:
        return 0.0
    jac = len(ta & tb) / len(ta | tb)
    seq = difflib.SequenceMatcher(None, khong_dau(a), khong_dau(b)).ratio()
    return 0.5 * jac + 0.5 * seq


def dong_bieu_mau(wb):
    ds = []
    for ten_sh in ('BC tháng 9', 'KH tháng 10'):
        if ten_sh not in wb.sheetnames:
            continue
        ws = wb[ten_sh]
        r0 = cot_nd = cot_sp = None
        for r in range(1, 12):
            ten = [khong_dau(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            if any('noi dung cong viec' in t for t in ten):
                r0 = r
                cot_nd = next(i for i, t in enumerate(ten, 1) if 'noi dung cong viec' in t)
                cot_sp = next((i for i, t in enumerate(ten, 1) if t.startswith('san pham')), None)
                break
        if not r0 or not cot_sp:
            continue
        for r in range(r0 + 1, ws.max_row + 1):
            tt, nd = ws.cell(r, 1).value, ws.cell(r, cot_nd).value
            # Chi dong nhiem vu cu the (TT dang 1.1, 2.3); bo dong Truc, dong muc La Ma
            if nd and re.match(r'^\d+\.\d+$', str(tt).strip()):
                ds.append((ten_sh, r, cot_sp, str(nd)))
    return ds


def khop(trich_yeu, ung_vien):
    xep = sorted(((diem(trich_yeu, u[3]), u) for u in ung_vien), key=lambda x: -x[0])
    if not xep or xep[0][0] < DIEM_KHOP:
        return None, xep[0][0] if xep else 0
    if len(xep) > 1 and xep[0][0] - xep[1][0] < CACH_BIET:
        return None, xep[0][0]
    return xep[0][1], xep[0][0]


# ---------- Chuong trinh chinh ----------

def main():
    ap = argparse.ArgumentParser(description='Cập nhật CSDL theo dõi nhiệm vụ từ kết quả tra cứu iOffice')
    ap.add_argument('dau_vao', nargs='*', help='Tệp kết quả tra cứu (.xlsx/.csv/.html)')
    ap.add_argument('--thu', action='store_true', help='Chạy thử: in kết quả, không ghi tệp nào')
    ap.add_argument('--nhan-file', metavar='THU_MUC', help='Thư mục tệp đã tải (vd. Downloads)')
    ap.add_argument('--thu-muc', default=THU_MUC, help='Thư mục chứa CSDL (mặc định 01-Dau-Moi-Nop)')
    a = ap.parse_args()

    bay_gio = datetime.now()
    duong_csdl = os.path.join(a.thu_muc, TEN_CSDL)
    thu_muc_file = os.path.join(a.thu_muc, THU_MUC_FILE)
    ghi_chu, can_xu_ly = [], []

    wb = openpyxl.load_workbook(duong_csdl)
    try:
        r0, cot = chuan_hoa_cau_truc(wb, ghi_chu)
    except ValueError as e:
        print('LỖI CSDL:', e)
        return 1
    ws = wb['CSDL']
    da_co = {khoa_so(ws.cell(r, cot['Số ký hiệu']).value) for r in cac_dong(ws, r0, cot)}

    # Buoc 2-3: doc, loc, chong trung, phan loai
    so_quet, moi, ten_goc = 0, [], {}
    for p in a.dau_vao:
        try:
            ban_ghi = doc_dau_vao(p)
        except (ValueError, OSError) as e:
            print('LỖI ĐẦU VÀO %s: %s' % (p, e))
            return 1
        so_quet += len(ban_ghi)
        for b in ban_ghi:
            if b.get('ten_file'):
                ten_goc[khoa_so(b['so_ky_hieu'])] = [t.strip() for t in str(b['ten_file']).split('|')
                                                   if t.strip()]
            ngay = doc_ngay(b.get('ngay'))
            k = khoa_so(b['so_ky_hieu'])
            if not ngay or ngay < NGAY_TU or k in da_co or not loai_vb(b['so_ky_hieu']):
                continue
            nhom, ly_do = phan_loai(b['trich_yeu'], b['so_ky_hieu'])
            da_co.add(k)
            if nhom == 'Không theo dõi':
                continue
            b.update(ngay=ngay, nhom=nhom, ly_do=ly_do)
            moi.append(b)
    moi.sort(key=lambda b: (b['ngay'], b['so_ky_hieu']))

    # Phan loai lai dong 'Can xem lai' theo luat hien hanh (luat duoc bo sung dan). Chi dong may
    # dat 'Can xem lai' — dong nguoi da doi nhom thi khong con o nhom nay nen khong bi dong toi.
    phan_lai = []
    for r in cac_dong(ws, r0, cot):
        if ws.cell(r, cot['Nhóm']).value != 'Cần xem lại':
            continue
        so = ws.cell(r, cot['Số ký hiệu']).value
        nhom, ly_do = phan_loai(ws.cell(r, cot['Trích yếu']).value, so)
        if nhom in ('Cần xem lại', 'Không theo dõi'):
            continue
        ws.cell(r, cot['Nhóm']).value = nhom
        ws.cell(r, cot['Lý do phân loại']).value = ly_do + ' (phân loại lại %s)' % bay_gio.strftime('%d/%m/%Y')
        for c in cot.values():
            ws.cell(r, c).fill = PatternFill('solid', fgColor=MAU[nhom])
        phan_lai.append('%s: Cần xem lại → %s (%s)' % (so, nhom, ly_do))
    if phan_lai:
        ghi_chu.append('Phân loại lại %d dòng' % len(phan_lai))

    # Buoc 5: ghi dong moi
    stt = sum(1 for _ in cac_dong(ws, r0, cot))
    mau_o = {c: ws.cell(r0 + 1, c)._style for c in cot.values()} if stt else {}
    for b in moi:
        r = ws.max_row + 1
        while r > r0 + 1 and not ws.cell(r - 1, cot['Số ký hiệu']).value:
            r -= 1
        stt += 1
        gt = {'STT': stt, 'Số ký hiệu': b['so_ky_hieu'], 'Loại': loai_vb(b['so_ky_hieu']),
              'Ngày ban hành': b['ngay'].strftime('%d/%m/%Y'), 'Trích yếu': b['trich_yeu'],
              'Đơn vị/người xử lý chính': b.get('xu_ly') or None, 'Độ khẩn': b.get('do_khan'),
              'Nhóm': b['nhom'], 'Lý do phân loại': b['ly_do'], 'Tình trạng': 'Mới ban hành',
              'Doc ID': b.get('doc_id') or None, 'Ngày ghi nhận': bay_gio.strftime('%d/%m/%Y')}
        for t, c in cot.items():
            o = ws.cell(r, c, gt.get(t))
            if c in mau_o:
                o._style = mau_o[c]
            o.fill = PatternFill('solid', fgColor=MAU[b['nhom']])
        if not b.get('doc_id'):
            can_xu_ly.append('%s: thiếu Doc ID (tệp đầu vào không có cột này)' % b['so_ky_hieu'])

    # Buoc 4: doi chieu file cho MOI dong da ghi CSDL (25/10/2026: moi VB deu tai ve de trich cong
    # viec sau) — file san co trong 2026-09, file moi tai trong --nhan-file
    da_tai = []
    for r in cac_dong(ws, r0, cot):
        if ws.cell(r, cot['File đã tải']).value:
            continue
        so = ws.cell(r, cot['Số ký hiệu']).value
        trich = ws.cell(r, cot['Trích yếu']).value
        san = tim_file(thu_muc_file, so)
        if san:  # da co trong 2026-09 (lan truoc hoac bo sung tay) -> ghi nhan tep uu tien
            ws.cell(r, cot['File đã tải']).value = chon_chinh(san)
            da_tai.append(chon_chinh(san) + ' (đã có sẵn)')
            continue
        tai = tim_file(a.nhan_file, so) if a.nhan_file else []
        if not tai and a.nhan_file:  # ten tep goc do skill ktc-ioffice ghi trong CSV
            tai = [t for t in ten_goc.get(khoa_so(so), [])
                   if os.path.isfile(os.path.join(a.nhan_file, t))]
        if not tai:
            can_xu_ly.append('%s: chưa có tệp — mở showDocDetail(%s,\'\',\'2\') để tải' %
                             (so, ws.cell(r, cot['Doc ID']).value))
            ghi_chu.append('Chưa tải %s' % so)
            continue
        # Chuyen TAT CA tep dinh kem cua van ban; tep chinh (.docx uu tien) mang ten chuan,
        # tep phu them hau to _2, _3... Khong de tep co san (tuong duong mv -n).
        chinh = chon_chinh(tai)
        goc_ten = os.path.splitext(ten_file_chuan(so, trich, ''))[0]
        thu_tu = [chinh] + [t for t in tai if t != chinh]
        for i, t in enumerate(thu_tu, 1):
            dich_ten = goc_ten + ('' if i == 1 else '_%d' % i) + os.path.splitext(t)[1].lower()
            dich = os.path.join(thu_muc_file, dich_ten)
            if os.path.exists(dich):
                can_xu_ly.append('%s: không chuyển %s, đích đã tồn tại %s' % (so, t, dich_ten))
                continue
            if not a.thu:
                shutil.move(os.path.join(a.nhan_file, t), dich)
            if i == 1:
                ws.cell(r, cot['File đã tải']).value = dich_ten
            da_tai.append(dich_ten)

    # Buoc 6: khop bieu mau QLDT tren BAN SAO
    duong_bm = os.path.join(a.thu_muc, BIEU_MAU)
    ban_sao = os.path.join(a.thu_muc, '2026-09', 'QLĐT_may-dien_%s.xlsx' % bay_gio.strftime('%Y%m%d'))
    da_khop, chua_khop = [], []
    if os.path.exists(duong_bm):
        wb_bm = openpyxl.load_workbook(ban_sao if os.path.exists(ban_sao) else duong_bm)
        ung_vien = dong_bieu_mau(wb_bm)
        for r in cac_dong(ws, r0, cot):
            if ws.cell(r, cot['Nhóm']).value != 'Nhiệm vụ' or \
                    ws.cell(r, cot['Nhiệm vụ tương ứng trong biểu mẫu']).value:
                continue
            so = ws.cell(r, cot['Số ký hiệu']).value
            u, d = khop(ws.cell(r, cot['Trích yếu']).value, ung_vien)
            if not u:
                ws.cell(r, cot['Nhiệm vụ tương ứng trong biểu mẫu']).value = 'Chưa khớp'
                chua_khop.append('%s (điểm cao nhất %.2f)' % (so, d))
                continue
            sh, rb, cs, nd = u
            o = wb_bm[sh].cell(rb, cs)
            o.value = '%s; %s' % (o.value, so) if o.value else so
            o.fill = PatternFill('solid', fgColor=VANG)
            o.comment = Comment('Do máy điền %s — cần người rà lại' % bay_gio.strftime('%d/%m/%Y %H:%M'),
                                'cap_nhat_csdl')
            ws.cell(r, cot['Nhiệm vụ tương ứng trong biểu mẫu']).value = \
                'QLĐT.xlsx · %s · dòng %d (máy khớp %.2f, cần rà)' % (sh, rb, d)
            da_khop.append('%s → %s dòng %d: %s' % (so, sh, rb, nd[:70]))
        if da_khop and not a.thu:
            wb_bm.save(ban_sao)
    else:
        ghi_chu.append('Không thấy %s' % BIEU_MAU)

    # Nhat ky + luu
    dem = {n: sum(1 for b in moi if b['nhom'] == n) for n in MAU}
    nk = wb['Nhật ký chạy']
    nk.append([bay_gio.strftime('%d/%m/%Y %H:%M'), so_quet, len(moi), dem['Nhiệm vụ'],
               dem['Vận hành đào tạo'], dem['Cần xem lại'], len(da_tai),
               '; '.join(['Chờ phân loại %d' % dem['Chờ phân loại']] + ghi_chu) or 'cap_nhat_csdl.py'])
    if not a.thu:
        sao_luu = os.path.join(a.thu_muc, '_sao-luu')
        os.makedirs(sao_luu, exist_ok=True)
        shutil.copy2(duong_csdl, os.path.join(sao_luu, 'CSDL_%s.xlsx' % bay_gio.strftime('%Y%m%d-%H%M%S')))
        wb.save(duong_csdl)

    # Buoc 7: bao cao
    print('=== %sKẾT QUẢ CẬP NHẬT CSDL — %s ===' % ('[CHẠY THỬ, KHÔNG GHI] ' if a.thu else '',
                                                   bay_gio.strftime('%d/%m/%Y %H:%M')))
    if not moi and not da_tai and not da_khop and not can_xu_ly and not phan_lai:
        print('Không có văn bản mới.')
        return 0
    print('Quét %d VB · mới %d (Nhiệm vụ %d · Vận hành đào tạo %d · Cần xem lại %d)' %
          (so_quet, len(moi), dem['Nhiệm vụ'], dem['Vận hành đào tạo'], dem['Cần xem lại']))
    print('  (Chờ phân loại %d — thông báo khác, tải về để trích công việc sau)' % dem['Chờ phân loại'])
    for b in moi:
        print('  + %-14s %-17s %s' % (b['so_ky_hieu'], b['nhom'], b['trich_yeu'][:70]))
    for tieu_de, ds in (('Phân loại lại theo luật mới', phan_lai), ('File đã ghi nhận', da_tai), ('Đã điền vào biểu mẫu (bản sao)', da_khop),
                        ('Chưa khớp biểu mẫu QLĐT (bình thường với VB ngoài Phòng QLĐT)', chua_khop),
                        ('Cần người xử lý', can_xu_ly +
                         ['%s: nhóm Cần xem lại — %s' % (b['so_ky_hieu'], b['ly_do'])
                          for b in moi if b['nhom'] == 'Cần xem lại'])):
        if ds:
            print('%s:' % tieu_de)
            for x in ds:
                print('  - ' + x)
    return 0


if __name__ == '__main__':
    sys.exit(main())
