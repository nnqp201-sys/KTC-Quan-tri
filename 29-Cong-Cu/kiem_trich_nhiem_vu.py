# -*- coding: utf-8 -*-
"""Buoc 2-3 quy trinh trich nhiem vu (skill ktc-trich-nhiem-vu): chuan bi van ban cho AI doc, roi
KIEM ket qua AI trich (JSON) va lap SO CHO DUYET. Chi phat hien, khong sua noi dung AI trich.

  python 29-Cong-Cu/kiem_trich_nhiem_vu.py --chuan-bi
      Liet ke van ban nhom 'Nhiem vu' co tep, chua trich; ghi toan van danh so dong (L001...) vao
      29-Cong-Cu/_trung_gian/van-ban/<so>.txt de AI doc va trich dan dung dong.
  python 29-Cong-Cu/kiem_trich_nhiem_vu.py [--json <thu muc>]
      Kiem moi JSON trong 30-Ket-Qua/<hom nay>/trich-nhiem-vu/ -> 21-Master-Task-Register/Cho-duyet_<ngay>.xlsx
      (giu nguyen cot "Quyết định" nguoi duyet da nhap neu so da ton tai).

Muc co: DO = phai sua truoc khi duyet (nghi bia trich dan, sai ma, sai cau truc); VANG = nguoi duyet can xem.
Ma thoat: 1 neu co co DO, nguoc lai 0.
"""
import argparse
import difflib
import glob
import json
import os
import re
import sys
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cap_nhat_csdl import THU_MUC, TEN_CSDL, THU_MUC_FILE, chon_chinh, tim_file  # noqa: E402
from doc_van_ban import MA_DON_VI, doc_toan_van, khong_dau  # noqa: E402

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TRUNG_GIAN = os.path.join(GOC, '29-Cong-Cu', '_trung_gian', 'van-ban')
KET_QUA = os.path.join(GOC, '30-Ket-Qua')
MASTER_DIR = os.path.join(GOC, '21-Master-Task-Register')
MASTER = os.path.join(MASTER_DIR, 'Master-Task-Register_20260913_v0.1.xlsx')

BAT_BUOC = ['Ma_Tam', 'Ten_Nhiem_Vu', 'San_Pham_Yeu_Cau', 'Can_Cu_Dieu_Khoan', 'Dong', 'Truc',
            'Noi_Ham', 'Trich_Dan', 'Do_Tin_Cay']
COT = ['Quyết định', 'Ghi chú duyệt', 'Mức', 'Cờ', 'Số ký hiệu', 'Ma_Tam', 'Parent', 'Ten_Nhiem_Vu',
       'Mo_Ta', 'San_Pham_Yeu_Cau', 'So_Luong', 'Don_Vi_Chu_Tri', 'Don_Vi_Chu_Tri_Goc', 'Don_Vi_Phoi_Hop',
       'Don_Vi_Phoi_Hop_Goc', 'Nguoi_Chi_Dao', 'Han_Hoan_Thanh', 'Ky_Ke_Hoach', 'Truc', 'Noi_Ham',
       'Ma_NV_Chuan', 'Can_Cu_Dieu_Khoan', 'Dong', 'Trich_Dan', 'Do_Tin_Cay', 'Ghi_Chu', 'Ngay_Ban_Hanh',
       'Co_Quan', 'Task_ID']
RONG = {'Quyết định': 12, 'Ghi chú duyệt': 20, 'Mức': 7, 'Cờ': 40, 'Ten_Nhiem_Vu': 45, 'Mo_Ta': 50,
        'San_Pham_Yeu_Cau': 35, 'Trich_Dan': 50, 'Noi_Ham': 30, 'Don_Vi_Chu_Tri_Goc': 25,
        'Don_Vi_Phoi_Hop_Goc': 25, 'Ghi_Chu': 30}
MAU = {'ĐỎ': 'F8CBAD', 'VÀNG': 'FFF2CC', '': 'E2EFDA'}


def slug(so):
    return re.sub(r'[^A-Za-z0-9]+', '-', khong_dau(so)).strip('-').upper()


def phang(s):
    """So khop trich dan: bo dau, bo moi ky tu khong phai chu/so."""
    return re.sub(r'[^a-z0-9]', '', khong_dau(s))


def doc_ngay(s):
    m = re.match(r'^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$', str(s or ''))
    if not m:
        return None
    try:
        return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def mo_csdl(thu_muc):
    wb = openpyxl.load_workbook(os.path.join(thu_muc, TEN_CSDL), read_only=True, data_only=True)
    ws = wb['CSDL']
    hang = list(ws.iter_rows(values_only=True))
    r0 = next(i for i, h in enumerate(hang[:12]) if str(h[0] or '').strip() == 'STT')
    ten = list(hang[r0])
    return [dict(zip(ten, h)) for h in hang[r0 + 1:] if h[ten.index('Số ký hiệu')]]


def tu_khoa(s):
    return {t for t in re.findall(r'[a-z0-9]+', khong_dau(s)) if len(t) > 2}


def khop_trich_yeu(dong, trich_yeu):
    """Ty le tu khoa cua trich yeu (CSDL) co mat trong 30 dong dau van ban — nhan dien van ban chinh."""
    ty = tu_khoa(trich_yeu) - {'hoach', 'thong', 'bao', 'nam', 'cua', 'cac', 'truong'}
    if not ty:
        return 1.0
    return len(ty & tu_khoa(' '.join(dong[:30]))) / len(ty)


LOAI_VB = ['ke hoach', 'thong bao', 'bao cao', 'phieu', 'bien ban', 'thong tu', 'nghi dinh', 'quyet dinh',
           'to trinh', 'giay moi', 'cong van', 'mau', 'danh sach', 'chuong trinh', 'de an']


def la_van_ban_chinh(dong, so, trich_yeu):
    """Tra '' neu dung la van ban chinh cua Truong, nguoc lai ly do. Bon dieu kien:
    co quan ban hanh la Truong; tieu de loai dung (KH -> KẾ HOẠCH, TB -> THÔNG BÁO); so ky hieu o dong
    'Số:' (neu co ghi so) dung so; trich yeu khop >= 50%. Bao cao ra soat 897 nhac lai ten ke hoach nhung
    tieu de la BÁO CÁO va so BC -> bi loai."""
    dau = [khong_dau(x) for x in dong[:25]]
    if not any('truong cao dang kon tum' in x for x in dau[:6]):
        return 'cơ quan ban hành không phải Trường'
    ky = '/KH-' if '/KH-' in so else '/TB-'
    so_dong = next((x for x in dau if re.search(r'\bso\s*:', x)), '')
    if so_dong and not re.search(r'/%s-' % ky.strip('/-').lower(), so_dong):
        return 'dòng "Số:" là loại khác (%s)' % so_dong.split('|')[0].strip()[:40]
    m = re.search(r'\bso\s*:\s*(\d+)\s*/', so_dong)
    if m and m.group(1) != re.match(r'\s*(\d+)', so).group(1):
        return 'số văn bản trong tệp là %s' % m.group(1)
    tieu_de = next((x.strip() for x in dau if not x.startswith('[') and
                    any(re.match(r'^%s\b' % l, x.strip()) for l in LOAI_VB) and len(x) < 60 and
                    dong[dau.index(x)].upper() == dong[dau.index(x)]), '')
    can = 'ke hoach' if ky == '/KH-' else 'thong bao'
    if not tieu_de.startswith(can) and not (ky == '/TB-' and tieu_de.startswith('ket luan')):
        return 'tiêu đề văn bản là "%s"' % (tieu_de.upper() or 'không xác định')
    diem = khop_trich_yeu(dong, trich_yeu)
    return '' if diem >= 0.5 else 'không khớp trích yếu (%.0f%%)' % (diem * 100)


def toan_van(thu_muc, so, tep_chinh=None, trich_yeu=''):
    """Dong toan van: tep CHINH (khop trich yeu) truoc, phu luc .xlsx sau. Tra (dong, ten tep chinh, loi).
    Tep tai ve co the lech (bao cao kem theo, phieu...) -> thu lan luot theo uu tien, chi nhan tep ma
    trich yeu xuat hien o phan dau; khong tep nao khop -> bao LECH TEP, khong de AI trich nham."""
    tm = os.path.join(thu_muc, THU_MUC_FILE)
    cac = tim_file(tm, so)
    if tep_chinh and tep_chinh not in cac and os.path.isfile(os.path.join(tm, tep_chinh)):
        cac.append(tep_chinh)
    phu_luc = [os.path.join(tm, t) for t in cac if t.lower().endswith(('.xlsx', '.xls'))] + \
              [os.path.join(thu_muc, t) for t in tim_file(thu_muc, so) if t.lower().endswith(('.xlsx', '.xls'))]
    doc_duoc = [t for t in cac if t.lower().endswith(('.docx', '.pdf'))]
    if not doc_duoc:
        return [], '', 'không có tệp .docx/.pdf'
    thu, chinh, dong = [], '', []
    con = list(doc_duoc)
    while con:
        t = chon_chinh(con)
        con.remove(t)
        try:
            d = doc_toan_van(os.path.join(tm, t))
        except Exception as e:
            thu.append('%s: không đọc được (%s)' % (t, e))
            continue
        if not d:
            thu.append('%s: PDF ảnh quét' % t)
            continue
        ly_do = la_van_ban_chinh(d, so, trich_yeu)
        if not ly_do:
            chinh, dong = t, d
            break
        thu.append('%s: %s' % (t, ly_do))
    if not chinh:
        return [], '', 'LỆCH TỆP — không tệp nào là văn bản chính: ' + '; '.join(thu)
    for p in phu_luc:
        try:
            dong += ['=== Phụ lục: %s ===' % os.path.basename(p)] + doc_toan_van(p)
        except Exception as e:
            dong.append('=== Phụ lục %s: không đọc được (%s) ===' % (os.path.basename(p), e))
    return dong, chinh, ''


def json_da_co():
    return {os.path.basename(p) for p in glob.glob(os.path.join(KET_QUA, '*', 'trich-nhiem-vu', '*.json'))}


# ---------- --chuan-bi ----------

def chuan_bi(thu_muc):
    os.makedirs(TRUNG_GIAN, exist_ok=True)
    da_co = json_da_co()
    can, loi = [], []
    for r in mo_csdl(thu_muc):
        so = r['Số ký hiệu']
        if r.get('Nhóm') != 'Nhiệm vụ' or slug(so) + '.json' in da_co:
            continue
        dong, chinh, ly_do = toan_van(thu_muc, so, r.get('File đã tải'), r.get('Trích yếu'))
        if ly_do:
            loi.append('%s: %s' % (so, ly_do))
            continue
        duong = os.path.join(TRUNG_GIAN, slug(so) + '.txt')
        with open(duong, 'w', encoding='utf-8') as f:
            f.write('# %s | %s | ngày %s | tệp %s\n' % (so, r.get('Trích yếu'), r.get('Ngày ban hành'), chinh))
            for i, x in enumerate(dong, 1):
                f.write('L%03d %s\n' % (i, x))
        can.append((so, os.path.relpath(duong, GOC), len(dong)))
    print('=== VĂN BẢN CHỜ AI TRÍCH: %d ===' % len(can))
    for so, p, n in can:
        print('  %-14s %4d dòng  %s' % (so, n, p))
    if loi:
        print('Không chuẩn bị được (%d):' % len(loi))
        for x in loi:
            print('  - ' + x)
    print('JSON ghi vào: %s' % os.path.relpath(os.path.join(KET_QUA, datetime.now().strftime('%Y-%m-%d'),
                                                            'trich-nhiem-vu'), GOC))
    return 0


# ---------- kiem ----------

def kiem_van_ban(d, dong, csdl_kh):
    """Tra (cac dong nhiem vu kem co, cac co cap van ban)."""
    so = d.get('So_Ky_Hieu', '?')
    ngay_bh = doc_ngay(d.get('Ngay_Ban_Hanh'))
    van = phang(' '.join(dong))
    kq, co_vb = [], []
    ma_tam = {n.get('Ma_Tam') for n in d.get('Nhiem_Vu', [])}
    for n in d.get('Nhiem_Vu', []):
        do, vang = [], []
        thieu = [k for k in BAT_BUOC if n.get(k) in (None, '', [])]
        if thieu:
            do.append('thiếu ' + ', '.join(thieu))
        if n.get('Trich_Dan') and phang(n['Trich_Dan']) not in van:
            do.append('NGHI BỊA: trích dẫn không có trong văn bản')
        sai_dong = [x for x in n.get('Dong') or [] if not isinstance(x, int) or not 1 <= x <= len(dong)]
        if sai_dong:
            do.append('số dòng ngoài văn bản: %s' % sai_dong)
        ma = n.get('Don_Vi_Chu_Tri') or ''
        if ma and ma not in MA_DON_VI:
            do.append('mã chủ trì sai: %s' % ma)
        if not ma:
            vang.append('chủ trì không có mã (%s)' % (n.get('Don_Vi_Chu_Tri_Goc') or 'trống'))
        for m in [x.strip() for x in str(n.get('Don_Vi_Phoi_Hop') or '').split(';') if x.strip()]:
            if m not in MA_DON_VI:
                do.append('mã phối hợp sai: %s' % m)
        if n.get('Parent') and n['Parent'] not in ma_tam:
            do.append('Parent %s không có trong văn bản' % n['Parent'])
        try:
            truc = int(n.get('Truc'))
            if not 1 <= truc <= 6 or not str(n.get('Noi_Ham') or '').startswith('Trục %d' % truc):
                do.append('Trục/Nội hàm không khớp')
        except (TypeError, ValueError):
            do.append('Trục không hợp lệ')
        han = n.get('Han_Hoan_Thanh') or ''
        if not han:
            vang.append('không có hạn')
        elif doc_ngay(han):
            if ngay_bh and doc_ngay(han) < ngay_bh:
                vang.append('hạn trước ngày ban hành')
        elif not re.match(r'^Tháng \d{1,2}/\d{4}$|^Quý [IV]+/\d{4}$', han):
            do.append('hạn sai định dạng: %s' % han)
        if not re.match(r'^\d{4}-(T(0[1-9]|1[0-2])|Q[1-4]|NAM|CD)$', str(n.get('Ky_Ke_Hoach') or '')):
            do.append('Kỳ kế hoạch sai: %s' % n.get('Ky_Ke_Hoach'))
        if n.get('San_Pham_Yeu_Cau') == 'THIEU_DU_LIEU':
            vang.append('chưa xác định sản phẩm')
        if n.get('Do_Tin_Cay') in ('Thấp', 'Trung bình'):
            vang.append('độ tin cậy %s' % n['Do_Tin_Cay'].lower())
        # Noi TB -> KH: nhiem vu "tham muu ke hoach ... <so hieu VB>" ma KH do da ban hanh sau ngay giao
        for ref in re.findall(r'\d+/\d{4}/[A-ZĐ][A-ZĐa-zđ-]*', '%s %s' % (n.get('San_Pham_Yeu_Cau'), n.get('Ten_Nhiem_Vu'))):
            for kh in csdl_kh:
                if phang(ref) in phang(kh['Trích yếu']) and kh['Số ký hiệu'] != so and \
                        (not ngay_bh or (doc_ngay(kh['Ngày ban hành']) or ngay_bh) >= ngay_bh):
                    vang.append('CÓ THỂ ĐÃ XONG: %s (%s)' % (kh['Số ký hiệu'], str(kh['Trích yếu'])[:60]))
        kq.append((n, do, list(dict.fromkeys(vang))))
    # Sot muc: dong danh so / gach dau trong pham vi giao viec chua duoc nhiem vu hay loai tru nao tro toi
    pv = d.get('Pham_Vi_Giao_Viec') or [1, len(dong)]
    phu = {x for n in d.get('Nhiem_Vu', []) for x in n.get('Dong') or []} | \
          {x for l in d.get('Loai_Tru', []) for x in l.get('Dong') or []}
    sot = [i for i in range(max(1, pv[0]), min(len(dong), pv[1]) + 1)
           if re.match(r'^(\d+(\.\d+)*[.)]|[-–+•*]|[a-zđ]\)|\[Bảng \d+\.(?!1\])\d+\])\s*', dong[i - 1])
           and not dong[i - 1].rstrip().endswith(':') and i not in phu]
    if sot:
        co_vb.append('CÓ THỂ SÓT dòng %s' % ', '.join('L%03d' % i for i in sot))
    if not d.get('Nhiem_Vu'):
        co_vb.append('không có nhiệm vụ nào')
    return kq, co_vb


def giong(a, b):
    """Ten gan giong VA cung so hieu (neu co): 'TT 76/2026' khac 'TT 79/2026' du chu gan nhu y het."""
    so_a, so_b = set(re.findall(r'\d+', a)), set(re.findall(r'\d+', b))
    return so_a == so_b and difflib.SequenceMatcher(None, a, b).ratio() > 0.85


def trung_lap(ds, master):
    """Gan co trung giua cac nhiem vu trong dot va voi Master (cung chu tri, ten gan giong)."""
    goc = [(m.get('Don_Vi_Chu_Tri'), khong_dau(m.get('Ten_Nhiem_Vu')), m.get('Task_ID')) for m in master]
    for i, (so, n, do, vang) in enumerate(ds):
        t = khong_dau(n.get('Ten_Nhiem_Vu'))
        for so2, n2, _, _ in ds[:i]:
            if n2.get('Don_Vi_Chu_Tri') == n.get('Don_Vi_Chu_Tri') and \
                    giong(t, khong_dau(n2.get('Ten_Nhiem_Vu'))):
                vang.append('CÓ THỂ TRÙNG %s %s' % (so2, n2.get('Ma_Tam')))
        for dv, t2, tid in goc:
            if dv == n.get('Don_Vi_Chu_Tri') and giong(t, t2):
                vang.append('CÓ THỂ TRÙNG Master %s' % tid)


def doc_master():
    if not os.path.exists(MASTER):
        return []
    ws = openpyxl.load_workbook(MASTER, read_only=True)['01-Master-Task-Register']
    hang = list(ws.iter_rows(values_only=True))
    return [dict(zip(hang[0], h)) for h in hang[3:] if h and h[0]]


def kiem(thu_muc, thu_muc_json, ra=MASTER_DIR):
    tep_json = sorted(glob.glob(os.path.join(thu_muc_json, '*.json')))
    if not tep_json:
        print('Không có JSON nào trong %s' % thu_muc_json)
        return 0
    csdl = mo_csdl(thu_muc)
    csdl_kh = [r for r in csdl if '/KH-' in str(r['Số ký hiệu'])]
    theo_so = {r['Số ký hiệu']: r for r in csdl}
    ds, tong_vb, loai_tru = [], [], []
    for p in tep_json:
        try:
            d = json.load(open(p, encoding='utf-8'))
        except ValueError as e:
            tong_vb.append([os.path.basename(p), 'ĐỎ', 'JSON hỏng: %s' % e, 0, 0])
            continue
        so = d.get('So_Ky_Hieu', '')
        r_csdl = theo_so.get(so) or {}
        dong, _, ly_do = toan_van(thu_muc, so, r_csdl.get('File đã tải'), r_csdl.get('Trích yếu'))
        if ly_do:
            tong_vb.append([so, 'ĐỎ', 'không đọc lại được văn bản gốc: ' + ly_do, 0, 0])
            continue
        kq, co_vb = kiem_van_ban(d, dong, csdl_kh)
        for n, do, vang in kq:
            n.update(Ngay_Ban_Hanh=d.get('Ngay_Ban_Hanh'), Co_Quan=d.get('Co_Quan'))
            ds.append((so, n, do, vang))
        for l in d.get('Loai_Tru', []):
            loai_tru.append([so, ', '.join('L%03d' % x for x in l.get('Dong') or []), l.get('Noi_Dung'), l.get('Ly_Do')])
        tong_vb.append([so, 'ĐỎ' if any(k[1] for k in kq) else ('VÀNG' if co_vb else ''),
                        '; '.join(co_vb), len(kq), len(d.get('Loai_Tru', []))])
    trung_lap(ds, doc_master())

    duong = os.path.join(ra, 'Cho-duyet_%s.xlsx' % datetime.now().strftime('%Y%m%d'))
    cu = {}
    if os.path.exists(duong):  # giu quyet dinh nguoi duyet da nhap
        ws = openpyxl.load_workbook(duong)['Chờ duyệt']
        h = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            cu[(r[h.index('Số ký hiệu')], r[h.index('Ma_Tam')])] = (r[0], r[1], r[h.index('Task_ID')])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Chờ duyệt'
    ws.append(COT)
    for so, n, do, vang in ds:
        muc = 'ĐỎ' if do else ('VÀNG' if vang else '')
        qd, gc, tid = cu.get((so, n.get('Ma_Tam')), (None, None, None))
        gt = dict(n, **{'Quyết định': qd, 'Ghi chú duyệt': gc, 'Mức': muc, 'Cờ': '; '.join(do + vang),
                        'Số ký hiệu': so, 'Task_ID': tid,
                        'Dong': ', '.join('L%03d' % x for x in n.get('Dong') or [] if isinstance(x, int))})
        ws.append([gt.get(c) for c in COT])
        for c in range(1, len(COT) + 1):
            o = ws.cell(ws.max_row, c)
            o.fill = PatternFill('solid', fgColor=MAU[muc])
            o.alignment = Alignment(wrap_text=True, vertical='top')
            o.font = Font(name='Times New Roman', size=12)
    for sh, tieu_de, hang in (('Văn bản', ['Số ký hiệu', 'Mức', 'Cờ văn bản', 'Số nhiệm vụ', 'Số loại trừ'], tong_vb),
                              ('Loại trừ', ['Số ký hiệu', 'Dòng', 'Nội dung', 'Lý do'], loai_tru)):
        w2 = wb.create_sheet(sh)
        w2.append(tieu_de)
        for h in hang:
            w2.append(h)
    for w in wb.worksheets:
        for j, t in enumerate([c.value for c in w[1]], 1):
            w.cell(1, j).font = Font(name='Times New Roman', size=12, bold=True)
        for hang in w.iter_rows(min_row=2):
            for o in hang:
                if not o.font.bold:
                    o.font = Font(name='Times New Roman', size=12)
            w.column_dimensions[openpyxl.utils.get_column_letter(j)].width = RONG.get(t, 14)
        w.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wb.save(duong)

    so_do = sum(1 for x in ds if x[2])
    print('=== KIỂM TRÍCH NHIỆM VỤ — %d văn bản · %d nhiệm vụ · ĐỎ %d · VÀNG %d ===' % (
        len(tong_vb), len(ds), so_do, sum(1 for x in ds if not x[2] and x[3])))
    for v in tong_vb:
        print('  %-14s %-5s nv=%-3s loại trừ=%-3s %s' % (v[0], v[1] or 'ổn', v[3], v[4], v[2]))
    for so, n, do, vang in ds:
        if do or vang:
            print('  [%s] %s %s: %s' % ('ĐỎ' if do else 'VÀNG', so, n.get('Ma_Tam'), '; '.join(do + vang)))
    print('Sổ chờ duyệt: %s' % os.path.relpath(duong, GOC))
    return 1 if so_do or any(v[1] == 'ĐỎ' for v in tong_vb) else 0


def main():
    ap = argparse.ArgumentParser(description='Chuẩn bị và kiểm kết quả AI trích nhiệm vụ')
    ap.add_argument('--chuan-bi', action='store_true', help='Xuất toàn văn có số dòng cho AI đọc')
    ap.add_argument('--json', help='Thư mục JSON (mặc định 30-Ket-Qua/<hôm nay>/trich-nhiem-vu)')
    ap.add_argument('--thu-muc', default=THU_MUC, help='Thư mục CSDL và tệp văn bản')
    ap.add_argument('--ra', default=MASTER_DIR, help='Nơi ghi sổ chờ duyệt')
    a = ap.parse_args()
    if a.chuan_bi:
        return chuan_bi(a.thu_muc)
    return kiem(a.thu_muc, a.json or os.path.join(KET_QUA, datetime.now().strftime('%Y-%m-%d'), 'trich-nhiem-vu'), a.ra)


if __name__ == '__main__':
    sys.exit(main())
