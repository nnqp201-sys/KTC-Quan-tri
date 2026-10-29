# -*- coding: utf-8 -*-
"""Trich danh muc San pham/Cong viec cua PHAN MEM KPI tinh (VNPT-DGCB) + anh xa sang QD 2119.

Tu 27/10/2026, Truong dang ky ke hoach cong tac quy tren phan mem KPI tinh Quang Ngai
(https://kpi.quangngai.dcs.vn). Phan mem dinh danh san pham bang ma rieng, quy tac
"Ma nhom SPCV_so thu tu tang dan" (HDSD VNPT-DGCB v2.0, muc IV.2.4.2.2, tr. 35):
  - N1..N5 = Nhom 1..5 theo QD 2119 (doi chieu 27/10/2026: khop 414/414 san pham, 0 lech nhom)
  - T1..T6 = sau nhom Truong khai them cho san pham chua co trong QD 2119 (39 san pham).
    Noi dung khop 6 Truc TB 817 — nhung CHUA co van ban xac nhan, nen cot he_so de trong
    va ghi THIEU_DU_LIEU. Khong tu dat he so (CLAUDE.md, QD 2119).

Nguon:
  1. Tep mau import cua phan mem: 11-Du-lieu-Cong-Viec/*Import*.xlsx, sheet `SPCV`
  2. Phu luc QD 2119 trong KTC-Database (chi doc) — lay he so, nhom, truc, noi ham, mo ta

Chay:  python 29-Cong-Cu/trich_danh_muc_pm_kpi.py [--nguon <tep>] [--ra <thu muc>]
Ra:    danh-muc-spcv-pm-kpi.csv + .nguon.txt (sha256 hai nguon, so dong, so dong thieu he so)
"""
import argparse
import csv
import glob
import hashlib
import os
import re
import sys
import unicodedata
import warnings

DU_AN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

TEN_CSV = "danh-muc-spcv-pm-kpi.csv"
MAC_DINH_RA = os.path.join(DU_AN, "11-Du-lieu-Cong-Viec", "DANH MUC SAN PHAM CONG VIEC")
MA_SPCV = re.compile(r"^([NT]\d)_(\d+)$")
COT = ["ma_spcv", "ten_spcv", "nhom_pm", "ma_qd2119", "nhom_qd2119", "he_so",
       "truc", "noi_ham", "ma_vb", "san_pham_minh_chung", "mo_ta", "ghi_chu"]


def chuan(s):
    """Chuan hoa ten san pham de doi chieu: bo dau cau, khoang trang thua, chu thuong."""
    s = unicodedata.normalize("NFC", str(s)).lower().strip()
    return re.sub(r"[\s\.,;:\(\)/\-–]+", " ", s).strip()


# Hai san pham QD 2119 duoc doi ten trong danh muc phan mem (doi chieu thu cong 27/10/2026)
DOI_TEN = {
    chuan("Nghị quyết Hội nghị viên chức và người lao động"):
        chuan("Nghị quyết Hội nghị CBVC"),
    chuan("Đề xuất, quyết định công nhận HSSV được hưởng chế độ, chính sách"):
        chuan("Đề xuất, quyết định công nhận SV được hưởng chế độ, chính sách"),
}


def tim_nguon_pm():
    ds = [p for p in glob.glob(os.path.join(DU_AN, "11-Du-lieu-Cong-Viec", "*.xlsx"))
          if "import" in os.path.basename(p).lower()]
    if not ds:
        raise FileNotFoundError(
            "Không thấy tệp mẫu import của phần mềm KPI (11-Du-lieu-Cong-Viec/*Import*.xlsx) "
            "— dừng, hỏi người dùng.")
    return sorted(ds)[-1]


def tim_nguon_2119():
    import duong_dan
    db = duong_dan.ktc_database()
    ds = glob.glob(os.path.join(db, "02-KTC-Regulations", "**", "PL-2119-QD-CDKT_*.xlsx"), recursive=True)
    if not ds:
        raise FileNotFoundError(
            "Không thấy phụ lục QĐ 2119 (PL-2119-QD-CDKT_*.xlsx) trong KTC-Database/02-KTC-Regulations "
            "— dừng, hỏi người dùng.")
    return sorted(ds)[-1]


def doc_pm(nguon):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(nguon, data_only=True, read_only=True)
    if "SPCV" not in wb.sheetnames:
        raise ValueError(f"Tệp {os.path.basename(nguon)} không có sheet 'SPCV' — dừng, hỏi người dùng.")
    ra = []
    for r in wb["SPCV"].iter_rows(min_row=3, values_only=True):
        if not r[1] or str(r[1]).strip() == "Mã sản phẩm công việc":
            continue
        ra.append((str(r[1]).strip(), str(r[2] or "").strip()))
    return ra


def doc_2119(nguon):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(nguon, data_only=True, read_only=True)
    ra = {}
    for r in wb["Danh_muc_san_pham"].iter_rows(min_row=5, values_only=True):
        if r[1] and r[2]:
            ra.setdefault(chuan(r[2]), r)
    return ra


def ghep(pm, pl):
    hang, loi = [], []
    for ma, ten in pm:
        m = MA_SPCV.match(ma)
        if not m:
            loi.append(f"mã sai dạng '{ma}' (phải là N1_241, T4_461…)")
            continue
        nhom_pm = m.group(1)
        k = chuan(ten)
        o = pl.get(k) or pl.get(DOI_TEN.get(k, ""))
        if o:
            ma2119 = str(o[1]).strip()
            pc = ma2119.split(".")
            truc, noi_ham, ma_vb = (pc + ["", "", ""])[:3]
            nhom = str(o[5] or "").strip()
            if nhom_pm.startswith("N") and nhom.replace("Nhóm ", "N") != nhom_pm:
                loi.append(f"{ma}: nhóm phần mềm {nhom_pm} ≠ {nhom} của QĐ 2119 ({ten[:40]}…)")
            hang.append({"ma_spcv": ma, "ten_spcv": ten, "nhom_pm": nhom_pm, "ma_qd2119": ma2119,
                         "nhom_qd2119": nhom, "he_so": str(o[6] or "").strip(),
                         "truc": truc, "noi_ham": noi_ham, "ma_vb": ma_vb,
                         "san_pham_minh_chung": str(o[4] or "").strip(), "mo_ta": str(o[3] or "").strip(),
                         "ghi_chu": "Đổi tên so với PL-2119" if k in DOI_TEN else ""})
        else:
            hang.append({"ma_spcv": ma, "ten_spcv": ten, "nhom_pm": nhom_pm, "ma_qd2119": "",
                         "nhom_qd2119": "", "he_so": "",
                         "truc": nhom_pm[1] if nhom_pm.startswith("T") else "",
                         "noi_ham": "", "ma_vb": "", "san_pham_minh_chung": "", "mo_ta": "",
                         "ghi_chu": "THIEU_DU_LIEU — chưa có trong PL-2119, chưa có hệ số"})
    return hang, loi


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--nguon", help="tệp mẫu import của phần mềm KPI (.xlsx, sheet SPCV)")
    ap.add_argument("--pl2119", help="phụ lục QĐ 2119 (.xlsx) — mặc định tìm trong KTC-Database")
    ap.add_argument("--ra", default=MAC_DINH_RA)
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    warnings.simplefilter("ignore")

    n_pm = a.nguon or tim_nguon_pm()
    n_pl = a.pl2119 or tim_nguon_2119()
    hang, loi = ghep(doc_pm(n_pm), doc_2119(n_pl))
    trung = sorted({h["ma_spcv"] for h in hang if sum(x["ma_spcv"] == h["ma_spcv"] for x in hang) > 1})
    if loi or trung:
        print(f"✗ Lỗi: {loi[:5]} · trùng mã: {trung[:5]}")
        return 1

    os.makedirs(a.ra, exist_ok=True)
    ra = os.path.join(a.ra, TEN_CSV)
    with open(ra, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COT)
        w.writeheader()
        w.writerows(hang)

    thieu = [h for h in hang if not h["he_so"]]
    with open(ra + ".nguon.txt", "w", encoding="utf-8", newline="") as f:
        f.write(f"nguon_phan_mem: 11-Du-lieu-Cong-Viec/{os.path.basename(n_pm)}\n"
                f"sha256: {hashlib.sha256(open(n_pm, 'rb').read()).hexdigest()}\n"
                f"nguon_he_so: KTC-Database/02-KTC-Regulations/.../{os.path.basename(n_pl)}\n"
                f"sha256: {hashlib.sha256(open(n_pl, 'rb').read()).hexdigest()}\n"
                f"so_dong: {len(hang)}\nso_dong_co_he_so: {len(hang) - len(thieu)}\n"
                f"so_dong_thieu_he_so: {len(thieu)}\n"
                "trang_thai: ma chinh thuc dung tren phan mem KPI tinh (kpi.quangngai.dcs.vn) tu 27/10/2026; "
                "he so lay theo QD 2119/QD-CDKT ngay 28/9/2026\n")
    print(f"✓ Đã trích {len(hang)} sản phẩm/công việc → {ra}")
    print(f"  {len(hang) - len(thieu)} dòng có hệ số QĐ 2119 · {len(thieu)} dòng THIEU_DU_LIEU "
          f"({', '.join(sorted({h['nhom_pm'] for h in thieu}))})")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
