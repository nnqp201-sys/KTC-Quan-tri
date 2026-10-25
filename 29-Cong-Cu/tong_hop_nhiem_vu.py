# -*- coding: utf-8 -*-
"""Buoc 6 quy trinh trich nhiem vu: tong hop Master Task Register cho giao ban tuan / bao cao thang.

Xuat 30-Ket-Qua/<ngay>/tong-hop/Tong-hop-nhiem-vu_<yyyymmdd>.xlsx, 4 sheet:
  Theo đơn vị · Đến hạn <N> ngày · Quá hạn · Theo Trục
Han dang "Tháng m/yyyy" tinh la ngay cuoi thang. Nhiem vu da xong = Trang_Thai thuoc DA_XONG.

  python 29-Cong-Cu/tong_hop_nhiem_vu.py [--ngay 7] [--master <tệp>]
"""
import argparse
import calendar
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from doc_van_ban import MA_DON_VI  # noqa: E402
from kiem_trich_nhiem_vu import GOC, MASTER, doc_ngay  # noqa: E402

DA_XONG = ('Đã hoàn thành', 'Đã kiểm tra', 'Đã báo cáo')
DUNG = ('Hủy',)


def han_ngay(han):
    d = doc_ngay(han)
    if d:
        return d
    m = re.match(r'^Tháng (\d{1,2})/(\d{4})$', str(han or ''))
    if m:
        y, mth = int(m.group(2)), int(m.group(1))
        return datetime(y, mth, calendar.monthrange(y, mth)[1])
    return None


def main():
    ap = argparse.ArgumentParser(description='Tổng hợp nhiệm vụ theo đơn vị, hạn, Trục')
    ap.add_argument('--ngay', type=int, default=7, help='Cửa sổ "sắp đến hạn" (ngày)')
    ap.add_argument('--master', default=MASTER)
    ap.add_argument('--ra', help='Thư mục ghi (mặc định 30-Ket-Qua/<hôm nay>/tong-hop)')
    a = ap.parse_args()

    hom_nay = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    ws = openpyxl.load_workbook(a.master, read_only=True)['01-Master-Task-Register']
    hang = list(ws.iter_rows(values_only=True))
    ds = [dict(zip(hang[0], h)) for h in hang[3:] if h and h[0]]
    ds = [n for n in ds if n.get('Trang_Thai_Phu') not in DUNG]

    dv = defaultdict(lambda: [0, 0, 0, 0, 0])   # tong, xong, dang lam, den han, qua han
    truc = defaultdict(lambda: [0, 0, 0])        # tong, xong, qua han
    den_han, qua_han = [], []
    for n in ds:
        k = n.get('Don_Vi_Chu_Tri') or '(chưa có mã)'
        xong = n.get('Trang_Thai') in DA_XONG
        h = han_ngay(n.get('Han_Hoan_Thanh'))
        dv[k][0] += 1
        truc[n.get('Truc')][0] += 1
        if xong:
            dv[k][1] += 1
            truc[n.get('Truc')][1] += 1
            continue
        dv[k][2] += 1
        if h and h < hom_nay:
            dv[k][4] += 1
            truc[n.get('Truc')][2] += 1
            qua_han.append((n, (hom_nay - h).days))
        elif h and h <= hom_nay + timedelta(days=a.ngay):
            dv[k][3] += 1
            den_han.append((n, (h - hom_nay).days))

    wb = openpyxl.Workbook()
    w = wb.active
    w.title = 'Theo đơn vị'
    w.append(['Mã', 'Đơn vị', 'Tổng', 'Đã xong', 'Chưa xong', 'Đến hạn %d ngày' % a.ngay, 'Quá hạn'])
    for k in sorted(dv, key=lambda x: (x not in MA_DON_VI, x)):
        w.append([k, MA_DON_VI.get(k, '')] + dv[k][:2] + [dv[k][2], dv[k][3], dv[k][4]])
    w.append(['', 'TỔNG'] + [sum(v[i] for v in dv.values()) for i in range(5)])
    cot = ['Task_ID', 'Don_Vi_Chu_Tri', 'Ten_Nhiem_Vu', 'San_Pham_Yeu_Cau', 'Han_Hoan_Thanh', 'Trang_Thai',
           'Phan_Tram_Tien_Do', 'Can_Cu_So_Hieu']
    for ten, nhom, nhan in (('Đến hạn %d ngày' % a.ngay, sorted(den_han, key=lambda x: x[1]), 'Còn (ngày)'),
                            ('Quá hạn', sorted(qua_han, key=lambda x: -x[1]), 'Quá (ngày)')):
        w2 = wb.create_sheet(ten)
        w2.append(cot + [nhan])
        for n, so_ngay in nhom:
            w2.append([n.get(c) for c in cot] + [so_ngay])
    w3 = wb.create_sheet('Theo Trục')
    w3.append(['Trục', 'Tổng', 'Đã xong', 'Quá hạn'])
    for t in sorted(truc, key=lambda x: (x is None, x)):
        w3.append(['Trục %s' % t if t else '(trống)'] + truc[t])
    for s in wb.worksheets:
        for c in s[1]:
            c.font = Font(bold=True)
        for j, rong in enumerate([18, 14, 50, 35, 13, 14, 10, 15, 10], 1):
            s.column_dimensions[openpyxl.utils.get_column_letter(j)].width = rong
        s.freeze_panes = 'A2'

    ra = a.ra or os.path.join(GOC, '30-Ket-Qua', hom_nay.strftime('%Y-%m-%d'), 'tong-hop')
    os.makedirs(ra, exist_ok=True)
    duong = os.path.join(ra, 'Tong-hop-nhiem-vu_%s.xlsx' % hom_nay.strftime('%Y%m%d'))
    wb.save(duong)
    print('=== TỔNG HỢP %d nhiệm vụ — đến hạn %d ngày: %d · quá hạn: %d ===' % (
        len(ds), a.ngay, len(den_han), len(qua_han)))
    for k in sorted(dv, key=lambda x: (x not in MA_DON_VI, x)):
        v = dv[k]
        print('  %-14s tổng %-3d xong %-3d chưa %-3d đến hạn %-3d quá hạn %d' % (k, v[0], v[1], v[2], v[3], v[4]))
    print('Tệp: %s' % os.path.relpath(duong, GOC))
    return 0


if __name__ == '__main__':
    sys.exit(main())
