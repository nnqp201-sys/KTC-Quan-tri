# -*- coding: utf-8 -*-
"""Kiểm nhất quán một hay nhiều văn bản (.docx) của dự án với du-an_*.yaml.

Dùng:  python 29-Cong-Cu/dau_tu_cong/kiem_nhat_quan_du_an.py <tep.docx> [<tep2.docx> ...] [--yaml du-an.yaml]

Bắt: (1) năm ngày tháng vô lý; (2) số hiệu văn bản có trong YAML nhưng ngày khác; (3) tên dự án trong dấu
ngoặc kép khác nguyên văn; (4) số tiền lạ (không có trong YAML) — Mức 3, cần người soát; (5) lỗi gõ đã gặp.
Mã thoát 1 khi có lỗi Mức 1. Không sửa tệp.
"""
import argparse
import os
import re
import sys

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from du_an import tai  # noqa: E402

LOI_GO = ["bỗ sung", "chỉ tiết", "hợp nhấc", "phê duyêt", "thẩm đinh", "Quản Ngãi", "Quảng Ngãi.,"]


def doc_dong(p):
    d = Document(p)
    dong = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            seen = set()
            for c in r.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                dong.extend(x.text for x in c.paragraphs)
    return [x for x in dong if x.strip()]


def so_tien_cho_phep(d):
    ok = set()

    def quet(x):
        if isinstance(x, dict):
            for v in x.values():
                quet(v)
        elif isinstance(x, (list, tuple)):
            for v in x:
                quet(v)
        elif isinstance(x, int) and x >= 1000:
            ok.add(f"{x:,}".replace(",", "."))
        elif isinstance(x, str):
            ok.update(re.findall(r"\d{1,3}(?:\.\d{3})+", x))

    quet(d)
    # các số hạng mục trong QĐ 2158 (giá gói KS gồm hai khoản mục)
    ok.update({"93.995.981", "758.349.265"})
    return ok


def kiem(p, d):
    kq = []
    dong = doc_dong(p)
    van = "\n".join(dong)
    ten = d["du_an"]["ten"]
    cho_phep = so_tien_cho_phep(d)

    for i, s in enumerate(dong, 1):
        for m in re.finditer(r"\b(\d{1,2})[/.](\d{1,2})[/.](\d{4})\b", s):
            nam = int(m.group(3))
            if nam < 1990 or nam > 2035:
                kq.append((1, "NK01", i, m.group(0), "Năm vô lý (nhiều khả năng gõ sai)"))
        for m in re.finditer(r"(?:số|Số)\s+(\d{1,5}/[\w\-ĐĐ/]+)\s*(?:\([^)]*\)\s*)?ngày\s+(\d{1,2}/\d{1,2}/\d{4}|\d{1,2} tháng \d{1,2} năm \d{4})", s):
            so, ngay = m.group(1), m.group(2)
            for ma, v in d["van_ban"].items():
                if v["so"] == so:
                    chuan = re.sub(r"^0", "", v["ngay"].split("/")[0]) + "/" + str(int(v["ngay"].split("/")[1])) + "/" + v["ngay"].split("/")[2]
                    g = ngay
                    if "tháng" in g:
                        a = re.findall(r"\d+", g)
                        g = f"{int(a[0])}/{int(a[1])}/{a[2]}"
                    else:
                        a = g.split("/")
                        g = f"{int(a[0])}/{int(a[1])}/{a[2]}"
                    if g != chuan:
                        kq.append((1, "NK02", i, m.group(0)[:80], f"Ngày khác YAML ({v['ngay']}) — kiểm {ma}"))
        for m in re.finditer(r"“(Đầu tư, cải tạo[^”]*)”", s):
            if m.group(1).strip().rstrip(".") != ten:
                kq.append((1, "NK03", i, m.group(1)[:80], "Tên dự án khác nguyên văn trong YAML (QĐ 805, NQ 46)"))
        for m in re.finditer(r"\d{1,3}(?:\.\d{3})+", s):
            if m.group(0) not in cho_phep and not re.match(r"^\d{1,3}\.\d{3}$", m.group(0)) :
                kq.append((3, "NK04", i, m.group(0), "Số tiền không có trong YAML — soát nguồn"))
        for loi in LOI_GO:
            if loi in s:
                kq.append((2, "NK05", i, loi, "Lỗi gõ đã gặp"))
    return sorted(kq, key=lambda x: (x[0], x[2]))


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("tep", nargs="+")
    ap.add_argument("--yaml", default=None)
    a = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")
    d = tai(a.yaml)
    co1 = False
    for p in a.tep:
        kq = kiem(p, d)
        print(f"\n=== {os.path.basename(p)} — {len(kq)} gợi ý ===")
        for muc, ma, dong, trich, mota in kq:
            print(f"  [Mức {muc}] {ma} dòng {dong}: {mota} — “{trich}”")
            co1 |= muc == 1
        if not kq:
            print("  ✓ Không phát hiện lệch so với YAML.")
    return 1 if co1 else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
