# -*- coding: utf-8 -*-
"""Tờ trình đề nghị phê duyệt KHLCNT giai đoạn CBĐT — ráp nội dung vào mẫu TTr của Trường (Downloads/TTr de nghi phe duyet KHLCNT.docx)."""
import copy, docx, docx.oxml.ns
from docx.text.paragraph import Paragraph

import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from du_an import tai_ve, onedrive_cbdt, ket_qua  # noqa: E402

MAU = tai_ve("TTr de nghi phe duyet KHLCNT.docx")
V9 = onedrive_cbdt("QD-KHLCNT-giai-doan-CBDT_Du-an-CLC-2026-2030_20261007_v9-sach.docx")
OUT = ket_qua("2026-10-07", "Soan-Thao-VB", "TTr-de-nghi-phe-duyet-KHLCNT-giai-doan-CBDT_20261007.docx")

d = docx.Document(MAU)
v9 = docx.Document(V9)
TEN = "Đầu tư, cải tạo các công trình phục vụ người học, nhà giáo và mua sắm bổ sung thiết bị đáp ứng mục tiêu xây dựng Trường Cao đẳng Kon Tum trở thành trường cao đẳng chất lượng cao giai đoạn 2026 - 2030"
VON = "Chương trình mục tiêu quốc gia xây dựng nông thôn mới, giảm nghèo bền vững và phát triển kinh tế - xã hội vùng đồng bào dân tộc thiểu số và miền núi giai đoạn 2026 - 2035, giai đoạn I: Từ năm 2026 đến năm 2030"

def set_text(par, text):
    from docx.oxml.ns import qn
    ts = [t for t in par._p.iter(qn("w:t"))
          if not any(a.tag in (qn("w:drawing"), qn("w:pict")) for a in t.iterancestors())]
    if not ts:
        par.add_run(text); return
    ts[0].text = text
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for t in ts[1:]:
        t.text = ""

def find(prefix, start=0):
    for i, p in enumerate(d.paragraphs):
        if i >= start and p.text.strip().startswith(prefix):
            return i
    raise ValueError(prefix)

P = d.paragraphs
# ---- vai trò đoạn (nhân bản trước khi xóa) ----
role = {k: copy.deepcopy(P[i]._p) for k, i in {
    "H": 7, "T": 8, "CC": 17, "KHONG": 28, "A": 34, "BUL": 36, "PLUS": 41, "TX": 40}.items()}

# ---- đầu trang ----
t0 = d.tables[0]
set_text(t0.rows[0].cells[0].paragraphs[1], "PHÒNG TỔNG HỢP - HÀNH CHÍNH VÀ QUẢN TRỊ")
p_so = t0.rows[1].cells[0].paragraphs[0]
r = p_so.add_run("Số: 16/TTr-THHCQT")
from docx.shared import Pt
r.font.size = Pt(13)
set_text(t0.rows[1].cells[1].paragraphs[0], "Quảng Ngãi, ngày 07 tháng 10 năm 2026")

# ---- tiêu đề ----
set_text(P[2], "đề nghị phê duyệt Kế hoạch lựa chọn nhà thầu giai đoạn chuẩn bị đầu tư")
set_text(P[3], "Dự án “" + TEN + "”")

# ---- thân: xóa 7..57, dựng lại trước đoạn "VII." ----
i_vii = find("VII.")
for p in list(d.paragraphs)[7:i_vii]:
    p._p.getparent().remove(p._p)
anchor = d.paragraphs[find("VII.")]

def add(roleName, text):
    new = copy.deepcopy(role[roleName])
    anchor._p.addprevious(new)
    par = Paragraph(new, anchor._parent)
    set_text(par, text)
    return par

add("H", "I. Mô tả tóm tắt dự án")
add("T", "1. Tên dự án: “" + TEN + "”.")
add("T", "2. Sơ bộ tổng mức đầu tư: 24.894.000.000 đồng (Hai mươi bốn tỷ tám trăm chín mươi bốn triệu đồng). Dự án thuộc Nhóm C theo quy định tại khoản 4 Điều 11 của Luật Đầu tư công.")
add("T", "3. Dự toán kinh phí chuẩn bị đầu tư: 1.519.778.268 đồng (Một tỷ năm trăm mười chín triệu bảy trăm bảy mươi tám nghìn hai trăm sáu mươi tám đồng), đã được Hiệu trưởng phê duyệt tại Quyết định số 2158/QĐ-CĐKT ngày 29/9/2026.")
add("T", "4. Tên chủ đầu tư (cơ quan tổ chức chuẩn bị dự án): Trường Cao đẳng Kon Tum;")
add("T", "5. Nguồn vốn: " + VON + " (ngân sách trung ương và ngân sách tỉnh); cơ cấu theo kế hoạch vốn do cấp có thẩm quyền giao;")
add("T", "6. Địa điểm thực hiện: Trong khuôn viên các cơ sở thuộc Trường Cao đẳng Kon Tum;")
add("T", "7. Thời gian thực hiện: Chuẩn bị đầu tư từ Quý IV năm 2026; thời gian thực hiện dự án dự kiến giai đoạn 2026 - 2030;")
add("T", "8. Số lượng gói thầu: 02.")
add("H", "II. Căn cứ pháp lý")
CC = [
 "Luật Đấu thầu ngày 23/6/2023 (hợp nhất tại Văn bản hợp nhất số 74/VBHN-VPQH ngày 25/3/2026 của Văn phòng Quốc hội);",
 "Luật Đầu tư công ngày 29/11/2024 (hợp nhất tại Văn bản hợp nhất số 03/VBHN-VPQH ngày 05/2/2026 của Văn phòng Quốc hội);",
 "Luật Xây dựng ngày 10/12/2025;",
 "Nghị định số 214/2025/NĐ-CP ngày 04/8/2025 của Chính phủ quy định chi tiết một số điều và biện pháp thi hành Luật Đấu thầu về lựa chọn nhà thầu (hợp nhất tại Văn bản hợp nhất số 36/2026/VBHN-NĐ-BTC ngày 23/9/2026 của Bộ Tài chính);",
 "Nghị định số 358/2025/NĐ-CP ngày 31/12/2025 của Chính phủ quy định cơ chế quản lý, tổ chức thực hiện các chương trình mục tiêu quốc gia;",
 "Nghị định số 217/2026/NĐ-CP ngày 19/6/2026 của Chính phủ quy định chi tiết một số điều của Luật Xây dựng về quản lý hoạt động xây dựng;",
 "Nghị quyết số 56/NQ-HĐND ngày 28/8/2026 của Hội đồng nhân dân tỉnh Quảng Ngãi về việc phân bổ kế hoạch đầu tư công và kinh phí thường xuyên năm 2026 thực hiện Chương trình mục tiêu quốc gia xây dựng nông thôn mới, giảm nghèo bền vững và phát triển kinh tế - xã hội vùng đồng bào dân tộc thiểu số và miền núi giai đoạn 2026 - 2035, giai đoạn I: Từ năm 2026 đến năm 2030 trên địa bàn tỉnh Quảng Ngãi;",
 "Quyết định số 805/QĐ-UBND ngày 11/9/2026 của Ủy ban nhân dân tỉnh Quảng Ngãi phân bổ kế hoạch vốn đầu tư công và kinh phí thường xuyên năm 2026 thực hiện Chương trình mục tiêu quốc gia xây dựng nông thôn mới, giảm nghèo bền vững và phát triển kinh tế - xã hội vùng đồng bào dân tộc thiểu số và miền núi giai đoạn 2026 - 2035, giai đoạn I: Từ năm 2026 đến năm 2030 trên địa bàn tỉnh Quảng Ngãi (Nguồn vốn: Ngân sách trung ương và ngân sách tỉnh);",
 "Quyết định số 2161/QĐ-CĐKT ngày 29/9/2026 của Hiệu trưởng Trường Cao đẳng Kon Tum phân bổ chi tiết kế hoạch đầu tư công năm 2026 thực hiện Chương trình mục tiêu quốc gia xây dựng nông thôn mới, giảm nghèo bền vững và phát triển kinh tế - xã hội vùng đồng bào dân tộc thiểu số và miền núi giai đoạn 2026 - 2035, giai đoạn I: Từ năm 2026 đến năm 2030;",
 "Quyết định số 2158/QĐ-CĐKT ngày 29/9/2026 của Hiệu trưởng Trường Cao đẳng Kon Tum phê duyệt nhiệm vụ và dự toán kinh phí chuẩn bị đầu tư Dự án “" + TEN + "”.",
]
for i, t in enumerate(CC, 1):
    add("CC", f"{i}. {t}")
add("H", "III. Phần công việc đã thực hiện")
add("KHONG", "Không")
add("H", "IV. Phần công việc không áp dụng được một trong các hình thức lựa chọn nhà thầu")
add("KHONG", "Không")
add("H", "V. Phần công việc thuộc kế hoạch lựa chọn nhà thầu")
add("T", "1. Bảng tổng hợp phần công việc thuộc kế hoạch lựa chọn nhà thầu: Chi tiết tại Phụ lục kèm theo.")
add("T", "2. Giải trình nội dung kế hoạch lựa chọn nhà thầu:")
add("A", "a) Cơ sở phân chia các gói thầu:")
add("TX", "Cơ sở phân chia dự án thành các gói thầu, căn cứ vào:")
add("BUL", "- Nhiệm vụ và dự toán kinh phí chuẩn bị đầu tư đã được phê duyệt tại Quyết định số 2158/QĐ-CĐKT;")
add("BUL", "- Điều 37 Luật Đấu thầu về nguyên tắc lập kế hoạch lựa chọn nhà thầu; Điều 46 Nghị định số 217/2026/NĐ-CP về thuê tư vấn quản lý dự án.")
add("TX", "Việc phân dự án thành các gói thầu căn cứ vào nội dung, tính chất của công việc, trình tự thực hiện theo thời gian và theo các nguyên tắc sau:")
add("PLUS", "+ Bảo đảm tính đồng bộ về mặt kỹ thuật, không chia công việc thành các gói thầu quá nhỏ làm mất sự thống nhất, đồng bộ;")
add("PLUS", "+ Bảo đảm tiến độ thực hiện công tác chuẩn bị đầu tư và thực hiện dự án;")
add("PLUS", "+ Bảo đảm quy mô hợp lý, phù hợp năng lực nhà thầu và thị trường tư vấn.")
add("TX", "Cụ thể: Gói thầu số 1 gộp hai khoản mục khảo sát xây dựng (93.995.981 đồng) và lập Báo cáo kinh tế - kỹ thuật (758.349.265 đồng) tại Phụ lục Quyết định số 2158/QĐ-CĐKT để một nhà thầu thực hiện, bảo đảm đồng bộ giữa số liệu khảo sát và hồ sơ thiết kế; việc gộp không nhằm tránh quy định về hình thức lựa chọn nhà thầu, tổng giá gói thầu 852.345.246 đồng. Gói thầu số 2 là tư vấn quản lý dự án, thực hiện từ giai đoạn chuẩn bị đầu tư đến khi quyết toán hoàn thành dự án.")
add("A", "b) Giá gói thầu: Giá gói thầu được xác định trên cơ sở:")
add("BUL", "- Dự toán kinh phí chuẩn bị đầu tư tại Phụ lục Quyết định số 2158/QĐ-CĐKT (tư vấn quản lý dự án 667.433.022 đồng; khảo sát xây dựng 93.995.981 đồng; lập Báo cáo kinh tế - kỹ thuật 758.349.265 đồng), xác định theo phương pháp, định mức nêu tại Quyết định nêu trên;")
add("BUL", "- Khoản 2 Điều 37 Luật Đấu thầu: gói thầu số 2 có thời gian thực hiện dài hơn 01 năm nên giá gói thầu được xác định trên cơ sở toàn bộ thời gian thực hiện gói thầu.")
add("A", "c) Nguồn vốn: " + VON + " (ngân sách trung ương và ngân sách tỉnh), theo Nghị quyết số 56/NQ-HĐND, Quyết định số 805/QĐ-UBND và Quyết định số 2161/QĐ-CĐKT; việc ký kết, thanh toán hợp đồng thực hiện trong phạm vi kế hoạch vốn do cấp có thẩm quyền giao.")
add("A", "d) Hình thức và phương thức lựa chọn nhà thầu: Chỉ định thầu theo quy trình rút gọn, thực hiện theo khoản 4 Điều 78 và Điều 80 Nghị định số 214/2025/NĐ-CP (gói thầu dịch vụ tư vấn thuộc dự án có giá gói thầu không quá 03 tỷ đồng; giá hai gói thầu lần lượt 852.345.246 đồng và 667.433.022 đồng). Kế hoạch do người đứng đầu đơn vị được giao nhiệm vụ chuẩn bị dự án phê duyệt theo điểm a khoản 1 Điều 38 và khoản 2 Điều 40 Luật Đấu thầu (gói thầu thực hiện trước khi phê duyệt dự án). Chỉ định thầu rút gọn không quy định phương thức lựa chọn nhà thầu.")
add("A", "đ) Thời gian tổ chức lựa chọn nhà thầu: Thực hiện theo khoản 5 Điều 39 Luật Đấu thầu; dự kiến không quá 07 ngày.")
add("A", "e) Thời gian bắt đầu tổ chức lựa chọn nhà thầu: Quý IV năm 2026.")
add("A", "g) Loại hợp đồng: Trọn gói.")
add("A", "h) Thời gian thực hiện gói thầu: Gói thầu số 1: 120 ngày kể từ ngày hợp đồng có hiệu lực. Gói thầu số 2: không quá 03 năm, từ giai đoạn chuẩn bị đầu tư đến khi quyết toán hoàn thành dự án, phù hợp thời gian bố trí vốn thực hiện dự án nhóm C theo điểm d khoản 2 Điều 57 Luật Đầu tư công.")
add("A", "i) Tùy chọn mua thêm: Không áp dụng.")
add("A", "k) Giám sát hoạt động đấu thầu: Không đề xuất.")
add("H", "VI. Phần công việc chưa đủ điều kiện lập kế hoạch lựa chọn nhà thầu")
add("KHONG", "Không có")
# xóa đoạn "VII." cũ không cần: giữ làm tiêu đề phần VII

# ---- bảng VII ----
t1 = d.tables[1]
def cell(c, text):
    set_text(c.paragraphs[0], text)
    for extra in c.paragraphs[1:]:
        for t in extra._p.iter(docx.oxml.ns.qn("w:t")):
            t.text = ""
vals = {1: "0", 2: "0", 3: "1.519.778.268", 4: "0"}
for ri, v in vals.items():
    cell(t1.rows[ri].cells[2], v)
cell(t1.rows[5].cells[2], "1.519.778.268")
cell(t1.rows[6].cells[0], "Dự toán kinh phí chuẩn bị đầu tư (không vượt sơ bộ tổng mức đầu tư 24.894.000.000 đồng)")
cell(t1.rows[6].cells[2], "1.519.778.268")

# ---- kiến nghị ----
i_kn = find("Trên cơ sở những nội dung")
set_text(d.paragraphs[i_kn], "Trên cơ sở những nội dung nêu trên, Phòng Tổng hợp - Hành chính và Quản trị kính trình Hiệu trưởng xem xét, phê duyệt Kế hoạch lựa chọn nhà thầu giai đoạn chuẩn bị đầu tư Dự án “" + TEN + "” (kèm theo dự thảo Quyết định và Phụ lục).")

# ---- chữ ký ----
t2 = d.tables[2]
ps = t2.rows[0].cells[1].paragraphs
set_text(ps[0], "KT. TRƯỞNG PHÒNG")
set_text(ps[1], "PHÓ TRƯỞNG PHÒNG")

# ---- Phụ lục ----
set_text(d.paragraphs[find("KẾ HOẠCH LỰA CHỌN NHÀ THẦU MUA THUỐC")], "KẾ HOẠCH LỰA CHỌN NHÀ THẦU GIAI ĐOẠN CHUẨN BỊ ĐẦU TƯ DỰ ÁN")
set_text(d.paragraphs[find("TẠI PHÒNG KHÁM ĐA KHOA")], "“" + TEN.upper() + "”")

tb = d.tables[3]
src = v9.tables[2]
def uniq(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc)); out.append(c)
    return out
def v9row(i):
    return [c.text.strip() for c in uniq(src.rows[i])]
rows_src = [v9row(2), v9row(3)]
for ri, vs in zip((2, 3), rows_src):
    cs = uniq(tb.rows[ri])
    # v9 có 14 cột: TT, CĐT, tên, tóm tắt, giá, nguồn, hình thức, phương thức, TG tổ chức, bắt đầu, loại HĐ, TG thực hiện, tùy chọn, giám sát
    mapping = [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11]
    for c, k in zip(cs, mapping):
        cell(c, vs[k])
cs = uniq(tb.rows[4])
cell(cs[-1] if len(cs) < 6 else cs[1], "1.519.778.268") if False else None
for c in uniq(tb.rows[4]):
    if c.text.strip().replace(".", "").isdigit():
        cell(c, "1.519.778.268")

# ghi chú dưới bảng
last = d.paragraphs[-1]
for txt in ("(*) " + VON + "; cơ cấu ngân sách trung ương/ngân sách tỉnh theo kế hoạch vốn do cấp có thẩm quyền giao.",
            "(**) Dự án nhóm C thì thời gian bố trí vốn thực hiện dự án không quá 03 năm (Điều 57 khoản 2 điểm d Luật Đầu tư công)."):
    new = copy.deepcopy(role["T"]); last._p.addprevious(new) if False else d.element.body.insert(len(d.element.body) - 1, new)
    set_text(Paragraph(new, last._parent), txt)

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
d.save(OUT)
print("saved", OUT)
