# -*- coding: utf-8 -*-
"""Hoan thien du thao Bao cao giai ngan von DTC thang 9/2026 theo QD 64/2026/QD-UBND (Track Changes).

Goc: (thu muc Downloads) Bao cao tinh hinh thuc hien va giai ngan VĐT cong KH nam 2026_Thang 9.docx
Ra : 30-Ket-Qua/2026-10-14/Soan-Thao/BC_Giai-ngan-von-DTC-thang-9-2026_20261014_v1_TrackChanges.docx
     + ban sach (chap nhan het thay doi) ..._v1_ban-sach.docx
"""
import copy
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx.oxml.ns import qn
from ktc_trackchanges import TrackChanges, kiem_tra, doi_chieu_goc, _runs_text

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from du_an import tai_ve, ket_qua  # noqa: E402

GOC = tai_ve("Bao cao tinh hinh thuc hien va giai ngan VĐT cong KH nam 2026_Thang 9.docx")
OUT_DIR = ket_qua("2026-10-14", "Soan-Thao")
os.makedirs(OUT_DIR, exist_ok=True)
OUT = OUT_DIR + "/BC_Giai-ngan-von-DTC-thang-9-2026_20261014_v1_TrackChanges.docx"
SACH = OUT_DIR + "/BC_Giai-ngan-von-DTC-thang-9-2026_20261014_v1_ban-sach.docx"

tc = TrackChanges(GOC)


def dem(cum):
    return sum(1 for p in tc._doan() if cum in _runs_text(p))


def T(cu, moi):
    assert dem(cu) == 1, (dem(cu), cu)
    tc.thay(cu, moi)


def bo_dam(p_moi):
    for ins in p_moi.findall(qn("w:ins")):
        for r in ins.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                for tag in ("w:b", "w:bCs"):
                    for e in rpr.findall(qn(tag)):
                        rpr.remove(e)


def them(moc, text):
    assert dem(moc) == 1, (dem(moc), moc)
    p = tc.them_sau(moc, text)
    bo_dam(p)
    return p


def them_nhan(moc, nhan, noi_dung):
    """Chen doan moi: nhan in dam + noi dung thuong (giong cac gach dau dong muc 1.1)."""
    assert dem(moc) == 1, (dem(moc), moc)
    p = tc.them_sau(moc, nhan)
    ins = p.find(qn("w:ins"))
    r1 = ins.find(qn("w:r"))
    r2 = copy.deepcopy(r1)
    for t in r2.findall(qn("w:t")):
        t.text = noi_dung
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    rpr = r2.find(qn("w:rPr"))
    if rpr is not None:
        for tag in ("w:b", "w:bCs"):
            for e in rpr.findall(qn(tag)):
                rpr.remove(e)
    r1.addnext(r2)
    return p


# --- Mo dau -----------------------------------------------------------------
T("25/05/2026", "25/5/2026")
T("Quảng Ngãi. ", "Quảng Ngãi, Trường Cao đẳng Kon Tum báo cáo tình hình thực hiện và giải ngân vốn đầu tư công tháng 9 năm 2026 như sau:")
tc.xoa_doan(" Trường Cao đẳng Kon Tum báo cáo tình hình")
# doan mo dau trong ban nhap de in nghieng; bo in nghieng cho dong nhat voi than van ban
for p in tc.doc.paragraphs:
    if p.text.startswith("Thực hiện Quyết định số 64/2026"):
        for r in p._element.iter(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                for tag in ("w:i", "w:iCs"):
                    for e in rpr.findall(qn(tag)):
                        rpr.remove(e)
        break

# --- Danh so muc theo mau 01/TTKHN (so La Ma cho phan lon) -------------------
T("1. Báo cáo định kỳ", "I. Báo cáo định kỳ")
T("giai đoạn I: từ năm 2026", "giai đoạn I: Từ năm 2026")
T("1.1. Thông tin dự án", "1. Thông tin dự án")
T("1.2. Giải ngân vốn đầu tư công kế hoạch năm:", "2. Giải ngân vốn đầu tư công kế hoạch năm 2026:")

# --- 1. Thong tin du an ------------------------------------------------------
T("Việc sử dụng vốn chỉ được thực hiện sau khi có quyết định giao kế hoạch vốn của cấp có thẩm quyền.",
  "Kế hoạch vốn năm 2026 do Ủy ban nhân dân tỉnh giao tại Quyết định số 805/QĐ-UBND ngày 11/9/2026 và được Trường phân bổ chi tiết tại Quyết định số 2161/QĐ-CĐKT ngày 29/9/2026.")
them_nhan("Thời gian thực hiện dự án dự kiến", "- Tình trạng dự án:",
          " Dự án đang trong giai đoạn chuẩn bị đầu tư, chưa được cấp có thẩm quyền phê duyệt dự án.")

# --- 2. Giai ngan -------------------------------------------------------------
T("Tổng số vốn được giao trong năm 2026: ",
  "Tổng số vốn được giao trong năm 2026: 260 triệu đồng, trong đó vốn kế hoạch năm trước được phép kéo dài thời hạn giải ngân: Không có; vốn kế hoạch năm: 260 triệu đồng, cụ thể: ")
T("đến hết tháng 9: 0 đồng", "đến hết tháng 8: 0 triệu đồng")
MOC21 = "Lũy kế giải ngân vốn từ đầu năm"
them(MOC21, "+ Giải ngân vốn kế hoạch năm: 0 triệu đồng, đạt 0% kế hoạch giao (bao gồm: (i) nguồn vốn được kiểm soát, giải ngân qua Kho bạc Nhà nước: 0 triệu đồng; (ii) nguồn vốn do các đơn vị khác được giao kiểm soát, giải ngân: Không có).")
them(MOC21, "+ Giải ngân vốn kế hoạch năm trước được phép kéo dài: Không có.")
T("đến hết tháng báo cáo: 0 đồng", "đến hết tháng 9 (tháng báo cáo): 0 triệu đồng")
T("tỉnh giao.", "tỉnh giao, trong đó:")
MOC22 = "Ước lũy kế giải ngân vốn từ đầu năm"
them(MOC22, "+ Giải ngân vốn kế hoạch năm: 0 triệu đồng, đạt 0% kế hoạch giao (bao gồm: (i) nguồn vốn được kiểm soát, giải ngân qua Kho bạc Nhà nước: 0 triệu đồng; (ii) nguồn vốn do các đơn vị khác được giao kiểm soát, giải ngân: Không có).")
them(MOC22, "+ Giải ngân vốn kế hoạch năm trước được phép kéo dài: Không có.")

# --- II. Nhan xet, danh gia ----------------------------------------------------
MOC24 = "Nhận xét, đánh giá"
them(MOC24, "- Đánh giá tình hình thực hiện, giải ngân vốn đầu tư công: Kế hoạch vốn năm 2026 của Dự án được Ủy ban nhân dân tỉnh giao tại Quyết định số 805/QĐ-UBND ngày 11/9/2026; đến hết tháng 9 chưa phát sinh giải ngân (đạt 0% kế hoạch). Dự án mới được giao vốn năm 2026 nên không có số liệu so sánh cùng kỳ năm trước. Dự án thuộc nhóm C; không thuộc dự án quan trọng quốc gia, dự án trọng điểm, liên vùng, dự án khoa học, công nghệ, đổi mới sáng tạo và chuyển đổi số quốc gia.")
them(MOC24, "- Số vốn chưa phân bổ; các quyết định điều chỉnh kế hoạch vốn: Không có. Toàn bộ 260 triệu đồng kế hoạch vốn năm 2026 đã được Trường phân bổ chi tiết tại Quyết định số 2161/QĐ-CĐKT ngày 29/9/2026.")
T("Đánh giá tình hình thực hiện các Chương trình mục tiêu quốc gia.",
  "Đánh giá tình hình thực hiện các Chương trình mục tiêu quốc gia: Dự án đang trong giai đoạn chuẩn bị đầu tư. Trường đã phê duyệt nhiệm vụ và dự toán kinh phí chuẩn bị đầu tư (Quyết định số 2158/QĐ-CĐKT ngày 29/9/2026), phê duyệt kế hoạch lựa chọn nhà thầu giai đoạn chuẩn bị đầu tư (Quyết định số 2236/QĐ-CĐKT ngày 07/10/2026) và đã đăng tải kế hoạch trên Hệ thống mạng đấu thầu quốc gia; tiếp theo thực hiện chỉ định thầu rút gọn hai gói thầu tư vấn (khảo sát xây dựng, lập Báo cáo kinh tế - kỹ thuật; quản lý dự án).")
them("Những hạn chế, khó khăn, vướng mắc", "+ Nguyên nhân về cơ chế chính sách: Không.")
T(" (triển khai phân bổ kế hoạch vốn, tổ chức bồi thường hỗ trợ tái định cư, tổ chức đấu thầu lựa chọn nhà thầu thi công, chậm thực hiện thủ tục giải ngân qua Kho bạc Nhà nước...).",
  ": Kế hoạch vốn năm 2026 được Hội đồng nhân dân tỉnh phân bổ tại Nghị quyết số 56/NQ-HĐND ngày 28/8/2026 và Ủy ban nhân dân tỉnh giao tại Quyết định số 805/QĐ-UBND ngày 11/9/2026, trong khi Dự án đang ở giai đoạn chuẩn bị đầu tư; kế hoạch lựa chọn nhà thầu tư vấn mới được phê duyệt ngày 07/10/2026 nên chưa ký kết hợp đồng, chưa phát sinh khối lượng nghiệm thu, thanh toán qua Kho bạc Nhà nước.")
T("+ Nguyên nhân khác.", "+ Nguyên nhân khác: Không.")

# --- III. Giai phap va kien nghi ----------------------------------------------
T("3. Giải pháp và kiến nghị", "III. Giải pháp và kiến nghị")
T("2. Nhận xét, đánh giá", "II. Nhận xét, đánh giá")
T("đẩy nhanh giải ngân vốn đầu tư công;", "đẩy nhanh giải ngân vốn đầu tư công:")
MOC30 = "Giải pháp của đơn vị"
them(MOC30, "+ Hoàn thiện hồ sơ nghiệm thu, thanh toán ngay khi có khối lượng hoàn thành; Phòng Tài chính - Kế toán phối hợp Phòng TH-HC&QT đối chiếu số liệu giải ngân với Kho bạc Nhà nước Khu vực XV hằng tháng.")
them(MOC30, "+ Tổ chức chỉ định thầu rút gọn, hoàn thiện và ký kết hợp đồng hai gói thầu tư vấn ngay sau khi thương thảo xong; theo dõi, đôn đốc nhà thầu hoàn thành khảo sát xây dựng và lập Báo cáo kinh tế - kỹ thuật đúng tiến độ.")
T("+ Kiến nghị về cơ chế chính sách.", "+ Kiến nghị về cơ chế chính sách: Không.")
T(" (kiến nghị đối với cơ quan quản lý nhà nước, chủ đầu tư…)./.",
  ": Đề nghị Sở Tài chính, Sở Nông nghiệp và Môi trường quan tâm hướng dẫn, hỗ trợ Trường về trình tự, thủ tục thẩm định, phê duyệt dự án khi Trường hoàn thành Báo cáo kinh tế - kỹ thuật, bảo đảm giải ngân hết kế hoạch vốn năm 2026 theo quy định./.")

# --- Hoan tat trinh bay: bo to vang (danh dau cho dien cua ban nhap), giu khoi ky ten cung trang ---
body_el = tc.doc.element.body
for hl in list(body_el.iter(qn("w:highlight"))):
    hl.getparent().remove(hl)
for shd in list(body_el.iter(qn("w:shd"))):
    if shd.getparent().tag == qn("w:rPr"):
        shd.getparent().remove(shd)
pars = [p for p in tc.doc.paragraphs if p.text.strip()]
for p in pars[-2:]:
    p.paragraph_format.keep_with_next = True
for tr in tc.doc.tables[-1].rows:
    trpr = tr._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        from docx.oxml import OxmlElement
        cs = OxmlElement("w:cantSplit")
        trpr.insert(0, cs)

tc.luu(OUT)
kq = kiem_tra(OUT)
print("kiem_tra:", kq["dat"], kq["loi"])
print(tc.bao_cao())

# --- Ban sach: chap nhan het thay doi ---------------------------------------
from docx import Document
d = Document(OUT)
body = d.element.body
for dele in list(body.iter(qn("w:del"))):
    par = dele.getparent()
    if par.tag == qn("w:rPr"):          # dau doan bi xoa -> go ca doan
        p = par.getparent().getparent()
        if p.getparent() is not None:
            p.getparent().remove(p)
        continue
    par.remove(dele)
for ins in list(body.iter(qn("w:ins"))):
    par = ins.getparent()
    idx = par.index(ins)
    for ch in list(ins):
        par.insert(idx, ch)
        idx += 1
    par.remove(ins)
d.save(SACH)
print("OK", OUT, SACH)
