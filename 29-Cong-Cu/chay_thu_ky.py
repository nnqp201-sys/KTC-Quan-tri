# -*- coding: utf-8 -*-
"""Chay thu mot ky va CHAM DIEM tu dong — TANG 2.

Chay: python 29-Cong-Cu/chay_thu_ky.py            # cham diem, in bang
      python 29-Cong-Cu/chay_thu_ky.py --ghi      # cham diem + ghi 1 dong vao BANG-CHI-SO.md

Doi chieu san pham sinh ra voi VAN BAN DA BAN HANH cung ky:
  BC-375 (bao cao thang 8)  ·  PL-375 (phu luc)  ·  KH-834 (ke hoach thang 9)

Danh muc nhan muc con duoc RUT THANG TU TEP MAU, khong go tay — de khi mau
doi thi phep cham doi theo.
"""
from __future__ import annotations

import datetime
import io
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
DU_AN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NGOAI = os.path.dirname(DU_AN)
sys.path.insert(0, os.path.join(DU_AN, "29-Cong-Cu"))

import openpyxl                                        # noqa: E402
from docx import Document                              # noqa: E402

KHO = os.path.join(NGOAI, "KTC-Database", "02-KTC-Regulations")
BC375 = os.path.join(KHO, "BC-375_Bao-cao-ket-qua-thang-8-2026_20260906_v1.docx")
PL375 = os.path.join(KHO, "PL-375_Phu-luc-chi-tiet-ket-qua-cong-tac-thang-8-2026_20260906_v1.xlsx")
KH834 = os.path.join(KHO, "KH-834_Ke-hoach-cong-tac-thang-9-2026_20260906_v1.xlsx")
MAU = os.path.join(DU_AN, "25-KTC-Bao-Cao", "00. Mau bao cao thang (cap Truong).docx")
BANG = os.path.join(DU_AN, "92-Kinh-Nghiem", "02-Regression", "BANG-CHI-SO.md")


# ------------------------------------------------------------------ do
def the_thuc(p):
    d = Document(p)
    s = d.sections[0]
    f = d.styles["Normal"].font
    p0 = next((x for x in d.paragraphs
               if x.paragraph_format.first_line_indent), None)
    return (round(s.page_width.cm), round(s.top_margin.cm), round(s.bottom_margin.cm),
            round(s.left_margin.cm), round(s.right_margin.cm), f.name,
            f.size.pt if f.size else None, len(d.tables),
            round(p0.paragraph_format.first_line_indent.cm, 2) if p0 else None)


def nhan_mau() -> set:
    """Rut danh muc nhan muc con TU TEP MAU chinh thuc."""
    d = Document(MAU)
    ra = set()
    for p in d.paragraphs:
        t = " ".join(p.text.split())
        m = re.match(r"^\*\s*([^:\[]{3,60}):", t)
        if m:
            ra.add(m.group(1).strip())
    return ra


def nhan_san_pham(p) -> set:
    d = Document(p)
    ra = set()
    for x in d.paragraphs:
        t = " ".join(x.text.split())
        m = re.match(r"^\*\s*([^:]{3,70}):", t)
        if m and not m.group(1).startswith("Nghị quyết"):
            ra.add(m.group(1).strip())
    return ra


def dem_nhiem_vu(p, cot_nd=2, cot_dv=4):
    ws = openpyxl.load_workbook(p, data_only=True).active
    n = 0
    for r in ws.iter_rows(values_only=True):
        nd = r[cot_nd - 1] if len(r) >= cot_nd else None
        dv = r[cot_dv - 1] if len(r) >= cot_dv else None
        if nd and dv and len(str(nd).strip()) > 20 and not re.match(r"Trục|Các nhiệm vụ", str(nd)):
            n += 1
    return n


def cong_thuc_quy_doi(p, c_diem, c_he_so):
    ws = openpyxl.load_workbook(p, data_only=True).active
    dung = tong = 0
    for r in ws.iter_rows(values_only=True):
        if len(r) < max(c_diem, c_he_so):
            continue
        a, b = r[c_diem - 1], r[c_he_so - 1]
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            tong += 1
            if abs(b - a * 0.01) < 0.005:
                dung += 1
    return dung, tong


# ------------------------------------------------------------------ cham
def cham(bc, pl, kh) -> list:
    """Tra ve [(ten, nguong, ket_qua, dat)]"""
    ra = []

    tt_moi, tt_goc = the_thuc(bc), the_thuc(BC375)
    ra.append(("Thể thức khớp BC-375", "trùng khớp tuyệt đối",
               "khớp" if tt_moi == tt_goc else f"lệch: {tt_moi} ≠ {tt_goc}",
               tt_moi == tt_goc))

    d = Document(bc)
    n_phan = sum(1 for x in d.paragraphs
                 if x.text.strip().startswith(("I. ", "II. ", "III. ")))
    ra.append(("Số phần I/II/III", "3", str(n_phan), n_phan == 3))

    mau, sp = nhan_mau(), nhan_san_pham(bc)
    ngoai_mau = sorted(sp - mau)
    ti_le = len(sp & mau) / len(sp) if sp else 0
    ra.append(("Nhãn mục con nằm trong danh mục mẫu", "100%",
               f"{len(sp & mau)}/{len(sp)} = {ti_le:.0%}"
               + (f"; ngoài mẫu: {', '.join(ngoai_mau[:3])}…" if ngoai_mau else ""),
               ti_le == 1.0))

    from vanphong import kiem_tra
    vp = sum(len(kiem_tra(x.text)) for x in d.paragraphs)
    ra.append(("Vi phạm văn phong cấp Trường", "0", str(vp), vp == 0))

    n_pl, n_goc = dem_nhiem_vu(pl), dem_nhiem_vu(PL375)
    lech = abs(n_pl - n_goc) / n_goc if n_goc else 1
    ra.append(("Quy mô phụ lục so với PL-375", "±20%",
               f"{n_pl} so với {n_goc} → lệch {lech:.0%}", lech <= 0.20))

    n_kh, n_kh_goc = dem_nhiem_vu(kh), dem_nhiem_vu(KH834)
    lech_kh = abs(n_kh - n_kh_goc) / n_kh_goc if n_kh_goc else 1
    ra.append(("Quy mô kế hoạch so với KH-834", "±20%",
               f"{n_kh} so với {n_kh_goc} → lệch {lech_kh:.0%}", lech_kh <= 0.20))

    dung, tong = cong_thuc_quy_doi(pl, 8, 9)
    ra.append(("Công thức hệ số = điểm × 1% (phụ lục)", "100% dòng",
               f"{dung}/{tong}", tong > 0 and dung == tong))

    for f, ten in ((pl, "phụ lục"), (kh, "kế hoạch")):
        ws = openpyxl.load_workbook(f).active
        r0 = next((i for i in range(1, ws.max_row + 1)
                   if str(ws.cell(i, 1).value).strip() == "(1)"), None)
        gop = len([m for m in ws.merged_cells.ranges if r0 and m.min_row > r0])
        ra.append((f"Vùng gộp ô trong vùng dữ liệu ({ten})", "0", str(gop), gop == 0))
    return ra


def tim_san_pham(ngay: str):
    d = os.path.join(DU_AN, "30-Ket-Qua", ngay)
    if not os.path.isdir(d):
        return None
    g = lambda p: next((os.path.join(d, f) for f in sorted(os.listdir(d))
                        if f.startswith(p)), None)
    return g("BC_"), g("PL_"), g("KH_")


def main():
    ngay = next((a for a in sys.argv[1:] if re.fullmatch(r"\d{4}-\d{2}-\d{2}", a)), None)
    if not ngay:
        ds = sorted(os.listdir(os.path.join(DU_AN, "30-Ket-Qua")))
        ngay = ds[-1] if ds else None
    sp = tim_san_pham(ngay) if ngay else None
    if not sp or not all(sp):
        print(f"✗ Không tìm đủ 3 sản phẩm trong 30-Ket-Qua/{ngay}/ — chạy pipeline trước:")
        print("    python 29-Cong-Cu/trich_tuong_thuat.py && python 29-Cong-Cu/build_bc2.py "
              "&& python 29-Cong-Cu/build_xl.py")
        return 1
    bc, pl, kh = sp
    print("=" * 78)
    print(f"CHẤM ĐIỂM KỲ — sản phẩm ngày {ngay}")
    print(f"  Báo cáo   : {os.path.basename(bc)}")
    print(f"  Phụ lục   : {os.path.basename(pl)}")
    print(f"  Kế hoạch  : {os.path.basename(kh)}")
    print(f"  Đáp án    : BC-375 · PL-375 · KH-834 (đã ban hành)")
    print("=" * 78)
    kq = cham(bc, pl, kh)
    print(f"\n{'Chỉ số':46s} {'Ngưỡng':14s} {'Kết quả':28s}")
    print("─" * 96)
    for ten, nguong, gt, dat in kq:
        print(f"{'✓' if dat else '✗'} {ten:44s} {nguong:14s} {gt[:60]}")
    n_dat = sum(1 for *_, d in kq if d)
    print("─" * 96)
    print(f"ĐẠT {n_dat}/{len(kq)} chỉ số")

    if "--ghi" in sys.argv:
        ghi_bang(ngay, kq, n_dat)
    else:
        print("\n(chạy với --ghi để ghi một dòng vào BANG-CHI-SO.md)")
    return 0 if n_dat == len(kq) else 1


def ghi_bang(ngay, kq, n_dat):
    hom_nay = datetime.date.today().strftime("%d/%m/%Y")
    ver = "—"
    p = os.path.join(DU_AN, "25-KTC-Bao-Cao", "ktc-bao-cao-v3.5.skill")
    if os.path.exists(p):
        ver = "v3.5"
    dong = (f"| {hom_nay} | {ver} | **{n_dat}/{len(kq)}** | "
            + " | ".join(("✓" if d else f"✗ {g[:26]}") for *_, g, d in
                         [(a, b, c, d) for a, b, c, d in kq]) + " |")
    s = io.open(BANG, encoding="utf-8").read() if os.path.exists(BANG) else ""
    neo = "<!-- DONG-MOI -->"
    if neo in s:
        s = s.replace(neo, dong + "\n" + neo)
        io.open(BANG, "w", encoding="utf-8").write(s)
        print(f"\n✓ Đã ghi 1 dòng vào {os.path.relpath(BANG, DU_AN)}")
    else:
        print(f"\n⚠ Không thấy mốc {neo} trong BANG-CHI-SO.md — chưa ghi")


if __name__ == "__main__":
    sys.exit(main())
