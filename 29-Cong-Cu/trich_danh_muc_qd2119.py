# -*- coding: utf-8 -*-
"""Trich Danh muc san pham, cong viec CHINH THUC kem QD 2119/QD-CDKT (28/9/2026) ra CSV cho skill KPI.

QD 2119/QD-CDKT ngay 28/9/2026 cua Hieu truong ban hanh Danh muc san pham, cong viec — THAY THE danh muc du thao kem
TB 1052/TB-CDKT (DL-20260928-001). Nguon: KTC-Database/02-KTC-Regulations/02-01-.../PL-2119-QD-CDKT_*.xlsx, sheet
`Danh_muc_san_pham` (tim qua duong_dan — khong ghi cung o dia). CSV kem sha256 tep nguon de biet khi nao trich lai.

He so lay DUNG theo tung san pham. Tap he so hop le theo Nhom (do tren PL-2119): Nhom 1 = 0,3 · 0,5 · 1,0;
Nhom 2 = 1,2 · 1,5 · 2,0; Nhom 3 = 2,5; Nhom 4 = 3,5; Nhom 5 = 4,5. Cot `lech_nhom` chi DANH DAU dong co he so
ngoai tap cua Nhom — khong sua (phu luc da ban hanh la can cu).

Chay:  python 29-Cong-Cu/trich_danh_muc_qd2119.py [--ra <thu muc>] [--nguon <tep>]   (mac dinh 28-KTC-KPI/references/data/)
Kiem:  ma san pham dung dang `Truc.NoiHam.MaVB.STT`, khong trung, khong dong thieu he so; so dong khop bang goc.
"""
import argparse
import csv
import glob
import hashlib
import os
import re
import sys
import warnings

DU_AN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

TEN_CSV = "he-so-san-pham-QD2119.csv"
HE_SO_NHOM = {"Nhóm 1": (0.3, 0.5, 1.0), "Nhóm 2": (1.2, 1.5, 2.0), "Nhóm 3": (2.5,), "Nhóm 4": (3.5,),
              "Nhóm 5": (4.5,)}
MA_SP = re.compile(r"^(\d+)\.(\d+)\.([A-Z]{2}\d{2})\.(\d+)$")
COT = ["ma_san_pham", "stt", "truc", "noi_ham", "ma_vb", "ten_san_pham", "mo_ta", "loai_san_pham", "san_pham_chuan",
       "nhom", "he_so", "linh_vuc", "lech_nhom"]


def tim_nguon():
    import duong_dan
    db = duong_dan.ktc_database()
    ds = glob.glob(os.path.join(db, "02-KTC-Regulations", "**", "PL-2119-QD-CDKT_*.xlsx"), recursive=True)
    if not ds:
        raise FileNotFoundError("Không thấy phụ lục QĐ 2119 (PL-2119-QD-CDKT_*.xlsx) trong KTC-Database/02-KTC-Regulations "
                                "— dừng, hỏi người dùng.")
    return sorted(ds)[-1]


def trich(nguon):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(nguon, data_only=True, read_only=True)
    ws = wb["Danh_muc_san_pham"]
    hang, linh_vuc, tong, loi = [], "", 0, []
    for r in ws.iter_rows(min_row=5, values_only=True):
        stt, ma, ten, mo_ta, sp, nhom, hs, chuan = (list(r) + [None] * 8)[:8]
        if stt is None and ma is None and ten is None:
            continue
        if ma is None and nhom is None and hs is None:     # dong tieu de noi ham (1..38)
            linh_vuc = str(stt).strip()
            continue
        tong += 1
        ma = str(ma or "").strip()
        m = MA_SP.match(ma)
        if not m:
            loi.append(f"STT {stt}: mã sản phẩm sai dạng '{ma}'")
            continue
        nhom = str(nhom or "").strip()
        tap = HE_SO_NHOM.get(nhom)
        lech = "" if (isinstance(hs, (int, float)) and tap and any(abs(hs - x) < 1e-9 for x in tap)) \
            else f"hệ số {hs} ngoài tập hệ số {nhom} {tap}"
        hang.append({"ma_san_pham": ma, "stt": str(stt).strip(), "truc": m.group(1), "noi_ham": m.group(2),
                     "ma_vb": m.group(3), "ten_san_pham": (ten or "").strip(), "mo_ta": (mo_ta or "").strip(),
                     "loai_san_pham": (sp or "").strip(), "san_pham_chuan": str(chuan or "").strip(),
                     "nhom": nhom, "he_so": hs, "linh_vuc": linh_vuc, "lech_nhom": lech})
    return hang, tong, loi


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--ra", default=os.path.join(DU_AN, "28-KTC-KPI", "references", "data"))
    ap.add_argument("--nguon")
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    warnings.simplefilter("ignore")
    nguon = a.nguon or tim_nguon()
    hang, tong, loi = trich(nguon)
    thieu = [h["ma_san_pham"] for h in hang if h["he_so"] in (None, "")]
    trung = sorted({h["ma_san_pham"] for h in hang if sum(x["ma_san_pham"] == h["ma_san_pham"] for x in hang) > 1})
    if loi or thieu or trung or len(hang) != tong:
        print(f"✗ Trích {len(hang)}/{tong} dòng · lỗi mã: {loi[:5]} · thiếu hệ số: {thieu[:5]} · trùng mã: {trung[:5]}")
        return 1
    os.makedirs(a.ra, exist_ok=True)
    ra = os.path.join(a.ra, TEN_CSV)
    with open(ra, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COT)
        w.writeheader()
        w.writerows(hang)
    sha = hashlib.sha256(open(nguon, "rb").read()).hexdigest()
    goc = os.path.dirname(os.path.dirname(os.path.dirname(nguon)))
    with open(ra + ".nguon.txt", "w", encoding="utf-8", newline="") as f:
        f.write(f"nguon: KTC-Database/{os.path.relpath(nguon, goc).replace(os.sep, '/')}\n"
                f"sha256: {sha}\nso_dong: {len(hang)}\n"
                f"so_dong_lech_nhom: {sum(1 for h in hang if h['lech_nhom'])}\n"
                "trang_thai: CHINH THUC — Quyet dinh so 2119/QD-CDKT ngay 28/9/2026 (thay the danh muc du thao kem TB 1052)\n")
    print(f"✓ Đã trích {len(hang)} sản phẩm (khớp {tong} dòng có mã trong phụ lục) → {ra}")
    print(f"  {sum(1 for h in hang if h['lech_nhom'])} dòng hệ số ngoài tập của Nhóm · sha256 nguồn {sha[:12]}…")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
