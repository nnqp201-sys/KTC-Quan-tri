# -*- coding: utf-8 -*-
"""Tu kiem phan CAN CU va VIEN DAN cua du thao van ban hanh chinh — DL-20260919-002.

Quy tac: 20-Chuan-Chung/17-Quy-Tac-Vien-Dan.md (NĐ 30/2020 Phu luc I Phan I Muc II khoan 6,
Phu luc II Muc V khoan 7; Phap lenh 01/2012/UBTVQH13 sua doi 2026 Dieu 4; quy uoc Truong qua 897).

Chi PHAT HIEN va GOI Y muc — khong sua tep. Muc chinh thuc do ra soat KTC-Ra-Soat-897 ket luan.
Khong tra hieu luc van ban (viec do phai doi chieu kho 01-04).

Chay:  python 29-Cong-Cu/kiem_vien_dan.py <tep .docx|.md|.txt> [...]
Ma thoat: 1 neu co goi y Muc 1, nguoc lai 0.
"""
import os
import re
import sys

# Van ban Truong da bi thay the (03-Phap-Ly.md cua 897, muc 3)
DA_THAY = [
    (r"\b49/QĐ-CĐKT", "QĐ 49/QĐ-CĐKT", "QĐ 1976/QĐ-CĐKT (14/9/2026)"),
    (r"\b988/QĐ-CĐKT", "QĐ 988/QĐ-CĐKT", "QĐ 1976/QĐ-CĐKT (14/9/2026)"),
    (r"\b215/QĐ-CĐKT", "QĐ 215/QĐ-CĐKT", "QĐ 389/QĐ-CĐKT (26/02/2026)"),
    (r"\b50/QĐ-CĐKT", "QĐ 50/QĐ-CĐKT", "QĐ 1299/QĐ-CĐKT (27/5/2026)"),
    (r"\b1060/QĐ-CĐKT", "QĐ 1060/QĐ-CĐKT", "QĐ 1400/QĐ-CĐKT (12/6/2026)"),
    (r"\b340/TB-CĐCĐ", "TB 340/TB-CĐCĐ", "TB 597/TB-CĐKT (19/5/2026)"),
    (r"\b109/CĐCĐ-HCQT", "CV 109/CĐCĐ-HCQT", "TB 597/TB-CĐKT (19/5/2026)"),
]

SO_LUAT = r"số\s+\d+/\d{4}/(?:QH|UBTVQH)\d+"
LA_CAN_CU = re.compile(r"^\s*[-–—•+*]?\s*Căn cứ\b")
GACH_DAU = re.compile(r"^\s*[-–—•+*]\s*Căn cứ\b")
LA_XET = re.compile(r"^\s*Xét\b")
DAU_VB = 60  # so doan dau van ban chua khoi can cu ban hanh


def doc_tep(p):
    """Tra ve danh sach doan van (paragraph) — .docx doc bang python-docx."""
    if p.lower().endswith(".docx"):
        import docx
        d = docx.Document(p)
        doan = [x.text for x in d.paragraphs]
        for bang in d.tables:           # tieu de van ban thuong nam trong bang
            for hang in bang.rows:
                for o in hang.cells:
                    doan.extend(x.text for x in o.paragraphs)
        return doan
    with open(p, encoding="utf-8") as f:
        return f.read().splitlines()


def _la_qd_hieu_truong(doan):
    """Quyet dinh cua Hieu truong: co ten loai QUYẾT ĐỊNH va tham quyen HIỆU TRƯỞNG."""
    van = "\n".join(doan)
    return bool(re.search(r"^\s*QUYẾT ĐỊNH\s*$", van, re.M)) and "HIỆU TRƯỞNG" in van


def kiem_tra(doan):
    """doan: list[str]. Tra ve list[(muc, ma, so_dong, trich, mo_ta)] — so_dong tinh tu 1."""
    kq = []

    def them(muc, ma, i, mo_ta):
        kq.append((muc, ma, i + 1, doan[i].strip()[:90], mo_ta))

    # Khoi can cu ban hanh: nam o DAU van ban (DAU_VB doan dau). "Căn cứ …" sau do la cau van trong noi
    # dung (vd. So tay "Căn cứ vào báo cáo…") — khong ap quy tac dong can cu. Thu that 19/9/2026 tren kho 02.
    khoi = [i for i, s in enumerate(doan[:DAU_VB]) if LA_CAN_CU.match(s) or LA_XET.match(s)]
    can_cu = [i for i in khoi if LA_CAN_CU.match(doan[i])]

    for i in can_cu:
        s = doan[i]
        la_luat = re.search(r"Căn cứ\s+(Bộ luật|Luật|Pháp lệnh)\s+[A-ZĐ]", s)
        if la_luat and not re.search(r"ngày\s+\d{1,2}(\s+tháng\s+\d{1,2}\s+năm\s+|/)\d", s):
            them(2, "VD02", i, "Căn cứ Luật/Pháp lệnh ghi tên và ngày ban hành — dòng này thiếu ngày "
                 "(897, 02-Noi-Dung mục 1 dòng 4)")
        if re.search(r"Căn cứ\s+(các\s+)?Văn bản hợp nhất", s):
            them(3, "VD03", i, "VBHN đứng làm căn cứ chính — ghi văn bản gốc trước, VBHN đặt trong ngoặc "
                 "“(hợp nhất tại Văn bản hợp nhất số …)”")
        if re.search(r"Luật số\s+\d+/\d{4}/QH\d+", s) and not re.search(r"Luật\s+[A-ZĐ]\w", s):
            them(2, "VD05", i, "Chỉ dẫn luật sửa đổi, thiếu tên luật gốc được sửa đổi, bổ sung")
        if GACH_DAU.match(s):
            them(2, "VD07", i, "Gạch đầu dòng trước “Căn cứ” là quy tắc văn bản Đảng — văn bản hành chính "
                 "không dùng")

    # VD06 — dau cuoi dong trong khoi can cu
    if khoi:
        # chi xet cum lien tiep dau tien (bo qua dong trong giua cac can cu)
        cum = [khoi[0]]
        for i in khoi[1:]:
            if all(not doan[j].strip() for j in range(cum[-1] + 1, i)):
                cum.append(i)
            else:
                break
        for k, i in enumerate(cum):
            cuoi = doan[i].rstrip()
            can = "." if k == len(cum) - 1 else ";"
            if not cuoi.endswith(can):
                them(2, "VD06", i, f"Cuối dòng căn cứ phải là “{can}” (dòng giữa “;”, dòng cuối “.”)")

    # VD08 — QD cua Hieu truong: can cu dau tien la QD 1976
    # Dan doi truoc cua QD 1976 (49, 988) o vi tri dau: dung vi tri, chi co the sai doi -> de VD09 bao
    if can_cu and _la_qd_hieu_truong(doan) and not re.search(
            r"\b(1976|988|49)/QĐ-CĐKT", doan[can_cu[0]]):
        them(1, "VD08", can_cu[0], "Quyết định của Hiệu trưởng: căn cứ đầu tiên phải là Quyết định số "
             "1976/QĐ-CĐKT ngày 14/9/2026 (897, 08-Quy-Uoc-Rieng-CDKT mục 1) — áp cho dự thảo ban hành "
             "từ 14/9/2026; văn bản cũ đối chiếu văn bản hiệu lực tại ngày ban hành")

    for i, s in enumerate(doan):
        if re.search(r"(Bộ luật|Luật|Pháp lệnh)\s+[A-ZĐ][^;\n]{0,120}?" + SO_LUAT, s):
            them(2, "VD01", i, "Văn bản hành chính: viện dẫn Luật/Pháp lệnh không ghi số hiệu, kể cả khi đã "
                 "có VBHN — chỉ ghi tên (và ngày ban hành ở phần căn cứ). NĐ 30 PL I, Phần I, Mục II, "
                 "khoản 6; 897 02-Noi-Dung dòng 4")
        if re.search(r"Pháp lệnh[^;.\n]{0,160}?của Quốc hội", s) and "Thường vụ" not in s:
            them(2, "VD04", i, "Pháp lệnh do Ủy ban Thường vụ Quốc hội ban hành, không phải Quốc hội")
        for mau, cu, moi in DA_THAY:
            # "…thay thế/bãi bỏ Quyết định số 215/QĐ-CĐKT" la dieu khoan bai bo, khong phai vien dan
            if re.search(mau, s) and not re.search(
                    r"(thay thế|bãi bỏ|hết hiệu lực)[^.;]{0,80}?" + mau, s):
                them(1, "VD09", i, f"{cu} đã bị thay thế bởi {moi} — chỉ hợp lệ nếu dự thảo ban hành "
                     "trước ngày thay thế")
        if LA_CAN_CU.match(s) and re.search(r"checklist|Checklist|\.md\b|\.docx\b|biểu mẫu nội bộ|"
                                           r"KTC-Ra-Soat|Skill-Library", s):
            them(1, "VD10", i, "Không dẫn mẫu/checklist/tệp nội bộ làm căn cứ pháp lý")
        for m in re.finditer(r"(?:khoản\s+\d+\s+|điểm\s+[a-zđ]\s+(?:khoản\s+\d+\s+)?)?"
                             r"\b(điều|chương|tiểu mục)\s+(\d+|[IVXL]+)\b", s):
            them(3, "VD11", i, f"Viết hoa “{m.group(1).capitalize()} {m.group(2)}” khi viện dẫn "
                 "(NĐ 30, Phụ lục II, Mục V, khoản 7)")
            break

    # VD12 — cung so ky hieu xuat hien lai kem trich yeu day du (sau lan dau)
    da_gap = {}
    for i, s in enumerate(doan):
        for m in re.finditer(r"(?:Quyết định|Nghị định|Thông tư|Thông báo|Kế hoạch)\s+số\s+"
                             r"(\d+/[\w\-/Đ]+)\s+ngày[^;.\n]{0,40}?(?:của|về|ban hành|quy định)", s):
            so = m.group(1)
            if so in da_gap and not LA_CAN_CU.match(s):
                them(4, "VD12", i, f"{so} đã viện dẫn đầy đủ ở dòng {da_gap[so] + 1} — lần sau chỉ ghi tên "
                     "loại và số, ký hiệu (NĐ 30, Phụ lục I, Phần I, Mục II, khoản 6 điểm b)")
            else:
                da_gap.setdefault(so, i)
    kq.sort(key=lambda x: (x[0], x[2]))
    return kq


def in_bao_cao(ten, kq):
    print(f"\n=== {ten} — {len(kq)} gợi ý ===")
    if not kq:
        print("  ✓ Không phát hiện lỗi viện dẫn theo 12 phép kiểm VD01–VD12.")
    for muc, ma, dong, trich, mo_ta in kq:
        print(f"  [Mức {muc}] {ma} dòng {dong}: {mo_ta}\n           “{trich}”")
    print("  Lưu ý: công cụ không tra hiệu lực văn bản — đối chiếu kho 01-04 và rà soát 897 vẫn bắt buộc.")


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    co_muc_1 = False
    for p in argv:
        if not os.path.isfile(p):
            print(f"✗ không thấy tệp {p}")
            return 2
        kq = kiem_tra(doc_tep(p))
        in_bao_cao(p, kq)
        co_muc_1 |= any(x[0] == 1 for x in kq)
    return 1 if co_muc_1 else 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
