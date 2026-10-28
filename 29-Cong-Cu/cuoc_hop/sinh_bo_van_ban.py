# -*- coding: utf-8 -*-
"""Sinh tron bo van ban sau cuoc hop gop y De an tu MOT bang y kien (yaml) — 28/10/2026.

    python 29-Cong-Cu/cuoc_hop/sinh_bo_van_ban.py <bang-y-kien.yaml> <mau_09.docx> <mau_bc_tiep_thu.docx> <thu_muc_ra>

  <mau_09.docx>          ban lam viec tao bang: kiem_the_thuc.py --tao "<KTC-Database>/03-Templates(1)/09-Bien-ban.dotx" <dich>
                         (chay LENH RIENG — hook guard chan lenh vua nhac kho vua ghi)
  <mau_bc_tiep_thu.docx> ban chep BC 344/BC-CDKT (kho 05, De an Dang kiem lan 3) — chep bang Copy-Item, lenh rieng
Dau ra: BB-Hop-gop-y-<ma>_<ngay>.docx (Truong ban hanh) + BC-To-XDDA_Tiep-thu-giai-trinh-<ma>_<ngay>.docx (To ban hanh).
O chua co du lieu (null trong yaml) ghi "…" va liet ke cuoi chuong trinh — hoi anh roi chay lai.
Mau yaml: 29-Cong-Cu/cuoc_hop/mau-bang-y-kien.yaml
"""
import copy
import os
import sys

import yaml
import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rap_mau import dat, xoa, nhan_ban, nhan_ban_dong, ke_duoi_doan, bo_trang_trang_cuoi  # noqa: E402

CHO = "…"
thieu = []
NHAN = {"TT": "Tiếp thu ý kiến:", "GT": "Giải trình:", "KTT": "Không tiếp thu:"}


def v(x, ten):
    if x in (None, ""):
        thieu.append(ten)
        return CHO
    return str(x)


def ten_nguoi(n):
    s = f"{n['xung_ho']} {n['ten']}"
    return s + (f", {n['don_vi']}" if n.get("don_vi") else "")


# ------------------------------------------------------------------ bien ban (mau 09)
def bien_ban(Y, mau, out):
    h = Y["cuoc_hop"]; de_an = Y["de_an"]
    d = docx.Document(mau)
    dd, mm, yy = h["ngay"].split("/")
    dat(d.tables[0].rows[1].cells[1].paragraphs[0], ["Quảng Ngãi", ", ngày ", dd, " ", "tháng ", mm, " năm 20", yy[2:]], False)
    P = d.paragraphs
    tieu_de = f"Họp góp ý dự thảo {de_an['ten']} ({de_an['lan']})"
    dat(P[1], [tieu_de], False)
    dat(P[3], ["Thời gian bắt đầu:", " ", f"{v(h.get('gio_bat_dau'), 'giờ bắt đầu')}, ngày {h['ngay']}", "."], False)
    dat(P[4], ["Địa điểm:", " ", v(h.get("dia_diem"), "địa điểm"), "."], False)
    tp = h.get("thanh_phan") or []
    sau = P[5]
    for t in tp:
        sau = nhan_ban(P[6], sau, ["- ", t, ";"])
    for q in P[6:9]:
        xoa(q)
    dat(P[9], ["*", " Chủ trì:", f" {h['chu_tri']}."])
    dat(P[10], ["*", " Thư ký: ", v(h.get("thu_ky"), "thư ký"), "."], False)
    dat(P[11], ["Nội dung cuộc họp:", " ", h["noi_dung"], ",", " cụ thể như sau:"])
    muc = 0
    sau = P[11]
    def tieu_muc(text):
        nonlocal muc, sau
        muc += 1
        sau = nhan_ban(P[12], sau, [f"{muc}.", " ", text, ":"])
    def dong(i, text, gach=False):
        nonlocal sau
        sau = nhan_ban(P[15], sau, ["- " if gach else f"({i}) ", text.rstrip("."), "."])
    if Y.get("bao_cao_tom_tat"):
        tieu_muc(f"{Y['bao_cao_tom_tat']['nguoi']} báo cáo tóm tắt dự thảo")
        for i, t in enumerate(Y["bao_cao_tom_tat"]["noi_dung"], 1):
            dong(i, t)
    tieu_muc("Các thành viên dự họp thảo luận, cho ý kiến về dự thảo, cụ thể như sau")
    for i, n in enumerate(Y["y_kien"], 1):
        ys = [y["noi_dung"].rstrip(".") for y in n["y"]]
        ys = [ys[0]] + [t[:1].lower() + t[1:] for t in ys[1:]]
        ten = ten_nguoi(n)
        dong(i, ten[:1].upper() + ten[1:] + ": " + "; ".join(ys))
    gt = [(n, y) for n in Y["y_kien"] for y in n["y"] if y["xu_ly"] == "GT"]
    if gt:
        tieu_muc(f"{Y['chu_the_tiep_thu']} giải trình")
        for i, (n, y) in enumerate(gt, 1):
            dong(i, y["tra_loi"])
    kl = Y.get("ket_luan") or []
    if kl:
        muc += 1
        sau = nhan_ban(P[24], sau, [f"{muc}", ".", " ", h["chu_tri"], ", chủ trì cuộc họp ", "kết luận nội dung cuộc họp như sau:"])
        for k in kl:
            dong(0, k["noi_dung"], gach=True)
    for q in P[12:27]:
        xoa(q)
    dat(P[27], ["Cuộc họp", " kết thúc vào lúc ", v(h.get("gio_ket_thuc"), "giờ kết thúc"), "", "", "", ", ngày ", h["ngay"],
                "", "", "", "", ".", "/."], False)
    S = d.tables[1]
    dat(S.rows[1].cells[0].paragraphs[1], [v(h.get("thu_ky"), "thư ký")], False)
    dat(S.rows[1].cells[1].paragraphs[1], [h.get("chu_tri_ho_ten") or v(None, "họ tên chủ trì")], False)
    ps = [p for p in S.rows[2].cells[0].paragraphs if p.text.startswith("-")]
    nn = h.get("noi_nhan_bb") or ["- Lãnh đạo Trường;", "- Thành viên Tổ xây dựng Đề án;", "- Lưu: VT, THHCQT."]
    for p, t in zip(ps, nn):
        dat(p, [t], False)
        for r in p.runs:
            r.bold = False
    bo_trang_trang_cuoi(d)
    d.save(out)


# ------------------------------------------------------------------ BC tiep thu (BC 344, bang 3 cot)
def bc_tiep_thu(Y, mau, out):
    h = Y["cuoc_hop"]; de_an = Y["de_an"]; to = Y["to"]
    d = docx.Document(mau)
    H = d.tables[0]
    dat(H.rows[0].cells[0].paragraphs[0], ["TRƯỜNG CAO ĐẲNG ", "KON TUM"], False)
    dat(H.rows[0].cells[0].paragraphs[1], ["", to["ten_hoa"]], False)
    ke_duoi_doan(H.rows[0].cells[0].paragraphs[1], rong_o_cm=8.5)
    dat(H.rows[1].cells[0].paragraphs[0], ["Số:  ", "      ", f" /{to['ky_hieu']}"], False)
    dat(H.rows[1].cells[2].paragraphs[0], ["Quảng Ngãi", ", ngày", "      ", "", " ", "tháng ", h["ngay"].split("/")[1], " năm 202",
                                           h["ngay"][-1]], False)
    P = d.paragraphs
    t1, t2 = de_an["ten_dong_1"], de_an["ten_dong_2"]
    dat(P[2], ["tiếp thu, giải trình ý kiến góp ý ", f"dự thảo {t1}"], False)
    dat(P[3], [t2], False)
    dat(P[4], [f"({de_an['lan']})"], False)
    i_kg = [i for i, p in enumerate(P) if p.text.startswith("Kính gửi")][0]
    p = P[i_kg]
    dat(p, ["Kính gửi: ", f"{Y['kinh_gui']}."], False)
    for q in P[i_kg + 1:i_kg + 8]:
        xoa(q)
    i_md = [i for i, p in enumerate(P) if p.text.startswith("Ngày ")][0]
    dat(P[i_md], [f"Tại cuộc họp ngày {h['ngay']}, các thành viên {to['ten']} (thành lập theo {to['qd']}) đã tham gia ý kiến góp ý "
                  f"dự thảo {de_an['ten']} ({de_an['lan']}); {h['chu_tri_ngan']} kết luận. Trên cơ sở ý kiến góp ý và kết luận của chủ trì, "
                  f"{to['ten']} báo cáo tiếp thu, giải trình cụ thể như sau:"])
    i_kt = [i for i, p in enumerate(P) if p.text.startswith("Trên đây")][0]
    dat(P[i_kt], [f"Trên đây là Báo cáo tiếp thu, giải trình ý kiến góp ý dự thảo {de_an['ten']} ({de_an['lan']}); {to['ten']} tiếp tục "
                  "chỉnh lý, hoàn thiện hồ sơ Đề án theo nội dung Báo cáo này, trình Hiệu trưởng xem xét, phê duyệt./."])
    T = d.tables[1]
    dat(T.rows[0].cells[2].paragraphs[0], [f"NỘI DUNG TIẾP THU, GIẢI TRÌNH CỦA {to['ten_hoa']}"], False)
    r_nhom, r_muc = copy.deepcopy(T.rows[1]._tr), copy.deepcopy(T.rows[2]._tr)
    for r in list(T.rows)[1:]:
        T._tbl.remove(r._tr)
    nhom_list = [(ten_nguoi(n) and f"Ý kiến của {ten_nguoi(n)}", n["y"]) for n in Y["y_kien"]]
    if Y.get("ket_luan"):
        nhom_list.append((f"Kết luận của {h['chu_tri_ngan']}", [dict(k, xu_ly="TT") for k in Y["ket_luan"]]))
    for i, (ten, ys) in enumerate(nhom_list, 1):
        cs = nhan_ban_dong(T, r_nhom)
        dat(cs[0].paragraphs[0], [str(i)])
        for c in cs[1:]:
            dat(c.paragraphs[0], [ten])
        for j, y in enumerate(ys, 1):
            cs = nhan_ban_dong(T, r_muc)
            dat(cs[0].paragraphs[0], [f"{i}.{j}"])
            for q in cs[1].paragraphs[:-1]:
                xoa(q)
            dat(cs[1].paragraphs[0], [y["noi_dung"]])
            ps = cs[2].paragraphs
            dat(ps[0], [NHAN[y["xu_ly"]]])
            dat(ps[1], [y.get("tra_loi") or v(None, f"trả lời ý {i}.{j}")])
            for r in ps[1].runs:
                r.bold = False
            for q in ps[2:]:
                xoa(q)
    S = d.tables[2].rows[0].cells
    ps = [p for p in S[0].paragraphs if p.text.startswith("-")]
    for k, p in enumerate(ps):
        if k < len(to["noi_nhan"]):
            dat(p, [to["noi_nhan"][k]], False)
        else:
            xoa(p)
    sp = S[1].paragraphs
    dat(sp[0], [f"TM. {to['ten_hoa']}"], False)
    e = copy.deepcopy(sp[0]._element); sp[0]._element.addnext(e)
    dat(docx.text.paragraph.Paragraph(e, sp[0]._parent), ["TỔ TRƯỞNG"], False)
    xoa(S[1].paragraphs[2])
    dat(S[1].paragraphs[-1], [to.get("to_truong") or v(None, "họ tên Tổ trưởng")], False)
    from docx.shared import Cm
    ST = d.tables[2]
    for gc, w in zip(ST._tbl.tblGrid.findall(qn("w:gridCol")), (7513, 5098)):
        gc.set(qn("w:w"), str(w))
    for c, w in zip(ST.rows[0].cells, (13.25, 9.0)):
        c.width = Cm(w)
    bo_trang_trang_cuoi(d)
    d.save(out)


def main(argv):
    if len(argv) != 4:
        print(__doc__); return 2
    Y = yaml.safe_load(open(argv[0], encoding="utf-8"))
    os.makedirs(argv[3], exist_ok=True)
    ma, ngay = Y["ma"], Y["cuoc_hop"]["ngay"].split("/")
    hau = f"{ngay[2]}{ngay[1]}{ngay[0]}"
    f1 = os.path.join(argv[3], f"BB-Hop-gop-y-{ma}_{hau}.docx")
    f2 = os.path.join(argv[3], f"BC-To-XDDA_Tiep-thu-giai-trinh-{ma}_{hau}.docx")
    bien_ban(Y, argv[1], f1)
    bc_tiep_thu(Y, argv[2], f2)
    print("Đã tạo:\n ", f1, "\n ", f2)
    if thieu:
        print("Ô còn trống (hỏi anh rồi chạy lại):", ", ".join(dict.fromkeys(thieu)))
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
