# -*- coding: utf-8 -*-
"""Xuat tep .xlsx dung Mau import cua PHAN MEM KPI tinh (VNPT-DGCB, kpi.quangngai.dcs.vn).

Dung cho buoc "Dang ky ke hoach cong tac" dau chu ky: Quan ly cong viec / Ca nhan CBCCVC /
Dang ky ke hoach cong tac → [Them tu file]. Mau do phan mem quy dinh, KHONG duoc doi so cot:
  dong 1    Don vi ap dung: <ten don vi>
  dong 2    Danh sach ke hoach thuc hien nhiem vu
  dong 3-6  ghi chu cua mau (du lieu dong 1-9 khong import vao he thong)
  dong 8-9  tieu de gop: TT(*) | San pham/cong viec(*) | So luong giao(*) | Mo ta
  dong 10+  du lieu; cot B phai ghi "<ma> - <ten>" dung het ky tu trong danh muc

Nguon ma: 11-Du-lieu-Cong-Viec/DANH MUC SAN PHAM CONG VIEC/danh-muc-spcv-pm-kpi.csv
(sinh boi 29-Cong-Cu/trich_danh_muc_pm_kpi.py). Ma khong co trong danh muc -> DUNG, khong tu dat ma.

Chay:
  # bang trong de don vi tu dien so luong (nhu mau goc)
  python 29-Cong-Cu/xuat_mau_import_kpi.py --don-vi "Khoa các Khoa học cơ bản" --ra <tep.xlsx>
  # tu ke hoach da co: CSV 3 cot ma_spcv,so_luong,mo_ta
  python 29-Cong-Cu/xuat_mau_import_kpi.py --don-vi "Phòng TH-HC&QT" --ke-hoach kh.csv --ra <tep.xlsx>
  # loc theo nhom hoac truc khi xuat bang trong
  python 29-Cong-Cu/xuat_mau_import_kpi.py --don-vi "..." --nhom N1,N2 --ra <tep.xlsx>
"""
import argparse
import csv
import os
import sys

DU_AN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DANH_MUC = os.path.join(DU_AN, "11-Du-lieu-Cong-Viec", "DANH MUC SAN PHAM CONG VIEC",
                        "danh-muc-spcv-pm-kpi.csv")
GHI_CHU = ["- Số cột trên hệ thống không thay đổi",
           "- Dữ liệu từ dòng 1 đến dòng 9 không được import vào hệ thống",
           "- Các cột thông tin có dấu (*) là bắt buộc nhập",
           "- Sản phẩm/công việc chọn từ danh sách mẫu"]
TIEU_DE = ["TT(*)", "Sản phẩm/công việc(*)", "Số lượng giao(*)", "Mô tả"]
RONG = {"A": 5.71, "B": 40.71, "C": 15.71, "D": 30.71}


def doc_danh_muc(duong_dan=DANH_MUC):
    if not os.path.isfile(duong_dan):
        raise FileNotFoundError(
            f"Không thấy danh mục {duong_dan} — chạy 29-Cong-Cu/trich_danh_muc_pm_kpi.py trước.")
    with open(duong_dan, encoding="utf-8-sig", newline="") as f:
        return {r["ma_spcv"]: r for r in csv.DictReader(f)}


def doc_ke_hoach(duong_dan):
    """CSV ke hoach cua don vi: cot ma_spcv (bat buoc), so_luong, mo_ta."""
    with open(duong_dan, encoding="utf-8-sig", newline="") as f:
        ds = list(csv.DictReader(f))
    if not ds or "ma_spcv" not in ds[0]:
        raise ValueError(f"{duong_dan} phải có cột 'ma_spcv' (và tùy chọn 'so_luong', 'mo_ta').")
    return [(r["ma_spcv"].strip(), (r.get("so_luong") or "").strip(), (r.get("mo_ta") or "").strip())
            for r in ds if (r.get("ma_spcv") or "").strip()]


def xuat(don_vi, dong, ra, tieu_de_bang="Danh sách kế hoạch thực hiện nhiệm vụ"):
    """dong: danh sach (ma, ten, so_luong, mo_ta) — ghi dung bo cuc mau cua phan mem."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Mẫu import"
    tnr = lambda **k: Font(name="Times New Roman", size=k.pop("sz", 11), **k)
    vien = Border(*[Side(style="thin")] * 4)

    ws["A1"] = f"Đơn vị áp dụng: {don_vi}"
    ws["A2"] = tieu_de_bang
    for i, g in enumerate(GHI_CHU, 3):
        ws.cell(i, 1, g)
    for r in range(1, 7):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(r, 1)
        c.font = tnr(bold=(r == 2), sz=16 if r == 2 else 11)
        c.alignment = Alignment(horizontal="center" if r == 2 else "left", wrap_text=True)

    for j, t in enumerate(TIEU_DE, 1):
        ws.merge_cells(start_row=8, start_column=j, end_row=9, end_column=j)
        for r in (8, 9):
            c = ws.cell(r, j)
            c.border = vien
        c = ws.cell(8, j, t)
        c.font = tnr()
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (ma, ten, sl, mo_ta) in enumerate(dong, 1):
        r = 9 + i
        for j, v in enumerate((i, f"{ma} - {ten}", sl, mo_ta), 1):
            c = ws.cell(r, j, v if v != "" else None)
            c.font = tnr()
            c.border = vien
            c.alignment = Alignment(horizontal="center" if j in (1, 3) else "left",
                                    vertical="center", wrap_text=(j != 1))

    for k, v in RONG.items():
        ws.column_dimensions[k].width = v
    os.makedirs(os.path.dirname(os.path.abspath(ra)) or ".", exist_ok=True)
    wb.save(ra)


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--don-vi", required=True, help="tên đơn vị áp dụng, ghi đúng như trên phần mềm")
    ap.add_argument("--ra", required=True, help="tệp .xlsx xuất ra")
    ap.add_argument("--ke-hoach", help="CSV kế hoạch của đơn vị: ma_spcv, so_luong, mo_ta")
    ap.add_argument("--nhom", help="lọc theo nhóm phần mềm khi xuất bảng trống, ví dụ N1,N2")
    ap.add_argument("--truc", help="lọc theo Trục, ví dụ 1,3")
    ap.add_argument("--danh-muc", default=DANH_MUC)
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    dm = doc_danh_muc(a.danh_muc)
    if a.ke_hoach:
        kh = doc_ke_hoach(a.ke_hoach)
        thieu = [m for m, _, _ in kh if m not in dm]
        if thieu:
            print(f"✗ {len(thieu)} mã không có trong danh mục phần mềm: {thieu[:10]} — "
                  "dừng, không tự đặt mã. Kiểm lại hoặc đề nghị P-TCCB bổ sung danh mục.")
            return 1
        dong = [(m, dm[m]["ten_spcv"], sl, mo_ta) for m, sl, mo_ta in kh]
    else:
        ds = list(dm.values())
        if a.nhom:
            giu = {x.strip().upper() for x in a.nhom.split(",")}
            ds = [r for r in ds if r["nhom_pm"] in giu]
        if a.truc:
            giu = {x.strip() for x in a.truc.split(",")}
            ds = [r for r in ds if r["truc"] in giu]
        if not ds:
            print("✗ Bộ lọc không còn sản phẩm nào — kiểm lại --nhom/--truc.")
            return 1
        dong = [(r["ma_spcv"], r["ten_spcv"], "", "") for r in ds]

    xuat(a.don_vi, dong, a.ra)

    thieu_hs = [m for m, _, _, _ in dong if not dm[m]["he_so"]]
    tong = 0.0
    for m, _, sl, _ in dong:
        try:
            tong += float(sl) * float(dm[m]["he_so"])
        except (TypeError, ValueError):
            pass
    print(f"✓ Đã xuất {len(dong)} dòng cho \"{a.don_vi}\" → {a.ra}")
    if a.ke_hoach:
        print(f"  Tổng điểm quy đổi (số lượng × hệ số QĐ 2119): {tong:.2f}")
    if thieu_hs:
        print(f"  ⚠ {len(thieu_hs)} mã chưa có hệ số (THIEU_DU_LIEU), không tính vào tổng: "
              f"{thieu_hs[:8]}{'…' if len(thieu_hs) > 8 else ''}")
    print("  Nạp lên: Quản lý công việc / Cá nhân CBCCVC / Đăng ký kế hoạch công tác → [Thêm từ file].")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
